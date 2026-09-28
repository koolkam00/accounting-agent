"""Tests for the synthetic deal-package generator (scripts/qoe_generate_deals.py, scripts/qoe_synth/).

The checks read generated files back with small independent readers (not
qoe/ingest.py) so the package is verified against the SPEC §3 formats
themselves. Unit tests use a tiny in-memory spec; the Meridian tests generate
dev deal 1 from its YAML spec into a temp dir.
"""

from __future__ import annotations

import copy
import csv
import importlib.util
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from qoe.money import D, ZERO, q2
from qoe.pdf_text import canonicalize_page_text, extract_pdf_pages
from qoe.periods import labels_for_month, month_range
from qoe.schemas import DealMeta, GroundTruth, PeriodDef, Treatment

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "scripts"
MERIDIAN_SPEC = ROOT / "data" / "qoe" / "specs" / "meridian_mechanical.yaml"
MERIDIAN_COMMITTED = ROOT / "data" / "qoe" / "dev" / "meridian_mechanical"
FOOTER = "SYNTHETIC — generated for QoE Evidence Review testing"

pytest.importorskip("reportlab")
import sys  # noqa: E402

sys.path.insert(0, str(SCRIPTS))
from qoe_synth import GenerationError, generate_deal, load_spec  # noqa: E402
from qoe_synth.spec import DealSpec  # noqa: E402


# ---------------------------------------------------------------------------
# Independent readers
# ---------------------------------------------------------------------------

INCOME_TYPES = {"income", "revenue", "sales", "other income", "otherincome"}
ACCOUNT_ROW = re.compile(r"^\s*(\d{3,6})\s*[·\-–:]?\s*(.+?)\s*$")
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


@dataclass(frozen=True)
class Row:
    account: str
    day: date
    amount: Decimal  # debit-positive
    num: str
    counterparty: str
    memo: str

    @property
    def month(self) -> str:
        return f"{self.day.year:04d}-{self.day.month:02d}"


def read_coa_types(path: Path) -> dict[str, str]:
    with path.open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    header = [h.strip().lstrip("*").lower() for h in rows[0]]
    num_i = next(i for i, h in enumerate(header) if h in ("account #", "number", "code"))
    type_i = next(i for i, h in enumerate(header) if h in ("type", "account type"))
    return {r[num_i]: r[type_i] for r in rows[1:] if r}


def read_qbo(path: Path, types: dict[str, str]) -> dict[int, Row]:
    out: dict[int, Row] = {}
    account = None
    with path.open(newline="", encoding="utf-8") as f:
        for line, rec in enumerate(csv.reader(f), start=1):
            if line <= 5 or not rec:
                continue
            if rec[0]:
                account = None if rec[0].startswith(("Total for ", "TOTAL")) else rec[0].split(" ", 1)[0]
                continue
            natural = D(rec[7])
            dp = -natural if types[account].strip().lower() in INCOME_TYPES else natural
            out[line] = Row(account, datetime.strptime(rec[1], "%m/%d/%Y").date(), dp, rec[3], rec[4], rec[5])
    return out


def read_netsuite(path: Path) -> dict[int, Row]:
    out: dict[int, Row] = {}
    with path.open(newline="", encoding="utf-8") as f:
        for line, rec in enumerate(csv.DictReader(f), start=2):
            out[line] = Row(
                rec["Account"].split(" ", 1)[0],
                datetime.strptime(rec["Date"], "%m/%d/%Y").date(),
                D(rec["Debit"]) - D(rec["Credit"]),
                rec["Document Number"],
                rec["Name"],
                rec["Memo"],
            )
    return out


def read_xero(path: Path) -> dict[int, Row]:
    ws = load_workbook(path, data_only=True)["Account Transactions"]
    out: dict[int, Row] = {}
    for r, rec in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        contact, _, text = str(rec[2]).partition(" - ")
        if not text:
            contact, text = "", contact
        out[r] = Row(str(rec[7]), rec[0].date(), D(rec[4]) - D(rec[5]), rec[3] or "", contact, text)
    return out


def read_gl(deal_dir: Path) -> dict[int, Row]:
    meta = yaml.safe_load((deal_dir / "deal.yaml").read_text())
    gl = deal_dir / meta["files"]["gl"]
    if gl.suffix == ".xlsx":
        return read_xero(gl)
    with gl.open(encoding="utf-8") as f:
        first = f.readline()
    if first.startswith("Internal ID"):
        return read_netsuite(gl)
    return read_qbo(gl, read_coa_types(deal_dir / meta["files"]["chart_of_accounts"]))


def _month_cell(value: object) -> str | None:
    if isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}"
    if not isinstance(value, str):
        return None
    m = re.fullmatch(r"([A-Za-z]{3})[a-z]*\s+(\d{4})", value.strip())
    if m and m.group(1).lower() in MONTHS:
        return f"{m.group(2)}-{MONTHS[m.group(1).lower()]:02d}"
    return value if re.fullmatch(r"\d{4}-\d{2}", value.strip()) else None


def read_pl(path: Path) -> dict[tuple[str, str], Decimal]:
    """(account, month) -> debit-positive amount, per the SPEC §3.5 rules."""
    ws = load_workbook(path, data_only=True).worksheets[0]
    months: list[str | None] | None = None
    credit = False
    out: dict[tuple[str, str], Decimal] = defaultdict(lambda: ZERO)
    for rec in ws.iter_rows(values_only=True):
        if months is None:
            parsed = [_month_cell(v) for v in rec]
            if sum(1 for p in parsed if p) >= 2:
                months = parsed
            continue
        label = rec[0]
        if not isinstance(label, str):
            continue
        if all(v is None for v in rec[1:]):
            name = label.lower()
            credit = "income" in name and "expense" not in name
            continue
        m = ACCOUNT_ROW.match(label)
        if not m:
            continue
        for month, value in zip(months[1:], rec[1:]):
            if month and value is not None:
                out[(m.group(1), month)] += -D(value) if credit else D(value)
    return dict(out)


def read_schedule(path: Path, labels: list[str]) -> dict[str, dict[str, Decimal]]:
    ws = load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if "Ref" in r and all(l in r for l in labels))
    cols = {l: rows[hdr].index(l) for l in labels}
    out: dict[str, dict[str, Decimal]] = {}
    for r in rows[hdr + 1 :]:
        ref, title = r[0], r[1]
        key = ref or (title or "").strip().lower()
        if key:
            out[key] = {l: D(r[c]) for l, c in cols.items()}
    return out


def by_period(rows: list[Row], periods: list[PeriodDef]) -> dict[str, Decimal]:
    out = {p.label: ZERO for p in periods}
    for r in rows:
        for label in labels_for_month(r.month, periods):
            out[label] += r.amount
    return out


def gl_ebitda(rows: dict[int, Row], types: dict[str, str], names: dict[str, str], periods: list[PeriodDef]) -> dict[str, Decimal]:
    """NI + interest + taxes + D&A from the GL rows (account names identify the addback classes)."""
    out = {p.label: ZERO for p in periods}
    for r in rows.values():
        name = names[r.account].lower()
        addback = any(w in name for w in ("interest", "depreciation", "amortization", "income tax"))
        for label in labels_for_month(r.month, periods):
            out[label] += (ZERO if addback else -r.amount)
    return out


def all_files(root: Path) -> dict[str, bytes]:
    return {str(p.relative_to(root)): p.read_bytes() for p in sorted(root.rglob("*")) if p.is_file()}


# ---------------------------------------------------------------------------
# A tiny in-memory deal spec
# ---------------------------------------------------------------------------


def tiny_spec(gl_format: str = "qbo_gl_csv") -> dict:
    return {
        "deal_id": "tiny_plumbing",
        "split": "dev",
        "seed": 7,
        "package_date": "2025-07-10",
        "company": {"name": "Tiny Plumbing Co, LLC", "industry": "Plumbing"},
        "periods": [
            {"label": "FY2024", "start": "2024-01", "end": "2024-12"},
            {"label": "TTM Jun-25", "start": "2024-07", "end": "2025-06"},
        ],
        "data_start": "2024-01",
        "data_end": "2025-06",
        "gl_format": gl_format,
        "sequences": {"ar": 1001},
        "accounts": [
            {"number": "1000", "name": "Checking", "type": "Bank"},
            {"number": "4000", "name": "Service Revenue", "type": "Income"},
            {"number": "4900", "name": "Sales Discounts", "type": "Income"},
            {"number": "5000", "name": "Materials", "type": "Cost of Goods Sold"},
            {"number": "6000", "name": "Wages", "type": "Expense"},
            {"number": "6400", "name": "Legal Fees", "type": "Expense"},
            {"number": "7000", "name": "Depreciation Expense", "type": "Expense"},
            {"number": "8000", "name": "Other Income", "type": "Other Income"},
            {"number": "8100", "name": "Interest Expense", "type": "Other Expense"},
        ],
        "background": [
            {"id": "sales", "account": "4000", "txn_type": "invoice", "schedule": {"count": [3, 5]},
             "amount": {"monthly": 90000, "growth": 0.05, "noise": 0.05},
             "counterparties": [{"name": "Acme Property Mgmt", "weight": 2}, "Birch Street HOA"],
             "memo": ["Plumbing service – job {job}", "Repipe – job {job}"], "num": {"sequence": "ar"}},
            {"id": "discounts", "account": "4900", "txn_type": "credit_memo", "direction": "debit",
             "schedule": {"days": [20]}, "amount": {"monthly": 900, "noise": 0.2},
             "memo": "Senior discounts – {mon} {yyyy}"},
            {"id": "materials", "account": "5000", "txn_type": "bill", "schedule": {"weekdays": [4]},
             "amount": {"monthly": 30000, "noise": 0.1},
             "counterparties": [{"name": "Pipe Supply Co", "num": "PS-{seq}", "seq_start": 500, "seq_step": [2, 9]}],
             "memo": "Weekly statement – week ending {mdy}"},
            {"id": "wages", "account": "6000", "txn_type": "payroll", "schedule": {"days": [15, -1], "adjust": "prior"},
             "amount": {"monthly": 25000, "noise": 0.01}, "memo": "Payroll – PPE {mdy}"},
            {"id": "dep", "account": "7000", "txn_type": "journal", "mode": "fixed", "schedule": {"days": [-1]},
             "amount": {"by_year": {2024: 1000, 2025: 1200}}, "memo": "Depreciation – {mon} {yyyy}"},
            {"id": "interest", "account": "8100", "txn_type": "journal", "mode": "fixed", "schedule": {"days": [-1]},
             "amount": {"fixed": 400}, "counterparties": ["First Test Bank"], "memo": "Loan interest – {mon} {yyyy}"},
        ],
        "planted": [
            {"key": "legal_{yyyy}{mm}", "repeat": {"start": "2024-10", "end": "2024-12", "days": [10]},
             "account": "6400", "counterparty": "Lawson Legal LLP", "num": "LL-{yy}{mm}", "amount": 5000,
             "memo": "Lawsuit defense – {mon} {yyyy}"},
            {"key": "recovery", "date": "2025-02-14", "account": "8000", "txn_type": "deposit",
             "counterparty": "Test Mutual Insurance", "amount": -4000, "memo": "Insurance reimbursement – legal costs"},
        ],
        "data_quality": {
            "duplicates": [{"key": "dup_legal_202411", "of": "legal_202411", "days_later": 2, "note": "Posted twice."}],
            "topside": [{"month": "2025-01", "account": "6000", "amount": 1500, "note": "Bonus accrual not in GL."}],
        },
        "parties": {"lawson": {"name": "Lawson Legal LLP", "address": ["1 Main Street", "Tampa, FL 33602"]}},
        "documents": [
            {"id": "inv", "filename": "1.1 Lawson Invoice LL-2410.pdf", "folder": "01 Legal", "template": "invoice",
             "supports": ["legal_202410"], "key_phrases": ["Invoice No.: LL-2410", "Total Due: $5,000.00"],
             "fields": {"issuer": "lawson", "bill_to": ["Tiny Plumbing Co, LLC"],
                        "meta": [["Invoice No.", "LL-2410"], ["Invoice Date", "October 10, 2024"]],
                        "lines": [{"description": "Defense of lawsuit", "qty": 10, "rate": 500}]}},
            {"id": "mail", "filename": "1.2 Owner email.txt", "folder": "01 Legal", "template": "email",
             "key_phrases": ["the suit is over"],
             "fields": {"from": "Owner <o@tiny.example>", "to": "CFO <c@tiny.example>",
                        "date": "Mon, 6 Jan 2025 09:00:00 -0500", "subject": "Lawsuit",
                        "body": "Good news — the suit is over.\n"}},
        ],
        "schedule": {"adjustments": [
            {"ref": "A-1", "title": "Lawsuit defense", "category": "Non-recurring", "accounts": "6400",
             "support": "DR 1.1", "amounts": {"FY2024": 15000, "TTM Jun-25": 15000}, "claim_keys": ["legal_2024*"]},
        ]},
        "ground_truth": {"authored_by": "unit test", "adjustments": [
            {"adj_id": "A-1", "case_type": "RECOVERY_OFFSET", "treatment": "REVISE",
             "amounts": {"FY2024": 15000, "TTM Jun-25": 11000}, "supporting": ["legal_2024*"],
             "recoveries": ["recovery"], "supporting_docs": ["inv", "mail"],
             "expected_flags": ["OFFSETTING_RECOVERY"], "rationale": "Defense costs net of the insurance recovery."},
        ]},
    }


FORMATS = ["qbo_gl_csv", "netsuite_csv", "xero_xlsx"]


@pytest.mark.parametrize("gl_format", FORMATS)
def test_tiny_deal_is_byte_identical_across_runs(tmp_path, gl_format):
    spec = DealSpec.model_validate(tiny_spec(gl_format))
    generate_deal(spec, tmp_path / "a")
    generate_deal(spec, tmp_path / "b")
    a, b = all_files(tmp_path / "a" / "tiny_plumbing"), all_files(tmp_path / "b" / "tiny_plumbing")
    assert a.keys() == b.keys()
    assert [k for k in a if a[k] != b[k]] == []


@pytest.mark.parametrize("gl_format", FORMATS)
def test_tiny_gl_round_trips_every_row(tmp_path, gl_format):
    result = generate_deal(DealSpec.model_validate(tiny_spec(gl_format)), tmp_path)
    rows = read_gl(result.deal_dir)
    ledger = result.ledger.by_key()
    assert len(rows) == len(result.key_rows) == len(result.ledger.gl_txns())
    for key, line in result.key_rows.items():
        t, r = ledger[key], rows[line]
        assert (r.account, r.day, r.amount, r.num, r.counterparty, r.memo) == (
            t.account, t.date, t.amount, t.num, t.counterparty, t.memo
        ), key


def test_tiny_format_layouts(tmp_path):
    qbo = generate_deal(DealSpec.model_validate(tiny_spec("qbo_gl_csv")), tmp_path / "q").deal_dir
    lines = (qbo / "gl" / "general_ledger.csv").read_text(encoding="utf-8").splitlines()
    assert lines[0] == '"Tiny Plumbing Co, LLC"' and lines[1] == "General Ledger"
    assert lines[2] == '"January 1, 2024 - June 30, 2025"' and lines[3] == ""
    assert lines[4] == ",Date,Transaction Type,Num,Name,Memo/Description,Split,Amount,Balance"
    assert lines[5] == "4000 Service Revenue,,,,,,,," and lines[-1] == "TOTAL,,,,,,,,"
    first = next(csv.reader([lines[6]]))
    assert first[2] == "Invoice" and first[6] == "Accounts Receivable (A/R)" and not first[7].startswith("-")
    discount = next(csv.reader([next(l for l in lines if "Senior discounts" in l)]))
    assert discount[7].startswith("-")  # contra revenue shows negative under the natural sign
    assert re.fullmatch(r'Total for 4000 Service Revenue,,,,,,,"\$[\d,]+\.\d\d",', next(l for l in lines if l.startswith("Total for 4000")))
    assert (qbo / "gl" / "chart_of_accounts.csv").read_text().splitlines()[0] == "Account #,Full name,Type,Detail type"

    ns = generate_deal(DealSpec.model_validate(tiny_spec("netsuite_csv")), tmp_path / "n").deal_dir
    ns_lines = (ns / "gl" / "general_ledger.csv").read_text().splitlines()
    assert ns_lines[0].startswith("Internal ID,Date,Period,Type,Document Number,Account,Name,Memo,Debit,Credit,Subsidiary")
    rec = next(csv.reader([ns_lines[1]]))
    assert re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", rec[1]) and re.fullmatch(r"[A-Z][a-z]{2} \d{4}", rec[2])
    assert bool(rec[8]) != bool(rec[9])

    xero = generate_deal(DealSpec.model_validate(tiny_spec("xero_xlsx")), tmp_path / "x").deal_dir
    wb = load_workbook(xero / "gl" / "general_ledger.xlsx")
    ws = wb["Account Transactions"]
    assert [ws.cell(row=r, column=1).value for r in (1, 2, 3)] == [
        "Account Transactions", "Tiny Plumbing Co, LLC", "For the period 1 January 2024 to 30 June 2025"]
    assert ws.cell(row=4, column=1).value is None
    assert [c.value for c in ws[5]] == ["Date", "Source", "Description", "Reference", "Debit", "Credit",
                                        "Running Balance", "Account Code", "Account"]
    assert yaml.safe_load((xero / "deal.yaml").read_text())["files"]["gl"] == "gl/general_ledger.xlsx"


def test_tiny_truth_monthly_pl_and_schedule(tmp_path):
    result = generate_deal(DealSpec.model_validate(tiny_spec()), tmp_path)
    deal = result.deal_dir
    gt = GroundTruth.model_validate_json((deal / "ground_truth.json").read_text())
    rows = read_gl(deal)
    key_rows = result.key_rows
    adj = gt.adjustments[0]
    assert adj.supporting_gl_rows == sorted(key_rows[f"legal_2024{m}"] for m in ("10", "11", "12"))
    assert adj.related_gl_rows == [key_rows["recovery"]]
    assert adj.supporting_docs == ["1.1 Lawson Invoice LL-2410.pdf", "1.2 Owner email.txt"]
    codes = {(d.code.value, d.month, d.account) for d in gt.data_quality}
    assert ("DUPLICATE_GL_ENTRY", "2024-11", "6400") in codes
    assert ("RECON_VARIANCE", "2025-01", "6000") in codes
    assert ("MGMT_EBITDA_DIFFERS_FROM_GL", None, None) in codes

    pl = read_pl(deal / "financials" / "monthly_pl.xlsx")
    gl = defaultdict(lambda: ZERO)
    for r in rows.values():
        gl[(r.account, r.month)] += r.amount
    diffs = {k: pl.get(k, ZERO) - gl.get(k, ZERO) for k in set(pl) | set(gl) if pl.get(k, ZERO) != gl.get(k, ZERO)}
    assert diffs == {("6000", "2025-01"): Decimal("1500.00")}

    sched = read_schedule(deal / "adjustments" / "management_adjusted_ebitda.xlsx", ["FY2024", "TTM Jun-25"])
    assert sched["reported ebitda"]["TTM Jun-25"] == D(gt.gl_ebitda["TTM Jun-25"]) - 1500
    assert sched["reported ebitda"]["FY2024"] == D(gt.gl_ebitda["FY2024"])
    assert D(gt.diligence_adjusted_ebitda["TTM Jun-25"]) == D(gt.gl_ebitda["TTM Jun-25"]) + 11000


def test_tiny_missing_gl_month(tmp_path):
    data = tiny_spec()
    data["data_quality"]["missing_gl_months"] = [{"month": "2024-03", "note": "March export missing."}]
    result = generate_deal(DealSpec.model_validate(data), tmp_path)
    rows = read_gl(result.deal_dir)
    assert not [r for r in rows.values() if r.month == "2024-03"]
    pl = read_pl(result.deal_dir / "financials" / "monthly_pl.xlsx")
    assert pl[("4000", "2024-03")] < 0  # the P&L still has the month
    gt = GroundTruth.model_validate_json((result.deal_dir / "ground_truth.json").read_text())
    assert any(d.code.value == "MISSING_PERIOD" and d.month == "2024-03" for d in gt.data_quality)


@pytest.mark.parametrize(
    "mutate, message",
    [
        (lambda d: d["ground_truth"]["adjustments"][0]["amounts"].update({"TTM Jun-25": 15000}), "GL rows give"),
        (lambda d: d["ground_truth"]["adjustments"][0].update({"supporting": ["nope_*"]}), "matches no GL row"),
        (lambda d: d["schedule"]["adjustments"][0]["amounts"].update({"FY2024": 14000}), "claim_keys total"),
        (lambda d: (d["documents"][0]["fields"]["lines"][0].update({"rate": 450}), d["documents"][0].update({"key_phrases": []})),
         "invoice total"),
        (lambda d: d["documents"][1].update({"key_phrases": ["not in the email"]}), "key phrase"),
        (lambda d: d["background"].append(
            {"id": "coffee", "account": "6000", "txn_type": "expense", "mode": "fixed", "schedule": {"weekdays": [0]},
             "amount": {"fixed": 85}, "counterparties": ["Joe's Coffee Service"], "memo": "Coffee service"}),
         "unplanned duplicate"),
        (lambda d: d["planted"][0].update({"account": "1000"}), "balance-sheet account"),
    ],
)
def test_spec_errors_are_caught(tmp_path, mutate, message):
    data = copy.deepcopy(tiny_spec())
    mutate(data)
    with pytest.raises(GenerationError, match=message):
        generate_deal(DealSpec.model_validate(data), tmp_path)


def test_cli_generates_from_yaml(tmp_path):
    spec_path = tmp_path / "tiny_plumbing.yaml"
    spec_path.write_text(yaml.safe_dump(tiny_spec(), allow_unicode=True, sort_keys=False), encoding="utf-8")
    cli = importlib.util.spec_from_file_location("qoe_generate_deals", SCRIPTS / "qoe_generate_deals.py")
    module = importlib.util.module_from_spec(cli)
    cli.loader.exec_module(module)
    assert module.main(["--spec", str(spec_path), "--out", str(tmp_path / "out")]) == 0
    meta = DealMeta.model_validate(yaml.safe_load((tmp_path / "out" / "tiny_plumbing" / "deal.yaml").read_text()))
    assert meta.gl_format == "auto" and meta.synthetic


# ---------------------------------------------------------------------------
# Dev deal 1: Meridian Mechanical Services (SPEC §10)
# ---------------------------------------------------------------------------

meridian_only = pytest.mark.skipif(not MERIDIAN_SPEC.exists(), reason="Meridian spec not present")

CLAIMS = {  # SPEC §10 management claims: FY2024 / FY2025 / TTM Jun-26
    "M-01": (0, 120000, 65500), "M-02": (360000, 360000, 360000), "M-03": (0, 48000, 48000),
    "M-04": (0, 45000, 15000), "M-05": (0, 0, 75000), "M-06": (0, 96000, 48000), "M-07": (0, 35000, 0),
    "M-08": (0, 62000, 62000), "M-09": (58000, 0, 0), "M-10": (0, 42000, 0), "M-11": (0, 80000, 80000),
    "M-12": (0, 0, 165000), "M-13": (0, 64000, 64000), "M-14": (0, 52000, 52000),
}
TRUTH = {  # SPEC §10 diligence amounts; None = REQUEST_INFO
    "M-01": ("REVISE", (0, 84500, 30000)), "M-02": ("REQUEST_INFO", None), "M-03": ("REVISE", (0, 31200, 31200)),
    "M-04": ("ACCEPT", (0, 45000, 15000)), "M-05": ("ACCEPT", (0, 0, 75000)), "M-06": ("REJECT", (0, 0, 0)),
    "M-07": ("REJECT", (0, 0, 0)), "M-08": ("REVISE", (0, 41000, 41000)), "M-09": ("REVISE", (58000, -40000, 0)),
    "M-10": ("REVISE", (-42000, 42000, 0)), "M-11": ("REVISE", (0, 80000, 0)), "M-12": ("REQUEST_INFO", None),
    "M-13": ("REJECT", (0, 0, 0)), "M-14": ("ACCEPT", (0, 52000, 52000)),
}
LABELS = ["FY2024", "FY2025", "TTM Jun-26"]


@pytest.fixture(scope="module")
def meridian(tmp_path_factory):
    if not MERIDIAN_SPEC.exists():
        pytest.skip("Meridian spec not present")
    return generate_deal(load_spec(MERIDIAN_SPEC), tmp_path_factory.mktemp("meridian"))


@pytest.fixture(scope="module")
def meridian_rows(meridian):
    return read_gl(meridian.deal_dir)


@pytest.fixture(scope="module")
def meridian_truth(meridian):
    return GroundTruth.model_validate_json((meridian.deal_dir / "ground_truth.json").read_text())


@pytest.fixture(scope="module")
def periods():
    return [PeriodDef(label="FY2024", start="2024-01", end="2024-12"), PeriodDef(label="FY2025", start="2025-01", end="2025-12"),
            PeriodDef(label="TTM Jun-26", start="2025-07", end="2026-06")]


@meridian_only
def test_meridian_is_byte_identical_across_runs(meridian, tmp_path):
    again = generate_deal(load_spec(MERIDIAN_SPEC), tmp_path)
    a, b = all_files(meridian.deal_dir), all_files(again.deal_dir)
    assert a.keys() == b.keys()
    assert [k for k in a if a[k] != b[k]] == []


@meridian_only
def test_committed_meridian_package_is_current(meridian):
    if not MERIDIAN_COMMITTED.exists():
        pytest.skip("data/qoe/dev/meridian_mechanical not generated")
    fresh, committed = all_files(meridian.deal_dir), all_files(MERIDIAN_COMMITTED)
    assert fresh.keys() == committed.keys()
    assert [k for k in fresh if fresh[k] != committed[k]] == [], "regenerate with scripts/qoe_generate_deals.py --all"


@meridian_only
def test_meridian_package_layout(meridian, meridian_truth):
    deal = meridian.deal_dir
    meta = DealMeta.model_validate(yaml.safe_load((deal / "deal.yaml").read_text()))
    assert meta.deal_id == "meridian_mechanical" and meta.target_name == "Meridian Mechanical Services, LLC"
    assert [p.label for p in meta.periods] == LABELS and (meta.data_start, meta.data_end) == ("2024-01", "2026-06")
    for rel in (meta.files.gl, meta.files.chart_of_accounts, meta.files.monthly_pl, meta.files.adjustments):
        assert (deal / rel).is_file(), rel
    assert "SYNTHETIC" in (deal / "README.txt").read_text()
    lines = (deal / meta.files.gl).read_text(encoding="utf-8").splitlines()
    assert lines[:5] == ['"Meridian Mechanical Services, LLC"', "General Ledger", '"January 1, 2024 - June 30, 2026"', "",
                         ",Date,Transaction Type,Num,Name,Memo/Description,Split,Amount,Balance"]
    assert lines[5] == "4000 Service Revenue - Commercial,,,,,,,," and lines[-1] == "TOTAL,,,,,,,,"
    names = [p.name for p in (deal / "documents").rglob("*") if p.is_file()]
    assert len(names) == len(set(names)) >= 40
    docs = set(names)
    assert {d for a in meridian_truth.adjustments for d in a.supporting_docs} <= docs
    assert meridian_truth.deal_id == "meridian_mechanical" and meridian_truth.split == "dev"


@meridian_only
def test_meridian_scale_and_levels(meridian, meridian_rows, meridian_truth, periods):
    assert 3000 <= len(meridian_rows) <= 8000
    revenue = by_period([r for r in meridian_rows.values() if r.account.startswith("4")], periods)
    for label in LABELS:
        assert Decimal("35000000") <= -revenue[label] <= Decimal("41000000"), label
        assert Decimal("4000000") <= D(meridian_truth.gl_ebitda[label]) <= Decimal("5000000"), label
    months = {r.month for r in meridian_rows.values()}
    assert months == set(month_range("2024-01", "2026-06"))
    assert all(not r.account.startswith(("1", "2", "3")) for r in meridian_rows.values())  # P&L accounts only


@meridian_only
def test_meridian_truth_matches_catalog(meridian_truth):
    got = {a.adj_id: a for a in meridian_truth.adjustments}
    assert list(got) == list(TRUTH)
    for adj_id, (treatment, amounts) in TRUTH.items():
        a = got[adj_id]
        assert a.treatment == Treatment(treatment), adj_id
        assert a.amounts == ({} if amounts is None else {l: f"{v}.00" for l, v in zip(LABELS, amounts)}), adj_id
        assert a.rationale and a.supporting_docs, adj_id
    assert {f.value for f in got["M-06"].expected_flags} == {"CONTRADICTORY_EVIDENCE", "CONTINUING_OBLIGATION", "RECURRING_PATTERN"}
    assert got["M-14"].ambiguity == "medium"


@meridian_only
def test_meridian_truth_amounts_tie_to_gl_rows(meridian, meridian_rows, meridian_truth, periods):
    """Truth amounts are pure pass-throughs of the supporting GL rows, except for
    documented effects: recoveries (M-09 insurance gain) and out-of-period moves (M-10)."""
    with_effects = set()
    for a in meridian_truth.adjustments:
        support = [meridian_rows[n] for n in a.supporting_gl_rows]
        assert len(support) == len(a.supporting_gl_rows) and not set(a.supporting_gl_rows) & set(a.related_gl_rows)
        if a.treatment == Treatment.REQUEST_INFO:
            assert a.amounts == {} and support
            continue
        computed = by_period(support, periods)
        eff = meridian.effects[a.adj_id]
        if eff.recovery_rows or eff.moves:
            with_effects.add(a.adj_id)
        for label, value in by_period([meridian_rows[n] for n in eff.recovery_rows], periods).items():
            computed[label] += value
        for rows, start, end in eff.moves:
            moved = sum((meridian_rows[n].amount for n in rows), ZERO)
            service = month_range(start, end)
            for p in periods:
                inside = sum(1 for m in service if p.start <= m <= p.end)
                computed[p.label] -= q2(moved * inside / len(service))
        assert {l: computed[l] for l in LABELS} == {l: D(a.amounts[l]) for l in LABELS}, a.adj_id
    assert with_effects == {"M-09", "M-10"}


@meridian_only
def test_meridian_schedule_foots(meridian):
    sched = read_schedule(meridian.deal_dir / "adjustments" / "management_adjusted_ebitda.xlsx", LABELS)
    assert [k for k in sched if k.startswith("M-")] == list(CLAIMS)
    for adj_id, claim in CLAIMS.items():
        assert [sched[adj_id][l] for l in LABELS] == [Decimal(v) for v in claim], adj_id
    for l in LABELS:
        total = sum((sched[a][l] for a in CLAIMS), ZERO)
        assert sched["total management adjustments"][l] == total
        assert sched["management adjusted ebitda"][l] == sched["reported ebitda"][l] + total
        parts = ("net income", "interest expense", "income tax expense", "depreciation and amortization")
        assert sched["reported ebitda"][l] == sum((sched[k][l] for k in parts), ZERO)


@meridian_only
def test_meridian_reported_ebitda_is_25k_below_gl(meridian, meridian_rows, meridian_truth, periods):
    deal = meridian.deal_dir
    with (deal / "gl" / "chart_of_accounts.csv").open(encoding="utf-8") as f:
        coa = list(csv.reader(f))[1:]
    names = {r[0]: r[1] for r in coa}
    ours = gl_ebitda(meridian_rows, read_coa_types(deal / "gl" / "chart_of_accounts.csv"), names, periods)
    assert {l: str(q2(v)) for l, v in ours.items()} == meridian_truth.gl_ebitda
    sched = read_schedule(deal / "adjustments" / "management_adjusted_ebitda.xlsx", LABELS)
    gap = {l: sched["reported ebitda"][l] - D(meridian_truth.gl_ebitda[l]) for l in LABELS}
    assert gap == {"FY2024": 0, "FY2025": Decimal("-25000"), "TTM Jun-26": Decimal("-25000")}
    finals = {l: sum((D(a.amounts[l]) for a in meridian_truth.adjustments if a.amounts), ZERO) for l in LABELS}
    for l in LABELS:
        assert D(meridian_truth.diligence_adjusted_ebitda[l]) == D(meridian_truth.gl_ebitda[l]) + finals[l]


@meridian_only
def test_meridian_monthly_pl_equals_gl_except_topside(meridian, meridian_rows):
    pl = read_pl(meridian.deal_dir / "financials" / "monthly_pl.xlsx")
    gl = defaultdict(lambda: ZERO)
    for r in meridian_rows.values():
        gl[(r.account, r.month)] += r.amount
    diffs = {k: pl.get(k, ZERO) - gl.get(k, ZERO) for k in set(pl) | set(gl) if pl.get(k, ZERO) != gl.get(k, ZERO)}
    assert diffs == {("6000", "2025-12"): Decimal("25000.00")}


@meridian_only
def test_meridian_planted_duplicate(meridian_rows, meridian_truth):
    dq = [d for d in meridian_truth.data_quality if d.code.value == "DUPLICATE_GL_ENTRY"]
    assert len(dq) == 1 and dq[0].month == "2025-05" and dq[0].account == "6200"
    first, second = (meridian_rows[n] for n in dq[0].gl_rows)
    assert (first.account, first.amount, first.num, first.counterparty) == (second.account, second.amount, second.num, second.counterparty)
    assert (second.day - first.day).days == 3 and first.counterparty == "Coastal Risk Insurance"
    codes = {d.code.value for d in meridian_truth.data_quality}
    assert codes == {"DUPLICATE_GL_ENTRY", "RECON_VARIANCE", "MGMT_EBITDA_DIFFERS_FROM_GL"}


@meridian_only
def test_meridian_documents_are_quotable(meridian):
    spec = load_spec(MERIDIAN_SPEC)
    docs_dir = meridian.deal_dir / "documents"
    by_name = {p.name: p for p in docs_dir.rglob("*") if p.is_file()}
    for doc in spec.documents:
        path = by_name[doc.filename]
        if path.suffix == ".pdf":
            pages = [canonicalize_page_text(p) for p in extract_pdf_pages(path)]
            assert pages and all(FOOTER in p for p in pages), doc.filename
            text = "\n".join(pages)
        else:
            text = canonicalize_page_text(path.read_text(encoding="utf-8"))
            assert all(f"\n{h}: " in "\n" + text for h in ("From", "To", "Date", "Subject")), doc.filename
            assert FOOTER in text
        for phrase in doc.key_phrases:
            assert phrase in text, (doc.filename, phrase)
    draft = canonicalize_page_text(extract_pdf_pages(by_name["5.1 Castellano Executive Employment Agreement DRAFT.pdf"])[0])
    assert "DRAFT" in draft and "By: ____" in draft and "/s/" not in draft
    msa = canonicalize_page_text("\n".join(extract_pdf_pages(by_name["6.1 Brightline Managed Services Agreement.pdf"])))
    assert "/s/ Richard Castellano" in msa
