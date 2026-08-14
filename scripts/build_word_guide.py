#!/usr/bin/env python3
"""One-shot generator for docs/Accounting_Agent_Guide.docx. Not part of the AP pipeline."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "docs" / "Accounting_Agent_Guide.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x4B)
    return p


def add_p(doc, text, *, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    set_run_font(run, name="Consolas", size=9)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            set_run_font(run, size=11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shading = hdr[i]._tePr if False else hdr[i]._tc.get_or_add_tcPr()
        # simple header fill
        from docx.oxml import OxmlElement

        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1B3A4B")
        shd.set(qn("w:val"), "clear")
        shading.append(shd)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=8)
    doc.add_paragraph()
    return table


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Accounting Agent")
    set_run_font(r, size=28, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B))

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Full repository guide")
    set_run_font(r, size=16, color=RGBColor(0x4A, 0x67, 0x75))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Deterministic accounts-payable three-way match  ·  Python control plane + Qwen3-8B extraction\n"
        "How the repo was built, what every folder and .py file does, GPU setup, and how to read outputs"
    )
    set_run_font(r, size=11, italic=True)

    add_p(
        doc,
        "This Word document is the readable tour of the accounting-agent repository. "
        "A Markdown twin lives at docs/how_it_works.md. Local tests do not call the GPU; "
        "they replay ground truth from expected.json. Live Qwen evaluation is a separate --live path.",
    )

    # ------------------------------------------------------------------
    add_heading(doc, "1. What this system is", 1)
    add_p(
        doc,
        "This is a prototype accounts-payable agent. A vendor invoice PDF goes in. "
        "A decision comes out: READY_FOR_DRAFT or HUMAN_REVIEW, plus a balanced journal "
        "proposal and an audit trail when the match is clean.",
    )
    add_p(
        doc,
        "The split is the product: the LLM only extracts structured JSON from invoice text. "
        "Python does every match, Decimal calculation, policy check, idempotency key, and decision. "
        "The model cannot override policy.",
    )
    add_bullets(
        doc,
        [
            "Invoice: the only document that arrives as a file (PDF).",
            "Purchase order and goods receipt: in production these live in an ERP (NetSuite, QuickBooks, SAP). Here they are JSON seeds loaded into SQLite.",
            "Three-way match: invoice vs PO vs receipts — same SKUs, prices within tolerance, quantity not above ordered or received.",
            "Fictional vendors only. Not production ERP-integrated.",
        ],
    )

    add_heading(doc, "2. How it was created", 1)
    add_p(doc, "The repo was built as an evaluation prototype in this order.", bold=True)
    add_heading(doc, "2.1 Control plane first", 2)
    add_p(
        doc,
        "Python modules under app/ were written as a bounded state machine: ingest, extract, validate, "
        "lookup, match, propose journal, create draft or review. Schemas are strict Pydantic v2 "
        "(extra fields forbidden). Money is decimal strings, never floats. Policy lives in one YAML file.",
    )
    add_heading(doc, "2.2 Fake ERP and fixtures", 2)
    add_p(
        doc,
        "scripts/generate_cases.py (seed 20260812) generates 50 synthetic cases with ReportLab. "
        "Each case is a folder with invoice.pdf, po.json, receipt.json, and expected.json. "
        "scripts/initialize_database.py loads POs, receipts, vendors, and duplicate seeds into SQLite. "
        "A LocalERPAdapter implements the same interface a real ERP adapter would.",
    )
    add_heading(doc, "2.3 First GPU measurement (easy invoices)", 2)
    add_p(
        doc,
        "The original PDFs had a labeled machine-readable block (Vendor:, LINE|sku|qty|price|total). "
        "On 2026-08-13, live Qwen3-8B on pinned vLLM scored 50/50 accuracy, schema-valid JSON, "
        "zero false READY_FOR_DRAFT. Repeat extracts at concurrency 1, 8, and 32 produced one hash per case. "
        "Those hashes matched a second H100 (PCIe US-KS-2 vs SXM AP-IN-1). "
        "That result is frozen in reports/. It was measured on the easy labeled PDFs.",
    )
    add_heading(doc, "2.4 Harder invoices (current fixtures)", 2)
    add_p(
        doc,
        "The generator was then changed so PDFs are still text-layer (no OCR) but much harder for the model: "
        "no LINE| block, inconsistent labels, remit-to and barcode distractors, two-column reading-order traps, "
        "multi-page totals, mixed money and date formats. expected.json still stores normalized fields; "
        "evidence quotes store the messy source form. Local mock tests still pass because they preseed "
        "from expected.json. Live Qwen accuracy on this harder set has not been re-run.",
    )

    # ------------------------------------------------------------------
    add_heading(doc, "3. End-to-end flow", 1)
    add_code(
        doc,
        "invoice.pdf\n"
        "  → INGEST     pypdf text layer → canonicalize → hash\n"
        "                (empty text → OCR_REQUIRED → HUMAN_REVIEW, no LLM)\n"
        "  → EXTRACT    Qwen or mock/preseed → ExtractedInvoice JSON\n"
        "  → VALIDATE   required fields, evidence quotes, line math\n"
        "  → LOOKUP     ERP PO + receipts + vendor + duplicate + idempotency key\n"
        "  → MATCH      invoice vs PO vs receipts vs policy.yaml\n"
        "  → JOURNAL    Dr expense GLs from PO lines, Cr Accounts Payable (if clean)\n"
        "  → DECISION   READY_FOR_DRAFT or HUMAN_REVIEW",
    )
    add_table(
        doc,
        ["Mode / path", "What happens"],
        [
            [
                "evaluate",
                "Run the pipeline; do not write a draft. Used by tests and accuracy eval.",
            ],
            [
                "create_draft",
                "If READY_FOR_DRAFT, persist a draft bill id. Same idempotency key returns the cached result.",
            ],
            [
                "Mock / preseed",
                "expected.json extraction is injected. make test, make evaluate-accuracy, Streamlit, run_case.py.",
            ],
            [
                "Live GPU",
                "Qwen3-8B via vLLM reads canonical text. evaluate_accuracy.py --live and evaluate_determinism.py --live.",
            ],
        ],
    )

    add_heading(doc, "4. Policy (what “match” means)", 1)
    add_bullets(
        doc,
        [
            "Currency must be USD.",
            "Unit price may differ from the PO by at most max(0.5% of PO price, $0.01).",
            "Invoice quantity may not exceed PO quantity or received quantity (tolerance 0).",
            "Invoice total must match line sum ± $0.01.",
            "Any non-zero tax → HUMAN_REVIEW.",
            "Missing or ambiguous required fields → HUMAN_REVIEW.",
            "Image-only PDFs (no text layer) → OCR_REQUIRED → HUMAN_REVIEW. This repo does not run OCR.",
        ],
    )

    add_heading(doc, "5. Folder map", 1)
    add_table(
        doc,
        ["Folder", "What is in it", "How it was created"],
        [
            [
                "app/",
                "The product: pipeline, schemas, matching, journal, LLM clients, SQLite ERP adapter, Streamlit UI.",
                "Hand-written Python. This is the control plane.",
            ],
            [
                "app/adapters/",
                "ERPAdapter protocol + LocalERPAdapter (SQLite).",
                "Hand-written so a real NetSuite/QBO adapter could replace SQLite later.",
            ],
            [
                "config/",
                "policy.yaml — match tolerances.",
                "Hand-written. Hashed into the idempotency key; changing it changes keys.",
            ],
            [
                "prompts/",
                "invoice_extraction_v1.txt — system prompt for Qwen.",
                "Hand-written. Also hashed into the idempotency key.",
            ],
            [
                "scripts/",
                "Generate fixtures, seed the DB, run one case, evaluate accuracy/determinism, build reports.",
                "Hand-written CLIs. generate_cases.py is what creates the 50 PDFs.",
            ],
            [
                "tests/",
                "pytest modules plus fixtures/ (50 synthetic cases).",
                "Tests hand-written. Fixtures generated by make generate-data.",
            ],
            [
                "tests/fixtures/development/",
                "case_001–case_030. Iterate and tune here.",
                "scripts/generate_cases.py, seed 20260812.",
            ],
            [
                "tests/fixtures/holdout/",
                "case_031–case_050. Score last. Do not hand-edit expected.json.",
                "Same generator. Holdout was not used for prompt/policy tuning.",
            ],
            [
                "deployment/",
                "Pinned vLLM start script, RunPod notes, CPU Dockerfile, compose file.",
                "Hand-written. Image digest and model revision must not be silently changed.",
            ],
            [
                "docs/",
                "how_it_works.md, technical_decisions.md, SECRET_SWEEP.md, this Word guide.",
                "Written as the project was measured and the harder PDFs were added.",
            ],
            [
                "reports/",
                "Accuracy JSON/CSV, determinism hashes, billing, final_report.md.",
                "Produced by evaluate_*.py, capture_environment.py, build_report.py on 2026-08-13 (easy PDFs).",
            ],
        ],
    )

    # ------------------------------------------------------------------
    add_heading(doc, "6. Every Python file", 1)
    add_p(
        doc,
        "Empty __init__.py files only mark packages (app/ and app/adapters/). "
        "They contain no logic.",
    )

    add_heading(doc, "6.1 app/ — runtime", 2)
    add_table(
        doc,
        [".py file", "What it does", "How it was created / used"],
        [
            [
                "app/schemas.py",
                "Pydantic models: ExtractedInvoice, PO, receipt, vendor, policy, journal, WorkflowResult, AuditManifest, exception-code list. extra=forbid.",
                "Hand-written contract. LLM JSON must validate against this. Schema version is hashed into idempotency.",
            ],
            [
                "app/settings.py",
                "Loads .env: vLLM URL/key, model pins, DATABASE_URL, AP account 2000, prompt/policy paths, temperature 0, seed 42.",
                "Hand-written. get_settings() is cached.",
            ],
            [
                "app/pdf_text.py",
                "pypdf PdfReader; extract_text(extraction_mode=plain) per page. If no page has text, ocr_required=True.",
                "Hand-written ingest. No OCR implementation — scanned PDFs stop here.",
            ],
            [
                "app/canonicalize.py",
                "NFC Unicode, newlines, collapse spaces, join pages with --- PAGE n ---, SHA-256 of PDF bytes and canonical text.",
                "Hand-written so hashes are stable across runs.",
            ],
            [
                "app/llm_client.py",
                "LLMClient ABC. MockLLMClient returns expected.json extraction (or a regex parser). VLLMLLMClient calls vLLM with json_schema and thinking disabled. INVOICE_JSON_SCHEMA lives here.",
                "Hand-written. Mock path is for local tests; live path is the GPU experiment.",
            ],
            [
                "app/validation.py",
                "Required fields, ambiguity, evidence quotes must be exact substrings of canonical text, Decimal arithmetic, normalize money/qty strings.",
                "Hand-written. Runs after every extract, mock or live.",
            ],
            [
                "app/matching.py",
                "Three-way match: duplicate, PO OPEN, vendor, currency, tax, per-SKU price tolerance, qty vs PO, qty vs receipts.",
                "Hand-written pure Python. This is the AP clerk.",
            ],
            [
                "app/journal.py",
                "Map invoice SKUs to PO gl_account, debit those accounts (freight on first GL), credit AP. Unbalanced → HUMAN_REVIEW.",
                "Hand-written. Only called when match exceptions are empty.",
            ],
            [
                "app/idempotency.py",
                "SHA-256 of pdf_hash | po_id | receipt_snapshot | prompt_hash | schema | policy | model_revision.",
                "Hand-written so create_draft is safe to retry.",
            ],
            [
                "app/audit.py",
                "In-memory event log (STATE, OCR_REQUIRED, COMPLETE). Optional SQLite persist.",
                "Hand-written audit trail for the UI and manifests.",
            ],
            [
                "app/pipeline.py",
                "State machine that calls everything above. run_case_dir loads a fixture folder and preseeds extraction from expected.json for mock runs.",
                "Hand-written orchestrator. This is the file to read first after schemas.",
            ],
            [
                "app/database.py",
                "SQLAlchemy tables: vendors, POs, PO lines, receipts, drafts, processed documents, duplicate seeds, audit events. SQLite foreign_keys ON.",
                "Hand-written fake ERP schema.",
            ],
            [
                "app/ui.py",
                "Streamlit demo: pick a fixture or upload a PDF, run mock pipeline, show decision, evidence, journal, hashes.",
                "Hand-written local demo. make run-ui. Does not call Qwen.",
            ],
            [
                "app/adapters/base.py",
                "Abstract ERPAdapter methods a real ERP integration would implement.",
                "Hand-written boundary. No QuickBooks/NetSuite code behind it yet.",
            ],
            [
                "app/adapters/local_erp.py",
                "SQLite implementation: ordered SELECTs, draft ids DRAFT- plus first 12 of the key, attach_source_document is a no-op.",
                "Hand-written stand-in for ERP APIs.",
            ],
        ],
    )

    add_heading(doc, "6.2 scripts/ — generate, seed, evaluate", 2)
    add_table(
        doc,
        [".py file", "What it does", "How it was created / used"],
        [
            [
                "scripts/generate_cases.py",
                "Builds 50 cases: 5 vendors, 8 SKUs, 5 PDF layouts, 9 match scenarios. Writes PDF + JSON + dataset_manifest.json. Re-extracts each PDF to bind evidence page numbers. ReportLab invariant=1 for byte-stable PDFs.",
                "Hand-written generator. make generate-data. This is how every fixture PDF was created.",
            ],
            [
                "scripts/initialize_database.py",
                "Creates SQLite tables, inserts V001–V005, loads every po.json and receipt.json, seeds duplicates from expected.json already_processed.",
                "Hand-written. make init-db. Eval scripts call it with --reset into a temp DB.",
            ],
            [
                "scripts/run_case.py",
                "CLI: --case case_001 --mode evaluate|create_draft. Prints WorkflowResult JSON using MockLLMClient.",
                "Hand-written smoke tool. make run-local.",
            ],
            [
                "scripts/evaluate_accuracy.py",
                "Scores each case vs expected.json decision and exception codes. Default mock+preseed. --live uses VLLMLLMClient and does not preseed. Writes reports/accuracy_mock.json or accuracy_gpu.json.",
                "Hand-written scorecard. This is the GPU accuracy test.",
            ],
            [
                "scripts/evaluate_determinism.py",
                "Always checks local PDF/canonical hashes. --live repeats extracts at concurrency 1/8/32 and hashes ExtractedInvoice JSON (volatile keys stripped).",
                "Hand-written. This is the GPU invariance test.",
            ],
            [
                "scripts/capture_environment.py",
                "Writes reports/environment.json: Python, packages, git commit, model pins.",
                "Hand-written. make capture-env.",
            ],
            [
                "scripts/build_report.py",
                "Assembles reports/final_report.md from accuracy, determinism, env, and billing JSON.",
                "Hand-written. make report.",
            ],
            [
                "scripts/build_word_guide.py",
                "Generates this Word document.",
                "One-shot helper. Not part of the AP pipeline.",
            ],
        ],
    )

    add_heading(doc, "6.3 tests/ — pytest", 2)
    add_table(
        doc,
        [".py file", "What it does"],
        [
            [
                "tests/test_pipeline.py",
                "Fixture cases through the full pipeline: exact match drafts, price-over review, duplicates, idempotent re-run, all 30 dev + 20 holdout decisions vs expected.json.",
            ],
            [
                "tests/test_matching.py",
                "Match rules without PDFs: exact, price within/over, qty vs receipt, duplicate, missing PO, nonzero tax, sorted exceptions.",
            ],
            [
                "tests/test_validation.py",
                "Required fields, evidence must match canonical text, math error, ambiguity.",
            ],
            [
                "tests/test_journal.py",
                "Balanced journal (debits sorted, AP credit) and unbalanced → review.",
            ],
            [
                "tests/test_idempotency.py",
                "Same inputs → same 64-char key; changing PO id changes the key.",
            ],
            [
                "tests/test_canonicalization.py",
                "NFC/newlines, whitespace policy, page separators, hash stability.",
            ],
            [
                "tests/test_pdf_reproducibility.py",
                "Regenerating a PDF is byte-identical; layouts have no LINE| machine block; every expected evidence quote appears in canonical text.",
            ],
        ],
    )

    # ------------------------------------------------------------------
    add_heading(doc, "7. Other important files (not Python)", 1)
    add_table(
        doc,
        ["File", "What it does"],
        [
            ["README.md", "Front door: measured results, harder-PDF note, local/GPU how-to, pins."],
            ["Makefile", "setup, generate-data, init-db, test, run-local, run-ui, evaluate-accuracy, evaluate-determinism, capture-env, report."],
            ["pyproject.toml", "Package metadata, Python 3.12 pin, dependencies, pytest pythonpath."],
            ["uv.lock", "Locked versions for uv sync. Do not hand-edit."],
            [".env.example", "Template env. Copy to .env (gitignored). Fake API key is fine locally."],
            [".gitignore", "Drops .venv, .env, *.db, caches. Allows a short list of reports/*.json."],
            ["config/policy.yaml", "Match tolerances. Changing it changes policy_hash and idempotency keys."],
            ["prompts/invoice_extraction_v1.txt", "Qwen system prompt: extract JSON, do not invent, exact evidence quotes."],
            ["deployment/start_vllm.sh", "docker run pinned vLLM image + Qwen3-8B + batch invariance + xgrammar."],
            ["deployment/runpod-template.md", "Same pins for renting an H100. Stop the pod after eval."],
            ["deployment/Dockerfile", "CPU app image (not the GPU server)."],
            ["deployment/docker-compose.yml", "App service; vLLM service commented out."],
            ["docs/how_it_works.md", "Markdown map of the repo, including every fixture case."],
            ["docs/technical_decisions.md", "Why temp=0 is not enough, GPU pins and sources, what was not proven."],
            ["docs/SECRET_SWEEP.md", "Do not commit .env, tokens, logs, or SQLite DBs."],
            ["tests/fixtures/dataset_manifest.json", "SHA-256 of every generated fixture file."],
        ],
    )

    add_heading(doc, "8. The 50 synthetic cases", 1)
    add_p(
        doc,
        "Each case directory contains exactly four files, all created by scripts/generate_cases.py:",
    )
    add_bullets(
        doc,
        [
            "invoice.pdf — ReportLab text-layer invoice (byte-reproducible). Five visual layouts rotate by case index. No LINE| machine block.",
            "po.json — purchase order seed loaded into SQLite (not how a real ERP stores POs).",
            "receipt.json — goods receipt seed loaded into SQLite.",
            "expected.json — ground truth: scenario, decision, exception codes, perfect extraction and evidence quotes.",
        ],
    )
    add_p(
        doc,
        "Layouts: classic, boxed, two_column (PO appears in extracted text before vendor), modern (lines page 1, totals page 2), compact. "
        "Distractors on every PDF: Harbor Street Holdings remit-to, INV-000000 barcode, PO-9999, coupon $8,888.00.",
    )
    add_table(
        doc,
        ["Cases", "Split", "Scenario", "Expected decision"],
        [
            ["001–015", "development", "exact_match", "READY_FOR_DRAFT"],
            ["016–020", "development", "price_within", "READY_FOR_DRAFT"],
            ["021–025", "development", "price_over", "HUMAN_REVIEW / PRICE_VARIANCE"],
            ["026–030", "development", "qty_exceeds_receipt", "HUMAN_REVIEW / QUANTITY_EXCEEDS_RECEIPT"],
            ["031–035", "holdout", "duplicate", "HUMAN_REVIEW / DUPLICATE_INVOICE"],
            ["036–040", "holdout", "missing_po", "HUMAN_REVIEW / PO_NOT_FOUND (PDF cites PO-MISSING-N)"],
            ["041–044", "holdout", "math_error", "HUMAN_REVIEW / INVOICE_MATH_ERROR"],
            ["045–047", "holdout", "vendor_mismatch", "HUMAN_REVIEW / VENDOR_MISMATCH"],
            ["048, 050", "holdout", "missing invoice number", "HUMAN_REVIEW / MISSING_INVOICE_NUMBER"],
            ["049", "holdout", "ambiguous vendor (Other Corp letterhead)", "HUMAN_REVIEW / AMBIGUOUS_FIELD"],
        ],
    )

    # ------------------------------------------------------------------
    add_heading(doc, "9. Local setup (no GPU)", 1)
    add_code(
        doc,
        "uv python install 3.12\n"
        "uv sync --python 3.12\n"
        "cp .env.example .env\n"
        "\n"
        "make generate-data   # regenerate the 50 PDFs if needed\n"
        "make init-db\n"
        "make test            # MockLLM, preseeds expected.json\n"
        "make evaluate-accuracy\n"
        "make run-ui          # Streamlit demo, still mock\n"
        "make report",
    )
    add_p(
        doc,
        "Open a PDF with open tests/fixtures/development/case_001/invoice.pdf. "
        "The Streamlit UI shows PO/receipt JSON and pipeline results; it does not render the PDF page.",
    )

    # ------------------------------------------------------------------
    add_heading(doc, "10. How to set up the GPU", 1)
    add_p(
        doc,
        "Do not silently change these pins. They are the measured configuration.",
        bold=True,
    )
    add_table(
        doc,
        ["Pin", "Value"],
        [
            ["Model", "Qwen/Qwen3-8B"],
            ["Revision", "b968826d9c46dd6066d109eabc6255188de91218"],
            ["Docker image", "vllm/vllm-openai:v0.27.1"],
            [
                "Digest",
                "sha256:0a51ea5b4ae2dc5d81890e5173f54203d2a3ae0cfffe51b8fd2afd4391bfd967",
            ],
            ["Batch invariance", "VLLM_BATCH_INVARIANT=1"],
            ["Thinking", "disabled (enable_thinking=false)"],
            ["Structured output", "xgrammar + OpenAI json_schema (not guided_*)"],
            ["Sampling", "temperature=0, top_p=1, seed=42, n=1"],
            ["Suggested GPU", "NVIDIA H100 (compute 9.0) — safest for batch-invariance docs"],
        ],
    )

    add_heading(doc, "10.1 Option A — local Docker GPU", 2)
    add_p(doc, "Requires NVIDIA drivers, Docker, and nvidia-container-toolkit.")
    add_code(doc, "bash deployment/start_vllm.sh")
    add_p(
        doc,
        "That pulls the image by digest and serves OpenAI-compatible chat on port 8000. "
        "Wait until the model is listed (first start downloads weights into the Hugging Face cache). "
        "In .env set:",
    )
    add_code(
        doc,
        "VLLM_BASE_URL=http://localhost:8000/v1\n"
        "VLLM_API_KEY=not-a-real-key\n"
        "MODEL_NAME=Qwen/Qwen3-8B",
    )

    add_heading(doc, "10.2 Option B — RunPod H100", 2)
    add_bullets(
        doc,
        [
            "Rent 1x H100 80GB (PCIe or SXM). Secure Cloud preferred.",
            "Use the exact image digest above. Set env VLLM_BATCH_INVARIANT=1.",
            "Start command: same flags as deployment/start_vllm.sh (thinking off, xgrammar backend).",
            "When healthy, set VLLM_BASE_URL in .env to the pod’s OpenAI-compatible proxy URL (usually …/v1).",
            "Auto-stop when eval finishes. Do not leave an idle GPU rented.",
            "Never commit the live API key. See docs/SECRET_SWEEP.md.",
        ],
    )

    add_heading(doc, "10.3 Run live evaluation", 2)
    add_p(
        doc,
        "Confirm the server is up, then from the repo root:",
    )
    add_code(
        doc,
        "uv run python scripts/evaluate_accuracy.py --live\n"
        "uv run python scripts/evaluate_determinism.py --live   # optional, longer",
    )
    add_p(
        doc,
        "--live does not preseed from expected.json. Qwen must extract from canonical PDF text. "
        "Score against expected decision and exception codes. Do not tune on holdout. "
        "Do not edit holdout expected.json after seeing failures. "
        "The published 50/50 was on the old labeled invoices; expect the harder PDFs to score lower.",
    )

    add_heading(doc, "10.4 What was not run last time", 2)
    add_bullets(
        doc,
        [
            "Negative control: same GPU with VLLM_BATCH_INVARIANT=0. That is the experiment that should break.",
            "Pod restart (process-lifetime invariance).",
            "Full 9,000-extract matrix (50 × 20 × {1,8,32} × 3 restarts).",
            "Live accuracy on the current harder PDFs.",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(doc, "11. How to read the outputs", 1)
    add_p(
        doc,
        "Pipeline output is a WorkflowResult. Eval scripts write files under reports/. "
        "The 2026-08-13 files describe the easy-invoice GPU run unless you overwrite them.",
    )

    add_heading(doc, "11.1 A single case (CLI or UI)", 2)
    add_code(doc, "make run-local\n# or: uv run python scripts/run_case.py --case case_001 --mode evaluate")
    add_p(doc, "JSON fields that matter:")
    add_table(
        doc,
        ["Field", "How to read it"],
        [
            ["decision", "READY_FOR_DRAFT (clean match + balanced journal) or HUMAN_REVIEW."],
            ["exception_codes", "Empty on a clean match. Any code forces review. Sorted unique strings."],
            ["extracted_invoice", "What the LLM (or preseed) produced. Compare to expected.json extraction."],
            ["control_checks", "Named pass/fail from validation and matching (PRICE, QTY_RCV, EVIDENCE, …)."],
            ["proposed_journal", "Null unless the match was clean. balanced must be true to draft."],
            ["idempotency_key", "64-char SHA-256. Same PDF+PO+receipts+prompt+policy+model → same key."],
            ["draft_bill_id", "Set only in create_draft mode on READY_FOR_DRAFT. Shape DRAFT- plus 12 hex chars."],
            ["ocr_required", "True only if pypdf found no text. Fixtures should be False."],
            ["messages", "Human-readable notes (cached idempotent hit, OCR stop, …)."],
        ],
    )

    add_heading(doc, "11.2 expected.json vs what the model returned", 2)
    add_p(
        doc,
        "expected.json is the answer key, not ERP data. Top-level decision and exception_codes are what accuracy scoring uses. "
        "extraction is the perfect extract: normalized amounts like \"250.00\" and ISO dates. "
        "evidence.quote is the messy substring from the PDF (\"$250.00\"). "
        "If live Qwen picks INV-000000 or Harbor Street Holdings, extraction is wrong even if JSON is schema-valid — "
        "the Python matcher will then likely HUMAN_REVIEW, which may still “pass” on a review case or fail a READY_FOR_DRAFT case.",
    )

    add_heading(doc, "11.3 reports/accuracy_gpu.json and accuracy_mock.json", 2)
    add_p(
        doc,
        "Written by evaluate_accuracy.py. Top-level keys: development, holdout, llm, gpu, preseeded_from_expected. "
        "Each split has total, correct, accuracy, schema_valid, false_ready_for_draft, and rows[].",
    )
    add_table(
        doc,
        ["Row field", "How to read it"],
        [
            ["pass", "True if decision matched expected and exception codes were OK."],
            ["expected_decision / got_decision", "Mismatch here is an accuracy failure. got=ERROR means the client threw."],
            ["expected_codes / got_codes", "For READY_FOR_DRAFT, got must be empty. For review, expected codes must be a subset of got."],
            ["false_ready_for_draft", "Count of cases the system would have auto-drafted but should not have. This is the dangerous error."],
            ["schema_valid", "ExtractedInvoice parsed. Schema-valid can still be the wrong vendor or total."],
        ],
    )
    add_p(
        doc,
        "Terminal output prints a running score: [development] case_001 pass=True got=READY_FOR_DRAFT expected=READY_FOR_DRAFT. "
        "A checkpoint file reports/accuracy_gpu.partial.json is written as cases finish so a kill does not lose progress.",
    )

    add_heading(doc, "11.4 reports/results.csv", 2)
    add_p(
        doc,
        "One row per case from the last live accuracy run. Columns: split, case_id, scenario, expected_decision, "
        "got_decision, expected_codes, got_codes, pass, llm, gpu. Filter pass=False first. "
        "The checked-in CSV is the 50/50 easy-invoice GPU run.",
    )

    add_heading(doc, "11.5 reports/determinism.json", 2)
    add_p(
        doc,
        "Local section: pdf_stable and canonical_stable should be true (same file hashed twice). "
        "GPU section: for each concurrency, unique hashes should equal the number of cases (one hash per case, not one per repeat). "
        "cases_with_>1_hash should be 0. cross_machine.pass true means two machines produced the same extract hashes. "
        "Status NOT_EXECUTED means that cell was skipped — do not treat it as a pass. "
        "Hashes are of ExtractedInvoice JSON with timestamps and request ids stripped.",
    )

    add_heading(doc, "11.6 reports/final_report.md", 2)
    add_p(
        doc,
        "Human summary produced by make report. Read Scope, Pins, Accuracy, Determinism table, Honest limitations. "
        "If GPU spend or hours look like estimates, the report is supposed to say so. "
        "The checked-in report is the easy-invoice 2026-08-13 run (~$38, ~13 GPU hours, pods p7tdoz42gtf92g and n56so1rkgs0ckh).",
    )

    add_heading(doc, "11.7 Other report files", 2)
    add_table(
        doc,
        ["File", "How to read it"],
        [
            [
                "environment.json",
                "Python version, package versions, git commit, dirty tree, model pins. Confirm you are looking at the commit you think you are.",
            ],
            [
                "gpu_runtime.json",
                "Which determinism cells ran, pod datacenter/GPU, skipped cells (negative control, restart, full 9000).",
            ],
            [
                "billing.json",
                "RunPod API billed USD and hours. Not a credit-card dump. Keys should never appear here.",
            ],
            [
                "determinism_records_r0.jsonl / r1.jsonl",
                "One JSON object per extract. r0 = PCIe US-KS-2, r1 = SXM AP-IN-1. Use if you need to debug a hash split.",
            ],
            [
                "determinism_r0_pcie_usks2.json",
                "Baseline snapshot from the first machine for cross-machine compare.",
            ],
            [
                "runpod_quote.md",
                "Pre-rental planning notes (written when spend was still $0). Historical.",
            ],
            [
                "failures/",
                "Empty README if no live failures. Failed --live cases would be dumped here. Check this folder after a hard-PDF GPU run.",
            ],
        ],
    )

    add_heading(doc, "11.8 Streamlit UI", 2)
    add_p(
        doc,
        "make run-ui. Select a fixture: you will see scenario, expected decision, PO JSON, receipt JSON. "
        "Run workflow: decision color (green draft / orange review), exception codes, tabs for extracted fields, "
        "evidence quotes, matching controls, proposed journal, audit hashes (PDF SHA-256, canonical hash). "
        "Upload tab uses the mock regex parser if there is no expected.json for that file. "
        "This UI is not the GPU path.",
    )

    # ------------------------------------------------------------------
    add_heading(doc, "12. What this repo is not", 1)
    add_bullets(
        doc,
        [
            "Not a production ERP integration (the adapter protocol exists; QuickBooks/NetSuite do not).",
            "Not an OCR product. Scanned / image-only PDFs stop at ingest.",
            "Not a vision model. Qwen reads pypdf text, not pixels.",
            "Not a general AP agent. Two decisions, one policy file, five fictional vendors.",
            "Local 50/50 accuracy does not mean Qwen still extracts the harder invoices. That needs --live.",
        ],
    )

    add_heading(doc, "13. Suggested reading order", 1)
    add_bullets(
        doc,
        [
            "This document (or docs/how_it_works.md) for the map.",
            "app/pipeline.py — the state machine.",
            "app/matching.py and config/policy.yaml — the clerk and the rules.",
            "app/llm_client.py — mock vs live.",
            "scripts/generate_cases.py — how the PDFs were drawn.",
            "Open case_001, case_003, and case_004 PDFs — classic vs reading-order trap vs two-page.",
            "docs/technical_decisions.md before changing GPU pins.",
            "Then start vLLM and run evaluate_accuracy.py --live.",
        ],
    )

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(18)
    r = footer.add_run(
        "Accounting Agent  ·  evaluation prototype  ·  generated from the repository, not from a live GPU run"
    )
    set_run_font(r, size=9, italic=True, color=RGBColor(0x66, 0x66, 0x66))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
