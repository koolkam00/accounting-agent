"""Capacity and margin model for QoE Evidence Review.

**Capacity is not value.** Freed salaried hours become value only when they
are redeployed, when a cost is actually avoided, or when delivery economics
change. A salaried associate costs the same whether the hours are used or not.
So the model keeps four views apart and never adds them together, because the
same freed hour can only be used once:

(a) Capacity only: hours freed and equivalent engagements. Hours, not dollars.
(b) Fixed-fee margin: on a fixed fee, fewer delivery hours lift engagement
    margin at standard cost. On hourly work the client keeps the saving, so
    billed fees fall. The firm's cash profit moves only through (c) or (d).
(c) Redeployment: value only for the share of freed hours actually converted
    to billable or other engagement work.
(d) Avoided incremental hire: only if the team would otherwise hire, capped at
    whole FTEs freed.

**No result is built in.** The preparer and reviewer time reductions are
MEASURED inputs that default to None. Until a timed comparison supplies them,
every output that depends on them is None here and "n/a" in the workbook.
Every other input is an ASSUMPTION with a note. The one example set
(``CommercialInputs.illustrative()``) is labeled ILLUSTRATIVE wherever it
appears.

``LINES`` is the single definition of each output: ``compute`` evaluates it in
Decimal and ``write_workbook`` renders the same line as an Excel formula, so
the Python model and the workbook cannot drift apart.
"""

from __future__ import annotations

from dataclasses import dataclass, fields
from decimal import ROUND_DOWN, Decimal
from pathlib import Path
from typing import Callable, Mapping, Optional

from qoe.money import D

ASSUMPTION = "ASSUMPTION"
MEASURED = "MEASURED"

CAPACITY_VS_VALUE = (
    "Freed salaried hours are capacity, not money. They become economic value only when they are "
    "redeployed, when a cost is actually avoided, or when delivery economics change."
)
NO_RESULT_NOTICE = (
    "No time saving has been measured. Scenario results appear only after measured preparer and "
    "reviewer time reductions are entered."
)

NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = "0.0%"

STATUS_AWAITING = "AWAITING_MEASUREMENT"
STATUS_INCOMPLETE = "INCOMPLETE"
STATUS_COMPLETE = "COMPLETE"


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class InputSpec:
    key: str
    ref: str
    label: str
    unit: str
    basis: str  # ASSUMPTION | MEASURED
    fraction: bool  # stored as a fraction (0.35), shown as a percentage
    number_format: str
    illustrative: Decimal
    note: str


INPUT_SPECS: tuple[InputSpec, ...] = (
    InputSpec(
        key="adjustments_per_engagement",
        ref="I-01",
        label="Management adjustments tested per engagement",
        unit="adjustments",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("12"),
        note="Seller-proposed EBITDA adjustments in a typical QoE. Use the median of your recent engagements.",
    ),
    InputSpec(
        key="preparer_hours_per_adjustment",
        ref="I-02",
        label="Baseline preparer hours per adjustment",
        unit="hours",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM1,
        illustrative=Decimal("6"),
        note=(
            "Associate / senior hours per adjustment today: GL pull, tie-out, document review, support schedule. "
            "Take it from time codes, not recollection."
        ),
    ),
    InputSpec(
        key="reviewer_hours_per_adjustment",
        ref="I-03",
        label="Baseline reviewer hours per adjustment",
        unit="hours",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM1,
        illustrative=Decimal("1.5"),
        note="Manager / senior manager review hours per adjustment today.",
    ),
    InputSpec(
        key="other_hours_per_engagement",
        ref="I-04",
        label="Other engagement hours (unaffected by the tool)",
        unit="hours",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("500"),
        note=(
            "Hours on the rest of the QoE: net working capital, net debt, databook, report, meetings. "
            "Needed to express freed hours as whole engagements."
        ),
    ),
    InputSpec(
        key="preparer_time_reduction",
        ref="I-05",
        label="Measured preparer time reduction",
        unit="% of baseline hours",
        basis=MEASURED,
        fraction=True,
        number_format=PCT,
        illustrative=Decimal("0.40"),
        note=(
            "From a timed comparison of the same adjustments with and without the tool. Leave blank until "
            "measured: the model shows no result without it."
        ),
    ),
    InputSpec(
        key="reviewer_time_reduction",
        ref="I-06",
        label="Measured reviewer time reduction",
        unit="% of baseline hours",
        basis=MEASURED,
        fraction=True,
        number_format=PCT,
        illustrative=Decimal("0.15"),
        note="Reviewer hours with vs without the tool, from the same timed comparison. Leave blank until measured.",
    ),
    InputSpec(
        key="engagements_per_year",
        ref="I-07",
        label="QoE engagements per year per team",
        unit="engagements",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("20"),
        note="Engagements the team delivers in a year.",
    ),
    InputSpec(
        key="preparer_cost_per_hour",
        ref="I-08",
        label="Blended preparer cost per hour",
        unit="USD / hour",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("90"),
        note="Fully loaded cost (salary, benefits, overhead) per available hour. Cost, not bill rate.",
    ),
    InputSpec(
        key="reviewer_cost_per_hour",
        ref="I-09",
        label="Blended reviewer cost per hour",
        unit="USD / hour",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("175"),
        note="Fully loaded cost per available hour for managers / senior managers.",
    ),
    InputSpec(
        key="fixed_fee_share",
        ref="I-10",
        label="Share of engagements on a fixed fee",
        unit="% of engagements",
        basis=ASSUMPTION,
        fraction=True,
        number_format=PCT,
        illustrative=Decimal("0.60"),
        note="The rest are billed hourly.",
    ),
    InputSpec(
        key="fixed_fee_per_engagement",
        ref="I-11",
        label="Fixed fee per QoE",
        unit="USD",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("150000"),
        note="Average fixed fee, excluding expenses.",
    ),
    InputSpec(
        key="net_bill_rate_per_hour",
        ref="I-12",
        label="Blended net realized bill rate",
        unit="USD / hour",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("300"),
        note="Standard rate x realization. Prices hourly engagements and values redeployed hours.",
    ),
    InputSpec(
        key="redeployment_rate",
        ref="I-13",
        label="Share of freed hours redeployed",
        unit="% of freed hours",
        basis=ASSUMPTION,
        fraction=True,
        number_format=PCT,
        illustrative=Decimal("0.50"),
        note="Share actually converted to billable or other engagement work. Depends on demand, not on the tool.",
    ),
    InputSpec(
        key="chargeable_hours_per_fte",
        ref="I-14",
        label="Chargeable hours per FTE per year",
        unit="hours / year",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("1500"),
        note="Converts freed hours into full-time equivalents.",
    ),
    InputSpec(
        key="hires_planned",
        ref="I-15",
        label="Incremental hires planned that freed capacity could replace",
        unit="hires",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("1"),
        note="0 if the team is not planning to hire. Capped at whole FTEs actually freed.",
    ),
    InputSpec(
        key="cost_per_hire",
        ref="I-16",
        label="Fully loaded annual cost per hire",
        unit="USD / year",
        basis=ASSUMPTION,
        fraction=False,
        number_format=NUM,
        illustrative=Decimal("140000"),
        note="Salary, benefits, overhead, and recruiting cost for the role that would have been hired.",
    ),
)

SPEC_BY_KEY = {s.key: s for s in INPUT_SPECS}
MEASURED_KEYS = tuple(s.key for s in INPUT_SPECS if s.basis == MEASURED)


@dataclass(frozen=True)
class CommercialInputs:
    """Model inputs; None means "not entered". Fractions are 0..1, never 35 for 35%."""

    adjustments_per_engagement: Optional[Decimal] = None
    preparer_hours_per_adjustment: Optional[Decimal] = None
    reviewer_hours_per_adjustment: Optional[Decimal] = None
    other_hours_per_engagement: Optional[Decimal] = None
    preparer_time_reduction: Optional[Decimal] = None
    reviewer_time_reduction: Optional[Decimal] = None
    engagements_per_year: Optional[Decimal] = None
    preparer_cost_per_hour: Optional[Decimal] = None
    reviewer_cost_per_hour: Optional[Decimal] = None
    fixed_fee_share: Optional[Decimal] = None
    fixed_fee_per_engagement: Optional[Decimal] = None
    net_bill_rate_per_hour: Optional[Decimal] = None
    redeployment_rate: Optional[Decimal] = None
    chargeable_hours_per_fte: Optional[Decimal] = None
    hires_planned: Optional[Decimal] = None
    cost_per_hire: Optional[Decimal] = None

    def __post_init__(self) -> None:
        for f in fields(self):
            value = getattr(self, f.name)
            if value is None:
                continue
            if not isinstance(value, Decimal):
                raise TypeError(f"{f.name} must be Decimal or None, got {type(value).__name__}")
            if value < 0:
                raise ValueError(f"{f.name} cannot be negative ({value})")
            if SPEC_BY_KEY[f.name].fraction and value > 1:
                raise ValueError(f"{f.name} is a fraction between 0 and 1 (enter 0.35 for 35%), got {value}")

    def as_dict(self) -> dict[str, Optional[Decimal]]:
        return {f.name: getattr(self, f.name) for f in fields(self)}

    @classmethod
    def from_mapping(cls, data: Mapping[str, object]) -> CommercialInputs:
        """Parse user-supplied values (numbers or strings); blank or None stays None."""
        unknown = sorted(set(data) - set(SPEC_BY_KEY))
        if unknown:
            raise KeyError(f"unknown commercial model inputs: {', '.join(unknown)}")
        parsed: dict[str, Optional[Decimal]] = {}
        for key, value in data.items():
            parsed[key] = None if value is None or (isinstance(value, str) and not value.strip()) else D(value)
        return cls(**parsed)

    @classmethod
    def illustrative(cls) -> CommercialInputs:
        """The one example set. ILLUSTRATIVE: not measured, not a forecast, not a benchmark."""
        return cls(**{s.key: s.illustrative for s in INPUT_SPECS})


# ---------------------------------------------------------------------------
# Output lines (single definition for Python and Excel)
# ---------------------------------------------------------------------------

Values = Mapping[str, Decimal]


@dataclass(frozen=True)
class Line:
    key: str
    ref: str
    section: str  # a | b | c | d
    label: str
    number_format: str
    deps: tuple[str, ...]  # input keys and earlier line keys
    excel: str  # expression over {dep} placeholders
    calc: Callable[[Values], Optional[Decimal]]
    how: str


def _div(num: Decimal, den: Decimal) -> Optional[Decimal]:
    return None if den == 0 else num / den


SECTIONS: dict[str, tuple[str, str]] = {
    "a": (
        "(a) Capacity only: hours, not dollars",
        "Salaried staff cost the same whether or not these hours are used. Equivalent engagements are "
        "capacity for more work only if there is demand to fill it.",
    ),
    "b": (
        "(b) Fixed-fee margin: delivery economics",
        "Engagement margin at standard cost (hours x cost rate). On a fixed fee the fee is unchanged, so "
        "fewer hours lift margin. On hourly work the client keeps the saving: billed fees fall with hours. "
        "The firm's cash profit changes only if the freed hours are redeployed (c) or a hire is avoided (d).",
    ),
    "c": (
        "(c) Redeployment: value only for hours actually redeployed",
        "Hours that are not redeployed are idle capacity with no economic value.",
    ),
    "d": (
        "(d) Avoided incremental hire",
        "Applies only if the team would otherwise hire. A fraction of an FTE does not avoid a hire.",
    ),
}

LINES: tuple[Line, ...] = (
    Line(
        key="baseline_review_hours",
        ref="a.1",
        section="a",
        label="Baseline adjustment-review hours per engagement",
        number_format=NUM,
        deps=("adjustments_per_engagement", "preparer_hours_per_adjustment", "reviewer_hours_per_adjustment"),
        excel="{adjustments_per_engagement}*({preparer_hours_per_adjustment}+{reviewer_hours_per_adjustment})",
        calc=lambda v: (
            v["adjustments_per_engagement"] * (v["preparer_hours_per_adjustment"] + v["reviewer_hours_per_adjustment"])
        ),
        how="I-01 x (I-02 + I-03)",
    ),
    Line(
        key="preparer_hours_freed",
        ref="a.2",
        section="a",
        label="Preparer hours freed per engagement",
        number_format=NUM1,
        deps=("adjustments_per_engagement", "preparer_hours_per_adjustment", "preparer_time_reduction"),
        excel="{adjustments_per_engagement}*{preparer_hours_per_adjustment}*{preparer_time_reduction}",
        calc=lambda v: (
            v["adjustments_per_engagement"] * v["preparer_hours_per_adjustment"] * v["preparer_time_reduction"]
        ),
        how="I-01 x I-02 x I-05",
    ),
    Line(
        key="reviewer_hours_freed",
        ref="a.3",
        section="a",
        label="Reviewer hours freed per engagement",
        number_format=NUM1,
        deps=("adjustments_per_engagement", "reviewer_hours_per_adjustment", "reviewer_time_reduction"),
        excel="{adjustments_per_engagement}*{reviewer_hours_per_adjustment}*{reviewer_time_reduction}",
        calc=lambda v: (
            v["adjustments_per_engagement"] * v["reviewer_hours_per_adjustment"] * v["reviewer_time_reduction"]
        ),
        how="I-01 x I-03 x I-06",
    ),
    Line(
        key="hours_freed_per_engagement",
        ref="a.4",
        section="a",
        label="Total hours freed per engagement",
        number_format=NUM1,
        deps=("preparer_hours_freed", "reviewer_hours_freed"),
        excel="{preparer_hours_freed}+{reviewer_hours_freed}",
        calc=lambda v: v["preparer_hours_freed"] + v["reviewer_hours_freed"],
        how="a.2 + a.3",
    ),
    Line(
        key="hours_freed_per_year",
        ref="a.5",
        section="a",
        label="Hours freed per team per year",
        number_format=NUM,
        deps=("hours_freed_per_engagement", "engagements_per_year"),
        excel="{hours_freed_per_engagement}*{engagements_per_year}",
        calc=lambda v: v["hours_freed_per_engagement"] * v["engagements_per_year"],
        how="a.4 x I-07",
    ),
    Line(
        key="engagement_hours_with_tool",
        ref="a.6",
        section="a",
        label="Total hours per engagement with the tool",
        number_format=NUM,
        deps=("baseline_review_hours", "other_hours_per_engagement", "hours_freed_per_engagement"),
        excel="{baseline_review_hours}+{other_hours_per_engagement}-{hours_freed_per_engagement}",
        calc=lambda v: v["baseline_review_hours"] + v["other_hours_per_engagement"] - v["hours_freed_per_engagement"],
        how="a.1 + I-04 - a.4",
    ),
    Line(
        key="equivalent_engagements",
        ref="a.7",
        section="a",
        label="Equivalent additional engagements per year (if demand exists)",
        number_format=NUM1,
        deps=("hours_freed_per_year", "engagement_hours_with_tool"),
        excel='IF({engagement_hours_with_tool}=0,"n/a",{hours_freed_per_year}/{engagement_hours_with_tool})',
        calc=lambda v: _div(v["hours_freed_per_year"], v["engagement_hours_with_tool"]),
        how="a.5 / a.6",
    ),
    Line(
        key="fixed_fee_engagements",
        ref="b.1",
        section="b",
        label="Fixed-fee engagements per year",
        number_format=NUM1,
        deps=("engagements_per_year", "fixed_fee_share"),
        excel="{engagements_per_year}*{fixed_fee_share}",
        calc=lambda v: v["engagements_per_year"] * v["fixed_fee_share"],
        how="I-07 x I-10",
    ),
    Line(
        key="cost_avoided_per_engagement",
        ref="b.2",
        section="b",
        label="Delivery cost avoided per engagement at standard cost (USD)",
        number_format=NUM,
        deps=("preparer_hours_freed", "preparer_cost_per_hour", "reviewer_hours_freed", "reviewer_cost_per_hour"),
        excel="{preparer_hours_freed}*{preparer_cost_per_hour}+{reviewer_hours_freed}*{reviewer_cost_per_hour}",
        calc=lambda v: (
            v["preparer_hours_freed"] * v["preparer_cost_per_hour"]
            + v["reviewer_hours_freed"] * v["reviewer_cost_per_hour"]
        ),
        how="a.2 x I-08 + a.3 x I-09",
    ),
    Line(
        key="margin_lift_per_fixed_fee_engagement",
        ref="b.3",
        section="b",
        label="Margin lift per fixed-fee engagement (% of fee)",
        number_format=PCT,
        deps=("cost_avoided_per_engagement", "fixed_fee_per_engagement"),
        excel='IF({fixed_fee_per_engagement}=0,"n/a",{cost_avoided_per_engagement}/{fixed_fee_per_engagement})',
        calc=lambda v: _div(v["cost_avoided_per_engagement"], v["fixed_fee_per_engagement"]),
        how="b.2 / I-11",
    ),
    Line(
        key="fixed_fee_margin_per_year",
        ref="b.4",
        section="b",
        label="Engagement-margin improvement on fixed-fee work (USD per year)",
        number_format=NUM,
        deps=("cost_avoided_per_engagement", "fixed_fee_engagements"),
        excel="{cost_avoided_per_engagement}*{fixed_fee_engagements}",
        calc=lambda v: v["cost_avoided_per_engagement"] * v["fixed_fee_engagements"],
        how="b.2 x b.1",
    ),
    Line(
        key="hourly_engagements",
        ref="b.5",
        section="b",
        label="Hourly-billed engagements per year",
        number_format=NUM1,
        deps=("engagements_per_year", "fixed_fee_share"),
        excel="{engagements_per_year}*(1-{fixed_fee_share})",
        calc=lambda v: v["engagements_per_year"] * (1 - v["fixed_fee_share"]),
        how="I-07 x (1 - I-10)",
    ),
    Line(
        key="hourly_fees_forgone_per_year",
        ref="b.6",
        section="b",
        label="Fees forgone on hourly engagements (USD per year)",
        number_format=NUM,
        deps=("hours_freed_per_engagement", "net_bill_rate_per_hour", "hourly_engagements"),
        excel="-{hours_freed_per_engagement}*{net_bill_rate_per_hour}*{hourly_engagements}",
        calc=lambda v: -(v["hours_freed_per_engagement"] * v["net_bill_rate_per_hour"] * v["hourly_engagements"]),
        how="-(a.4 x I-12 x b.5)",
    ),
    Line(
        key="hourly_cost_avoided_per_year",
        ref="b.7",
        section="b",
        label="Delivery cost avoided on hourly engagements at standard cost (USD per year)",
        number_format=NUM,
        deps=("cost_avoided_per_engagement", "hourly_engagements"),
        excel="{cost_avoided_per_engagement}*{hourly_engagements}",
        calc=lambda v: v["cost_avoided_per_engagement"] * v["hourly_engagements"],
        how="b.2 x b.5",
    ),
    Line(
        key="hourly_net_margin_per_year",
        ref="b.8",
        section="b",
        label="Net engagement-margin effect on hourly work (USD per year)",
        number_format=NUM,
        deps=("hourly_fees_forgone_per_year", "hourly_cost_avoided_per_year"),
        excel="{hourly_fees_forgone_per_year}+{hourly_cost_avoided_per_year}",
        calc=lambda v: v["hourly_fees_forgone_per_year"] + v["hourly_cost_avoided_per_year"],
        how="b.6 + b.7",
    ),
    Line(
        key="hours_redeployed_per_year",
        ref="c.1",
        section="c",
        label="Freed hours redeployed per year",
        number_format=NUM,
        deps=("hours_freed_per_year", "redeployment_rate"),
        excel="{hours_freed_per_year}*{redeployment_rate}",
        calc=lambda v: v["hours_freed_per_year"] * v["redeployment_rate"],
        how="a.5 x I-13",
    ),
    Line(
        key="hours_not_redeployed_per_year",
        ref="c.2",
        section="c",
        label="Freed hours not redeployed (no economic value)",
        number_format=NUM,
        deps=("hours_freed_per_year", "hours_redeployed_per_year"),
        excel="{hours_freed_per_year}-{hours_redeployed_per_year}",
        calc=lambda v: v["hours_freed_per_year"] - v["hours_redeployed_per_year"],
        how="a.5 - c.1",
    ),
    Line(
        key="redeployment_value_per_year",
        ref="c.3",
        section="c",
        label="Value of redeployed hours (USD per year)",
        number_format=NUM,
        deps=("hours_redeployed_per_year", "net_bill_rate_per_hour"),
        excel="{hours_redeployed_per_year}*{net_bill_rate_per_hour}",
        calc=lambda v: v["hours_redeployed_per_year"] * v["net_bill_rate_per_hour"],
        how="c.1 x I-12",
    ),
    Line(
        key="fte_freed",
        ref="d.1",
        section="d",
        label="Full-time equivalents freed",
        number_format=NUM2,
        deps=("hours_freed_per_year", "chargeable_hours_per_fte"),
        excel='IF({chargeable_hours_per_fte}=0,"n/a",{hours_freed_per_year}/{chargeable_hours_per_fte})',
        calc=lambda v: _div(v["hours_freed_per_year"], v["chargeable_hours_per_fte"]),
        how="a.5 / I-14",
    ),
    Line(
        key="hires_avoided",
        ref="d.2",
        section="d",
        label="Hires avoided (capped at whole FTEs freed)",
        number_format=NUM,
        deps=("hires_planned", "fte_freed"),
        excel="MIN({hires_planned},ROUNDDOWN({fte_freed},0))",
        calc=lambda v: min(v["hires_planned"], v["fte_freed"].quantize(Decimal("1"), rounding=ROUND_DOWN)),
        how="MIN(I-15, whole FTEs in d.1)",
    ),
    Line(
        key="hire_cost_avoided_per_year",
        ref="d.3",
        section="d",
        label="Hiring cost avoided (USD per year)",
        number_format=NUM,
        deps=("hires_avoided", "cost_per_hire"),
        excel="{hires_avoided}*{cost_per_hire}",
        calc=lambda v: v["hires_avoided"] * v["cost_per_hire"],
        how="d.2 x I-16",
    ),
)

LINE_BY_KEY = {line.key: line for line in LINES}


def status(inputs: CommercialInputs) -> str:
    values = inputs.as_dict()
    if any(values[k] is None for k in MEASURED_KEYS):
        return STATUS_AWAITING
    if any(v is None for v in values.values()):
        return STATUS_INCOMPLETE
    return STATUS_COMPLETE


def compute(inputs: CommercialInputs) -> dict[str, Optional[Decimal]]:
    """Every output line; None where an input it depends on is missing (or a denominator is 0)."""
    values: dict[str, Optional[Decimal]] = dict(inputs.as_dict())
    out: dict[str, Optional[Decimal]] = {}
    for line in LINES:
        deps = [values[d] for d in line.deps]
        result = None if any(d is None for d in deps) else line.calc(values)  # type: ignore[arg-type]
        values[line.key] = result
        out[line.key] = result
    return out


def by_scenario(results: Mapping[str, Optional[Decimal]]) -> dict[str, dict[str, Optional[Decimal]]]:
    grouped: dict[str, dict[str, Optional[Decimal]]] = {s: {} for s in SECTIONS}
    for line in LINES:
        grouped[line.section][line.key] = results[line.key]
    return grouped


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

ASSUMPTIONS_SHEET = "Assumptions"
SCENARIOS_SHEET = "Scenarios"
YOUR_COL = "C"
ILLUSTRATIVE_COL = "F"
FIRST_INPUT_ROW = 13
SCENARIO_YOUR_COL = "C"
SCENARIO_ILLUSTRATIVE_COL = "D"


def input_row(key: str) -> int:
    return FIRST_INPUT_ROW + [s.key for s in INPUT_SPECS].index(key)


def write_workbook(
    path: Path, inputs: Optional[CommercialInputs] = None, source_notes: Optional[Mapping[str, str]] = None
) -> dict[str, int]:
    """Write the Assumptions + Scenarios workbook; returns the Scenarios row of each line key.

    ``inputs`` fills the "Your value" column (default: all blank). The
    illustrative column always carries ``CommercialInputs.illustrative()``.
    Every Scenarios value is a formula on the Assumptions sheet.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    inputs = inputs or CommercialInputs()
    source_notes = source_notes or {}
    your = inputs.as_dict()
    example = CommercialInputs.illustrative().as_dict()

    def font(**kw: object) -> Font:
        return Font(name="Arial", size=kw.pop("size", 10), **kw)  # type: ignore[arg-type]

    blue = "0000FF"
    yellow = PatternFill("solid", start_color="FFFF00", end_color="FFFF00")
    grey = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
    header_fill = PatternFill("solid", start_color="1F3864", end_color="1F3864")
    section_fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws = wb.active
    ws.title = ASSUMPTIONS_SHEET
    last_input_row = FIRST_INPUT_ROW + len(INPUT_SPECS) - 1

    ws["A1"] = "QoE Evidence Review: commercial model inputs"
    ws["A1"].font = font(size=14, bold=True)
    ws["A2"] = "ILLUSTRATIVE MODEL. " + NO_RESULT_NOTICE
    ws["A2"].font = font(bold=True, color="C00000")
    ws["A3"] = CAPACITY_VS_VALUE
    ws["A3"].font = font(italic=True)

    ws["A5"] = "Legend"
    ws["A5"].font = font(bold=True)
    legend: list[tuple[object, Font, Optional[PatternFill], str]] = [
        (123, font(color=blue), None, "Blue text: an input you can change."),
        (
            None,
            font(color=blue),
            yellow,
            "Yellow fill: replace with your firm's own data (or a measured value) before relying on any output.",
        ),
        ("=1+1", font(), None, "Black text: a formula. Do not overwrite."),
        (
            123,
            font(color=blue),
            grey,
            "Grey column: ILLUSTRATIVE example values only. Not measured, not a forecast, not a benchmark.",
        ),
    ]
    for i, (sample, sample_font, fill, text) in enumerate(legend, start=6):
        cell = ws.cell(row=i, column=2, value=sample)
        cell.font = sample_font
        cell.border = box
        if fill is not None:
            cell.fill = fill
        ws.cell(row=i, column=3, value=text).font = font()
    ws["A10"] = (
        "Percentages are stored as fractions (0.35 = 35%). Basis ASSUMPTION = your estimate; "
        "MEASURED = from a timed comparison only."
    )
    ws["A10"].font = font(italic=True)

    headers = ["Ref", "Input", "Your value", "Unit", "Basis (your value)", "Illustrative example", "Note"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=FIRST_INPUT_ROW - 1, column=col, value=text)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    fraction_dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    fraction_dv.error = "Enter a fraction between 0 and 1 (0.35 for 35%)."
    fraction_dv.errorTitle = "Fraction expected"
    number_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    number_dv.error = "Enter a number of zero or more."
    number_dv.errorTitle = "Non-negative number expected"
    ws.add_data_validation(fraction_dv)
    ws.add_data_validation(number_dv)

    for spec in INPUT_SPECS:
        r = input_row(spec.key)
        ws.cell(row=r, column=1, value=spec.ref).font = font()
        ws.cell(row=r, column=2, value=spec.label).font = font()
        value = your[spec.key]
        c = ws.cell(row=r, column=3, value=value)
        c.font = font(color=blue)
        c.fill = yellow
        c.number_format = spec.number_format
        c.border = box
        (fraction_dv if spec.fraction else number_dv).add(c.coordinate)
        ws.cell(row=r, column=4, value=spec.unit).font = font()
        if spec.basis == MEASURED:
            basis = ws.cell(row=r, column=5, value=f'=IF(ISNUMBER({YOUR_COL}{r}),"MEASURED","MEASUREMENT PENDING")')
            basis.font = font(bold=True)
        else:
            ws.cell(row=r, column=5, value=ASSUMPTION).font = font()
        ex = ws.cell(row=r, column=6, value=example[spec.key])
        ex.font = font(color=blue)
        ex.fill = grey
        ex.number_format = spec.number_format
        ex.border = box
        note = spec.note
        if spec.basis == MEASURED:
            note += " Illustrative value is NOT a measured result."
        if source_notes.get(spec.key):
            note += f" Your source: {source_notes[spec.key]}"
        n = ws.cell(row=r, column=7, value=note)
        n.font = font()
        n.alignment = wrap

    check_row = last_input_row + 2
    ws.cell(row=check_row, column=1, value="Input checks").font = font(bold=True)
    measured_refs = ",".join(f"{{col}}{input_row(k)}" for k in MEASURED_KEYS)
    checks = [
        ("Time reductions measured?", f"=COUNT({measured_refs})={len(MEASURED_KEYS)}"),
        ("All inputs entered?", f"=COUNT({{col}}{FIRST_INPUT_ROW}:{{col}}{last_input_row})={len(INPUT_SPECS)}"),
    ]
    check_rows: list[int] = []
    for i, (label, template) in enumerate(checks, start=1):
        r = check_row + i
        check_rows.append(r)
        ws.cell(row=r, column=2, value=label).font = font()
        for col in (YOUR_COL, ILLUSTRATIVE_COL):
            cell = ws[f"{col}{r}"]
            cell.value = template.format(col=col)
            cell.font = font()
    measured_check, complete_check = check_rows

    for col, width in zip("ABCDEFG", (8, 52, 14, 20, 24, 14, 90), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"C{FIRST_INPUT_ROW}"

    # --- Scenarios --------------------------------------------------------
    sc = wb.create_sheet(SCENARIOS_SHEET)
    sc["A1"] = "QoE Evidence Review: scenarios (capacity vs economic value)"
    sc["A1"].font = font(size=14, bold=True)
    sc["A2"] = CAPACITY_VS_VALUE + " The four views below are alternatives: do not add them together."
    sc["A2"].font = font(bold=True, color="C00000")
    sc["A3"] = (
        "Every value on this sheet is a formula on the Assumptions sheet. 'n/a' means an input it needs is blank."
    )
    sc["A3"].font = font(italic=True)
    header_row = 5
    for col, text in enumerate(
        ["Ref", "Line", "Your inputs", "Illustrative example (not measured)", "How computed"], start=1
    ):
        cell = sc.cell(row=header_row, column=col, value=text)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    status_row = header_row + 1
    sc.cell(row=status_row, column=2, value="Status").font = font(bold=True)
    for sc_col, a_col, prefix in (
        (SCENARIO_YOUR_COL, YOUR_COL, ""),
        (SCENARIO_ILLUSTRATIVE_COL, ILLUSTRATIVE_COL, '"ILLUSTRATIVE, not measured: "&'),
    ):
        cell = sc[f"{sc_col}{status_row}"]
        cell.value = (
            f"={prefix}IF(NOT({ASSUMPTIONS_SHEET}!{a_col}{measured_check}),"
            '"Awaiting measured time reduction: no result",'
            f'IF({ASSUMPTIONS_SHEET}!{a_col}{complete_check},"All inputs entered","Some inputs blank: see n/a"))'
        )
        cell.font = font(bold=True)
        cell.alignment = wrap

    rows: dict[str, int] = {}
    r = status_row + 2
    for section, (title, note) in SECTIONS.items():
        for col in range(1, 6):
            sc.cell(row=r, column=col).fill = section_fill
        sc.cell(row=r, column=1, value=title).font = font(bold=True)
        r += 1
        for line in (ln for ln in LINES if ln.section == section):
            rows[line.key] = r
            sc.cell(row=r, column=1, value=line.ref).font = font()
            sc.cell(row=r, column=2, value=line.label).font = font()
            for sc_col, a_col in ((SCENARIO_YOUR_COL, YOUR_COL), (SCENARIO_ILLUSTRATIVE_COL, ILLUSTRATIVE_COL)):
                refs = {
                    dep: (f"{sc_col}{rows[dep]}" if dep in rows else f"{ASSUMPTIONS_SHEET}!${a_col}${input_row(dep)}")
                    for dep in line.deps
                }
                expr = line.excel.format(**refs)
                guard = ",".join(refs[d] for d in line.deps)
                cell = sc[f"{sc_col}{r}"]
                cell.value = f'=IF(COUNT({guard})={len(line.deps)},{expr},"n/a")'
                cell.font = font()
                cell.number_format = line.number_format
                cell.alignment = Alignment(horizontal="right")
            sc.cell(row=r, column=5, value=line.how).font = font(color="595959")
            r += 1
        n = sc.cell(row=r, column=2, value=note)
        n.font = font(italic=True, color="595959")
        n.alignment = wrap
        sc.row_dimensions[r].height = 42 if len(note) > 120 else 28
        r += 2

    for col, width in zip("ABCDE", (7, 64, 16, 20, 30), strict=True):
        sc.column_dimensions[col].width = width
    sc.freeze_panes = f"C{header_row + 1}"

    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return rows


def read_back(path: Path, rows: Mapping[str, int]) -> dict[str, dict[str, object]]:
    """Cached Scenarios values after a recalculation: {"your": {...}, "illustrative": {...}}."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    try:
        sc = wb[SCENARIOS_SHEET]
        return {
            "your": {k: sc[f"{SCENARIO_YOUR_COL}{r}"].value for k, r in rows.items()},
            "illustrative": {k: sc[f"{SCENARIO_ILLUSTRATIVE_COL}{r}"].value for k, r in rows.items()},
        }
    finally:
        wb.close()


def compare_with_model(
    cached: Mapping[str, object], expected: Mapping[str, Optional[Decimal]], tolerance: Decimal = Decimal("0.005")
) -> list[str]:
    """Lines where the recalculated workbook disagrees with ``compute`` (empty list = agree)."""
    problems: list[str] = []
    for key, want in expected.items():
        got = cached.get(key)
        if want is None:
            if got != "n/a":
                problems.append(f"{key}: workbook {got!r}, model n/a")
        elif not isinstance(got, (int, float)) or isinstance(got, bool):
            problems.append(f"{key}: workbook {got!r}, model {want}")
        elif abs(D(got) - want) > tolerance * max(Decimal(1), abs(want)):
            problems.append(f"{key}: workbook {got}, model {want}")
    return problems


__all__ = [
    "ASSUMPTION",
    "CAPACITY_VS_VALUE",
    "INPUT_SPECS",
    "LINES",
    "MEASURED",
    "SECTIONS",
    "CommercialInputs",
    "by_scenario",
    "compare_with_model",
    "compute",
    "read_back",
    "status",
    "write_workbook",
]
