#!/usr/bin/env python3
"""Recalculate clay-icp-tables.xlsx with LibreOffice and check it against the CSVs and the Python view logic.

1. Baseline: every formula column LibreOffice can compute (score, readiness, view helpers) equals the CSV.
2. Scenario: change status, do_not_contact, known_to_founder, hold_before_outreach and founder_ok_to_queue
   on a copy, recalculate, and check the helpers equal what stage3_merge.view() gives for the same edits.
3. The daily-view tabs: LibreOffice 24.2 has no SORT or FILTER, so the test reads each tab's formula,
   checks it filters the right column and sorts on daily_view_rank, applies that FILTER and SORT to the
   recalculated values, and checks the result equals daily-view-*.csv (baseline) and the scenario's view.
Writes its scratch files under _work/test_workbook/. Prints PASS or the first differences.
    python3 -B clay-icp-tables/_build/test_workbook.py
"""
import os
import re
import shutil
import subprocess
import sys

import pandas as pd
from openpyxl import load_workbook

import common as C
import config as K
import stage3_merge as M

SCR = os.path.join(C.WORK, "test_workbook")
XLSX = os.path.join(C.A_DIR, "clay-icp-tables.xlsx")
CHECK = ["score", "ready_for_outreach", "in_flight", "group_blocked", "eligible", "view_pick", "daily_view_rank",
         "in_daily_view", "row_order"]
FAIL = []


def recalc(src, tag):
    d = os.path.join(SCR, tag)
    shutil.rmtree(d, ignore_errors=True)
    os.makedirs(d)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", d, src],
                   capture_output=True, timeout=600)
    out = os.path.join(d, os.path.basename(src))
    if not os.path.exists(out):
        sys.exit("LibreOffice produced no file")
    return load_workbook(out, data_only=True)


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{round(v, 1)}"
    return str(v)


def sheet_df(ws):
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    return pd.DataFrame([[norm(v) for v in r] for r in rows[1:]], columns=head)


def compare(tag, got, exp, cols):
    exp = exp.copy()
    for c in cols:
        a = got[c].tolist()
        b = [norm(float(x)) if c == "score" and x != "" else x for x in exp[c].tolist()]
        bad = [(exp.row_id.iloc[i], a[i], b[i]) for i in range(len(b)) if a[i] != b[i]]
        if bad:
            FAIL.append(f"{tag}: {c} differs on {len(bad)} rows, first {bad[:3]}")


def emulate_view(wb_formulas, wb_values, tab, src_tab):
    """Apply the tab's own SORT(FILTER()) to the recalculated values."""
    f = wb_formulas[tab]["A2"].value
    text = f.text if hasattr(f, "text") else str(f)
    m = re.search(r"FILTER\('([^']+)'!\$([A-Z]+)\$2:\$([A-Z]+)\$\d+,'[^']+'!\$([A-Z]+)\$2:\$[A-Z]+\$\d+=\"Y\"\),(\d+),TRUE\)", text)
    if not m:
        FAIL.append(f"{tab}: formula not recognised: {text[:120]}")
        return None
    sheet, c1, c2, condcol, k = m.group(1), m.group(2), m.group(3), m.group(4), int(m.group(5))
    ws = wb_values[sheet]
    from openpyxl.utils import column_index_from_string as ci
    head = [c.value for c in ws[1]]
    if head[ci(condcol) - 1] != "in_daily_view":
        FAIL.append(f"{tab}: FILTER condition column is {head[ci(condcol) - 1]}, not in_daily_view")
    block = head[ci(c1) - 1: ci(c2)]
    if block[k - 1] != "daily_view_rank":
        FAIL.append(f"{tab}: SORT column {k} is {block[k - 1]}, not daily_view_rank")
    tab_head = [c.value for c in wb_formulas[tab][1]]
    if tab_head != block:
        FAIL.append(f"{tab}: header row does not match the filtered block")
    df = sheet_df(ws)
    v = df[df.in_daily_view == "Y"][block].copy()
    v["_k"] = v.daily_view_rank.astype(int)
    return v.sort_values("_k").drop(columns="_k").reset_index(drop=True)


def rescore(df, side):
    """What the workbook formulas do to score and readiness after the scenario's edits."""
    df = df.copy()
    if side == "office":
        df["score_known"] = df.known_to_founder.map(lambda x: K.W_KNOWN if x == "Y" else 0)
        df["score_title"] = df.title_match.map(lambda x: K.P_TITLE.get(x, 0))
        parts = ["score_onsite", "score_size", "score_attendance", "score_address", "score_title", "score_known"]
    else:
        df["score_known"] = df.known_to_founder.map(lambda x: K.B_KNOWN if x == "Y" else 0)
        parts = ["score_buy", "score_ny", "score_why", "score_rev", "score_fit", "score_contact", "score_known",
                 "score_review"]
    tot = sum(pd.to_numeric(df[p], errors="coerce").fillna(0) for p in parts).round(1)
    df["score"] = [(f"{t:.1f}" if side == "office" else str(int(t))) if g == "pass" else "0"
                   for t, g in zip(tot, df.gate_result)]
    df = M.readiness(df, side)
    return M.view(df)


def main():
    os.makedirs(SCR, exist_ok=True)
    off = pd.read_csv(os.path.join(C.A_DIR, "offices.csv"), dtype=str, keep_default_na=False)
    br = pd.read_csv(os.path.join(C.A_DIR, "brands.csv"), dtype=str, keep_default_na=False)
    dvo = pd.read_csv(os.path.join(C.A_DIR, "daily-view-offices.csv"), dtype=str, keep_default_na=False)
    dvb = pd.read_csv(os.path.join(C.A_DIR, "daily-view-brands.csv"), dtype=str, keep_default_na=False)
    wf = load_workbook(XLSX)
    # 1. baseline
    wv = recalc(XLSX, "baseline")
    go, gb = sheet_df(wv["Offices"]), sheet_df(wv["Brands"])
    compare("baseline offices", go, off, CHECK)
    compare("baseline brands", gb, br, CHECK)
    for tab, src, dv in (("Daily view - Offices", "Offices", dvo), ("Daily view - Brands", "Brands", dvb)):
        v = emulate_view(wf, wv, tab, src)
        if v is not None and v.row_id.tolist() != dv.row_id.tolist():
            FAIL.append(f"baseline {tab}: {v.row_id.tolist()} vs CSV {dv.row_id.tolist()}")
    # 2. scenario
    edits = {
        "Offices": [("rank", "1", "status", "contacted"), ("rank", "2", "do_not_contact", "Y"),
                    ("company", "Signal AI", "founder_ok_to_queue", "Y"), ("company", "Orbital", "hold_before_outreach", "N"),
                    ("company", "Wellcom Worldwide", "known_to_founder", "Y"),
                    ("company", "Carbyne", "status", "queued")],  # Carbyne queued blocks Prepared 911 (same parent)
        "Brands": [("rank", "1", "status", "queued"), ("company", "Bruce Cost Ginger Ale", "status", "contacted"),
                   ("company", "LesserEvil", "do_not_contact", "Y"), ("rank", "40", "known_to_founder", "Y"),
                   ("rank", "5", "status", "wrong_person")],
    }
    wb = load_workbook(XLSX)
    exp = {}
    for tab, df, side in (("Offices", off, "office"), ("Brands", br, "brand")):
        ws = wb[tab]
        head = [c.value for c in ws[1]]
        d2 = df.copy()
        for key, kv, col, new in edits[tab]:
            idx = d2.index[d2[key] == kv]
            assert len(idx) == 1, (tab, key, kv)
            i = idx[0]
            d2.at[i, col] = new
            ws.cell(row=i + 2, column=head.index(col) + 1, value=new)
        exp[tab] = rescore(d2, side)
    scen = os.path.join(SCR, "scenario.xlsx")
    wb.save(scen)
    ws2 = recalc(scen, "scenario")
    compare("scenario offices", sheet_df(ws2["Offices"]), exp["Offices"], CHECK)
    compare("scenario brands", sheet_df(ws2["Brands"]), exp["Brands"], CHECK)
    wf2 = load_workbook(scen)
    for tab, src in (("Daily view - Offices", "Offices"), ("Daily view - Brands", "Brands")):
        v = emulate_view(wf2, ws2, tab, src)
        e = exp[src]
        e = e[e.in_daily_view == "Y"].copy()
        e["_k"] = e.daily_view_rank.astype(int)
        e = e.sort_values("_k")
        if v is not None and v.row_id.tolist() != e.row_id.tolist():
            FAIL.append(f"scenario {tab}: {v.row_id.tolist()} vs expected {e.row_id.tolist()}")
        if v is not None:
            print(f"scenario {tab}: {', '.join(v.company.tolist())}")
    # spot facts the scenario must show
    eo, eb = exp["Offices"], exp["Brands"]
    def g(df, comp, col):
        return df[df.company == comp][col].iloc[0]
    facts = [("Prepared 911 blocked by Carbyne in flight", g(eo, "Prepared 911", "group_blocked") == "Y"),
             ("Signal AI ready after founder_ok_to_queue", g(eo, "Signal AI", "ready_for_outreach") == "Y"),
             ("Moshi blocked by Bruce Cost in flight", g(eb, "Moshi", "group_blocked") == "Y"),
             ("PRETZELIZED blocked by LesserEvil opt-out", g(eb, "PRETZELIZED", "group_blocked") == "Y")]
    for name, ok in facts:
        if not ok:
            FAIL.append("scenario fact failed: " + name)
    print("FAIL:\n  " + "\n  ".join(FAIL) if FAIL else "PASS: baseline and scenario match; daily-view formulas "
          "filter in_daily_view and sort on daily_view_rank")
    return 0 if not FAIL else 1


if __name__ == "__main__":
    sys.exit(main())
