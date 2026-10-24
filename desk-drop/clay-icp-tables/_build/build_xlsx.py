#!/usr/bin/env python3
"""Build clay-icp-tables/clay-icp-tables.xlsx: the Google Sheets version of both tables, with live views.

Tabs: Offices, Brands (the tables of record), Daily view - Offices, Daily view - Brands, Walk list - Offices,
Founder review - Brands, Enrichment (where Clay's or Hunter's email results are pasted), Dictionary, Scoring,
Routine.

Every computed column is a formula, written out to row 1000 so a row added at the bottom (a referral) computes
too: the contact the views show (primary, or the backup once promote_backup = Y), the channel (warm-intro once
known_to_founder = Y, email once the Enrichment tab holds a verified address), the score and its one-sentence
reason, readiness, the do-not-contact and company-group blocks, and each view's slot. The view tabs are plain
INDEX/MATCH lookups on those slots (no SORT or FILTER), so they compute the same way in Google Sheets, Excel and
LibreOffice. A guard blanks every view if the table's formulas have been pasted over. stage3_merge.derive() is
the same logic in Python; _build/test_workbook.py recalculates this file and checks the two agree.
No macros, no scripts.
"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import common as C
import config as K
import dictionary as DD
import stage3_merge as M

OUT = os.path.join(C.A_DIR, "clay-icp-tables.xlsx")
LAST = 1000  # formulas and validations reach row 1000 so added rows compute
EN_LAST_ROW, EN_LAST_COL = 2000, "GZ"  # the Enrichment tab: up to 2,000 rows and 208 columns of a Clay export
BRAND_NAME = "Desk Drop"   # the working name: change it here or in Routine!B1 (company_name)
TAGLINE = "Emerging brands. Delivered to your workday."
EN = "Enrichment"

FILL = {"founder": PatternFill("solid", fgColor="FFF2CC"), "formula": PatternFill("solid", fgColor="D9D9D9"),
        "E / founder": PatternFill("solid", fgColor="DDEBF7"), "build": PatternFill("solid", fgColor="FFFFFF")}
RED = PatternFill("solid", fgColor="F4CCCC")
HEAD = Font(bold=True)
NUMERIC = {"score", "score_raw", "score_onsite", "score_size", "score_attendance", "score_address", "score_title",
           "score_known", "score_buy", "score_ny", "score_why", "score_rev", "score_fit", "score_contact",
           "score_review", "rank", "onsite_low", "onsite_high", "onsite_central", "onsite_fit_share", "lat", "lon",
           "blocks_outside", "why_now_age_days", "graded_retailer_count_within_ten",
           "nearest_retailer_blocks_outside", "daily_view_rank", "view_slot", "walk_slot", "review_slot", "row_order"}
VIEW_NUMERIC = {"score", "daily_view_rank", "rank", "walk_slot", "review_slot"}
IN_FLIGHT_ARR = "{" + ",".join(f'"{s}"' for s in M.IN_FLIGHT) + "}"
FORMULA_COLS = {
    "office": set(M.SWAP) | {"channel", "score", "score_reason", "score_title", "score_known", "daily_view_rank",
                             "ready_for_outreach", "ready_reason", "primary_work_email", "primary_email_status",
                             "backup_work_email", "backup_email_status"} | set(M.HELPERS["office"]),
    "brand": (set(M.SWAP) - {"title_match"}) | {"channel", "score", "score_reason", "score_known", "daily_view_rank",
                                                "ready_for_outreach", "ready_reason", "primary_work_email",
                                                "primary_email_status", "backup_work_email", "backup_email_status"}
    | set(M.HELPERS["brand"]),
}


def load(name):
    return pd.read_csv(os.path.join(C.A_DIR, name), dtype=str, keep_default_na=False)


def val(col, v):
    if col in NUMERIC and v not in ("", None):
        try:
            f = float(v)
            return int(f) if f.is_integer() and col not in ("lat", "lon", "onsite_fit_share") else f
        except ValueError:
            return v
    return v


def en_lookup(r, field):
    """The value of column `field` for this row's row_id on the Enrichment tab (headers matched by name)."""
    rng = f"'{EN}'!$A$2:${EN_LAST_COL}${EN_LAST_ROW}"
    hdr = f"'{EN}'!$A$1:${EN_LAST_COL}$1"
    return (f'IFERROR(INDEX({rng},MATCH($A{r},INDEX({rng},0,MATCH("row_id",{hdr},0)),0),'
            f'MATCH("{field}",{hdr},0))&"","")')


def formulas(side, L, r):
    """Formula text for row r. L maps column name -> letter."""
    def c(name):
        return f"{L[name]}{r}"

    def rng(name):
        return f"${L[name]}$2:${L[name]}${LAST}"

    def blank(expr):
        return f'=IF($A{r}="","",{expr})'
    f = {}
    for pre, ek, sk in (("primary", "work_email", "email_status"), ("backup", "backup_work_email",
                                                                     "backup_email_status")):
        f[f"{pre}_work_email"] = blank(en_lookup(r, ek))
        x = en_lookup(r, sk)
        f[f"{pre}_email_status"] = blank(f'IF({c(pre + "_work_email")}="","none",IF(OR({x}="verified",{x}="unverified"),'
                                         f'{x},"unverified"))')
    promo = c("promote_backup")
    for out, src in M.SWAP.items():
        if out == "title_match" and side == "brand":
            continue
        f[out] = blank(f'IF({promo}="Y",{c("backup_" + src)}&"",{c("primary_" + src)}&"")')
    f["channel"] = blank(f'IF({c("known_to_founder")}="Y","warm-intro",IF({c("email_status")}="verified","email",'
                         f'IF({c("first_name")}<>"","linkedin",{c("channel_base")}&"")))')
    who = f'IF({promo}="Y",{c("backup_reason_who")}&"",{c("primary_reason_who")}&"")'
    if side == "office":
        f["score_title"] = blank(f'IF({c("title_match")}="exact",{K.P_TITLE["exact"]},IF({c("title_match")}="adjacent",'
                                 f'{K.P_TITLE["adjacent"]},0))')
        f["score_known"] = blank(f'IF({c("known_to_founder")}="Y",{K.W_KNOWN},0)')
        parts = "+".join(f"N({c(x)})" for x in ("score_onsite", "score_size", "score_attendance", "score_address",
                                                  "score_title", "score_known"))
        f["score"] = blank(f'IF({c("gate_result")}="pass",ROUND({parts},1),0)')
        f["score_reason"] = blank(f'IF({c("gate_result")}<>"pass",{c("score_reason_core")}&"",{c("score_reason_core")}'
                                  f'&"; "&{who}&IF({c("known_to_founder")}="Y","; known to the founder (+{K.W_KNOWN})","")&".")')
        f["ready_for_outreach"] = blank(f'IF(AND({c("gate_result")}="pass",{c("hold_before_outreach")}<>"Y",'
                                        f'{c("first_name")}<>"",OR({c("onsite_founder_call")}<>"Y",'
                                        f'{c("founder_ok_to_queue")}="Y")),"Y","N")')
        f["ready_reason"] = blank(f'_xlfn.TEXTJOIN("; ",TRUE,IF({c("gate_result")}<>"pass","gated out",""),'
                                  f'IF({c("hold_before_outreach")}="Y","on hold: "&{c("hold_reason")},""),'
                                  f'IF({c("first_name")}="","no contact named on an opened page",""),'
                                  f'IF(AND({c("onsite_founder_call")}="Y",{c("founder_ok_to_queue")}<>"Y"),'
                                  f'"Tuesday estimate band "&{c("onsite_band")}&": founder\'s call",""))')
    else:
        f["score_known"] = blank(f'IF({c("known_to_founder")}="Y",{K.B_KNOWN},0)')
        parts = "+".join(f"N({c(x)})" for x in ("score_buy", "score_ny", "score_why", "score_rev", "score_fit",
                                                  "score_contact", "score_known", "score_review"))
        f["score"] = blank(f'IF({c("gate_result")}="pass",ROUND({parts},0),0)')
        f["score_reason"] = blank(f'IF({c("gate_result")}<>"pass",{c("score_reason_core")}&"",{c("score_reason_core")}'
                                  f'&", "&{who}&{c("score_reason_tail")}&IF({c("known_to_founder")}="Y",'
                                  f'"; known to the founder (+{K.B_KNOWN})","")&".")')
        f["ready_for_outreach"] = blank(f'IF(AND({c("gate_result")}="pass",{c("hold_before_outreach")}<>"Y",'
                                        f'{c("first_name")}<>"",OR({c("review_blocking")}<>"Y",'
                                        f'{c("founder_review_cleared")}="Y")),"Y","N")')
        f["ready_reason"] = blank(f'_xlfn.TEXTJOIN("; ",TRUE,IF({c("gate_result")}<>"pass","gated out",""),'
                                  f'IF({c("hold_before_outreach")}="Y","on hold: "&{c("hold_reason")},""),'
                                  f'IF({c("first_name")}="","no contact named on an opened page",""),'
                                  f'IF(AND({c("review_blocking")}="Y",{c("founder_review_cleared")}<>"Y"),'
                                  f'"founder review pending (Founder review tab)",""))')
    f["in_flight"] = blank(f'IF(ISNUMBER(MATCH({c("status")},{IN_FLIGHT_ARR},0)),"Y","N")')
    f["dnc_any"] = blank(f'IF(OR({c("do_not_contact")}="Y",{c("status")}="do_not_contact",'
                         f'{c("opt_out_date")}&""<>""),"Y","N")')
    f["group_key"] = blank(f'IF({c("company_group")}&""<>"",{c("company_group")}&"",$A{r})')
    f["group_blocked"] = blank(f'IF(COUNTIFS({rng("group_key")},{c("group_key")},{rng("in_flight")},"Y")'
                               f'+COUNTIFS({rng("group_key")},{c("group_key")},{rng("dnc_any")},"Y")>0,"Y","N")')
    f["eligible"] = blank(f'IF(AND({c("status")}="new",{c("do_not_contact")}="N",{c("dnc_any")}="N",'
                          f'{c("ready_for_outreach")}="Y",{c("group_blocked")}="N"),"Y","N")')
    f["view_pick"] = blank(f'IF({c("eligible")}<>"Y","N",IF(COUNTIFS({rng("group_key")},{c("group_key")},'
                           f'{rng("eligible")},"Y",{rng("score")},">"&{c("score")})+COUNTIFS({rng("group_key")},'
                           f'{c("group_key")},{rng("eligible")},"Y",{rng("score")},{c("score")},{rng("row_order")},'
                           f'"<"&{c("row_order")})=0,"Y","N"))')
    f["daily_view_rank"] = (f'=IF({c("view_pick")}<>"Y","",COUNTIFS({rng("view_pick")},"Y",{rng("score")},'
                            f'">"&{c("score")})+COUNTIFS({rng("view_pick")},"Y",{rng("score")},{c("score")},'
                            f'{rng("row_order")},"<"&{c("row_order")})+1)')
    f["in_daily_view"] = blank(f'IF(AND({c("eligible")}="Y",{c("view_pick")}="Y",{c("daily_view_rank")}&""<>""),'
                               f'IF({c("daily_view_rank")}<={K.DAILY_VIEW_SIZE},"Y","N"),"N")')
    f["view_slot"] = f'=IF({c("in_daily_view")}="Y",{c("daily_view_rank")},"")'
    if side == "office":
        ok, slot, size = "walk_ok", "walk_slot", M.WALK_SIZE
        f[ok] = blank(f'IF(AND({c("channel")}="walk-in",{c("gate_result")}="pass",{c("status")}="new",'
                      f'{c("dnc_any")}="N",{c("group_blocked")}="N",{c("hold_before_outreach")}<>"Y"),"Y","N")')
    else:
        ok, slot, size = "review_ok", "review_slot", M.REVIEW_SIZE
        f[ok] = blank(f'IF(AND({c("gate_result")}="pass",{c("review_blocking")}="Y",{c("founder_review_cleared")}<>"Y",'
                      f'{c("status")}="new",{c("dnc_any")}="N"),"Y","N")')
    k = (f'COUNTIFS({rng(ok)},"Y",{rng("score")},">"&{c("score")})+COUNTIFS({rng(ok)},"Y",{rng("score")},'
         f'{c("score")},{rng("row_order")},"<"&{c("row_order")})+1')
    f[slot] = f'=IF({c(ok)}<>"Y","",IF({k}<={size},{k},""))'
    f["intact"] = f'=IF(AND(_xlfn.ISFORMULA({c("first_name")}),_xlfn.ISFORMULA({c("score")})),"Y","N")'
    f["row_order"] = "=ROW()-1"
    return f


def header_fill(col, ddict):
    who = ddict.get(col, {}).get("edited_by", "build")
    return FILL.get(who, FILL["build"])


VALIDATION = {
    "status": None, "do_not_contact": "Y,N", "known_to_founder": "Y,N", "hold_before_outreach": "Y,N",
    "founder_ok_to_queue": "Y", "promote_backup": "Y", "founder_review_cleared": "Y",
}


def table_sheet(wb, title, df, side, ddict):
    ws = wb.create_sheet(title)
    cols = list(df.columns)
    L = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = HEAD
        cell.fill = header_fill(c, ddict)
    recs = df.to_dict("records")
    for i in range(2, LAST + 1):
        fm = formulas(side, L, i)
        rec = recs[i - 2] if i - 2 < len(recs) else None
        for j, c in enumerate(cols, start=1):
            if c in fm:
                ws.cell(row=i, column=j, value=fm[c])
            elif rec is not None and rec[c] != "":
                ws.cell(row=i, column=j, value=val(c, rec[c]))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"
    for c in cols:
        w = 14
        if c in ("company", "title", "check_before_send", "score_reason", "why_now", "contact_note", "b_notes",
                 "notes", "founder_review_reason", "gate_reason", "ready_reason", "hold_reason", "onsite_basis"):
            w = 40
        ws.column_dimensions[L[c]].width = w
    lists = dict(VALIDATION)
    lists["status"] = ",".join(M.STATUSES[side])
    if "email_status" in L:
        pass
    for c, lst in lists.items():
        if c not in L:
            continue
        dv = DataValidation(type="list", formula1=f'"{lst}"', allow_blank=True, showErrorMessage=True,
                            errorStyle="stop", errorTitle="Not an allowed value",
                            error=f"{c} takes only: {lst.replace(',', ', ')}")
        dv.add(f"{L[c]}2:{L[c]}{LAST}")
        ws.add_data_validation(dv)
    # status, do_not_contact and opt_out_date must agree (R1-05): a mismatch turns the three cells red
    s, d, o = L["status"], L["do_not_contact"], L["opt_out_date"]
    mis = (f'AND($A2<>"",OR(AND(${s}2="do_not_contact",${d}2<>"Y"),AND(${d}2="Y",${s}2<>"do_not_contact"),'
           f'AND(${o}2&""<>"",${d}2<>"Y")))')
    for col in (s, d, o):
        ws.conditional_formatting.add(f"{col}2:{col}{LAST}", FormulaRule(formula=[mis], fill=RED))
    for c in ("status", "do_not_contact", "promote_backup"):
        ws[f"{L[c]}1"].comment = Comment(
            {"status": "CONTRACTS section 4 words only (the cell refuses anything else). An opt-out: status = "
                       "do_not_contact, do_not_contact = Y and opt_out_date, all three, the same day. A red cell "
                       "means the three disagree.",
             "do_not_contact": "Y the same day anyone opts out. Never set back to N. Any one of do_not_contact = Y, "
                               "status = do_not_contact or a filled opt_out_date keeps the row and its whole "
                               "company_group out of every view.",
             "promote_backup": "Y after a wrong_person reply: the backup becomes the contact everywhere (the view, "
                               "title_match, score, channel, check_before_send). Then set status = new."}[c],
            "workstream A")
    return ws, L


def view_sheet(wb, title, src, L, cols, slot, size, empty_msg, comment):
    """A live view: row k shows the source row whose `slot` is k. A guard (cell in the last column + 2) blanks
    the view if any row of the source table has lost its formulas (a paste over the tab)."""
    ws = wb.create_sheet(title)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = HEAD
    gcol = get_column_letter(len(cols) + 2)
    ws[f"{gcol}1"] = "guard"
    ws[f"{gcol}2"] = f"=COUNTIF('{src}'!${L['intact']}$2:${L['intact']}${LAST},\"N\")"
    slot_rng = f"'{src}'!${L[slot]}$2:${L[slot]}${LAST}"
    for k in range(1, size + 1):
        for j, c in enumerate(cols, start=1):
            col_rng = f"'{src}'!${L[c]}$2:${L[c]}${LAST}"
            look = f"INDEX({col_rng},MATCH({k},{slot_rng},0))"
            if c not in VIEW_NUMERIC:
                look += '&""'
            if k == 1 and j == 1:
                fallback = f'"{empty_msg}"'
                guard = '"The table tab was pasted over and its formulas are gone: undo the paste or restore it ' \
                        '(File > Version history). Paste Clay exports on the Enrichment tab only."'
            else:
                fallback = '""'
                guard = '""'
            ws.cell(row=k + 1, column=j, value=f"=IF(${gcol}$2>0,{guard},IFERROR({look},{fallback}))")
    ws["A1"].comment = Comment(comment, "workstream A")
    ws.freeze_panes = "C2"
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.column_dimensions["B"].width = 26
    for name, w in (("score_reason", 50), ("check_before_send", 60), ("founder_review_reason", 60)):
        if name in cols:
            ws.column_dimensions[get_column_letter(cols.index(name) + 1)].width = w
    ws.column_dimensions[gcol].hidden = True
    return ws


def enrichment_sheet(wb, enrich):
    ws = wb.create_sheet(EN)
    cols = list(enrich.columns) if len(enrich.columns) else M.ENRICH_COLS
    for j, c in enumerate(cols, start=1):
        ws.cell(row=1, column=j, value=c).font = HEAD
    for i, rec in enumerate(enrich.to_dict("records"), start=2):
        for j, c in enumerate(cols, start=1):
            ws.cell(row=i, column=j, value=rec.get(c, ""))
    ws["A1"].comment = Comment(
        "Paste the Clay export here (select A1, paste, header row included), or type rows by hand: row_id plus "
        "work_email and email_status (verified or unverified), and backup_work_email and backup_email_status. "
        "Columns are found by their header names, so order and extra columns do not matter. Only these four "
        "columns are read: nothing pasted here can change status, do_not_contact or any founder control.",
        "workstream A")
    for j in range(1, max(len(cols), 6) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22
    ws.freeze_panes = "A2"
    return ws


def text_sheet(wb, title, rows, widths, wrap_cols=()):
    ws = wb.create_sheet(title)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            if i == 1 or (isinstance(v, str) and v.isupper() and len(v) < 40):
                cell.font = HEAD
            if j in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


SCORING_ROWS = [
    ["table", "part", "points", "rule in plain words"],
    ["Offices", "HARD GATES", "", "Applied before any enrichment spend. A row fails if any one holds; a failed row scores 0."],
    ["Offices", "gate", "fail", "Outside the cluster: B's in_boundary = N against cluster-1-flatiron/boundary.geojson."],
    ["Offices", "gate", "fail", "Inside a coworking space (B's in_coworking = Y, or still a coworking tenant until a move B reports), or no staffed office at the address."],
    ["Offices", "gate", "fail", "Tuesday estimate wholly under 30 or wholly over 250. A range that crosses 30 or 250 passes with a flag."],
    ["Offices", "gate", "fail", "Coworking site whose operator is not shown to control the pantry, or a site possibly closed."],
    ["Offices", "SCORE", "0 to 100", "How likely this office fits the drop and can say yes, from what is known today."],
    ["Offices", "score_onsite", "0 to 30", "25 x the share of the Tuesday range inside 40 to 250 (ratio scale), plus 5 when the count is for the New York office; unknown headcount 7."],
    ["Offices", "score_size", "-4 to 15", "500 or fewer staff 15; 501 to 1,000 6; 1,001 to 4,999 2; 5,000+ 0; unknown 8; minus 4 when a controlling parent is named."],
    ["Offices", "score_attendance", "-5 to 15", "On-site 5 days 15; hybrid 4 days 12; 3 days 10; days not stated 7; 2 or fewer 5; no signal 6; remote -5; plus 3 when the policy names Tuesday (cap 15)."],
    ["Offices", "score_address", "0 to 15", "Occupied 10, occupancy-unconfirmed 5, signed-not-occupied 2; 0 with an address conflict; plus source confidence high 5, medium 3, low 0."],
    ["Offices", "score_title", "0 to 10", "The contact shown (primary, or the backup once promote_backup = Y) in the five target titles (or their direct forms), on an opened page, in New York: 10. Any other named person: 5. Nobody: 0."],
    ["Offices", "score_known", "0 or 15", "known_to_founder = Y (the channel also becomes warm-intro)."],
    ["Offices", "ties", "", "Contacts named on the employer's own page or in dated press before names only The Org lists, then the share of the range inside 40 to 250, source confidence, name."],
    ["Offices", "Tuesday estimate", "", "Sourced count x New York share (100% for a New York office count; 50% to 100% for a company-wide count with a New York HQ; 5% to 100% otherwise) x the in-office rate (67% for firms of 500 or fewer, 57% for 501 to 4,999, 46% for 5,000+; Partnership for New York City, March 2025, via plan.md). Low to high, central = square root of low x high."],
    ["Offices", "space check", "", f"The estimate never exceeds what the row's own leased square footage seats: high at most sf / {K.SF_PER_PERSON_DENSE}, central at most sf / {K.SF_PER_PERSON_TYPICAL}. Where the headcount math exceeds the space, the range comes from the space (sf / {K.SF_PER_PERSON_ROOMY} to sf / {K.SF_PER_PERSON_DENSE}). For a company headquartered outside New York with one square-footage figure, the space replaces the 5% to 100% guess, within the headcount math. onsite_basis says which rule applied. An estimate until intake replaces it."],
    ["Brands", "HARD GATES", "", "A row fails if any one holds; a failed row scores 0."],
    ["Brands", "gate", "fail", "Excluded category: alcohol, cannabis or hemp, supplements, anything needing preparation."],
    ["Brands", "gate", "fail", "An explicit health claim anywhere on the brand's own site: homepage, product pages and descriptions (Shopify catalog), About, FAQ and similar pages. Structure or function wording ('supports a calm mood', 'helps support gut health', 'fights cavities', 'immunity boosting'), disease wording, or the FDA supplement disclaimer. Nutrient-content wording ('10g protein') passes with founder review."],
    ["Brands", "gate", "fail", "Not shelf-stable: refrigerated, frozen, under 60 days, or shelf_stable = N."],
    ["Brands", "gate", "fail", "No buy path: no graded store (store_selector or brand_claim) within ten blocks of the cluster and no online store."],
    ["Brands", "SCORE", "0 to 100", "How likely this brand fits a founding campaign now: buyable near the offices, New York first, with a reason to sample now and a reachable owner."],
    ["Brands", "score_buy", "0 to 25", "Store page inside the cluster 25; store page within ten blocks 20; brand locator inside 17; brand locator within ten 13; online store only 6; none 0."],
    ["Brands", "score_ny", "0 to 20", "HQ New York City 20; rest of New York State 13; Northeast 8; elsewhere 3; unknown 4; minus 3 when a New York HQ is unverified."],
    ["Brands", "score_why", "0 to 15", "Press in the last 180 days 15, 181 to 365 days 8; own new listing in 180 days 8, 181 to 365 days 4; none, undated or unlinked 0."],
    ["Brands", "score_rev", "0 to 10", "$500k to $10M 10; unknown 5; under $500k 2; over $10M 0."],
    ["Brands", "score_fit", "0 to 10", "Shelf-stable Y 4, likely 2, unknown 1; individually packaged Y 3, unknown 1; online store Y 3, unknown 1."],
    ["Brands", "score_contact", "0 to 10", "Founder named on an opened page 6 (first name only or snippet 3); marketing lead on an opened page 4 (snippet 2)."],
    ["Brands", "score_known", "0 or 10", "known_to_founder = Y (the channel also becomes warm-intro)."],
    ["Brands", "score_review", "-35 to 0", "The most severe founder-review type: revenue over $10M sourced -35; owned by a large company -25; scale signals -12; dormant -12; heritage -10; perishable risk -8; under $500k -5; alcohol-adjacent -5; functional-ingredient wording -4; nutrient wording -4; pack format -4; other -2."],
    ["Brands", "founder review", "", "review_blocking = Y for owner, scale, heritage, dormant, perishable, under $500k, alcohol-adjacent, pack format, functional-ingredient or implied-effect wording, and other open questions: the row waits on the Founder review tab until you set founder_review_cleared = Y (or status = no). Nutrient wording, an unstated shelf life and a store-only buy path are notes, not blocks."],
    ["Brands", "ties", "", "Contacts named on the brand's own page or in dated press before The Org-only names, then store grade, why-now points, name."],
    ["Both", "DAILY VIEW", "", "The ten best-scored rows with status = new, do_not_contact = N, no opt-out, ready_for_outreach = Y, one row per company_group, skipping any group with a row in flight (queued to materials_sent) or any do-not-contact signal (the flag, the status or an opt-out date)."],
    ["Both", "ready_for_outreach", "Y/N", "Passed the gates, not on hold, a contact named on an opened page, and (offices) a Tuesday band of 40 to 250 or founder_ok_to_queue = Y; (brands) no founder review pending."],
]


ROUTINE_ROWS = [
    ["minute", "step", "where", "what to do"],
    ["0:00-3:00", "1. Log yesterday", "Offices and Brands tabs",
     "For every reply, bounce and opt-out since yesterday set status (the cell takes CONTRACTS words only) and E's columns. Any opt-out in any words: status = do_not_contact, do_not_contact = Y, opt_out_date = today, all three (the cells turn red if they disagree). Never reverse it. A wrong_person reply: promote_backup = Y and status = new; the backup becomes the contact in every view."],
    ["3:00-4:00", "2. Open both daily views", "Daily view - Offices, Daily view - Brands",
     "They have already updated. Take the top five of each (ten a day, the plan's pace). Rows 6 to 10 are spares."],
    ["4:00-11:00", "3. Forty seconds a row, ten rows", "check_before_send column",
     "Read score_reason and check_before_send. Open the LinkedIn URL (pasted on Friday) and confirm the person still holds the title. If a row still says 'find ... on LinkedIn', do it or skip the row to a spare. A brand's store line: use it only if retailer_checked_date is this week. If the person has left, set status = wrong_person and promote_backup = Y (it comes back tomorrow with the backup)."],
    ["11:00-12:00", "4. Export", "File > Download > CSV (each daily view tab)",
     "Google names the files '<workbook> - Daily view - Offices.csv' and '... - Brands.csv'. Rename them daily-view-offices.csv and daily-view-brands.csv and save them in clay-icp-tables/live/ for E's draft maker (python3 messages/_render/fill.py drafts --daily ../../clay-icp-tables/live/daily-view-offices.csv ...). Export before step 5: a queued row leaves the live view at once."],
    ["12:00-13:00", "5. Queue", "status column",
     "Set status = queued on the rows you will send today (about five per side). They leave the daily view and their company_group is blocked, so no second person from the same company is in flight. Rows you skipped stay new for tomorrow."],
    ["", "", "", ""],
    ["FRIDAY PREP", "", "", ""],
    ["Friday, 30 to 40 min", "1. LinkedIn batch", "Offices and Brands tabs, filter eligible = Y",
     "Work down the next 25 eligible rows per side (the tabs are already in score order). Where primary_linkedin_url is blank, find the person on LinkedIn and paste the profile URL; where it is filled from another page, open it and confirm the name and title. This is the research that used to happen at send time."],
    ["", "2. Store check", "Brands tab, retailer_check_url",
     "For the next 25 eligible brands with a store_selector line, click retailer_check_url (the store's own catalog answer, no store picker). If the product is still listed, type today's date in retailer_checked_date; if not, leave it blank and do not use the store line."],
    ["", "3. Founder review batch", "Founder review - Brands",
     "Read each founder_review_reason. Clear the row with founder_review_cleared = Y, or set status = no. Only cleared rows (or rows with no blocking review) reach the daily view."],
    ["", "4. Walks and holds", "Walk list - Offices; Offices tab",
     "Plan the week's lobby walk from the Walk list tab (offices with nobody named). Copy B's walk notes into hold_before_outreach (N once the lobby directory shows the company)."],
    ["", "5. Known, referrals, emails", "Offices and Brands tabs; Enrichment tab",
     "Mark known_to_founder for anyone you know (the channel becomes warm-intro). Add a row at the bottom for each referral: a new row_id (AO-R-### or AB-R-###), company, primary_first_name, primary_last_name, primary_title, status = new, referral_from = the referrer's row_id, and company_group = the referrer's company_group when it is the same company (the formulas are already in every row to 1000). Paste the week's Clay or Hunter email results on the Enrichment tab. Never paste anything over the Offices or Brands tabs."],
    ["", "6. Save the live tables", "File > Download > CSV",
     "Download the Offices, Brands and Enrichment tabs into clay-icp-tables/live/ as offices-live.csv, brands-live.csv and enrichment-live.csv, so a rebuild carries everything over (clay-build-sheet.md, section 4)."],
    ["", "", "", ""],
    ["COLOR KEY", "", "", ""],
    ["yellow header", "you edit", "", "known_to_founder, hold_before_outreach, founder_ok_to_queue, promote_backup, primary_linkedin_url, backup_linkedin_url, founder_review_cleared, retailer_checked_date"],
    ["blue header", "outreach edits", "", "status, do_not_contact, opt_out_date and E's tracking columns, written the same day"],
    ["grey header", "formula", "", "never type here: the contact shown, emails (from the Enrichment tab), channel, score, score_reason, readiness and the view helpers"],
    ["white header", "build", "", "researched facts (including primary_* and backup_* people); change them by fixing the source in _build/ and rebuilding"],
]


def main():
    ddict = {r["column"]: r for r in DD.dictionary_rows()}
    off, br = load("offices.csv"), load("brands.csv")
    DD.check(list(off.columns), "offices")
    DD.check(list(br.columns), "brands")
    enrich = pd.read_csv(os.path.join(C.WORK, "enrichment_build.csv"), dtype=str, keep_default_na=False) \
        if os.path.exists(os.path.join(C.WORK, "enrichment_build.csv")) else pd.DataFrame(columns=M.ENRICH_COLS)
    wb = Workbook()
    wb.remove(wb.active)
    _, Lo = table_sheet(wb, "Offices", off, "office", ddict)
    _, Lb = table_sheet(wb, "Brands", br, "brand", ddict)
    view_sheet(wb, "Daily view - Offices", "Offices", Lo, M.STABLE + M.DV_EXTRA["office"], "view_slot",
               K.DAILY_VIEW_SIZE, "No row is ready today: see the Routine tab",
               "Live view: the ten best-scored Offices rows with status = new, do_not_contact = N, no opt-out, "
               "ready_for_outreach = Y and nobody in their company_group in flight or opted out; one row per group. "
               "It updates when you edit those columns.")
    view_sheet(wb, "Daily view - Brands", "Brands", Lb, M.STABLE + M.DV_EXTRA["brand"], "view_slot",
               K.DAILY_VIEW_SIZE, "No row is ready today: see the Routine tab",
               "Live view: the ten best-scored Brands rows with status = new, do_not_contact = N, no opt-out, "
               "ready_for_outreach = Y (founder review cleared) and nobody in their company_group in flight or "
               "opted out; one row per group.")
    view_sheet(wb, "Walk list - Offices", "Offices", Lo, M.WALK_VIEW, "walk_slot", M.WALK_SIZE,
               "No office is waiting for a walk-in",
               "Offices in the top 50 with nobody named (channel = walk-in), passed, status new, no hold, no "
               "do-not-contact: the lobby walk list, best score first.")
    view_sheet(wb, "Founder review - Brands", "Brands", Lb, M.REVIEW_VIEW, "review_slot", M.REVIEW_SIZE,
               "No brand is waiting for your review",
               "Brands that passed the gates but carry a blocking founder review: read founder_review_reason, then "
               "set founder_review_cleared = Y on the Brands tab (or status = no). Best score first.")
    enrichment_sheet(wb, enrich)
    rows = [["column", "tables", "group", "edited_by", "values", "meaning"]] + [
        [r["column"], r["tables"], r["group"], r["edited_by"], r["values"], r["meaning"]] for r in DD.dictionary_rows()]
    text_sheet(wb, "Dictionary", rows, [28, 9, 22, 12, 40, 100], wrap_cols=(5, 6))
    text_sheet(wb, "Scoring", SCORING_ROWS, [9, 20, 10, 140], wrap_cols=(4,))
    rt = text_sheet(wb, "Routine", [["company_name", BRAND_NAME, "", "The working name lives in this one cell."],
                                     ["TAGLINE", TAGLINE, "", ""],
                                     ["", "", "", ""]] + [[None] * 4] + ROUTINE_ROWS, [18, 30, 34, 120], wrap_cols=(4,))
    rt["A4"] = '="15-minute morning routine, "&B1&" outreach tables"'
    rt["A4"].font = HEAD
    for cell in rt[5]:
        cell.font = HEAD
    rt.freeze_panes = None
    wb.defined_names["company_name"] = DefinedName("company_name", attr_text="Routine!$B$1")
    wb.save(OUT)
    pd.DataFrame(DD.dictionary_rows()).to_csv(os.path.join(C.A_DIR, "column-dictionary.csv"), index=False)
    print("wrote", OUT, "and column-dictionary.csv")


if __name__ == "__main__":
    main()
