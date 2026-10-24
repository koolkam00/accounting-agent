#!/usr/bin/env python3
"""Behaviour tests: edit a copy of the workbook, recalculate with LibreOffice, assert.

Every test writes to a temp copy; the deliverable is never modified.
SAMPLE rows used here start with 'SAMPLE' (CONTRACTS s2) and never ship.
Cells are found through ref_map.json (written by build_model.py), so a layout change does not break the tests.
Tests tagged [R1-xx] cover the fixes for the independent review, round 1 (_review/round-1.md); [R2-xx] cover round 2 (_review/round-2.md);
[R3-xx] cover round 3 (_review/round-3.md).
Usage: python3 test_model.py   (exit 1 on any failure)
"""
import datetime as dt
import json
import os
import re
import sys
import tempfile

import openpyxl

import recalc

HERE = os.path.dirname(os.path.abspath(__file__))
REF = json.load(open(os.path.join(HERE, "ref_map.json")))
SRC = recalc.DEFAULT
D = dt.date
RESULTS = []


def cell(ref):
    sh, c = ref.split("!")
    return sh, c.replace("$", "")


def row_of(key):
    return int(re.sub(r"[^0-9]", "", cell(REF[key])[1]))


def run(edits, label):
    """edits: list of (sheet, cell, value). Returns recalculated workbook, integrity/schedule/data failures, error cells."""
    tmp = tempfile.mkdtemp(prefix="ddtest_")
    path = os.path.join(tmp, "t.xlsx")
    wb = openpyxl.load_workbook(SRC)
    for sh, c, v in edits:
        wb[sh][c].value = v
    wb.save(path)
    out, _ = recalc.recalc(path)
    rows, fails, errs = recalc.report(out, quiet=True)
    integ = [f for f in fails if f[1] in ("Integrity", "Schedule", "Data")]
    return out, integ, errs


def v(wb, key):
    sh, c = cell(REF[key])
    return wb[sh][c].value


def at(wb, sheet, c):
    return wb[sheet][c].value


def check(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + label + (f"  [{detail}]" if detail else ""))


def inp(key, value):
    sh, c = cell(REF[key])
    return (sh, c, value)


def near(a, b, tol=0.01):
    return isinstance(a, (int, float)) and abs(a - b) <= tol


def gate_tables(wb):
    sh, rng = REF["dash_gate_rows"].split("!")
    a, b = rng.replace("$", "").split(":")
    r0, r1 = int(a[1:]), int(b[1:])
    ws = wb[sh]
    status = {ws[f"B{r}"].value: ws[f"G{r}"].value for r in range(r0, r1 + 1) if ws[f"B{r}"].value}
    sh, rng = REF["dash_gate_overall"].split("!")
    a, b = rng.replace("$", "").split(":")
    gates = {ws[f"A{r}"].value: ws[f"C{r}"].value for r in range(int(a[1:]), int(b[1:]) + 1)}
    return status, gates


# ---------------------------------------------------------------- baseline
base, integ, errs = run([], "baseline")
check("baseline: no integrity failures and no error cells", not integ and not errs, f"{integ[:2]} {errs[:2]}")
b_low, b_end, b_ss = v(base, "cf_low"), v(base, "cf_end"), v(base, "ss_ph")
b_next_ph = v(base, "dash_next_ph")
print(f"   baseline low {b_low}, end {b_end}, steady per hour {b_ss:.2f}, next drop per hour {b_next_ph:.2f}")

# [R1-02] publication default is the plan's high end
pub_r = row_of("su_pub")
check("[R1-02] publication default = plan high end $2,000, low $250, high $2,000",
      at(base, "Inputs", f"C{pub_r}") == 2000 and at(base, "Inputs", f"E{pub_r}") == 250 and at(base, "Inputs", f"F{pub_r}") == 2000,
      (at(base, "Inputs", f"C{pub_r}"), at(base, "Inputs", f"E{pub_r}"), at(base, "Inputs", f"F{pub_r}")))
check("[R1-02][R3-02] low point -$4,476 in week 13 (trademark at Gate 2), capital to commit $5,476, gap to $6,510 $1,034",
      near(b_low, -4476) and v(base, "cf_low_week") == 13 and near(v(base, "cf_commit"), 5476) and near(v(base, "cf_gap"), 1034),
      (b_low, v(base, "cf_low_week"), v(base, "cf_commit"), v(base, "cf_gap")))
check("[R1-02] one-brand campaigns to recover startup cash = 5", v(base, "be_n_one") == 5, v(base, "be_n_one"))
# [R1-18c] a drop with no actual data shows no allocation and a blank contribution
check("[R1-18c] Drop_Actuals: drop without actuals shows a blank contribution and $0 allocations",
      at(base, "Drop_Actuals", "S8") in (None, "") and at(base, "Drop_Actuals", "I8") == 0, (at(base, "Drop_Actuals", "S8"), at(base, "Drop_Actuals", "I8")))
# [R1-11] insurance prepaid beyond the window
check("[R1-11] insurance prepaid beyond the six months = $700 (annual premium, five cover months inside)", near(v(base, "cf_prepaid_ins"), 700), v(base, "cf_prepaid_ins"))
check("[R1-11] all-in rate uses expensed overhead: $2,010 paid less $700 prepaid", near(v(base, "cf_over_exp"), 1310), v(base, "cf_over_exp"))
# [R1-07] two views of the six months
check("[R1-07] week-26 cash, roadmap only, is below the view with decision-window drops", v(base, "dash_cash_road") < v(base, "dash_cash_all"),
      (v(base, "dash_cash_road"), v(base, "dash_cash_all")))
check("[R1-07] all-in rate, roadmap only, is below the view with decision-window drops", v(base, "dash_allin_road") < v(base, "dash_allin_all"),
      (round(v(base, "dash_allin_road"), 2), round(v(base, "dash_allin_all"), 2)))
check("[R1-07] Schedule needs 12 new brand wins (60 qualified conversations) by DRP-005's sign-by date",
      at(base, "Schedule", "AB11") == 12 and near(at(base, "Schedule", "AC11"), 60), (at(base, "Schedule", "AB11"), at(base, "Schedule", "AC11")))
# [R1-04b] DRP-005 reports after the decision
check("[R1-04] DRP-005 debrief (T+10) is 2027-04-20 and flagged after the decision; Dashboard counts 1",
      str(at(base, "Schedule", "AF11"))[:10] == "2027-04-20" and at(base, "Schedule", "AG11") == "after the decision" and v(base, "dash_after_dec") == 1,
      (at(base, "Schedule", "AF11"), at(base, "Schedule", "AG11"), v(base, "dash_after_dec")))
# [R1-16] Gate 3 rate with a fatigue-free roster
check("[R1-16] Gate 3 rate with a roster of one drop per office a month = $51.4 (8 offices added, 82 h on run 2)",
      near(v(base, "dash_g3_roster"), 7660 / 149, 0.01) and v(base, "dash_g3_roster") < v(base, "dash_g3_rate"),
      (v(base, "dash_g3_roster"), v(base, "dash_g3_rate")))
# [R1-14] timing defaults follow their gates
ins_r = row_of("insurance_bind")
check("[R1-14] Gate 2 follows campaign zero (2027-01-22) and insurance binds at Gate 1 (2026-12-11)",
      str(v(base, "gate2_date"))[:10] == "2027-01-22" and str(v(base, "insurance_bind"))[:10] == "2026-12-11", (v(base, "gate2_date"), v(base, "insurance_bind")))
cert_r = row_of("su_cert")
check("[R1-14][R2-07] publication timing 2026-12-04 and certificate 2027-01-15 (formulas of the filing date, Thu 2026-11-19)",
      str(at(base, "Inputs", f"L{pub_r}"))[:10] == "2026-12-04" and str(at(base, "Inputs", f"L{cert_r}"))[:10] == "2027-01-15",
      (at(base, "Inputs", f"L{pub_r}"), at(base, "Inputs", f"L{cert_r}")))
# [R1-12] one cell holds the company name (formula view of the workbook)
wf = openpyxl.load_workbook(SRC)
name_cells = [f"{ws.title}!{c.coordinate}" for ws in wf.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and "Desk Drop" in c.value]
check("[R1-12] the company name sits in one cell only", name_cells == [REF["company_name"].replace("$", "")], name_cells)
# [R1-13] planning-estimate label on every tab built on defaults
missing = []
for t in ["Dashboard", "Inputs", "Scenarios", "Schedule", "Calculator", "Drop_PnL", "Cash_Flow", "Steady_State", "Sensitivity", "Plan_Check", "Capital", "Cash_Ledger",
          "Drop_Actuals"]:
    txt = " ".join(str(wf[t].cell(r, c).value) for r in range(1, 5) for c in range(1, 4) if wf[t].cell(r, c).value)
    if "planning estimate" not in txt.lower():
        missing.append(t)
check("[R1-13][R2-09d][R3-11a] 'Planning estimates' label in rows 1 to 4 of every tab built on defaults, Cash_Ledger and Drop_Actuals included", not missing, missing)
# [R1-18a] Plan_Check colours the expected difference
cf_rules = [r.formula for rng in wf["Plan_Check"].conditional_formatting for r in rng.rules]
check("[R1-18a] Plan_Check has a colour rule for 'differs, as expected'", any('"differs, as expected"' in str(f) for f in cf_rules))
# [R1-15] founding renewals explained
readme_txt = " ".join(str(c.value) for row in wf["README"].iter_rows() for c in row if c.value)
check("[R1-15] README explains why founding renewals default to 0 and what each one changes", "No founding-rate renewal is scheduled by default" in readme_txt)
check("[R1-05] README names the Dashboard rows that move with the price", "Price in force" in readme_txt and "A standard drop at that price" in readme_txt)

# ---------------------------------------------------------------- price
wb, integ, errs = run([inp("price_std_office", 180)], "price 180")
check("price $180: standard runs per hour rises", v(wb, "dash_g3_rate") > v(base, "dash_g3_rate"),
      f"{v(base, 'dash_g3_rate'):.2f} -> {v(wb, 'dash_g3_rate'):.2f}")
check("price $180: cash at week 26 rises", v(wb, "cf_end") > b_end, f"{b_end} -> {v(wb, 'cf_end')}")
check("price $180: inputs-changed count = 1", v(wb, "changed_count") == 1, f"changed {v(wb, 'changed_count')}")
check("price $180: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")
check("[R1-05] price $180: Dashboard shows the price in force", v(wb, "dash_price") == 180, v(wb, "dash_price"))
check("[R1-05] price $180: 'A standard drop at that price' moves from $73.7 to $90.4 an hour",
      near(v(base, "dash_std_ph"), 3830 / 52, 0.01) and near(v(wb, "dash_std_ph"), 90.4, 0.05), (v(base, "dash_std_ph"), v(wb, "dash_std_ph")))
check("[R1-03] price $180: break-even per brand on a standard drop moves (it no longer reads the fixed plan row)",
      v(wb, "be_c_three") > v(base, "be_c_three") + 1, (round(v(base, "be_c_three"), 2), round(v(wb, "be_c_three"), 2)))

# ---------------------------------------------------------------- processing off
wb, integ, errs = run([inp("processing_on", 0)], "processing off")
check("bank transfer: campaign zero contribution = $960", abs(at(wb, "Drop_PnL", "AJ7") - 960) < 0.01, at(wb, "Drop_PnL", "AJ7"))
check("bank transfer: low point improves", v(wb, "cf_low") > b_low, f"{b_low} -> {v(wb, 'cf_low')}")
check("bank transfer: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")

# ---------------------------------------------------------------- home storage
wb, integ, errs = run([inp("storage_choice", "home")], "home storage")
check("home storage: standard run hard costs = $595", abs(at(wb, "Drop_PnL", "AA8") - 595) < 0.01, at(wb, "Drop_PnL", "AA8"))
check("home storage: no storage paid in the cash flow", at(wb, "Cash_Flow", "AC22") == 0, at(wb, "Cash_Flow", "AC22"))
check("home storage: steady-state overhead = $160", abs(v(wb, "ss_over_paid") - 160) < 0.01, v(wb, "ss_over_paid"))
wb, integ, errs = run([inp("processing_on", 0), inp("storage_choice", "home")], "processing off and home storage")
check("[R1-03] bank transfer plus home storage: break-even per brand on a standard drop moves", v(wb, "be_c_three") > v(base, "be_c_three") + 1,
      (round(v(base, "be_c_three"), 2), round(v(wb, "be_c_three"), 2)))
check("[R1-03] bank transfer plus home storage: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")

# ---------------------------------------------------------------- insurance monthly
wb, integ, errs = run([inp("insurance_billing", "monthly")], "insurance monthly")
check("insurance monthly: five payments of $100 inside the window", abs(at(wb, "Cash_Flow", "AC21") + 500) < 0.01, at(wb, "Cash_Flow", "AC21"))
check("insurance monthly: low point improves", v(wb, "cf_low") > b_low, f"{b_low} -> {v(wb, 'cf_low')}")
check("insurance monthly: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")
check("[R1-11] insurance monthly: nothing prepaid, all-in rate unchanged", near(v(wb, "cf_prepaid_ins"), 0) and near(v(wb, "dash_allin_all"), v(base, "dash_allin_all")),
      (v(wb, "cf_prepaid_ins"), v(wb, "dash_allin_all"), v(base, "dash_allin_all")))

# ---------------------------------------------------------------- close rate
wb, integ, errs = run([inp("close_rate", 0.1)], "close rate 10%")
check("close rate 10%: hours per win = 9.75", abs(v(wb, "h_win") - 9.75) < 0.001, v(wb, "h_win"))
check("close rate 10%: standard runs per hour falls", v(wb, "dash_g3_rate") < v(base, "dash_g3_rate"))
check("[R1-07] close rate 10%: DRP-005 cumulative conversations implied = 120", near(at(wb, "Schedule", "AC11"), 120), at(wb, "Schedule", "AC11"))

# ---------------------------------------------------------------- drop size drivers
wb, integ, errs = run([inp("std_brands_drop", 2)], "two brands per standard drop")
check("two brands per standard drop: runs 1 and 2 carry 2 brands and per hour falls",
      at(wb, "Schedule", "K8") == 2 and at(wb, "Schedule", "K9") == 2 and v(wb, "dash_g3_rate") < v(base, "dash_g3_rate"),
      (at(wb, "Schedule", "K8"), v(wb, "dash_g3_rate")))
check("two brands per standard drop: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")
check("[R1-03] two brands per standard drop: break-even campaigns move with brands per drop", v(wb, "be_c_three") != v(base, "be_c_three"),
      (v(base, "be_c_three"), v(wb, "be_c_three")))
wb, integ, errs = run([inp("goal_offices", 20)], "goal of 20 offices")
check("goal of 20 offices: run 2 recruits 10 new offices and its hours rise", at(wb, "Schedule", "J9") == 10 and at(wb, "Drop_PnL", "AI9") > at(base, "Drop_PnL", "AI9"),
      (at(wb, "Schedule", "J9"), at(wb, "Drop_PnL", "AI9")))
check("goal of 20 offices: the run-2 fatigue warning clears", at(wb, "Schedule", "W9") == "ok", at(wb, "Schedule", "W9"))
check("[R1-16] goal of 20 offices: no offices to add, so the two Gate 3 rates agree", near(v(wb, "dash_g3_roster"), v(wb, "dash_g3_rate")),
      (v(wb, "dash_g3_roster"), v(wb, "dash_g3_rate")))

# ---------------------------------------------------------------- scenarios
for sc in ["A", "B", "C", "D"]:
    wb, integ, errs = run([inp("active_scenario", sc)], f"scenario {sc}")
    k4, k5 = at(wb, "Schedule", "K10"), at(wb, "Schedule", "K11")
    l4, l5 = at(wb, "Schedule", "L10"), at(wb, "Schedule", "L11")
    check(f"scenario {sc}: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")
    print(f"   scenario {sc}: D4 brands {k4} renewals {l4}; D5 brands {k5} renewals {l5}; end cash {v(wb, 'cf_end')}; "
          f"steady contrib {v(wb, 'ss_contrib')}; steady net {v(wb, 'ss_net')}; steady per hour {v(wb, 'ss_ph')}")
    if sc == "A":
        check("A: steady state at $180 per office earns more than base", v(wb, "ss_contrib") > v(base, "ss_contrib"))
    if sc == "B":
        check("B: no renewals scheduled and office-paid revenue $500 shown", l4 == 0 and l5 == 0 and near(v(wb, "ss_op_rev"), 500), v(wb, "ss_op_rev"))
        check("[R1-10] B: unpriced office-paid revenue is left out of the net (net = contribution less unabsorbed overhead)",
              near(v(wb, "ss_net"), v(wb, "ss_contrib") - v(wb, "ss_unabs")) and v(wb, "ss_op_in") == "no: cost to serve needs quote",
              (v(wb, "ss_net"), v(wb, "ss_op_in")))
        check("[R1-10] B: the Dashboard says office-paid revenue is left out", "left out of the net" in str(at(wb, "Dashboard", "D" + str(row_of("dash_ss")))))
    if sc == "C":
        check("C: only one decision-window drop runs", k4 == 3 and k5 == 0)
    if sc == "D":
        check("D: no decision-window drops and zero steady state", k4 == 0 and k5 == 0 and v(wb, "ss_contrib") == 0)
        check("D: cash at week 26 below base", v(wb, "cf_end") < b_end)
        check("[R1-07] D: with no decision-window drops the two six-month views agree (all-in and week-26 cash)",
              near(v(wb, "dash_allin_road"), v(wb, "dash_allin_all")) and near(v(wb, "dash_cash_road"), v(wb, "dash_cash_all"))
              and near(v(wb, "dash_cash_all"), v(base, "dash_cash_road")),
              (v(wb, "dash_allin_road"), v(wb, "dash_allin_all"), v(wb, "dash_cash_road"), v(base, "dash_cash_road")))
sc_cost_r = row_of("sc_op_cost")
wb, integ, errs = run([inp("active_scenario", "B"), ("Scenarios", f"E{sc_cost_r}", 50)], "scenario B with a quoted cost")
check("[R1-10] B with a $50 cost to serve: the office tier counts in the net (+$500 - $250)",
      v(wb, "ss_op_in") == "yes" and near(v(wb, "ss_net"), v(wb, "ss_contrib") - v(wb, "ss_unabs") + 250), (v(wb, "ss_op_in"), v(wb, "ss_net")))

# ---------------------------------------------------------------- edge cases on the schedule
wb, integ, errs = run([("Schedule", "K8", 0)], "run 1 with zero brands")
check("zero brands on run 1: the drop does not run, cost 0", at(wb, "Drop_PnL", "AA8") == 0 and at(wb, "Drop_PnL", "AK8") == "n/a")
check("zero brands on run 1: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")
check("[R1-09] zero brands on run 1: it recruits no offices and run 2 fills the roster instead", at(wb, "Schedule", "J8") == 0 and at(wb, "Schedule", "J9") == 7,
      (at(wb, "Schedule", "J8"), at(wb, "Schedule", "J9")))
wb, integ, errs = run([inp("c0_brands", 0), inp("std_brands_drop", 0)], "zero brands everywhere")
check("[R1-09] zero brands everywhere: no drop recruits an office and 'active offices planned' is 0",
      [at(wb, "Schedule", f"J{r}") for r in range(7, 12)] == [0, 0, 0, 0, 0] and v(wb, "dash_goal_goal_offices") == 0,
      ([at(wb, "Schedule", f"J{r}") for r in range(7, 12)], v(wb, "dash_goal_goal_offices")))
check("[R1-09] zero brands everywhere: no error cells", not errs, errs[:3])
wb, integ, errs = run([("Schedule", "I8", 1), ("Schedule", "J8", 1), ("Schedule", "K8", 1)], "one office")
check("one office on run 1: finite per hour", isinstance(at(wb, "Drop_PnL", "AK8"), (int, float)), at(wb, "Drop_PnL", "AK8"))
wb, integ, errs = run([("Schedule", "L10", 3)], "100% renewal on D4")
check("three renewals on D4: brand sales 6 h", abs(at(wb, "Drop_PnL", "AB10") - 6) < 0.001, at(wb, "Drop_PnL", "AB10"))
wb, integ, errs = run([("Schedule", "E8", D(2027, 2, 12))], "Friday drop")
check("a Friday drop is flagged and fails the schedule check", at(wb, "Schedule", "P8") == "not Tue/Wed" and any("schedule rules" in f[2] for f in integ))
wb, integ, errs = run([("Schedule", "E8", D(2026, 12, 22))], "blackout drop")
check("a drop in the December blackout is flagged", at(wb, "Schedule", "R8") == "clash")
wb, integ, errs = run([("Schedule", "A12", "DRP-006"), ("Schedule", "B12", "SAMPLE extra drop"), ("Schedule", "E12", D(2027, 4, 20)),
                       ("Schedule", "G12", "direct-ship"), ("Schedule", "H12", "standard"), ("Schedule", "I12", 10), ("Schedule", "J12", 8),
                       ("Schedule", "K12", 3), ("Schedule", "L12", 0), ("Schedule", "M12", 0)], "added drop")
check("an added drop flows into the cash flow (deposit inside the window)", v(wb, "cf_end") > b_end, f"{b_end} -> {v(wb, 'cf_end')}")
check("an added drop shows as a changed schedule row", v(wb, "dash_sched_changed") >= 1, v(wb, "dash_sched_changed"))
check("added drop: no integrity failures (per-drop cash still adds up to the cash flow)", not integ and not errs, f"{integ[:2]} {errs[:2]}")

# ---------------------------------------------------------------- [R1-14] dates follow their gates
lic_r = row_of("su_license")
wb, integ, errs = run([inp("gate1_date", D(2026, 12, 18))], "Gate 1 slips a week")
check("[R1-14] Gate 1 slips to 2026-12-18: insurance bind and licence follow, and only one input shows as changed",
      str(v(wb, "insurance_bind"))[:10] == "2026-12-18" and str(at(wb, "Inputs", f"L{lic_r}"))[:10] == "2026-12-18" and v(wb, "changed_count") == 1,
      (v(wb, "insurance_bind"), at(wb, "Inputs", f"L{lic_r}"), v(wb, "changed_count")))
wb, integ, errs = run([inp("gate0_date", D(2026, 11, 20))], "Gate 0 slips a week")
llc_r = row_of("su_llc")
check("[R1-14][R2-07] Gate 0 slips to 2026-11-20: filing moves to Fri 2026-11-27 (4 business days, skipping Thanksgiving) and publication to Fri 2026-12-11",
      str(at(wb, "Inputs", f"L{llc_r}"))[:10] == "2026-11-27" and str(at(wb, "Inputs", f"L{pub_r}"))[:10] == "2026-12-11",
      (at(wb, "Inputs", f"L{llc_r}"), at(wb, "Inputs", f"L{pub_r}")))
wb, integ, errs = run([("Schedule", "E7", D(2027, 1, 13))], "campaign zero moves a day")
check("[R1-14] campaign zero moves to Wed 2027-01-13: Gate 2 follows to 2027-01-25", str(v(wb, "gate2_date"))[:10] == "2027-01-25", v(wb, "gate2_date"))
check("[R1-14] campaign zero moves: no integrity failures", not integ and not errs, f"{integ[:2]} {errs[:2]}")

# ---------------------------------------------------------------- [R1-06] paste layouts are checked
wb, integ, errs = run([("Actuals_Campaigns", "AA4", "extra_column"), ("Actuals_Campaigns", "AA5", 1)], "extra column on campaign actuals")
labels = [f[2] for f in integ]
check("[R1-06] an extra column pasted after 'survey_responses_min_office' is flagged, and no helper is overwritten",
      any("Actuals_Campaigns: header row differs" in l for l in labels) and at(wb, "Actuals_Campaigns", "AL4") == "drop actuals in?", labels)
wb, integ, errs = run([("Pipeline", "E4", "event_type"), ("Hours_Log", "G5", 2), ("Actuals_DropCosts", "J4", "vendor")], "renamed and extra columns")
labels = [f[2] for f in integ]
check("[R1-06] a renamed Pipeline header, a value beside Hours_Log and an extra drop-cost column are each flagged",
      all(any(k in l for l in labels) for k in ("Pipeline: header row differs", "Hours_Log: header row differs", "Actuals_DropCosts: header row differs")), labels)

# ---------------------------------------------------------------- [R1-08] renewal answers
pl8 = [(D(2027, 3, 1), "brand", "SAMPLE-B-5", "SAMPLE brand 5", "renewal_answer", 1500, "yes"),
       (D(2027, 3, 2), "brand", "SAMPLE-B-5", "SAMPLE brand 5", "renewal_answer", 1500, "yes"),
       (D(2027, 3, 3), "brand", "SAMPLE-B-6", "SAMPLE brand 6", "renewal_answer", 1500, "maybe")]
ed8 = []
for i, row in enumerate(pl8):
    for j, val in enumerate(row):
        ed8.append(("Pipeline", f"{'ABCDEFG'[j]}{5 + i}", val))
wb, integ, errs = run(ed8, "renewal answers")
check("[R1-08] the same brand's 'yes' logged twice counts once, and a 'maybe' with $1,500 does not count", v(wb, "pl_g3_renewals") == 1, v(wb, "pl_g3_renewals"))

# ---------------------------------------------------------------- SAMPLE actuals and pipeline, as of 2027-01-29
ed = [inp("as_of_override", D(2027, 1, 29)), ("Schedule", "N7", "Y")]
pl = [  # date, side, org_id, org_name, event, value_number, value_text, controls_budget
    (D(2026, 10, 28), "office", "SAMPLE-O-1", "SAMPLE office 1", "conversation", None, None, None),
    (D(2026, 10, 29), "office", "SAMPLE-O-2", "SAMPLE office 2", "conversation", None, None, None),
    (D(2026, 10, 30), "office", "SAMPLE-O-3", "SAMPLE office 3", "conversation", None, None, None),
    (D(2026, 11, 2), "office", "SAMPLE-O-4", "SAMPLE office 4", "conversation", None, None, None),
    (D(2026, 11, 3), "office", "SAMPLE-O-5", "SAMPLE office 5", "conversation", None, None, None),
    (D(2026, 11, 4), "office", "SAMPLE-O-1", "SAMPLE office 1", "yes_in_principle", None, None, None),
    (D(2026, 11, 5), "office", "SAMPLE-O-2", "SAMPLE office 2", "yes_in_principle", None, None, None),
    (D(2026, 11, 6), "office", "SAMPLE-O-3", "SAMPLE office 3", "yes_in_principle", None, None, None),
    (D(2026, 11, 9), "office", "SAMPLE-O-4", "SAMPLE office 4", "yes_in_principle", None, None, None),
    (D(2026, 11, 10), "office", "SAMPLE-O-5", "SAMPLE office 5", "yes_in_principle", None, None, None),
    (D(2026, 11, 10), "office", "SAMPLE-O-5", "SAMPLE office 5", "yes_in_principle", None, "SAMPLE duplicate row", None),
    (D(2026, 11, 3), "brand", "SAMPLE-B-1", "SAMPLE brand 1", "budget_probe", 800, "SAMPLE quote", "Y"),
    (D(2026, 11, 4), "brand", "SAMPLE-B-2", "SAMPLE brand 2", "budget_probe", 600, "SAMPLE quote", "Y"),
    (D(2026, 11, 5), "brand", "SAMPLE-B-3", "SAMPLE brand 3", "budget_probe", 1000, "SAMPLE quote", "Y"),
    (D(2026, 11, 6), "brand", "SAMPLE-B-4", "SAMPLE brand 4", "budget_probe", 1500, "SAMPLE quote", "N"),
    (D(2026, 11, 20), "office", "SAMPLE-O-1", "SAMPLE office 1", "signed", 60, None, None),
    (D(2026, 11, 23), "office", "SAMPLE-O-2", "SAMPLE office 2", "signed", 55, None, None),
    (D(2026, 11, 24), "office", "SAMPLE-O-3", "SAMPLE office 3", "signed", 110, None, None),
    (D(2026, 11, 30), "office", "SAMPLE-O-4", "SAMPLE office 4", "signed", 95, None, None),
    (D(2026, 12, 2), "office", "SAMPLE-O-5", "SAMPLE office 5", "signed", 190, None, None),
    (D(2026, 12, 9), "brand", "SAMPLE-B-1", "SAMPLE brand 1", "paid", 600, None, "Y"),
    (D(2026, 12, 10), "brand", "SAMPLE-B-2", "SAMPLE brand 2", "paid", 600, None, "Y"),
    (D(2026, 11, 3), "office", "SAMPLE-O-1", "SAMPLE office 1", "budget_probe", 150, "SAMPLE quote", None),
]
for i, row in enumerate(pl):
    for j, val in enumerate(row):
        if val is not None:
            ed.append(("Pipeline", f"{'ABCDEFGH'[j]}{5 + i}", val))
camp = [  # campaign rows for DRP-001
    ["SAMPLE-CMP-001", "DRP-001", D(2027, 1, 12), "founder", "founding", "SAMPLE-BR-001", 5, 275, 250, 245, 240, 5, 160, 31, 24, 40, 3, 600, D(2026, 12, 9), 0, D(2027, 1, 21), "yes", "pending", None, None, 5],
    ["SAMPLE-CMP-002", "DRP-001", D(2027, 1, 12), "founder", "founding", "SAMPLE-BR-002", 5, 275, 250, 238, 210, 0, 140, 27, 20, 33, 1, 600, D(2026, 12, 10), 0, D(2027, 1, 22), "maybe", "pending", None, None, 4],
]
for i, row in enumerate(camp):
    for j, val in enumerate(row):
        if val is not None:
            ed.append(("Actuals_Campaigns", f"{openpyxl.utils.get_column_letter(j + 1)}{5 + i}", val))
for j, val in enumerate(["DRP-001", D(2027, 1, 12), "founder", 5, 2, 45, 95, 20, "SAMPLE"]):
    ed.append(("Actuals_DropCosts", f"{openpyxl.utils.get_column_letter(j + 1)}5", val))
hl = [(D(2026, 11, 2), None, "brand_sales", 6), (D(2026, 11, 25), "DRP-001", "brand_sales", 14), (D(2026, 12, 1), "DRP-001", "office_coordination", 18),
      (D(2027, 1, 7), "DRP-001", "receiving_kitting", 3.5), (D(2027, 1, 12), "DRP-001", "drop_day", 5.5), (D(2027, 1, 19), "DRP-001", "checks_pickup", 3),
      (D(2027, 1, 22), "DRP-001", "reporting_debriefs", 4), (D(2027, 1, 22), "DRP-001", "admin_invoicing", 2)]
for i, (d_, drop, act, h) in enumerate(hl):
    ed += [("Hours_Log", f"A{5 + i}", d_), ("Hours_Log", f"D{5 + i}", act), ("Hours_Log", f"E{5 + i}", h), ("Hours_Log", f"F{5 + i}", "SAMPLE")]
    if drop:
        ed.append(("Hours_Log", f"B{5 + i}", drop))
# Toggl: header row 9 and two entries
for j, h in enumerate(["User", "Project", "Description", "Start date", "Duration", "Tags"]):
    ed.append(("Toggl_Import", f"{'ABCDEF'[j]}9", h))
ed += [("Toggl_Import", "A10", "SAMPLE"), ("Toggl_Import", "B10", "DRP-001"), ("Toggl_Import", "D10", "2027-01-12"), ("Toggl_Import", "E10", "01:30:00"), ("Toggl_Import", "F10", "drop_day"),
       ("Toggl_Import", "A11", "SAMPLE"), ("Toggl_Import", "B11", "admin"), ("Toggl_Import", "D11", "2027-01-13"), ("Toggl_Import", "E11", "00:45:00"), ("Toggl_Import", "F11", "admin_invoicing")]
# [R1-18] an actual bank balance in week 13
ed.append(("Cash_Flow", "O30", -3000))
wb, integ, errs = run(ed, "SAMPLE actuals")
check("SAMPLE: no integrity or data failures", not integ and not errs, f"{integ[:3]} {errs[:3]}")
check("[R3-01] SAMPLE run with real dates and numbers: the banner reads 'all integrity and data checks pass'", at(wb, "Dashboard", "G5") == "all integrity and data checks pass",
      at(wb, "Dashboard", "G5"))
status, gates = gate_tables(wb)
print("   gate rows:", json.dumps(status, default=str)[:900])
print("   gate overall:", gates)
check("SAMPLE: Gate 0 passes (5 offices yes, 3 brands at $500+ controlling budget; duplicate row and non-controller ignored)",
      gates.get("Gate 0") == "pass", gates.get("Gate 0"))
check("SAMPLE: Gate 1 passes (5 signed, 510 headcount, 2 paid)", gates.get("Gate 1") == "pass", gates.get("Gate 1"))
check("SAMPLE: Gate 2 misses (one campaign under 5 responses per office)", gates.get("Gate 2") == "miss", gates.get("Gate 2"))
check("SAMPLE: Gate 3 not yet", gates.get("Gate 3") == "not yet", gates.get("Gate 3"))
check("SAMPLE: Drop_Actuals revenue $1,200 and hard costs $160 + $80 allocations",
      at(wb, "Drop_Actuals", "D7") == 1200 and abs(at(wb, "Drop_Actuals", "J7") - 240) < 0.01, (at(wb, "Drop_Actuals", "D7"), at(wb, "Drop_Actuals", "J7")))
check("SAMPLE: actual hours for DRP-001 = 50 hand + 1.5 Toggl", abs(at(wb, "Drop_Actuals", "R7") - 51.5) < 0.001, at(wb, "Drop_Actuals", "R7"))
check("SAMPLE: placement within a day = 450/500", abs(at(wb, "Drop_Actuals", "AG7") - 0.9) < 0.0001, at(wb, "Drop_Actuals", "AG7"))
check("SAMPLE: live basis switches to actual", at(wb, "Drop_PnL", "BF7") == "actual", at(wb, "Drop_PnL", "BF7"))
check("SAMPLE: actual receipts replace plan receipts for DRP-001 (both $1,200, in week 7)",
      at(wb, "Cash_Flow", "I10") == 1200 and at(wb, "Cash_Flow", "I9") == 0, (at(wb, "Cash_Flow", "I9"), at(wb, "Cash_Flow", "I10")))
check("SAMPLE: actual drop costs hit the cash flow on the drop date (week 12)", abs(at(wb, "Cash_Flow", "N17") + 160) < 0.01, at(wb, "Cash_Flow", "N17"))
check("SAMPLE: actual cash for DRP-001 inside the window = $1,200 - $160 = $1,040", near(at(wb, "Drop_Actuals", "AM7"), 1040), at(wb, "Drop_Actuals", "AM7"))
check("SAMPLE: Toggl hours read = 2.25", abs(v(wb, "TG_total") - 2.25) < 0.001, v(wb, "TG_total"))
check("SAMPLE: weekly hours log sums all hand and Toggl hours", abs(v(wb, "HL_week_total") - (56 + 2.25)) < 0.001, v(wb, "HL_week_total"))
check("[R1-18d] SAMPLE: with placement data but no renewal answers the pointer says 'not enough data: no renewal answers yet'",
      v(wb, "dash_pointer") == "not enough data: no renewal answers yet", v(wb, "dash_pointer"))
check("SAMPLE: median office probe = 150", v(wb, "pl_med_office") == 150, v(wb, "pl_med_office"))
check("[R1-18] SAMPLE: a typed bank balance in week 13 gives the latest week and its variance",
      v(wb, "cf_bank_week") == 13 and near(v(wb, "cf_bank_var"), -3000 - at(wb, "Cash_Flow", "O29")), (v(wb, "cf_bank_week"), v(wb, "cf_bank_var")))

# ---------------------------------------------------------------- [R1-04] decision rule after the decision date
def decision_case(n_offices, renewed, extra_actuals, as_of):
    e = [inp("as_of_override", as_of)]
    for i in range(n_offices):
        for j, val in enumerate([D(2026, 11, 20), "office", f"SAMPLE-O-{i}", f"SAMPLE office {i}", "budget_probe", 150, "SAMPLE answer"]):
            e.append(("Pipeline", f"{'ABCDEFG'[j]}{5 + i}", val))
    for i, rv in enumerate(renewed):   # campaign rows with a renewal answer, on DRP-001 (no 'Actuals in' needed for the rate)
        for j, val in enumerate([f"SAMPLE-CMP-{i}", "DRP-001", D(2027, 1, 12), "founder", "founding", f"SAMPLE-BR-{i}", 5]):
            e.append(("Actuals_Campaigns", f"{openpyxl.utils.get_column_letter(j + 1)}{5 + i}", val))
        e.append(("Actuals_Campaigns", f"W{5 + i}", rv))
    for r in extra_actuals:
        e.append(("Schedule", f"N{r}", "Y"))
    return run(e, "decision case")


wb, integ, errs = decision_case(5, ["N", "N"], [], D(2027, 4, 12))
check("[R1-04] after the decision date, renewal 0%, five offices at $100+, no extra drop with actuals: the rule waits instead of saying Stop",
      str(v(wb, "dash_rule")).startswith("not yet: the switch test needs 2 more extra drop(s)"), v(wb, "dash_rule"))
wb, integ, errs = decision_case(5, ["N", "N"], [10, 11], D(2027, 4, 26))
check("[R1-04] renewal 0%, five offices at $100+, two extra drops with actuals: Switch to office-paid", v(wb, "dash_rule") == "Switch to office-paid", v(wb, "dash_rule"))
wb, integ, errs = decision_case(4, ["N", "N"], [10, 11], D(2027, 4, 26))
check("[R1-04] renewal 0%, four offices, two extra drops with actuals, after the date: Stop (both tests fail)", v(wb, "dash_rule") == "Stop: both tests fail", v(wb, "dash_rule"))
wb, integ, errs = decision_case(0, ["Y", "N", "N"], [], D(2027, 4, 12))
check("[R1-04] renewal between 30% and 40%: the rule stays silent", str(v(wb, "dash_rule")).startswith("Rule silent"), v(wb, "dash_rule"))


# ================================================================ review round 2
import audit_formulas
# [R2-01] INDEX over one row names its row
_aud = audit_formulas.audit(SRC)
check("[R2-01] audit: no INDEX over a one-row range without a row argument, and at least 6,000 written as INDEX(range, 1, column)",
      _aud[7] == [] and _aud[8] >= 6000, (len(_aud[7]), _aud[8]))
check("[R2-01] the README's post-import cells read as expected: Scenarios!H10 = 2 and Cash_Flow!C35 = 13 (week of the low point)",
      at(base, "Scenarios", "H10") == 2 and at(base, "Cash_Flow", "C35") == 13 and REF["cf_low_week"] == "Cash_Flow!$C$35",
      (at(base, "Scenarios", "H10"), at(base, "Cash_Flow", "C35")))
wb, integ, errs = run([inp("active_scenario", "A")], "scenario A, INDEX column pick")
check("[R2-01] scenario A: the selector picks column D (post-decision price $180) and the low-point week is still found",
      v(wb, "sc_price_after") == 180 and v(wb, "cf_low_week") == 13, (v(wb, "sc_price_after"), v(wb, "cf_low_week")))

# [R2-07] legal-pack timing and new lines
att_r, acc_r, llc_r = row_of("su_attorney"), row_of("su_accountant"), row_of("su_llc")
check("[R2-07] attorney Wed 2026-11-18, LLC filing Thu 2026-11-19, accountant Fri 2026-11-20",
      [str(at(base, "Inputs", f"L{r_}"))[:10] for r_ in (att_r, llc_r, acc_r)] == ["2026-11-18", "2026-11-19", "2026-11-20"],
      [at(base, "Inputs", f"L{r_}") for r_ in (att_r, llc_r, acc_r)])
mb_r, ra_r = row_of("su_mailbox"), row_of("su_reg_agent")
check("[R2-07] mailbox (Include Y) and registered agent (Include N) lines are blank and flagged 'needs quote'",
      at(base, "Inputs", f"C{mb_r}") in (None, "") and at(base, "Inputs", f"N{mb_r}") == "Y" and "needs quote" in str(at(base, "Inputs", f"G{mb_r}"))
      and at(base, "Inputs", f"C{ra_r}") in (None, "") and at(base, "Inputs", f"N{ra_r}") == "N" and "needs quote" in str(at(base, "Inputs", f"G{ra_r}")),
      (at(base, "Inputs", f"G{mb_r}"), at(base, "Inputs", f"N{ra_r}")))
ins_r2 = row_of("insurance_monthly")
check("[R2-07] insurance status reads 'plan estimate (legal pack)'", str(at(base, "Inputs", f"G{ins_r2}")).startswith("plan estimate"), at(base, "Inputs", f"G{ins_r2}"))
wb, integ, errs = run([("Inputs", f"C{mb_r}", 120)], "mailbox quote typed")
check("[R2-07] a $120 mailbox quote flows into week 1 and deepens the low point by $120; capital tie-out still holds",
      near(v(wb, "cf_low"), b_low - 120) and near(at(wb, "Cash_Flow", "C25"), at(base, "Cash_Flow", "C25") - 120) and not integ,
      (v(wb, "cf_low"), at(wb, "Cash_Flow", "C25"), integ[:2]))

# [R2-08] cash at each stop point
check("[R2-08] stop points: Gate 0 paid out -$160, Gate 1 paid out -$3,120 with $1,200 received that day, Gate 1 week end -$3,856 (Cash_Flow week 7)",
      near(v(base, "dash_cr_out_gate0_date"), -160) and near(v(base, "dash_cr_out_gate1_date"), -3120) and near(v(base, "dash_cr_rec_gate1_date"), 1200)
      and near(v(base, "dash_cr_cum_gate1_date"), at(base, "Cash_Flow", "I27")) and near(v(base, "dash_cr_cum_gate1_date"), -3856),
      (v(base, "dash_cr_out_gate0_date"), v(base, "dash_cr_out_gate1_date"), v(base, "dash_cr_rec_gate1_date"), v(base, "dash_cr_cum_gate1_date")))
check("[R2-08] stop points: at the decision date paid out plus received equals week-26 cumulative cash ($11,409)",
      near(v(base, "dash_cr_out_decision_date") + v(base, "dash_cr_rec_decision_date"), v(base, "cf_end")) and near(v(base, "dash_cr_cum_decision_date"), v(base, "cf_end")),
      (v(base, "dash_cr_out_decision_date"), v(base, "dash_cr_rec_decision_date"), v(base, "cf_end")))

# [R2-09] polish
def readme_price_text(wb_):
    for row_ in wb_["README"].iter_rows():
        for c in row_:
            if isinstance(c.value, str) and "Standard price per office' (now" in c.value:
                return c.value
    return ""
wb, integ, errs = run([inp("price_std_office", 180)], "price 180 for README text")
check("[R2-09a] README price text is live: '(now $150)' at defaults, '(now $180)' after the edit",
      "(now $150)" in readme_price_text(base) and "(now $180)" in readme_price_text(wb), (readme_price_text(base)[:60], readme_price_text(wb)[:60]))
wb, integ, errs = run([("Actuals_Campaigns", "AK5", 1)], "value pasted in AK")
check("[R2-09b] a value in Actuals_Campaigns AK (last buffer column) is flagged", any("Actuals_Campaigns: header row differs" in f[2] for f in integ), [f[2] for f in integ])
check("[R2-09c][R3-02] break-even: 4 brand campaigns, which is 2 standard drops at three brands (startup outlays now include the $350 trademark)",
      v(base, "be_n_three") == 4 and v(base, "be_n_std_drops") == 2,
      (v(base, "be_n_three"), v(base, "be_n_std_drops")))

# [R2-06] no standard drop runs: no storage unit
wb, integ, errs = run([inp("std_brands_drop", 0)], "zero standard brands")
check("[R2-06] zero standard brands: storage start is blank, no storage is paid and nothing errors",
      v(wb, "storage_start") in (None, "") and at(wb, "Cash_Flow", "AC22") == 0 and not errs and not integ, (v(wb, "storage_start"), at(wb, "Cash_Flow", "AC22"), errs[:2], integ[:2]))

# [R2-03] hours logged outside a drop
def unlinked(hours_each, n_rows):
    e = []
    for i in range(n_rows):
        e += [("Hours_Log", f"A{5 + i}", D(2026, 10, 27) + dt.timedelta(days=7 * i)), ("Hours_Log", f"D{5 + i}", "admin_invoicing"),
              ("Hours_Log", f"E{5 + i}", hours_each), ("Hours_Log", f"F{5 + i}", "SAMPLE unlinked")]
    return e
wb, integ, errs = run(unlinked(10, 10), "100 unlinked hours")
road_r = row_of("dash_allin_road")
check("[R2-03] 100 hours logged with a blank drop_id lower both all-in rates and the detail names them",
      v(wb, "HL_unlinked_6m") == 100 and v(wb, "dash_allin_road") < v(base, "dash_allin_road") and v(wb, "dash_allin_all") < v(base, "dash_allin_all")
      and "100 logged hours not tied to a drop" in str(at(wb, "Dashboard", f"D{road_r}")),
      (v(wb, "HL_unlinked_6m"), round(v(wb, "dash_allin_road"), 2), round(v(wb, "dash_allin_all"), 2), str(at(wb, "Dashboard", f"D{road_r}"))[:90]))
wb, integ, errs = run(unlinked(2, 10), "20 unlinked hours")
check("[R2-03] 20 logged hours (fewer than the planned 45) leave the all-in rates on the plan's 45 hours",
      near(v(wb, "dash_allin_road"), v(base, "dash_allin_road")) and "45 planned discovery hours" in str(at(wb, "Dashboard", f"D{road_r}")),
      (v(wb, "dash_allin_road"), v(base, "dash_allin_road")))

# [R2-04] Gate 1 'paid in full' and [R2-05] probe medians
def pl_rows(rows):
    e = []
    for i, row_ in enumerate(rows):
        for j, val in enumerate(row_):
            if val is not None:
                e.append(("Pipeline", f"{'ABCDEFGHI'[j]}{5 + i}", val))
    return e
pay = [(D(2026, 11, 20), "brand", "SAMPLE-B-1", "SAMPLE brand 1", "paid", 100, None, "Y", None),
       (D(2026, 11, 25), "brand", "SAMPLE-B-2", "SAMPLE brand 2", "paid", 300, None, "Y", None),
       (D(2026, 12, 1), "brand", "SAMPLE-B-2", "SAMPLE brand 2", "paid", 300, None, "Y", None),
       (D(2026, 12, 2), "brand", "SAMPLE-B-1", "SAMPLE brand 1", "paid", 50, None, "Y", None)]
wb, integ, errs = run(pl_rows(pay) + [inp("as_of_override", D(2026, 12, 14))], "part payments")
check("[R2-04] $100 + $50 is not paid in full; $300 + $300 is: Gate 1 counts 1 brand paid in full, 2 with any payment",
      v(wb, "pl_g1_paid") == 1 and v(wb, "pl_paid_all") == 2, (v(wb, "pl_g1_paid"), v(wb, "pl_paid_all")))
probes = [(D(2026, 11, 2), "brand", "SAMPLE-B-1", "SAMPLE brand 1", "budget_probe", 400, "SAMPLE quote", "Y", None),
          (D(2026, 11, 2), "brand", "SAMPLE-B-1", "SAMPLE brand 1", "budget_probe", 1500, "SAMPLE quote", "Y", "after_1500"),
          (D(2026, 11, 3), "brand", "SAMPLE-B-2", "SAMPLE brand 2", "budget_probe", 800, "SAMPLE quote", "Y", None),
          (D(2026, 11, 5), "brand", "SAMPLE-B-2", "SAMPLE brand 2", "budget_probe", 800, "SAMPLE quote", "Y", None),
          (D(2026, 11, 6), "brand", "SAMPLE-B-3", "SAMPLE brand 3", "budget_probe", 1000, "SAMPLE quote", "N", "after_1500"),
          (D(2026, 11, 3), "office", "SAMPLE-O-1", "SAMPLE office 1", "budget_probe", 150, "SAMPLE quote", None, None),
          (D(2026, 11, 4), "office", "SAMPLE-O-1", "SAMPLE office 1", "budget_probe", 150, "SAMPLE quote", None, None),
          (D(2026, 11, 5), "office", "SAMPLE-O-2", "SAMPLE office 2", "budget_probe", 50, "SAMPLE quote", None, None)]
wb, integ, errs = run(pl_rows(probes), "probe medians")
check("[R2-05] medians use one before-anchor answer per organisation: brands $600 (400, 800), after-anchor $1,250 (1,500, 1,000), offices $100 (150, 50)",
      v(wb, "pl_med_brand") == 600 and v(wb, "pl_med_brand_after") == 1250 and v(wb, "pl_med_office") == 100,
      (v(wb, "pl_med_brand"), v(wb, "pl_med_brand_after"), v(wb, "pl_med_office")))
check("[R2-05] Gate 0 still counts a brand whose after-anchor answer clears $500 and who controls the budget (2: B-1 and B-2; B-3 does not control it)",
      v(wb, "pl_g0_brand_yes") == 2, v(wb, "pl_g0_brand_yes"))

# [R2-02] Gate 3 cannot pass on part of the evidence
def gate3_case(drp3_in):
    e = [inp("as_of_override", D(2027, 3, 15)), ("Schedule", "N8", "Y"), ("Schedule", "N9", "Y" if drp3_in else "N")]
    e += pl_rows([(D(2027, 3, 1), "brand", "SAMPLE-B-7", "SAMPLE brand 7", "renewal_answer", 1500, "yes", "Y", None)])
    for i in range(3):
        row_ = [f"SAMPLE-CMP-{10 + i}", "DRP-002", D(2027, 2, 9), "courier", "standard", f"SAMPLE-BR-{10 + i}", 10, 550, 500, 480, 450, 0, 300, 60, 40, 30, 2,
                1500, D(2027, 2, 10), 0, D(2027, 2, 18), "yes", "pending", 750, D(2027, 1, 26), 6]
        for j, val in enumerate(row_):
            e.append(("Actuals_Campaigns", f"{openpyxl.utils.get_column_letter(j + 1)}{5 + i}", val))
    for j, val in enumerate(["DRP-002", D(2027, 2, 9), "courier", 10, 3, 80, 300, 0, "SAMPLE"]):
        e.append(("Actuals_DropCosts", f"{openpyxl.utils.get_column_letter(j + 1)}5", val))
    for i, (act, h) in enumerate([("brand_sales", 18), ("office_coordination", 10), ("receiving_kitting", 4), ("drop_day", 8), ("checks_pickup", 4),
                                  ("reporting_debriefs", 6), ("admin_invoicing", 2)]):
        e += [("Hours_Log", f"A{5 + i}", D(2027, 2, 9)), ("Hours_Log", f"B{5 + i}", "DRP-002"), ("Hours_Log", f"D{5 + i}", act), ("Hours_Log", f"E{5 + i}", h)]
    return run(e, "gate 3 coverage")
wb, integ, errs = gate3_case(False)
status, gates = gate_tables(wb)
g3row = status.get("Contribution per founder hour, standard runs")
check("[R2-02] DRP-002 in, DRP-003 not (as of 2027-03-15): the Gate 3 rate reads 'not yet: 1 standard run(s) without actuals' and Gate 3 is not a pass",
      str(g3row).startswith("not yet: 1 standard run") and gates.get("Gate 3") == "not yet", (g3row, gates.get("Gate 3")))
check("[R2-02] the pointer and coverage line name the drops with results due that have no actuals (1 of 3 in)",
      v(wb, "dash_due_cov") == "1 of 3" and "2 drop(s) with results due have no actuals yet" in str(v(wb, "dash_pointer")), (v(wb, "dash_due_cov"), v(wb, "dash_pointer")))
check("[R2-02] SAMPLE Gate 3 run: no integrity or data failures", not integ and not errs, (integ[:2], errs[:2]))
check("[R2-02] the six-month goal 'contribution per founder hour, standard runs' follows the same rule (not yet while DRP-003 has no actuals)",
      str(v(wb, "dash_goal_status_g3_per_hour")).startswith("not yet: 1 standard run"), v(wb, "dash_goal_status_g3_per_hour"))
wb, integ, errs = gate3_case(True)
status, gates = gate_tables(wb)
check("[R2-02] with both standard runs in, the rate is judged on all of them and Gate 3 passes (renewal at $1,500 yes, rate above $50)",
      status.get("Contribution per founder hour, standard runs") == "pass" and gates.get("Gate 3") == "pass",
      (status.get("Contribution per founder hour, standard runs"), gates.get("Gate 3"), v(wb, "dash_g3_rate")))

# ---------------------------------------------------------------- bad data is caught
wb, integ, errs = run([("Hours_Log", "A5", D(2027, 1, 5)), ("Hours_Log", "D5", "kitting"), ("Hours_Log", "E5", 2),
                       ("Pipeline", "A5", D(2026, 11, 1)), ("Pipeline", "B5", "office"), ("Pipeline", "C5", "SAMPLE-O-9"), ("Pipeline", "E5", "meeting")], "bad data")
labels = [f[2] for f in integ]
check("bad activity key and bad event are flagged as data failures", any("invalid activity" in l for l in labels) and any("bad event" in l for l in labels), labels)

# ================================================================ review round 3
BANNER_OK = "all integrity and data checks pass"


def checks_rows(wb_):
    ws_ = wb_["Checks"]
    out_ = {}
    r_ = 5
    while ws_[f"A{r_}"].value is not None:
        out_[ws_[f"C{r_}"].value] = (ws_[f"B{r_}"].value, ws_[f"D{r_}"].value, ws_[f"E{r_}"].value)
        r_ += 1
    return out_


def text_fail(integ_, tab):
    return any(f[2].startswith(tab + ":") and "stored as text" in f[2] for f in integ_)


check("[R3-01] baseline: the Dashboard banner reads 'all integrity and data checks pass'", at(base, "Dashboard", "G5") == BANNER_OK, at(base, "Dashboard", "G5"))


# [R3-01] text in a date or money column: the same campaign-zero rows pasted with real dates and with the dates as text
def c0_case(text_dates=False, text_revenue=False):
    e = [inp("as_of_override", D(2027, 1, 29)), ("Schedule", "N7", "Y")]
    for i, row_ in enumerate(camp):
        for j, val in enumerate(row_):
            if val is None:
                continue
            if text_dates and isinstance(val, dt.date):
                val = val.isoformat()
            if text_revenue and j == 17:
                val = str(val)
            e.append(("Actuals_Campaigns", f"{openpyxl.utils.get_column_letter(j + 1)}{5 + i}", val))
    for j, val in enumerate(["DRP-001", D(2027, 1, 12), "founder", 5, 2, 45, 95, 20, "SAMPLE"]):
        if text_dates and isinstance(val, dt.date):
            val = val.isoformat()
        e.append(("Actuals_DropCosts", f"{openpyxl.utils.get_column_letter(j + 1)}5", val))
    return run(e, "campaign zero actuals")


wb_r, integ_r, errs_r = c0_case()
check("[R3-01] real dates: receipts $1,200 in week 7, no data failure, banner green",
      at(wb_r, "Cash_Flow", "I10") == 1200 and not integ_r and not errs_r and at(wb_r, "Dashboard", "G5") == BANNER_OK,
      (at(wb_r, "Cash_Flow", "I10"), integ_r[:2], at(wb_r, "Dashboard", "G5")))
wb, integ, errs = c0_case(text_dates=True)
check("[R3-01] the same rows with text dates: receipts vanish (the failure the check exists for), Checks flags both paste tabs, the banner names data failures",
      at(wb, "Cash_Flow", "I10") == 0 and text_fail(integ, "Actuals_Campaigns") and text_fail(integ, "Actuals_DropCosts")
      and str(at(wb, "Dashboard", "G5")).endswith("data check(s) failing: pasted rows are ignored or misread (see Checks)"),
      (at(wb, "Cash_Flow", "I10"), [f[2][:50] for f in integ], at(wb, "Dashboard", "G5")))
wb, integ, errs = c0_case(text_revenue=True)
check("[R3-01] revenue typed as text ('600'): flagged on Checks and on the banner", text_fail(integ, "Actuals_Campaigns") and at(wb, "Dashboard", "G5") != BANNER_OK,
      ([f[2][:50] for f in integ], at(wb, "Dashboard", "G5")))
txt_rows = [("Pipeline", "A5", "2026-11-04"), ("Pipeline", "B5", "office"), ("Pipeline", "C5", "SAMPLE-O-1"), ("Pipeline", "E5", "yes_in_principle"),
            ("Pipeline", "A6", D(2026, 11, 20)), ("Pipeline", "B6", "office"), ("Pipeline", "C6", "SAMPLE-O-1"), ("Pipeline", "E6", "signed"), ("Pipeline", "F6", "60"),
            ("Hours_Log", "A5", "2026-11-02"), ("Hours_Log", "D5", "brand_sales"), ("Hours_Log", "E5", 2),
            ("Hours_Log", "A6", D(2026, 11, 3)), ("Hours_Log", "D6", "brand_sales"), ("Hours_Log", "E6", "1.5")]
wb, integ, errs = run(txt_rows, "text in pipeline and hours")
ck_ = checks_rows(wb)
check("[R3-01] a text date and a text headcount on Pipeline, and a text date and text hours on Hours_Log, are each counted (2 and 2)",
      ck_["Pipeline: dates or value_number stored as text (columns A and F)"][1] == 2 and ck_["Hours_Log: dates or hours stored as text (columns A and E)"][1] == 2
      and text_fail(integ, "Pipeline") and text_fail(integ, "Hours_Log"),
      (ck_["Pipeline: dates or value_number stored as text (columns A and F)"], ck_["Hours_Log: dates or hours stored as text (columns A and E)"]))

# [R3-02] trademark in the six months at Gate 2
tm_r = row_of("su_tm")
check("[R3-02] trademark $350 included (default Y), timed on the Gate 2 day (2027-01-22)",
      at(base, "Inputs", f"N{tm_r}") == "Y" and at(base, "Inputs", f"S{tm_r}") == "Y" and str(at(base, "Inputs", f"R{tm_r}"))[:10] == "2027-01-22"
      and at(base, "Inputs", f"Q{tm_r}") == 350, (at(base, "Inputs", f"N{tm_r}"), at(base, "Inputs", f"R{tm_r}"), at(base, "Inputs", f"Q{tm_r}")))
wb, integ, errs = run([("Inputs", f"N{tm_r}", "N")], "trademark excluded")
check("[R3-02] Include = N restores the round-2 figures: low -$4,126 in week 12, commit $5,126, and the line shows as changed",
      near(v(wb, "cf_low"), -4126) and v(wb, "cf_low_week") == 12 and near(v(wb, "cf_commit"), 5126) and v(wb, "changed_count") == 1,
      (v(wb, "cf_low"), v(wb, "cf_low_week"), v(wb, "cf_commit"), v(wb, "changed_count")))

# [R3-03] a drop added mid-sequence (Tue 2027-03-16, between DRP-003 and DRP-004) in the next spare row
mid = [("Schedule", "B12", "SAMPLE added drop"), ("Schedule", "E12", D(2027, 3, 16)), ("Schedule", "G12", "courier"), ("Schedule", "H12", "standard"),
       ("Schedule", "I12", 10), ("Schedule", "J12", 0), ("Schedule", "K12", 3), ("Schedule", "L12", 0), ("Schedule", "M12", 0), inp("as_of_override", D(2027, 3, 10))]
wb, integ, errs = run(mid, "drop added mid-sequence")
ck_ = checks_rows(wb)
info = [k for k in ck_ if k and "out of date order" in k][0]
check("[R3-03] a mid-sequence drop: no integrity failure; the out-of-order row is information only (value 1, 'info')",
      not integ and not errs and ck_[info][0] == "Info" and ck_[info][1] == 1 and ck_[info][2] == "info" and at(wb, "Dashboard", "G5") == BANNER_OK,
      (integ[:2], ck_[info], at(wb, "Dashboard", "G5")))
check("[R3-03] as of 2027-03-10 the next scheduled drop is the added DRP-006 (Tue 2027-03-16), not DRP-004 (Tue 2027-03-23)",
      v(wb, "dash_next") == "DRP-006" and "2027-03-16" in str(at(wb, "Dashboard", "D" + str(row_of("dash_next")))),
      (v(wb, "dash_next"), at(wb, "Dashboard", "D" + str(row_of("dash_next")))))
wb, integ, errs = run([inp("as_of_override", D(2027, 4, 30))], "no drop left")
check("[R3-03] with no upcoming drop the Dashboard says 'none scheduled' and errors nowhere", v(wb, "dash_next") == "none scheduled" and not errs, (v(wb, "dash_next"), errs[:2]))

# [R3-04] a lever a scenario does not name follows the plan base
wb, integ, errs = run([inp("std_brands_drop", 2), inp("active_scenario", "A")], "two brands per drop under scenario A")
scq = [at(wb, "Scenarios", f"Q{r_}") for r_ in range(10, 22)]
check("[R3-04] two brands per standard drop under scenario A: the decision-window drops and the steady state carry 2 brands; no scenario cell shows 'changed'",
      at(wb, "Schedule", "K10") == 2 and at(wb, "Schedule", "K11") == 2 and v(wb, "sc_ss_brands") == 2 and "changed" not in scq and not integ,
      (at(wb, "Schedule", "K10"), v(wb, "sc_ss_brands"), scq))
check("[R3-04] the A steady state moves with it: two-brand drops at $180 earn less than its three-brand $9,406", v(wb, "ss_contrib") < 9406 - 1, v(wb, "ss_contrib"))
wb, integ, errs = run([("Inputs", f"C{row_of('plan_ss_drops')}", 3), inp("active_scenario", "B")], "three drops a month under B")
check("[R3-04] three steady-state drops a month on Inputs reach scenario B (and its hire threshold still follows the plan base)",
      v(wb, "sc_ss_drops") == 3 and at(wb, "Scenarios", "H21") == 2, (v(wb, "sc_ss_drops"), at(wb, "Scenarios", "H21")))

# [R3-05] cash at each stop point: refunds owed and the net
check("[R3-05] Gate 1 net equals paid out (founding cash received is refundable): -$3,120",
      near(v(base, "dash_cr_net_gate1_date"), -3120) and near(v(base, "dash_cr_ref_gate1_date"), 1200), (v(base, "dash_cr_net_gate1_date"), v(base, "dash_cr_ref_gate1_date")))
check("[R3-05] Gate 2 net -$4,126 (paid out -$5,326 plus $1,200 received, nothing to refund)",
      near(v(base, "dash_cr_net_gate2_date"), -4126) and near(v(base, "dash_cr_ref_gate2_date"), 0), (v(base, "dash_cr_net_gate2_date"), v(base, "dash_cr_ref_gate2_date")))
check("[R3-05] Gate 3: DRP-004's $2,250 of deposits are owed back, so the net is paid out + received - $2,250",
      near(v(base, "dash_cr_ref_gate3_date"), 2250)
      and near(v(base, "dash_cr_net_gate3_date"), v(base, "dash_cr_out_gate3_date") + v(base, "dash_cr_rec_gate3_date") - 2250) and v(base, "dash_cr_net_gate3_date") > 0,
      (v(base, "dash_cr_out_gate3_date"), v(base, "dash_cr_rec_gate3_date"), v(base, "dash_cr_ref_gate3_date"), v(base, "dash_cr_net_gate3_date")))

# [R3-06] unpriced lines are named beside the cash need
com_r = row_of("dash_commit")
hn_r = row_of("su_hnoa")
check("[R3-06] the capital-to-commit line says it counts 5 unpriced cost lines as $0", v(base, "unpriced_count") == 5 and "counts 5 unpriced cost line(s)" in str(at(base, "Dashboard", f"D{com_r}")),
      (v(base, "unpriced_count"), at(base, "Dashboard", f"D{com_r}")))
wb, integ, errs = run([("Inputs", f"C{hn_r}", 200)], "auto add-on quoted")
check("[R3-06] a $200 auto add-on quote: 4 unpriced lines left and the low point deepens by $200", v(wb, "unpriced_count") == 4 and near(v(wb, "cf_low"), b_low - 200),
      (v(wb, "unpriced_count"), v(wb, "cf_low")))

# [R3-07] renewal rate counts brands
rows7 = [("SAMPLE-CMP-1", "SAMPLE-BR-1", "Y"), ("SAMPLE-CMP-2", "SAMPLE-BR-1", "N"), ("SAMPLE-CMP-3", "SAMPLE-BR-2", "N")]
e7 = []
for i, (cid, bid, w) in enumerate(rows7):
    e7 += [("Actuals_Campaigns", f"A{5 + i}", cid), ("Actuals_Campaigns", f"B{5 + i}", "DRP-001"), ("Actuals_Campaigns", f"C{5 + i}", D(2027, 1, 12)),
           ("Actuals_Campaigns", f"F{5 + i}", bid), ("Actuals_Campaigns", f"G{5 + i}", 5), ("Actuals_Campaigns", f"W{5 + i}", w)]
wb, integ, errs = run(e7, "renewal per brand")
check("[R3-07] brand 1 renewed then declined, brand 2 declined: the renewal rate is 1 of 2 brands (50%), not 1 of 3 campaigns",
      near(v(wb, "dash_renewal"), 0.5, 0.0001), v(wb, "dash_renewal"))

# [R3-08] scenario A grows the roster
wb, integ, errs = run([inp("active_scenario", "A")], "scenario A roster")
check("[R3-08] scenario A: each decision-window drop recruits 4 new offices (roster 12 to 20) and carries 12 more office hours",
      at(wb, "Schedule", "J10") == 4 and at(wb, "Schedule", "J11") == 4 and at(wb, "Schedule", "V11") == 20
      and near(at(wb, "Drop_PnL", "AC10") - at(base, "Drop_PnL", "AC10"), 12), (at(wb, "Schedule", "J10"), at(wb, "Schedule", "V11"), at(wb, "Drop_PnL", "AC10")))

# [R3-09] week 1 moved on its own
wb, integ, errs = run([inp("week1_start", D(2026, 11, 9))], "week 1 moved alone")
check("[R3-09] week 1 moved two weeks with every gate and drop date left: a schedule check fails and the banner turns red",
      any("Week 1 moved" in f[2] for f in integ) and at(wb, "Dashboard", "G5") != BANNER_OK, ([f[2][:40] for f in integ], at(wb, "Dashboard", "G5")))
wb, integ, errs = run([inp("week1_start", D(2026, 11, 9)), inp("gate0_date", D(2026, 11, 27))], "week 1 and Gate 0 moved")
check("[R3-09] week 1 and Gate 0 moved together: the check is quiet", not any("Week 1 moved" in f[2] for f in integ), [f[2][:40] for f in integ])

# [R3-11] polish
wb, integ, errs = run([inp("c0_brands", 0), inp("std_brands_drop", 0)], "zero brands, headline rates")
check("[R3-11d] zero brands everywhere: both Gate 3 headline rates read 'n/a'", v(wb, "dash_g3_rate") == "n/a" and v(wb, "dash_g3_roster") == "n/a",
      (v(wb, "dash_g3_rate"), v(wb, "dash_g3_roster")))
check("[R3-11b] README tab: verified legal-pack fees are named in the opening paragraph", "or the legal pack" in readme_txt)
check("[R3-11c] one import route: the README tab gives File, Import, 'Replace spreadsheet'", "'Replace spreadsheet'" in readme_txt)

# ---------------------------------------------------------------- [R1-17] clean CSV exports
import csv as _csv
import glob as _glob
csv_dir = os.path.join(os.path.dirname(HERE), "csv")
bad_csv = []
files = sorted(_glob.glob(os.path.join(csv_dir, "*.csv")))
for f in files:
    rows_ = list(_csv.reader(open(f, encoding="utf-8")))
    if not rows_ or any(not re.fullmatch(r"[a-z0-9_]+", h) for h in rows_[0]) or len({len(r) for r in rows_}) != 1:
        bad_csv.append(os.path.basename(f))
check("[R1-17] every file in csv/ has one snake_case header row and rows of equal length (snapshots sit in csv/snapshots/)",
      len(files) >= 15 and not bad_csv and os.path.isdir(os.path.join(csv_dir, "snapshots")), (len(files), bad_csv))

n_fail = sum(1 for r in RESULTS if not r[1])
print(f"\n{len(RESULTS) - n_fail} of {len(RESULTS)} tests pass")
sys.exit(1 if n_fail else 0)
