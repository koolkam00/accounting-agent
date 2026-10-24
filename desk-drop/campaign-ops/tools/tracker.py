"""Read a tracker .xlsx (downloaded from Google Sheets, or built here) and recompute every number
from the logged rows, independently of the sheet's formulas.

Every generator (report, exports, route sheet, cards, links, invoices) uses this module, so the
numbers they print come from rows, never from retyped figures. compare() checks these numbers
against the sheet's own formula results when the file carries them (a Google Sheets download
does; a file built here does after tools/recalc.py).
"""
import datetime as dt
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import opscal  # noqa: E402
import trackerdef as T  # noqa: E402

T.define()


def _d(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    s = str(v).strip()
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        return None


def _t(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, (int, float)):
        secs = round(float(v) % 1 * 86400)
        return (dt.datetime(2000, 1, 1) + dt.timedelta(seconds=secs)).time()
    m = re.match(r"(\d{1,2}):(\d{2})", str(v))
    return dt.time(int(m.group(1)), int(m.group(2))) if m else None


def xround(x, n=0):
    """Spreadsheet ROUND: half away from zero (Python's round() is half to even)."""
    if x is None:
        return None
    m = 10 ** n
    v = abs(x) * m + 0.5 + 1e-9
    import math
    return math.copysign(math.floor(v) / m, x)


def _n(v):
    if v in (None, ""):
        return 0
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v)
    # Website rows: strip the one leading apostrophe the intake script may add (website/HANDOFF.md).
    if len(s) > 1 and s[0] == "'" and s[1] in "=+-@":
        s = s[1:]
    return s


class Tracker:
    def __init__(self, path):
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.rows = {}
        self.cached = {}
        for name, tab in T.TABS.items():
            if name not in self.wb.sheetnames:
                self.rows[name] = []
                continue
            ws = self.wb[name]
            hdr = [c.value for c in ws[tab.header_row]]
            idx = {h: i for i, h in enumerate(hdr) if h}
            out, cache = [], []
            skipped = 0
            for row in ws.iter_rows(min_row=tab.first(), max_row=ws.max_row, values_only=True):
                rec = {}
                for col in tab.cols:
                    i = idx.get(col.name)
                    v = row[i] if i is not None and i < len(row) else None
                    rec[col.name] = v
                key0 = tab.cols[0].name
                inputs = [c.name for c in tab.cols if c.kind != "A"] or [c.name for c in tab.cols]
                if all(rec[n] in (None, "") for n in inputs):
                    skipped += 1
                    continue
                keep = {k: v for k, v in rec.items() if tab.cols[tab.index[k]].kind != "A"}
                keep["_row"] = tab.first() + len(out) + skipped
                out.append(keep)
                cache.append({k: v for k, v in rec.items() if tab.cols[tab.index[k]].kind == "A"})
            self.rows[name] = out
            self.cached[name] = cache
        self.config = {}
        ws = self.wb["config"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                self.config[row[0]] = row[1]
        self.company_name = _s(self.config.get("company_name"))
        self.today = _d(self.config.get("today_override")) or dt.date.today()
        self._index()

    # ------------------------------------------------------------ lookups
    def _index(self):
        self.drops = {_s(r["drop_id"]): r for r in self.rows["drops"]}
        self.brands = {_s(r["brand_id"]): r for r in self.rows["brands"]}
        self.offices = {_s(r["office_id"]): r for r in self.rows["offices"]}
        self.campaigns = {_s(r["campaign_id"]): r for r in self.rows["campaigns"]}
        self.signed = self._paste("in_signed_offices", "office_id")
        self.intake = {}
        for tab in ("brand_intake", "brand_intake_founder"):
            if tab == "brand_intake_founder":
                cache = self.cached.get(tab, [])
                for r in cache:
                    if r.get("submitted_at"):
                        self.intake.setdefault(_s(r["submitted_at"]), r)
            else:
                for r in self.rows[tab]:
                    if r.get("submitted_at"):
                        self.intake.setdefault(_s(r["submitted_at"]), r)

    def _paste(self, tab, key):
        if tab not in self.wb.sheetnames:
            return {}
        ws = self.wb[tab]
        hdr = [c.value for c in ws[1]]
        out = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = {h: v for h, v in zip(hdr, row) if h}
            if rec.get(key) not in (None, ""):
                out.setdefault(_s(rec[key]), rec)
        return out

    def drop_date(self, drop_id):
        r = self.drops.get(drop_id)
        if not r:
            return None
        d = _d(r["drop_date"])
        if d is None or opscal.problems(d):
            return None
        return d

    def office(self, oid):
        """Office record: E's pasted row by header name, plus F's own columns."""
        r = dict(self.signed.get(oid, {}))
        f = self.offices.get(oid, {})
        r.update({k: v for k, v in f.items()})
        r["industry"] = _s(f.get("industry_override")) or _s(r.get("industry"))
        ph = _s(f.get("pantry_owner_phone_override")) or _s(r.get("pantry_owner_phone"))
        if not ph and _s(r.get("pantry_owner_is_contact")) == "Y":
            ph = _s(r.get("contact_phone"))
        r["owner_phone"] = ph
        h = r.get("tuesday_headcount")
        h = _n(h) if h not in (None, "") else None
        r["tuesday_headcount"] = h
        r["headcount_band"] = band(h)
        r["default_units"] = default_units(h, self.config)
        return r

    def brand(self, bid):
        b = dict(self.brands.get(bid, {}))
        src = self.intake.get(_s(b.get("intake_key")), {})
        for k, v in src.items():
            b.setdefault(k, v)
            if b.get(k) in (None, ""):
                b[k] = v
        return b

    def campaign(self, cid):
        c = dict(self.campaigns[cid])
        b = self.brand(_s(c["brand_id"]))
        c["brand_name"] = _s(b.get("brand_name"))
        c["category"] = _s(b.get("category"))
        c["drop_date"] = self.drop_date(_s(c["drop_id"]))
        c["delivery_model"] = _s(self.drops.get(_s(c["drop_id"]), {}).get("delivery_model"))
        c["price_expected_usd"] = expected_price(c, self.config)
        c["price_usd"] = _n(c["price_override_usd"]) if c.get("price_override_usd") not in (None, "") else c["price_expected_usd"]
        return c

    # ------------------------------------------------------------ the ladder, per office per campaign
    def approvals_for(self, cid):
        return [r for r in self.rows["approvals"] if _s(r["campaign_id"]) == cid]

    def office_drop_number(self, oid, dd):
        n = 0
        for s in self.rows["stops"]:
            if _s(s["office_id"]) != oid:
                continue
            sd = self.drop_date(_s(s["drop_id"]))
            if sd and dd and sd < dd and _s(s.get("outcome")) != "not_delivered":
                n += 1
        return n + 1

    def response_rows(self, cid):
        if cid not in self.wb.sheetnames:
            return None
        ws = self.wb[cid]
        out = []
        for i, row in enumerate(ws.iter_rows(min_row=2, max_col=7, values_only=True), start=2):
            if row[0] in (None, ""):
                continue
            ts = row[0]
            if isinstance(ts, (int, float)):
                ts = dt.datetime(1899, 12, 30) + dt.timedelta(days=float(ts))
            out.append({"_row": i, "ts": ts, "code": _s(row[1]), "tried": _s(row[2]), "rating": row[3], "buy": _s(row[4]),
                        "why": _s(row[5]), "comment": _s(row[6])})
        return out

    def ladder(self, cid):
        """Per-office rows plus campaign totals for one campaign, from raw rows only."""
        c = self.campaign(cid)
        T0 = c["drop_date"]
        win_end = opscal.add_bd(T0, 4) if T0 else None
        cutoff = opscal.add_bd(T0, 1) if T0 else None
        inv = [r for r in self.rows["inventory"] if _s(r["campaign_id"]) == cid]
        cnts = [r for r in self.rows["counts"] if _s(r["campaign_id"]) == cid]
        phs = [r for r in self.rows["photos"] if _s(r["campaign_id"]) == cid]
        resp = self.response_rows(cid) or []
        inwin = [x for x in resp if T0 and _d(x["ts"]) and T0 <= _d(x["ts"]) <= win_end] if T0 else []
        scans = self.rows["scans"]
        did = _s(c["drop_id"])
        stop_out = {_s(x["office_id"]): _s(x.get("outcome")) for x in self.rows["stops"] if _s(x["drop_id"]) == did}
        frozen = {_s(x["office_id"]): x for x in self.rows.get("survey_frozen", []) if _s(x["campaign_id"]) == cid}
        offices = []
        seq = 0
        for a in self.approvals_for(cid):
            oid = _s(a["office_id"])
            o = self.office(oid)
            n_drop = self.office_drop_number(oid, T0)
            mode = "explicit" if n_drop <= int(_n(self.config["approvals_explicit_drops"])) else "opt-out"
            rep = _s(a.get("reply"))
            ask = _d(a.get("asked_date"))
            if mode == "explicit":
                approved = "Y" if rep == "yes" else "N" if rep == "no" else "pending"
            else:
                approved = "N" if rep == "no" else ("Y" if ask and T0 and ask <= opscal.add_bd(T0, -5) else "pending")
            units_alloc = _n(a["units_override"]) if a.get("units_override") not in (None, "") else o["default_units"]
            mine = lambda rows, **kw: [r for r in rows if _s(r["office_id"]) == oid and all(_s(r[k]) == v for k, v in kw.items())]
            dl = mine(inv, movement="delivered")
            delivered = sum(_n(r["qty"]) for r in dl)
            dts = [dt.datetime.combine(_d(r["date"]), _t(r["time"]) or dt.time(0)) for r in dl if _d(r["date"])]
            delivered_at = max(dts) if dts else None
            dsrc = _s(dl[0]["count_source"]) if dl else ""
            # On the drop (review round 1, R1-02): approved, and not turned away at the door.
            refused = stop_out.get(oid) in REFUSED
            on_drop = "N" if approved == "N" else ("pending" if approved != "Y" else ("N" if (refused and delivered == 0) else "Y"))
            # Letters run over offices that took the brand (the sheet's office_seq does the same).
            if on_drop != "N":
                seq += 1
            usable = lambda r: _s(r["no_people_or_screens"]) == "Y" and (_s(r["photo_type"]) == "later" or _s(r["card_visible"]) == "Y")
            pp = [r for r in mine(phs, photo_type="placement") if usable(r)]
            pp0 = pp[0] if pp else None
            pdate = _d(pp0["taken_date"]) if pp0 else None
            pc_rows = mine(cnts, count_type="placed")
            placed_count = sum(_n(r["units"]) for r in pc_rows) if pc_rows else None
            placed = 0 if pdate is None else (placed_count or 0)
            within = "Y" if (pdate and cutoff and pdate <= cutoff) else "N"
            lag = None
            if delivered_at and pdate:
                pt = _t(pp0.get("taken_time")) or dt.time(0)
                da = delivered_at if isinstance(delivered_at, dt.datetime) else dt.datetime.combine(delivered_at, dt.time(0))
                lag = xround((dt.datetime.combine(pdate, pt) - da).total_seconds() / 3600, 1)
            topped = sum(_n(r["qty"]) for r in mine(inv, movement="topped_up"))
            r3 = mine(cnts, count_type="remaining_day3")
            r5 = mine(cnts, count_type="remaining_day5")
            rem3 = sum(_n(r["units"]) for r in r3) if r3 else None
            rem5 = sum(_n(r["units"]) for r in r5) if r5 else None
            depleted = max(0, placed + topped - rem5) if rem5 is not None else None
            lp = [r for r in mine(phs, photo_type="later") if usable(r) and _d(r["taken_date"])]
            ldate = max((_d(r["taken_date"]) for r in lp), default=None)
            lfile = next((_s(r["file_name"]) for r in lp if _d(r["taken_date"]) == ldate), "") if ldate else ""
            lday = opscal.window_day(T0, ldate) if (ldate and T0) else None
            code = "%s.%s" % (cid, oid)
            ro = [x for x in inwin if x["code"] == code]
            count = lambda f: sum(1 for x in ro if f(x))
            ratings = [x["rating"] for x in ro if isinstance(x["rating"], (int, float)) and x["rating"] >= 1]
            qslug, pslug = slug(cid, oid, "q"), slug(cid, oid, "p")

            def delta(sl):
                rs = [s for s in scans if _s(s["slug"]) == sl]
                close = [_n(s["clicks_to_date"]) for s in rs if _s(s["reading_type"]) == "close"]
                base = [_n(s["clicks_to_date"]) for s in rs if _s(s["reading_type"]) == "baseline"]
                if not close:
                    return None
                return max(close) - (max(base) if base else 0)

            label = chr(64 + seq) if on_drop != "N" else ""
            if label and _s(self.config.get("report_office_names")) == "names_if_public" and _s(o.get("public_naming_ok")) == "Y":
                label = _s(o.get("company"))
            offices.append({
                "office_id": oid, "company": _s(o.get("company")), "seq": seq if on_drop != "N" else "", "office_label": label,
                "industry": _s(o.get("industry")), "headcount_band": o["headcount_band"], "tuesday_headcount": o["tuesday_headcount"],
                "office_drop_number": n_drop, "approval_mode": mode, "approved": approved, "on_drop": on_drop,
                "stop_outcome": stop_out.get(oid, ""), "turned_away": approved == "Y" and on_drop == "N",
                "units_allocated": units_alloc, "card_code": code,
                "units_delivered": delivered, "delivered_at": delivered_at, "delivered_count_source": dsrc,
                "placement_photo_file": _s(pp0["file_name"]) if pp0 else "", "placement_photo_date": pdate,
                "placed_count": placed_count, "units_placed": placed, "placed_within_1bd": within, "placement_lag_hours": lag,
                "units_topped_up": topped, "remaining_day3": rem3, "remaining_day5": rem5, "units_depleted_day5": depleted,
                "later_photo_file": lfile, "later_photo_date": ldate, "later_photo_day": lday,
                "both_photos": "Y" if (pp0 and lfile) else "N",
                "survey_responses": len(ro), "tried_yes": count(lambda x: x["tried"] == "Yes"),
                "tried_no_didnt_see": count(lambda x: x["why"] == "Didn't see it"),
                "tried_no_not_for_me": count(lambda x: x["why"] == "Not for me"),
                "tried_no_gone": count(lambda x: x["why"] == "Gone by the time I looked"),
                "rating_n": len(ratings), "rating_sum": sum(ratings),
                "buy_yes": count(lambda x: x["buy"] == "Yes"), "buy_maybe": count(lambda x: x["buy"] == "Maybe"),
                "buy_no": count(lambda x: x["buy"] == "No"), "comments": count(lambda x: x["comment"].strip() != ""),
                "qr_scans": delta(qslug), "post_link_opens": delta(pslug), "qr_slug": qslug, "post_slug": pslug,
                "survey_frozen": "Y" if oid in frozen else "",
            })
            if oid in frozen and T0:
                # The raw answers were deleted after 12 months: the frozen totals stand in (R1-12).
                fz = frozen[oid]
                for k in FROZEN_COLS:
                    offices[-1][k] = int(_n(fz.get(k)))
        tot = lambda k: sum((o[k] or 0) for o in offices)
        approved = [o for o in offices if o["approved"] == "Y"]
        ondrop = [o for o in offices if o["on_drop"] == "Y"]
        deleted = _d(c.get("answers_deleted_date")) is not None
        rn = tot("rating_n")
        inv_rows = [r for r in self.rows["invoices"] if _s(r["campaign_id"]) == cid]
        reds = [r for r in self.rows["redemptions"] if _s(r["campaign_id"]) == cid and _d(r["as_of_date"])]
        asof = max((_d(r["as_of_date"]) for r in reds), default=None)
        red_to_date = sum(_n(r["redemptions_to_date"]) for r in reds if _d(r["as_of_date"]) == asof) if asof else None
        d30 = [r for r in reds if T0 and (_d(r["as_of_date"]) - T0).days >= 30]
        d30d = min((_d(r["as_of_date"]) for r in d30), default=None)
        red30 = sum(_n(r["redemptions_to_date"]) for r in d30 if _d(r["as_of_date"]) == d30d) if d30d else None
        paid_dates = lambda types: [_d(r["paid_date"]) for r in inv_rows if _s(r["invoice_type"]) in types and _d(r["paid_date"])]
        deposit_rows = [r for r in inv_rows if _s(r["invoice_type"]) == "standard_deposit" and _d(r["paid_date"])]
        depleted_all = [o["units_depleted_day5"] for o in ondrop]
        camp = {
            "campaign": c, "offices": offices, "drop_date": T0, "window_end": win_end,
            "offices_on_route": len(offices), "offices_approved": len(approved), "offices_on_drop": len(ondrop),
            "turned_away": [o for o in offices if o["turned_away"]],
            "units_allocated": sum(o["units_allocated"] or 0 for o in approved),
            "units_received": sum(_n(r["qty"]) for r in inv if _s(r["movement"]) == "received"),
            "units_rejected": sum(_n(r["qty"]) for r in inv if _s(r["movement"]) == "rejected"),
            "units_delivered": tot("units_delivered"), "units_placed": tot("units_placed"),
            "units_placed_within_1bd": sum(o["units_placed"] for o in offices if o["placed_within_1bd"] == "Y"),
            "units_topped_up": tot("units_topped_up"),
            "units_depleted_day5": None if any(x is None for x in depleted_all) else sum(depleted_all),
            "survey_responses": tot("survey_responses"),
            "survey_responses_min_office": min((o["survey_responses"] for o in ondrop), default=None),
            "survey_tried_yes": tot("tried_yes"), "tried_no_didnt_see": tot("tried_no_didnt_see"),
            "tried_no_not_for_me": tot("tried_no_not_for_me"), "tried_no_gone": tot("tried_no_gone"),
            "survey_rating_n": rn, "survey_rating_avg": xround(tot("rating_sum") / rn, 1) if rn else None,
            "survey_buy_yes": tot("buy_yes"), "survey_buy_maybe": tot("buy_maybe"), "survey_buy_no": tot("buy_no"),
            "comments_total": tot("comments"),
            "responses_in_window": len(inwin), "responses_total": len(resp),
            "responses_outside_window": None if deleted else len(resp) - len(inwin),
            "responses_unmatched_code": None if deleted else len(inwin) - tot("survey_responses"),
            "qr_scans": sum(o["qr_scans"] or 0 for o in offices), "post_link_opens": sum(o["post_link_opens"] or 0 for o in offices),
            "code_redemptions_to_date": red_to_date, "code_redemptions_as_of": asof, "code_redemptions_day30": red30,
            "revenue_usd": round(sum(_n(r["amount_usd"]) for r in inv_rows), 2),
            "deposit_usd": round(sum(_n(r["paid_amount_usd"]) for r in deposit_rows), 2) if deposit_rows else None,
            "deposit_paid_date": max((_d(r["paid_date"]) for r in deposit_rows), default=None),
            "invoice_paid_date": max(paid_dates(("founding_full", "standard_balance")), default=None),
            "processing_fee_usd": round(sum(_n(r["processing_fee_usd"]) for r in inv_rows), 2),
            "reserve_on_hand": sum(reserve_delta(r) for r in inv), "leftover_on_hand": sum(leftover_delta(r) for r in inv),
        }
        camp["placed_share_1bd"] = (camp["units_placed_within_1bd"] / camp["units_delivered"]) if camp["units_delivered"] else None
        camp["photo_compliance"] = (sum(1 for o in ondrop if o["both_photos"] == "Y") / len(ondrop)) if ondrop else None
        camp.update(self.money(cid, c, camp))
        return camp

    def money(self, cid, c, camp):
        did = _s(c["drop_id"])
        same_drop = [x for x in self.campaigns if _s(self.campaigns[x]["drop_id"]) == did]
        approved_by = {x: sum(1 for a in self.approvals_for(x) if self._approved(x, a)) for x in same_drop}
        tot_off = sum(approved_by.values())
        share = approved_by.get(cid, 0) / tot_off if tot_off else 0
        costs = [r for r in self.rows["costs"] if _s(r["drop_id"]) == did]
        direct = sum(_n(r["amount_usd"]) for r in costs if _s(r["campaign_id"]) == cid)
        shared_pool = sum(_n(r["amount_usd"]) for r in costs if _s(r["campaign_id"]) == "")
        hrs = [r for r in self.rows["hours"] if _s(r["drop_id"]) == did]
        h_direct = sum(_n(r["hours"]) for r in self.rows["hours"] if _s(r["campaign_id"]) == cid)
        h_pool = sum(_n(r["hours"]) for r in hrs if _s(r["campaign_id"]) == "")
        shared = xround(share * shared_pool, 2)
        h_shared = xround(share * h_pool, 2)
        contribution = camp["revenue_usd"] - camp["processing_fee_usd"] - direct - shared
        hours = h_direct + h_shared
        return {"direct_cost_usd": round(direct, 2), "shared_cost_usd": shared, "hours_direct": h_direct, "hours_shared": h_shared,
                "contribution_usd": xround(contribution, 2), "contribution_per_hour_usd": xround(contribution / hours, 2) if hours else None,
                "cost_share": share}

    def brand_qualified(self, bid):
        """The brands tab's qualified column, recomputed: all seven checks a plain yes and not an excluded product."""
        if bid not in self.brands:
            return False, "brand_id not on brands"
        b = self.brand(bid)
        qs = ["q1_sealed_shelf_stable_60d", "q2_labeled_ingredients_allergens", "q3_buy_path_near_cluster",
              "q4_can_ship_275_in_10bd", "q5_product_liability_insurance", "q6_contact_can_approve_600"]
        fails = [str(i + 1) for i, q in enumerate(qs) if _s(b.get(q)) != "Y"]
        if not _s(b.get("q7_why_now")).strip():
            fails.append("7")
        ok = not fails and _s(b.get("excluded_category_confirmed")) == "Y"
        return ok, ("Y (as answered)" if ok else "N: check " + " ".join(fails) + (" excluded?" if _s(b.get("excluded_category_confirmed")) != "Y" else ""))

    def release_status(self, cid, lad=None):
        """The sheet's ship_release rule, recomputed: RELEASED, or HOLD with the reasons."""
        lad = lad or self.ladder(cid)
        c = lad["campaign"]
        paid = sum(_n(r["paid_amount_usd"]) for r in self.rows["invoices"] if _s(r["campaign_id"]) == cid
                   and _s(r["invoice_type"]) in ("founding_full", "standard_deposit") and _d(r.get("paid_date")))
        price = c["price_usd"] or 0
        due = price if _s(c["tier"]) == "founding" else xround(price * _n(self.config["deposit_share"]), 2)
        b = self.brands.get(_s(c["brand_id"]), {})
        why = []
        if not due or paid < due:
            why.append("payment (%s of %s)" % (int(paid), int(due or 0)))
        if not _d(c.get("agreement_signed_date")):
            why.append("agreement")
        if not (_d(c.get("coi_received_date")) and _s(b.get("coi_check_ok")) == "Y"):
            why.append("certificate")
        if not self.brand_qualified(_s(c["brand_id"]))[0]:
            why.append("qualification")
        if lad["drop_date"] is None:
            why.append("drop date")
        pending = sum(1 for o in lad["offices"] if o["approved"] == "pending")
        if pending:
            why.append("approvals (%d pending)" % pending)
        if lad["offices_approved"] == 0:
            why.append("no office approved")
        return "RELEASED" if not why else "HOLD: " + " ".join(why)

    def _approved(self, cid, a):
        c = self.campaigns[cid]
        T0 = self.drop_date(_s(c["drop_id"]))
        n_drop = self.office_drop_number(_s(a["office_id"]), T0)
        rep = _s(a.get("reply"))
        if n_drop <= int(_n(self.config["approvals_explicit_drops"])):
            return rep == "yes"
        ask = _d(a.get("asked_date"))
        return rep != "no" and ask is not None and T0 is not None and ask <= opscal.add_bd(T0, -5)

    def drop_money(self, did):
        camps = [x for x in self.campaigns if _s(self.campaigns[x]["drop_id"]) == did]
        rev = sum(self.ladder(x)["revenue_usd"] for x in camps)
        fee = sum(self.ladder(x)["processing_fee_usd"] for x in camps)
        costs = [r for r in self.rows["costs"] if _s(r["drop_id"]) == did]
        by = lambda t: round(sum(_n(r["amount_usd"]) for r in costs if _s(r["cost_type"]) == t), 2)
        hours = sum(_n(r["hours"]) for r in self.rows["hours"] if _s(r["drop_id"]) == did)
        stops = [s for s in self.rows["stops"] if _s(s["drop_id"]) == did]
        direct = by("kit") + by("transport") + by("other_direct")
        contrib = rev - fee - direct
        lads = [self.ladder(x) for x in camps]
        dl = sum(l["units_delivered"] for l in lads)
        resp = [o["survey_responses"] for l in lads for o in l["offices"] if o["on_drop"] == "Y"]
        g2 = {"gate2_placed_share": (sum(l["units_placed_within_1bd"] for l in lads) / dl) if dl else None,
              "gate2_min_responses": min(resp) if resp else None,
              "gate2_buy_again_yes": sum(1 for x in camps if _s(self.campaigns[x].get("brand_would_buy_again")) == "yes")}
        return {**g2, "revenue_usd": rev, "kit_cost_usd": by("kit"), "transport_cost_usd": by("transport"),
                "other_direct_cost_usd": by("other_direct"), "direct_cost_usd": round(direct, 2), "stops": len(stops),
                "cost_per_stop_usd": xround(direct / len(stops), 2) if stops else None, "hours": hours,
                "contribution_usd": xround(contrib, 2), "contribution_per_hour_usd": xround(contrib / hours, 2) if hours else None,
                "campaigns": camps}


# ------------------------------------------------------------ shared rules (same as the formulas)
def band(h):
    if h is None:
        return "unknown"
    return "under-40" if h < 40 else "40-75" if h <= 75 else "76-150" if h <= 150 else "151-250" if h <= 250 else "251-500" if h <= 500 else "over-500"


def mround(x, m):
    return m * round(x / m + 1e-9) if x >= 0 else -m * round(-x / m + 1e-9)


def default_units(h, cfg):
    if h is None:
        return None
    raw = _n(cfg["units_base"]) + _n(cfg["units_per_person"]) * h
    # MROUND rounds half away from zero.
    m = _n(cfg["units_round"])
    q = raw / m
    r = int(q + 0.5) if q >= 0 else -int(-q + 0.5)
    return int(min(_n(cfg["units_max"]), max(_n(cfg["units_min"]), r * m)))


def expected_price(c, cfg):
    n = c.get("offices_contracted")
    if n in (None, ""):
        return None
    n = _n(n)
    tier, lock = _s(c.get("tier")), _s(c.get("founding_lock"))
    ex = _n(cfg["extra_office_usd"])
    if tier == "founding":
        if lock == "Y":
            lm = _n(cfg["founding_lock_max_offices"])
            return _n(cfg["founding_lock_per_office_usd"]) * min(n, lm) + ex * max(0, n - lm)
        return _n(cfg["founding_price_usd"]) + ex * max(0, n - _n(cfg["founding_offices"]))
    if tier == "standard":
        return _n(cfg["standard_price_usd"]) + ex * max(0, n - _n(cfg["standard_offices"]))
    return None


def slug(cid, oid, kind):
    pre = "SAMPLE-" if cid.startswith("SAMPLE-") else ""
    c = cid.replace("SAMPLE-", "").replace("CMP-", "c").lower()
    o = oid.replace("SAMPLE-", "").replace("O-", "o").lower()
    return "%s%s-%s-%s" % (pre, c, o, kind)


PILE = ("returned_to_brand", "given_to_roster_office", "released_to_office", "disposed")
REFUSED = ("declined", "building_refused", "not_delivered")   # stop outcomes that leave an office off the drop
FROZEN_COLS = ("survey_responses", "tried_yes", "tried_no_didnt_see", "tried_no_not_for_me", "tried_no_gone", "rating_n",
               "rating_sum", "buy_yes", "buy_maybe", "buy_no", "comments")


def reserve_delta(r):
    m, q, b = _s(r["movement"]), _n(r["qty"]), _s(r.get("bucket")) or "reserve"
    if m == "received":
        return q
    if m in ("rejected", "delivered", "topped_up"):
        return -q
    if m in PILE and b == "reserve":
        return -q
    if m == "correction" and b == "reserve":
        return q
    return 0


def leftover_delta(r):
    m, q, b = _s(r["movement"]), _n(r["qty"]), _s(r.get("bucket")) or "reserve"
    if m == "collected":
        return q
    if m in PILE and b == "leftover":
        return -q
    if m == "correction" and b == "leftover":
        return q
    return 0


# ------------------------------------------------------------ comparison with the sheet's formulas
APPROVAL_COMPARE = ["approved", "on_drop", "survey_frozen", "units_allocated", "units_delivered", "units_placed", "placed_within_1bd", "units_topped_up",
                    "units_depleted_day5", "survey_responses", "tried_yes", "rating_n", "buy_yes", "buy_maybe", "buy_no",
                    "comments", "qr_scans", "post_link_opens", "both_photos", "later_photo_day", "office_label",
                    "placement_photo_file", "later_photo_file", "placement_lag_hours"]
CAMPAIGN_COMPARE = ["units_received", "units_delivered", "units_placed", "units_placed_within_1bd", "units_topped_up",
                    "units_depleted_day5", "survey_responses", "survey_responses_min_office", "survey_tried_yes",
                    "survey_rating_n", "survey_rating_avg", "survey_buy_yes", "survey_buy_maybe", "survey_buy_no",
                    "comments_total", "qr_scans", "post_link_opens", "code_redemptions_to_date", "code_redemptions_day30",
                    "revenue_usd", "deposit_usd", "processing_fee_usd", "direct_cost_usd", "shared_cost_usd",
                    "contribution_usd", "contribution_per_hour_usd", "reserve_on_hand", "leftover_on_hand", "photo_compliance",
                    "responses_outside_window", "responses_unmatched_code", "offices_approved", "offices_on_drop", "units_allocated"]


def _eq(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, (dt.date, dt.datetime)) or isinstance(b, (dt.date, dt.datetime)):
        return _d(a) == _d(b)
    try:
        return abs(float(a) - float(b)) < 0.006
    except (TypeError, ValueError):
        return _s(a) == _s(b)


def compare(tr):
    """List every disagreement between this module's numbers and the sheet's cached formula values."""
    bad = []
    if not any(any(v not in (None, "") for v in r.values()) for r in tr.cached.get("campaigns", [])):
        return ["no cached formula values: recalculate first (tools/recalc.py) or download from Google Sheets"]
    cam_cache = {(_s(r["campaign_id"])): cc for r, cc in zip(tr.rows["campaigns"], tr.cached["campaigns"])}
    ap_cache = [(r, cc) for r, cc in zip(tr.rows["approvals"], tr.cached["approvals"])]
    for cid in tr.campaigns:
        lad = tr.ladder(cid)
        cc = cam_cache[cid]
        for k in CAMPAIGN_COMPARE:
            if not _eq(lad.get(k), cc.get(k)):
                bad.append("campaigns %s %s: python=%r sheet=%r" % (cid, k, lad.get(k), cc.get(k)))
        # The ship release: the same verdict both ways (the kit and the ship-date email use the Python one).
        py_rel, sh_rel = tr.release_status(cid, lad), _s(cc.get("ship_release"))
        if py_rel.startswith("RELEASED") != sh_rel.startswith("RELEASED"):
            bad.append("campaigns %s ship_release: python=%r sheet=%r" % (cid, py_rel, sh_rel))
        byo = {o["office_id"]: o for o in lad["offices"]}
        for r, acc in ap_cache:
            if _s(r["campaign_id"]) != cid:
                continue
            o = byo[_s(r["office_id"])]
            for k in APPROVAL_COMPARE:
                if not _eq(o.get(k), acc.get(k)):
                    bad.append("approvals %s/%s %s: python=%r sheet=%r" % (cid, o["office_id"], k, o.get(k), acc.get(k)))
    dcache = {_s(r["drop_id"]): cc for r, cc in zip(tr.rows["drops"], tr.cached["drops"])}
    for did in tr.drops:
        m = tr.drop_money(did)
        for k in ("revenue_usd", "kit_cost_usd", "transport_cost_usd", "other_direct_cost_usd", "direct_cost_usd",
                  "cost_per_stop_usd", "hours", "contribution_usd", "contribution_per_hour_usd", "gate2_placed_share",
                  "gate2_min_responses", "gate2_buy_again_yes"):
            if not _eq(m.get(k), dcache[did].get(k)):
                bad.append("drops %s %s: python=%r sheet=%r" % (did, k, m.get(k), dcache[did].get(k)))
    return bad
