#!/usr/bin/env python3
"""Every form of a SAMPLE dry run, as the CSV rows each form produced (dry runs only).

  python3 tools/dump_forms.py RECALCULATED_TRACKER.xlsx SCENARIO OUT_DIR

Writes OUT_DIR/ (one header row each, UTF-8):
  office_intake-website.csv      rows the website's office form wrote (CONTRACTS 6.3 columns)
  office_form-google.csv         the founder-side Google Form's own response rows (its question titles)
  office_intake_founder.csv      the same answers as the tracker maps them to the 6.3 columns
  brand_intake-website.csv, brand_form-google.csv, brand_intake_founder.csv   the same for brands
  survey-<campaign_id>.csv       each campaign's survey responses (Google's columns A to G)
  SAMPLE-newsletter-signup-<campaign_id>.csv   the sign-up form's own responses, for brands that chose the
                                 newsletter. In real use these never leave Google Forms (forms/survey-form.md,
                                 Part 3); the dry run writes them here only to show the T+5 pass-on and deletion.
  SAMPLE-newsletter-pass-on-<campaign_id>.txt   exactly what template 09 sends: ticked addresses, sorted, nothing else.
Refuses anything that is not SAMPLE data.
"""
import csv
import datetime as dt
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common  # noqa: E402
import opscal  # noqa: E402
import trackerdef as T  # noqa: E402


def val(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if (v.hour, v.minute, v.second) != (0, 0, 0) else v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def tab_rows(ws):
    hdr = []
    for c in ws[1]:
        if c.value in (None, ""):
            break   # a note may sit after a blank column (the founder-side form tabs): not a column
        hdr.append(c.value)
    rows = []
    for r in ws.iter_rows(min_row=2, max_col=len(hdr), values_only=True):
        if any(v not in (None, "") for v in r):
            rows.append([val(v) for v in r])
    return hdr, rows


def write(path, hdr, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(hdr)
        w.writerows(rows)
    return len(rows)


def main(argv):
    if len(argv) < 4:
        print(__doc__)
        return 1
    path, name, out = argv[1], argv[2], argv[3]
    wb = openpyxl.load_workbook(path, data_only=True)
    T.define()
    cids = [r[0] for r in wb["campaigns"].iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
    if not cids or any(not str(c).startswith("SAMPLE") for c in cids):
        raise SystemExit("REFUSED: dump_forms.py is for SAMPLE dry runs only.")
    os.makedirs(out, exist_ok=True)
    made = []
    for tab, fn in (("office_intake", "office_intake-website.csv"), ("office_form", "office_form-google.csv"),
                    ("office_intake_founder", "office_intake_founder.csv"), ("brand_intake", "brand_intake-website.csv"),
                    ("brand_form", "brand_form-google.csv"), ("brand_intake_founder", "brand_intake_founder.csv")):
        hdr, rows = tab_rows(wb[tab])
        made.append((fn, write(os.path.join(out, fn), hdr, rows)))
    for sn in wb.sheetnames:
        if sn.startswith("SAMPLE-CMP-"):
            hdr, rows = tab_rows(wb[sn])
            made.append(("survey-%s.csv" % sn, write(os.path.join(out, "survey-%s.csv" % sn), hdr, rows)))
    # Newsletter sign-up: SAMPLE responses consistent with the campaign row (ticked = newsletter_passed_count, plus
    # two unticked and one with no address), so the pass-on and both deletions can be followed in the dry run.
    camp = T.TABS["campaigns"]
    ws = wb["campaigns"]
    idx = {c.name: camp.index[c.name] + 1 for c in camp.cols}
    for r in range(camp.first(), camp.last() + 1):
        cid = ws.cell(row=r, column=1).value
        if not cid or ws.cell(row=r, column=idx["newsletter_option"]).value != "Y":
            continue
        if not str(cid).startswith("SAMPLE"):
            raise SystemExit("REFUSED: not SAMPLE data")
        n = int(ws.cell(row=r, column=idx["newsletter_passed_count"]).value or 0)
        brand = str(ws.cell(row=r, column=idx["brand_name"]).value or "[BRAND]")
        drops = T.TABS["drops"]
        dws = wb["drops"]
        did = ws.cell(row=r, column=idx["drop_id"]).value
        dcol = drops.index["drop_date"] + 1
        tday = next(dws.cell(row=k, column=dcol).value for k in range(drops.first(), drops.last() + 1) if dws.cell(row=k, column=1).value == did)
        tday = tday.date() if isinstance(tday, dt.datetime) else tday
        consent = "Share my email with %s under its privacy policy." % brand
        rows = []
        tag = str(cid).lower().replace("sample-", "")
        for i in range(n):
            ts = dt.datetime.combine(opscal.add_bd(tday, i % 4), dt.time(10 + i % 6, 7 * i % 60))
            rows.append([val(ts), "sample-reader-%s-%02d@example.com" % (tag, n - i), consent])
        for j, (email, ok) in enumerate((("sample-reader-%s-90@example.com" % tag, ""), ("sample-reader-%s-91@example.com" % tag, ""), ("", ""))):
            ts = dt.datetime.combine(opscal.add_bd(tday, 1 + j), dt.time(14, 10 * j))
            rows.append([val(ts), email, ok])
        rows.sort(key=lambda x: x[0])
        fn = "SAMPLE-newsletter-signup-%s.csv" % cid
        made.append((fn, write(os.path.join(out, fn), ["Timestamp", "Email (optional)", "Consent (optional)"], rows)))
        ticked = sorted(x[1] for x in rows if x[1] and x[2])
        passed = ws.cell(row=r, column=idx["newsletter_passed_date"]).value
        deleted = ws.cell(row=r, column=idx["newsletter_deleted_date"]).value
        un = ws.cell(row=r, column=idx["newsletter_unticked_deleted_date"]).value
        txt = ("SAMPLE. What template 09 sends to %s on %s: the ticked addresses only, sorted, nothing else.\n\n%s\n\n"
               "Unticked responses deleted in the form on %s (%d of them). Every response, the form and the sent email deleted on %s.\n"
               "The tracker holds only the dates and the count (%d).\n") % (
            brand, val(passed) or "[T+5]", "\n".join(ticked), val(un) or "[T+5]", len(rows) - len(ticked), val(deleted) or "[T+5 + 30 days]", n)
        fn2 = "SAMPLE-newsletter-pass-on-%s.txt" % cid
        common.write(os.path.join(out, fn2), txt)
        made.append((fn2, len(ticked)))
    print("forms: %s -> %s" % (", ".join("%s %d" % m for m in made), out))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
