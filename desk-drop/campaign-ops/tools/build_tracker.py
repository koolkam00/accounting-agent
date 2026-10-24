#!/usr/bin/env python3
"""Build the campaign tracker workbook (Google-Sheets-compatible .xlsx, formulas intact, no macros).

  python3 tools/build_tracker.py                 # writes tracker/campaign-tracker.xlsx (blank, ready to import)
  python3 tools/build_tracker.py --out X.xlsx    # somewhere else

The dry runs import this module and call build(sample=...) to add SAMPLE rows and
Google-Forms-shaped response tabs, then recalculate with LibreOffice.

The company name is never written by this script: it is read from
desk-drop-identity/brand.json (BRAND_NAME) into the one config cell, company_name.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import formulas as F  # noqa: E402
import opscal  # noqa: E402
import trackerdef as T  # noqa: E402

ROOT = T.ROOT
DESK = os.path.dirname(ROOT)
BRAND_JSON = os.path.join(DESK, "desk-drop-identity", "brand.json")
D_REPORT_JS = os.path.join(DESK, "desk-drop-identity", "system", "dd-report.js")

GREY = PatternFill("solid", fgColor="EFEFEF")
CREAM = PatternFill("solid", fgColor="FFFBEA")
HEAD_I = PatternFill("solid", fgColor="1F3864")
HEAD_A = PatternFill("solid", fgColor="7F7F7F")
HEAD_P = PatternFill("solid", fgColor="BF8F00")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
RED = Font(color="C00000", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")

ORDER = ["README", "dashboard", "survey_dashboard", "report", "config", "calendar", "drops", "campaigns", "brands",
         "offices", "approvals", "stops", "inventory", "counts", "photos", "links", "scans", "redemptions", "issues",
         "hours", "costs", "invoices", "office_intake", "brand_intake", "office_form", "brand_form",
         "office_intake_founder", "brand_intake_founder", "in_signed_offices", "in_paid_brands", "responses_template"]


def company_name():
    return json.load(open(BRAND_JSON))["BRAND_NAME"]


def limitations_text():
    """D's fixed limitations paragraph, read from desk-drop-identity/system/dd-report.js so it cannot drift."""
    js = open(D_REPORT_JS, encoding="utf-8").read()
    m = re.search(r"R\.LIMITATIONS\s*=\s*((?:'(?:[^'\\]|\\.)*'\s*\+?\s*)+);", js)
    parts = re.findall(r"'((?:[^'\\]|\\.)*)'", m.group(1))
    return "".join(parts).replace("\\'", "'")


def survey_headers(brand="[BRAND]"):
    s = json.load(open(os.path.join(ROOT, "spec", "forms.json")))["forms"]["survey"]["response_columns"]
    return [h.replace("{BRAND}", brand) for h in s]


def form_headers(key):
    form = json.load(open(os.path.join(ROOT, "spec", "forms.json")))["forms"][key]
    return ["Timestamp"] + [it["title"] for it in form["items"]]


def write_tab(wb, tab):
    ws = wb.create_sheet(tab.name)
    ws.sheet_properties.tabColor = tab.color
    h = tab.header_row
    for i, col in enumerate(tab.cols, start=1):
        cell = ws.cell(row=h, column=i, value=col.name)
        cell.font = WHITE_BOLD
        cell.fill = {"I": HEAD_I, "A": HEAD_A, "P": HEAD_P}[col.kind]
        cell.alignment = Alignment(vertical="center")
        note = col.help or {"A": "Formula. Never type over it.", "P": "Pasted or written by a form or script."}.get(col.kind, "")
        if note:
            cell.comment = Comment(note, "tracker")
        ws.column_dimensions[get_column_letter(i)].width = col.width
    ws.freeze_panes = tab.freeze
    first, last = tab.first(), tab.last()
    for i, col in enumerate(tab.cols, start=1):
        L = get_column_letter(i)
        if col.kind == "A":
            for r in range(first, last + 1):
                c = ws.cell(row=r, column=i, value="=" + col.f.replace("{r}", str(r)))
                c.fill = GREY
                if col.fmt:
                    c.number_format = col.fmt
        else:
            if col.fmt:
                for r in range(first, last + 1):
                    ws.cell(row=r, column=i).number_format = col.fmt
            if col.name in ("intake_key",):
                for r in range(first, last + 1):
                    ws.cell(row=r, column=i).number_format = "@"
            if tab.name in ("office_intake", "brand_intake"):
                pass
        if col.dv:
            kind, f1 = col.dv
            dv = DataValidation(type=kind, formula1=f1, allow_blank=True, showErrorMessage=True,
                                errorTitle="Not on the list", error="Pick a value from the list.")
            ws.add_data_validation(dv)
            dv.add("%s%d:%s%d" % (L, first, L, last))
    if "checks" in tab.index:
        L = tab.letter("checks")
        ws.conditional_formatting.add("%s%d:%s%d" % (L, first, L, last),
                                      FormulaRule(formula=['LEN(%s%d)>0' % (L, first)], font=Font(color="C00000")))
    return ws


def drop_date_validation(ws):
    t = T.TABS["drops"]
    L = t.letter("drop_date")
    f = ('AND(OR(WEEKDAY({L}2)=3,WEEKDAY({L}2)=4),COUNTIF(INDIRECT("calendar!$A$2:$A$80"),{L}2)=0,'
         'COUNTIFS(INDIRECT("calendar!$D$2:$D$20"),"<="&WORKDAY({L}2,4,INDIRECT("calendar!$A$2:$A$80")),INDIRECT("calendar!$E$2:$E$20"),">="&{L}2)=0)').format(L=L)
    dv = DataValidation(type="custom", formula1=f, allow_blank=True, showErrorMessage=True, errorStyle="stop",
                        errorTitle="Not a valid drop day",
                        error="Drop days are Tuesday or Wednesday business days, and the five-day window must be clear of 18-31 Dec, 18-31 Aug and the 4 July week. Run tools/opscal.py next to pick one.")
    ws.add_data_validation(dv)
    dv.add("%s%d:%s%d" % (L, t.first(), L, t.last()))
    ws.conditional_formatting.add("%s%d:%s%d" % (t.letter("date_check"), t.first(), t.letter("date_check"), t.last()),
                                  FormulaRule(formula=['LEFT(%s%d,7)="REFUSED"' % (t.letter("date_check"), t.first())], font=RED))


def write_config(wb):
    ws = wb.create_sheet("config")
    ws.sheet_properties.tabColor = "DDEBF7"
    for c, w in zip("ABC", (28, 44, 90)):
        ws.column_dimensions[c].width = w
    for i, h in enumerate(("key", "value", "note"), start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font, cell.fill = WHITE_BOLD, HEAD_I
    for k, v, note in T.CONFIG:
        r = T.CFG_ROW[k]
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=company_name() if k == "company_name" else v)
        ws.cell(row=r, column=3, value=note).alignment = WRAP
        if k == "today_override":
            ws.cell(row=r, column=2).number_format = T.DATE
    ws.cell(row=T.CFG_ROW["company_name"], column=2).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "A2"


def write_calendar(wb):
    ws = wb.create_sheet("calendar")
    ws.sheet_properties.tabColor = "DDEBF7"
    heads = ["holiday", "name", "", "blackout_start", "blackout_end", "reason"]
    for i, h in enumerate(heads, start=1):
        if h:
            c = ws.cell(row=1, column=i, value=h)
            c.font, c.fill = WHITE_BOLD, HEAD_A
    r = 2
    for y in range(2026, 2031):
        for d, name in sorted(opscal.holidays(y).items()):
            ws.cell(row=r, column=1, value=d).number_format = T.DATE
            ws.cell(row=r, column=2, value=name)
            r += 1
    r = 2
    for y in range(2026, 2031):
        for b0, b1, why in opscal.blackouts(y):
            ws.cell(row=r, column=4, value=b0).number_format = T.DATE
            ws.cell(row=r, column=5, value=b1).number_format = T.DATE
            ws.cell(row=r, column=6, value=why)
            r += 1
    ws.cell(row=1, column=8, value="Source: CONTRACTS section 3 (derived, founder to confirm). US federal holidays by rule, observed dates; written by tools/opscal.py. Do not edit: rebuild with tools/build_tracker.py, or add a row by hand only to add a new year.").alignment = WRAP
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 34
    ws.column_dimensions["H"].width = 60


def write_form_tab(wb, name, headers, rows=0, note=""):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "FFF9E5"
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = WHITE_BOLD, HEAD_P
        ws.column_dimensions[get_column_letter(i)].width = 18 if i > 1 else 20
    if note:
        ws.cell(row=1, column=len(headers) + 2, value=note).alignment = WRAP
        ws.column_dimensions[get_column_letter(len(headers) + 2)].width = 70
    ws.freeze_panes = "B2"
    return ws


def mirror(tabname, colname, k):
    """INDEX of the k-th data row of a column (for dashboard mirrors)."""
    return 'INDEX(%s,%d)' % (T.R(tabname, colname), k)


def blank_if(cond_ref, x):
    return '=IF(%s="","",%s)' % (cond_ref, x)


def write_dashboard(wb):
    ws = wb.create_sheet("dashboard")
    ws.sheet_properties.tabColor = "C00000"
    ws["A1"] = '="Campaign tracker · "&%s' % T.cfg("company_name")
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = ('=IF(COUNTIF(%s,"SAMPLE*")+COUNTIF(%s,"SAMPLE*")>0,"SAMPLE DATA. Illustrative placeholders, not results.","")'
                % (T.R("campaigns", "campaign_id"), T.R("drops", "drop_id")))
    ws["A2"].font = Font(bold=True, size=14, color="C00000")
    ws["A3"] = "Today"
    ws["B3"] = "=" + T.TODAY
    ws["B3"].number_format = "ddd yyyy-mm-dd"
    ws["A3"].font = BOLD
    # Checks
    r = 5
    ws.cell(row=r, column=1, value="Checks to fix (0 everywhere = nothing to fix; lines starting info: are facts, not counted)").font = Font(bold=True, size=12)
    checks = [
        ("Drops with a refused date", 'COUNTIF(%s,"REFUSED*")' % T.R("drops", "date_check"), "drops, date_check"),
        ("Campaigns with checks to read", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("campaigns", "checks"), T.R("campaigns", "checks")), "campaigns, checks"),
        ("Campaigns on ship HOLD", 'COUNTIF(%s,"HOLD*")' % T.R("campaigns", "ship_release"), "campaigns, ship_release"),
        ("Office rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("approvals", "checks"), T.R("approvals", "checks")), "approvals, checks"),
        ("Stops with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("stops", "checks"), T.R("stops", "checks")), "stops, checks"),
        ("Inventory rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("inventory", "checks"), T.R("inventory", "checks")), "inventory, checks"),
        ("Count rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("counts", "checks"), T.R("counts", "checks")), "counts, checks"),
        ("Photo rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("photos", "checks"), T.R("photos", "checks")), "photos, checks"),
        ("Link rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("links", "checks"), T.R("links", "checks")), "links, checks"),
        ("Scan readings with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("scans", "checks"), T.R("scans", "checks")), "scans, checks"),
        ("Issues open", 'SUMPRODUCT((LEN(%s)>0)*(LEN(%s)=0))' % (T.R("issues", "issue_id"), T.R("issues", "closed_date")), "issues"),
        ("Issue rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("issues", "checks"), T.R("issues", "checks")), "issues, checks"),
        ("Hours rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("hours", "checks"), T.R("hours", "checks")), "hours, checks"),
        ("Cost rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("costs", "checks"), T.R("costs", "checks")), "costs, checks"),
        ("Invoices overdue", 'COUNTIF(%s,"overdue")' % T.R("invoices", "status"), "invoices, status"),
        ("Invoice rows with checks", 'COUNTIF(%s,"?*")-COUNTIF(%s,"info:*")' % (T.R("invoices", "checks"), T.R("invoices", "checks")), "invoices, checks"),
        ("Brands in E's paid_brands.csv not on brands", 'SUMPRODUCT((%s<>"")*(COUNTIF(%s,%s)=0))' % (
            "in_paid_brands!$A$2:$A$%d" % (T.TABS["in_paid_brands"].last()), T.R("brands", "brand_id"),
            "INDEX(in_paid_brands!$A$2:$BZ$%d,0,MATCH(\"brand_id\",in_paid_brands!$A$1:$BZ$1,0))" % T.TABS["in_paid_brands"].last()), "in_paid_brands"),
        ("Offices in E's signed_offices.csv not on offices", 'SUMPRODUCT((%s<>"")*(COUNTIF(%s,%s)=0))' % (
            "in_signed_offices!$A$2:$A$%d" % (T.TABS["in_signed_offices"].last()), T.R("offices", "office_id"),
            "INDEX(in_signed_offices!$A$2:$BZ$%d,0,MATCH(\"office_id\",in_signed_offices!$A$1:$BZ$1,0))" % T.TABS["in_signed_offices"].last()), "in_signed_offices"),
    ]
    ws.cell(row=r + 1, column=1, value="check").font = BOLD
    ws.cell(row=r + 1, column=2, value="count").font = BOLD
    ws.cell(row=r + 1, column=3, value="where to look").font = BOLD
    for i, (label, f, where) in enumerate(checks):
        rr = r + 2 + i
        ws.cell(row=rr, column=1, value=label)
        c = ws.cell(row=rr, column=2, value="=IFERROR(%s,0)" % f)
        ws.cell(row=rr, column=3, value=where)
        ws.conditional_formatting.add("B%d" % rr, FormulaRule(formula=["B%d>0" % rr], font=RED))
    r = r + 3 + len(checks)
    # Drops
    ws.cell(row=r, column=1, value="Drops").font = Font(bold=True, size=12)
    dcols = [("drop_id", "drop"), ("drop_date", "date"), ("delivery_model", "model"), ("status", "status"),
             ("date_check", "date check"), ("next_step", "next step"), ("next_step_date", "on"), ("stops", "stops"),
             ("campaigns", "brands"), ("units_allocated", "units"), ("revenue_usd", "revenue"), ("direct_cost_usd", "direct cost"),
             ("cost_per_stop_usd", "cost per stop"), ("hours", "founder hours"), ("contribution_usd", "contribution"),
             ("contribution_per_hour_usd", "per founder hour")]
    for j, (cn, lab) in enumerate(dcols, start=1):
        ws.cell(row=r + 1, column=j, value=lab).font = BOLD
    for k in range(1, 16):
        rr = r + 1 + k
        for j, (cn, lab) in enumerate(dcols, start=1):
            x = mirror("drops", cn, k)
            cell = ws.cell(row=rr, column=j, value='=IF(%s="","",%s)' % (mirror("drops", "drop_id", k), x))
            col = T.TABS["drops"].cols[T.TABS["drops"].index[cn]]
            if col.fmt:
                cell.number_format = col.fmt
    r = r + 18
    ws.cell(row=r, column=1, value="Campaigns (every rung kept separate)").font = Font(bold=True, size=12)
    ccols = [("campaign_id", "campaign"), ("brand_name", "brand"), ("drop_date", "drop"), ("tier", "tier"),
             ("ship_release", "ship release"), ("offices_approved", "offices approved"), ("offices_contracted", "contracted"),
             ("units_delivered", "delivered"), ("units_placed_within_1bd", "placed within 1 bd"), ("placed_share_1bd", "share"),
             ("units_topped_up", "topped up"), ("units_depleted_day5", "depleted by day 5"), ("survey_responses", "responded"),
             ("survey_responses_min_office", "fewest at one office"), ("survey_tried_yes", "said they tried"),
             ("survey_rating_avg", "avg rating"), ("qr_scans", "QR scans"), ("code_redemptions_to_date", "code redemptions"),
             ("revenue_usd", "revenue"), ("contribution_per_hour_usd", "per founder hour"), ("gate2_placed", "G2 placed"),
             ("gate2_responses", "G2 responses"), ("gate2_report_on_time", "G2 report"), ("brand_would_buy_again", "buy again"),
             ("checks", "checks")]
    for j, (cn, lab) in enumerate(ccols, start=1):
        ws.cell(row=r + 1, column=j, value=lab).font = BOLD
    for k in range(1, 31):
        rr = r + 1 + k
        for j, (cn, lab) in enumerate(ccols, start=1):
            x = mirror("campaigns", cn, k)
            cell = ws.cell(row=rr, column=j, value='=IF(%s="","",%s)' % (mirror("campaigns", "campaign_id", k), x))
            col = T.TABS["campaigns"].cols[T.TABS["campaigns"].index[cn]]
            if col.fmt:
                cell.number_format = col.fmt
    ws.cell(row=r + 33, column=1, value="Incremental purchases are not measured. No column here adds the rungs together.").font = Font(italic=True)
    r = r + 35
    ws.cell(row=r, column=1, value="Founder hours by activity").font = Font(bold=True, size=12)
    ws.cell(row=r + 1, column=1, value="activity").font = BOLD
    ws.cell(row=r + 1, column=2, value="all hours").font = BOLD
    ws.cell(row=r + 1, column=3, value="not tied to a drop").font = BOLD
    for j in range(1, 7):
        ws.cell(row=r + 1, column=3 + j, value='=IF(%s="","",%s)' % (mirror("drops", "drop_id", j), mirror("drops", "drop_id", j))).font = BOLD
    for i, a in enumerate(T.ACTIVITIES):
        rr = r + 2 + i
        ws.cell(row=rr, column=1, value=a)
        ws.cell(row=rr, column=2, value='=SUMIFS(%s,%s,"%s")' % (T.R("hours", "hours"), T.R("hours", "activity"), a)).number_format = "0.0"
        ws.cell(row=rr, column=3, value='=SUMIFS(%s,%s,"%s",%s,"")' % (T.R("hours", "hours"), T.R("hours", "activity"), a, T.R("hours", "drop_id"))).number_format = "0.0"
        for j in range(1, 7):
            hdr = "%s%d" % (get_column_letter(3 + j), r + 1)
            ws.cell(row=rr, column=3 + j, value='=IF(%s="","",SUMIFS(%s,%s,"%s",%s,%s))' % (hdr, T.R("hours", "hours"), T.R("hours", "activity"), a, T.R("hours", "drop_id"), hdr)).number_format = "0.0"
    rr = r + 2 + len(T.ACTIVITIES)
    ws.cell(row=rr, column=1, value="total").font = BOLD
    for j in range(2, 10):
        L = get_column_letter(j)
        ws.cell(row=rr, column=j, value='=IF(%s%d="","",SUM(%s%d:%s%d))' % (L, r + 1, L, r + 2, L, rr - 1) if j > 3 else "=SUM(%s%d:%s%d)" % (L, r + 2, L, rr - 1)).number_format = "0.0"
    for c, w in zip("ABCDEFGHIJKLMNOPQRSTUVWXY", (30, 18, 14, 10, 30, 30, 12, 9, 9, 9, 10, 11, 11, 10, 11, 10, 10, 10, 10, 10, 9, 9, 9, 9, 40)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A4"


def write_survey_dashboard(wb):
    ws = wb.create_sheet("survey_dashboard")
    ws.sheet_properties.tabColor = "70AD47"
    ws["A1"] = "Survey responses"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "Campaign"
    ws["A2"].font = BOLD
    ws["B2"] = '=IF(%s="","",%s)' % (mirror("campaigns", "campaign_id", 1), mirror("campaigns", "campaign_id", 1))
    ws["C2"] = "Type a campaign_id over B2 to switch. Counts are live from the campaign's form tab, inside the window T to T+4 only."
    dv = DataValidation(type="list", formula1="=%s" % T.R("campaigns", "campaign_id"), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B2")
    ws["A3"] = ('=IF(LEFT($B$2,6)="SAMPLE","SAMPLE DATA. Illustrative placeholders, not results.","")')
    ws["A3"].font = Font(bold=True, color="C00000")
    sel = "$B$2"
    camp = lambda col: 'IFERROR(INDEX(%s,MATCH(%s,%s,0)),"")' % (T.R("campaigns", col), sel, T.R("campaigns", "campaign_id"))
    ws["A4"] = '=IF($B$2="","",%s&" · drop "&TEXT(%s,"ddd yyyy-mm-dd")&" · window closes "&TEXT(%s,"ddd yyyy-mm-dd"))' % (
        camp("brand_name"), camp("drop_date"), 'IFERROR(INDEX(%s,MATCH(%s,%s,0)),"")' % (T.R("drops", "t_plus_4"), camp("drop_id"), T.R("drops", "drop_id")))
    heads = ["office", "company", "responses", "said they tried", "didn't see it", "not for me", "gone when they looked",
             "ratings", "avg rating", "would buy: yes", "maybe", "no", "comments", "status"]
    cols = ["office_label", "company", "survey_responses", "tried_yes", "tried_no_didnt_see", "tried_no_not_for_me",
            "tried_no_gone", "rating_n", None, "buy_yes", "buy_maybe", "buy_no", "comments", None]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=6, column=j, value=h).font = BOLD
    for k in range(1, 13):
        rr = 6 + k
        m = 'MATCH(%s&"|"&%d,%s,0)' % (sel, k, T.R("approvals", "seq_key"))
        for j, cn in enumerate(cols, start=1):
            if cn:
                f = '=IFERROR(INDEX(%s,%s),"")' % (T.R("approvals", cn), m)
            elif heads[j - 1] == "avg rating":
                f = '=IFERROR(IF(INDEX(%s,%s)=0,"",ROUND(INDEX(%s,%s)/INDEX(%s,%s),1)),"")' % (
                    T.R("approvals", "rating_n"), m, T.R("approvals", "rating_sum"), m, T.R("approvals", "rating_n"), m)
            else:
                f = '=IFERROR(IF(INDEX(%s,%s)>=%s,"OK","under "&%s&": ask the owner to re-post"),"")' % (
                    T.R("approvals", "survey_responses"), m, T.cfg("min_responses_per_office"), T.cfg("min_responses_per_office"))
            ws.cell(row=rr, column=j, value=f)
    ws.cell(row=19, column=1, value="total").font = BOLD
    for j in (3, 4, 5, 6, 7, 8, 10, 11, 12, 13):
        L = get_column_letter(j)
        ws.cell(row=19, column=j, value="=SUM(%s7:%s18)" % (L, L)).font = BOLD
    ws.cell(row=19, column=9, value='=IFERROR(%s,"")' % camp("survey_rating_avg")).font = BOLD
    ws.cell(row=21, column=1, value="Not counted").font = BOLD
    ws.cell(row=22, column=1, value="outside the window (before T or after T+4)")
    ws.cell(row=22, column=3, value="=" + camp("responses_outside_window"))
    ws.cell(row=23, column=1, value="card code matches no office on this campaign")
    ws.cell(row=23, column=3, value="=" + camp("responses_unmatched_code"))
    ws.cell(row=25, column=1, value=("Who sees what (legal pack file 13, D6 and D9): this sheet is the founder's only. The brand gets totals and at most three "
                                     "comments word for word; each office gets its own summary, never answers row by row. No email address is ever in this file: "
                                     "newsletter sign-ups live only in the campaign's separate sign-up form.")).alignment = WRAP
    ws.merge_cells("A25:N25")
    ws.row_dimensions[25].height = 48
    for c, w in zip("ABCDEFGHIJKLMN", (16, 22, 10, 10, 10, 10, 12, 8, 9, 10, 8, 8, 9, 30)):
        ws.column_dimensions[c].width = w


def write_report(wb):
    ws = wb.create_sheet("report")
    ws.sheet_properties.tabColor = "1F3864"
    sel = "$B$2"
    camp = lambda col: 'IFERROR(INDEX(%s,MATCH(%s,%s,0)),"")' % (T.R("campaigns", col), sel, T.R("campaigns", "campaign_id"))
    drop = lambda col: 'IFERROR(INDEX(%s,MATCH(%s,%s,0)),"")' % (T.R("drops", col), camp("drop_id"), T.R("drops", "drop_id"))
    ws["A1"] = '=IF(LEFT($B$2,6)="SAMPLE","SAMPLE DATA. Illustrative placeholders, not results.","")'
    ws["A1"].font = Font(bold=True, size=14, color="C00000")
    ws["A2"] = "Campaign"
    ws["A2"].font = BOLD
    ws["B2"] = '=IF(%s="","",%s)' % (mirror("campaigns", "campaign_id", 1), mirror("campaigns", "campaign_id", 1))
    dv = DataValidation(type="list", formula1="=%s" % T.R("campaigns", "campaign_id"), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B2")
    ws["D2"] = "Type a campaign_id over B2. Everything below is formulas over the logged rows: nothing is retyped. The PDF comes from tools/make_report.py."
    ws["A4"] = ('=IF($B$2="","",%s&IF(%s<>""," · "&%s,"")&" · "&%s&" · Drop "&TEXT(%s,"ddd yyyy-mm-dd")&" · Window "&TEXT(%s,"ddd yyyy-mm-dd")&" to "&TEXT(%s,"ddd yyyy-mm-dd"))'
                % (camp("brand_name"), camp("product_sku"), camp("product_sku"), T.cfg("cluster_name"), camp("drop_date"), camp("drop_date"), drop("t_plus_4")))
    ws["A4"].font = Font(bold=True, size=13)
    rungs = [("Delivered", camp("units_delivered"), '"handed over on the drop day"'),
             ("Placed", camp("units_placed"), '"within one business day: "&%s&" of "&%s&IF(%s="",""," ("&TEXT(%s,"0%%")&")")' % (
                 camp("units_placed_within_1bd"), camp("units_delivered"), camp("placed_share_1bd"), camp("placed_share_1bd"))),
             ("Depleted", camp("units_depleted_day5"), '"gone by day 5; plus "&%s&" day-3 top-ups from reserve"' % camp("units_topped_up")),
             ("Responded", camp("survey_responses"), '"survey responses; "&%s&" said they tried it"' % camp("survey_tried_yes")),
             ("Scanned", camp("qr_scans"), '"QR scans, T to T+4"'),
             ("Purchased", camp("code_redemptions_to_date"), '"code redemptions reported by the brand"&IF(%s="",""," as of "&TEXT(%s,"yyyy-mm-dd"))' % (camp("code_redemptions_as_of"), camp("code_redemptions_as_of"))),
             ("Incremental", '"not measured"', '"this report does not show whether the campaign caused any sale"')]
    for j, (lab, val, note) in enumerate(rungs, start=1):
        ws.cell(row=6, column=j, value=lab).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=6, column=j).fill = HEAD_I
        ws.cell(row=7, column=j, value="=" + val).font = Font(bold=True, size=14)
        ws.cell(row=8, column=j, value="=IF($B$2=\"\",\"\",%s)" % note).alignment = WRAP
    ws.row_dimensions[8].height = 44
    heads = ["Office", "Industry", "Tuesday headcount", "Units delivered", "Units placed", "Gone by day 5", "Responses", "QR scans", "Placement photo", "Later photo"]
    cols = ["office_label", None, None, "units_delivered", "units_placed", "units_depleted_day5", "survey_responses", "qr_scans", "placement_photo_date", "later_photo_date"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=10, column=j, value=h)
        c.font = BOLD
        c.border = Border(bottom=THIN)
    for k in range(1, 13):
        rr = 10 + k
        m = 'MATCH(%s&"|"&%d,%s,0)' % (sel, k, T.R("approvals", "seq_key"))
        off = 'INDEX(%s,%s)' % (T.R("approvals", "office_id"), m)
        for j, cn in enumerate(cols, start=1):
            if cn == "office_label":
                f = '=IFERROR(INDEX(%s,%s)&IF(INDEX(%s,%s)="N","*",""),"")' % (T.R("approvals", "office_label"), m, T.R("approvals", "placed_within_1bd"), m)
            elif j == 2:
                f = '=IFERROR(%s,"")' % lk_off("industry", off)
            elif j == 3:
                f = '=IFERROR(SUBSTITUTE(%s,"-"," to "),"")' % lk_off("headcount_band", off)
            elif cn in ("placement_photo_date", "later_photo_date"):
                f = '=IFERROR(IF(INDEX(%s,%s)="","none received",TEXT(INDEX(%s,%s),"yyyy-mm-dd")),"")' % (T.R("approvals", cn), m, T.R("approvals", cn), m)
            else:
                f = '=IFERROR(INDEX(%s,%s),"")' % (T.R("approvals", cn), m)
            ws.cell(row=rr, column=j, value=f)
    ws.cell(row=23, column=1, value="Total").font = BOLD
    for j in (4, 5, 6, 7, 8):
        L = get_column_letter(j)
        ws.cell(row=23, column=j, value="=SUM(%s11:%s22)" % (L, L)).font = BOLD
    ws.cell(row=24, column=1, value='=IF(COUNTIF(A11:A22,"*~*")>0,"* placed after the first business day","")')
    n = camp("survey_responses")
    pct = lambda x, d: 'IF(N(%s)=0,"-",TEXT(%s/%s,"0%%")&" ("&%s&")")' % (d, x, d, x)
    buyn = '(N(%s)+N(%s)+N(%s))' % (camp("survey_buy_yes"), camp("survey_buy_maybe"), camp("survey_buy_no"))
    ws["A26"] = ('=IF($B$2="","","Survey (n = "&%s&"): tried "&%s&" · average rating "&%s&" of 5 ("&%s&" ratings) · would buy, of "&%s&" who answered: yes "&%s&", maybe "&%s&", no "&%s)'
                 % (n, pct(camp("survey_tried_yes"), n), camp("survey_rating_avg"), camp("survey_rating_n"), buyn,
                    pct(camp("survey_buy_yes"), buyn), pct(camp("survey_buy_maybe"), buyn), pct(camp("survey_buy_no"), buyn)))
    ws["A26"].font = BOLD
    ws["A27"] = '=IF($B$2="","","Did not try: didn\'t see it "&%s&", not for me "&%s&", gone by the time they looked "&%s)' % (
        camp("tried_no_didnt_see"), camp("tried_no_not_for_me"), camp("tried_no_gone"))
    for i in range(1, 4):
        ws["A%d" % (27 + i)] = '=IF(%s="","","“"&%s&"”")' % (camp("comment_%d" % i), camp("comment_%d" % i))
        ws["A%d" % (27 + i)].alignment = WRAP
        ws.merge_cells("A%d:J%d" % (27 + i, 27 + i))
        ws.row_dimensions[27 + i].height = 30
    ws["A31"] = '=IF($B$2="","","Three of "&%s&" comments, word for word.")' % camp("comments_total")
    ws["A33"] = ('=IF($B$2="","","Engagement: "&N(%s)&" QR scans across "&%s&" offices · "&IF(%s="","no code redemptions reported yet",%s&" code redemptions to date"))'
                 % (camp("qr_scans"), camp("offices_on_route"), camp("code_redemptions_to_date"), camp("code_redemptions_to_date")))
    ws["A33"].font = BOLD
    ws["A35"] = "What we learned"
    ws["A35"].font = BOLD
    ws["A36"] = '=%s' % camp("learned")
    ws["A36"].alignment = WRAP
    ws.merge_cells("A36:J36")
    ws.row_dimensions[36].height = 60
    ws["A38"] = "Recommendation"
    ws["A38"].font = BOLD
    ws["A39"] = '=%s' % camp("recommendation")
    ws["A39"].alignment = WRAP
    ws.merge_cells("A39:J39")
    ws.row_dimensions[39].height = 32
    ws["A41"] = "Limitations"
    ws["A41"].font = BOLD
    ws["A42"] = limitations_text()
    ws["A42"].alignment = WRAP
    ws.merge_cells("A42:J42")
    ws.row_dimensions[42].height = 96
    ws["A44"] = '="Prepared by "&%s&" · "&%s&" · "&%s' % (T.cfg("founder_name"), T.cfg("company_name"), T.cfg("founder_email"))
    for c, w in zip("ABCDEFGHIJ", (12, 18, 14, 12, 12, 12, 11, 10, 15, 15)):
        ws.column_dimensions[c].width = w


def lk_off(col, off):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (T.R("offices", col), off, T.R("offices", "office_id"))


README_LINES = [
    ("Campaign tracker", "title"),
    ("Every number in a brand report comes from a row in this file or a photo. Type only in white cells. Grey cells are formulas: never type over them. Cream tabs are written by forms, scripts or a pasted file.", ""),
    ("Rename: change config!B2 (company_name). Nothing else in this file holds the name.", ""),
    ("", ""),
    ("Tabs in the order a campaign touches them", "h"),
    ("office_intake, brand_intake: website form rows (CONTRACTS 6.3), written by the website's script when it is bound to this file.", ""),
    ("office_form, brand_form: the founder-side Google Forms' response tabs; office_intake_founder, brand_intake_founder turn them into the same 6.3 columns.", ""),
    ("in_signed_offices: paste E's signed_offices.csv. offices: type the office_id; the rest fills by header name.", ""),
    ("brands: type brand_id and intake_key (the brand's submitted_at); the seven checks fill in. in_paid_brands: paste E's paid_brands.csv to check it.", ""),
    ("drops: one row per route. The date is refused unless it is a Tuesday or Wednesday business day with a window clear of blackouts.", ""),
    ("campaigns: one row per brand per drop: terms, card text, survey links, report inputs, and every ladder number.", ""),
    ("invoices: issue, then log paid_date. The ship date is released only when the pre-ship payment is in, the agreement is signed, the certificate checked and the offices approved.", ""),
    ("approvals: one row per office per campaign: the office's yes, its units, and its whole ladder (formulas).", ""),
    ("stops: one row per office per drop: the route.", ""),
    ("links: made by formulas. Two short links per office per campaign: qr (on the card) and post (the Slack or email post).", ""),
    ("inventory: every unit that moves. counts: shelf counts. photos: one row per photo.", ""),
    ("<campaign_id> tabs: each campaign's survey form writes its own tab, named with the campaign_id (forms/survey-form.md).", ""),
    ("scans: link totals read from the short-link tool. redemptions: what the brand reports. issues: every deviation.", ""),
    ("hours, costs: the founder hours log and the direct costs. Exports for the financial model come from tools/export_actuals.py.", ""),
    ("dashboard, survey_dashboard, report: read only.", ""),
    ("", ""),
    ("Rules this file keeps", "h"),
    ("The ladder rungs stay in separate columns: delivered, placed, depleted, responded, scanned, purchased. Incremental is not measured. Nothing adds them up.", ""),
    ("Placed means visible in a usable placement photo with the card. Placed within one business day means that photo is dated on or before T+1.", ""),
    ("Survey counts only include answers inside the window, T to the end of T+4, whose card code matches an office on the campaign.", ""),
    ("No employee data. No email address ever enters this file.", ""),
    ("SAMPLE: any ID or text starting with SAMPLE is test data; the dashboard and report show a SAMPLE banner.", ""),
    ("Google Sheets: File > Settings > Locale United States, Time zone America/New_York.", ""),
]


def write_readme(wb):
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 150
    for i, (line, kind) in enumerate(README_LINES, start=1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True)
        if kind == "title":
            c.value = '="Campaign tracker · "&config!$B$2'
            c.font = Font(bold=True, size=16)
        elif kind == "h":
            c.font = Font(bold=True, size=12)


def build(out, sample=None, response_tabs=None, today=None, config=None):
    """sample: {tab: [ {col: value} ]} rows written into input columns.
    response_tabs: {campaign_id: (brand_name, [ [timestamp, card_code, tried, rating, buy, why, comment] ])}."""
    F.build_all()
    wb = openpyxl.Workbook()
    write_readme(wb)
    write_dashboard(wb)
    write_survey_dashboard(wb)
    write_report(wb)
    write_config(wb)
    write_calendar(wb)
    for name in ORDER:
        if name in T.TABS:
            ws = write_tab(wb, T.TABS[name])
            if name == "drops":
                drop_date_validation(ws)
    write_form_tab(wb, "office_form", form_headers("office_intake"), note="Template of the founder-side office form's response tab. When you link the Google Form to this file, Google adds a new tab: delete this one and rename the new tab office_form (forms/office-intake-form.md).")
    write_form_tab(wb, "brand_form", form_headers("brand_intake"), note="Template of the founder-side brand form's response tab. When you link the Google Form, delete this tab and rename Google's new tab brand_form.")
    write_form_tab(wb, "responses_template", survey_headers(), note="What each campaign's survey tab looks like (columns A to G, in this order). Google writes it; you rename it to the campaign_id. Never insert or move columns in a survey tab.")
    if response_tabs:
        for cid, (brand, rows) in response_tabs.items():
            ws = write_form_tab(wb, cid, survey_headers(brand))
            for r, row in enumerate(rows, start=2):
                for j, v in enumerate(row, start=1):
                    c = ws.cell(row=r, column=j, value=v)
                    if j == 1:
                        c.number_format = "yyyy-mm-dd hh:mm:ss"
    # order tabs
    wb._sheets = [wb[n] for n in ORDER if n in wb.sheetnames] + [ws for ws in wb._sheets if ws.title not in ORDER]
    if today:
        wb["config"].cell(row=T.CFG_ROW["today_override"], column=2, value=today).number_format = T.DATE
    for k, v in (config or {}).items():
        wb["config"].cell(row=T.CFG_ROW[k], column=2, value=v)
    if sample:
        fill(wb, sample)
    wb.save(out)
    return out


def fill(wb, sample):
    for tab, rows in sample.items():
        ws = wb[tab]
        t = T.TABS.get(tab)
        if t is None:
            # free-form tab with a header row
            hdr = [c.value for c in ws[1]]
            for r, row in enumerate(rows, start=2):
                for k, v in row.items():
                    ws.cell(row=r, column=hdr.index(k) + 1, value=v)
            continue
        for r, row in enumerate(rows, start=t.first()):
            for k, v in row.items():
                if k not in t.index:
                    raise KeyError("%s has no column %s" % (tab, k))
                col = t.cols[t.index[k]]
                if col.kind == "A":
                    raise KeyError("%s.%s is a formula column" % (tab, k))
                c = ws.cell(row=r, column=t.index[k] + 1, value=v)
                if isinstance(v, dt.time):
                    c.number_format = T.TIME
                elif isinstance(v, (dt.date, dt.datetime)):
                    c.number_format = T.DATE


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(ROOT, "tracker", "campaign-tracker.xlsx"))
    a = ap.parse_args()
    build(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
