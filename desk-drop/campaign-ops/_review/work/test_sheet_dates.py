import sys, os, datetime as dt
sys.path.insert(0, 'tools')
import openpyxl, recalc as RC, opscal
src='tracker/campaign-tracker.xlsx'
wb=openpyxl.load_workbook(src)
ws=wb['drops']
dates=['2026-12-22','2027-07-06','2026-08-25','2027-01-11','2027-01-15','2027-01-12','2026-12-16','2027-01-14','2026-11-11','2027-08-11','2027-08-12','2027-06-30']
for i,dd in enumerate(dates, start=2):
    ws.cell(i,1,'TEST-DRP-%03d'%i); c=ws.cell(i,2,dt.datetime.fromisoformat(dd)); c.number_format='yyyy-mm-dd'
out='../../work/dates-in.xlsx'; wb.save(out)
r=RC.recalc(out,'../../work/dates-out.xlsx')
wb2=openpyxl.load_workbook(r,data_only=True); w=wb2['drops']
hdr=[c.value for c in w[1]]
ok=True
for i,dd in enumerate(dates, start=2):
    row={hdr[j]:w.cell(i,j+1).value for j in range(len(hdr))}
    py=opscal.problems(dd)
    s=opscal.sop_dates(dd) if not py else None
    sheet=row['date_check']
    agree = (sheet=='OK')==(not py)
    det=''
    if s:
        for k,col in (('ship_by','t_minus_10'),('ship_by_requested','ship_by_requested'),('receive','t_minus_5'),('kit_print','t_minus_3'),('confirm','t_minus_1'),('placement_confirm','t_plus_1'),('check_day3','t_plus_2'),('close','t_plus_4'),('pass_on','t_plus_5'),('report_due','t_plus_7'),('debrief','t_plus_10'),('return_label_deadline','t_plus_14'),('redemption_check','t_plus_30_cal')):
            v=row[col]; v=v.date() if hasattr(v,'date') else v
            if v!=s[k]: det+=' MISMATCH %s sheet=%s py=%s;'%(col,v,s[k]); agree=False
    else:
        if any(row[c] not in (None,'') for c in ('t_minus_10','t_plus_4','t_plus_7')): det+=' derived dates NOT blank'; agree=False
    ok&=agree
    print(dd, '| sheet:', sheet, '| py:', 'OK' if not py else 'REFUSED', '|', 'AGREE' if agree else 'DISAGREE', det)
print('ALL AGREE' if ok else 'SOME DISAGREE')
