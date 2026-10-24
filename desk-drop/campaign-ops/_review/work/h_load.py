"""Reviewer's independent load of F's exports into a COPY of H's workbook (H's file is only read).
Usage: python3 h_load.py <exports_dir> <out_dir>"""
import sys, os, csv, shutil, hashlib, datetime as dt, re, subprocess, tempfile
import openpyxl
H = '/home/user/accounting-agent/desk-drop/financial-model/Desk Drop financial model v1.xlsx'
exp, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
sha0 = hashlib.sha256(open(H, 'rb').read()).hexdigest()
dst = os.path.join(out, 'model-copy.xlsx'); shutil.copy(H, dst)
wb = openpyxl.load_workbook(dst)
def conv(k, v):
    if v == '': return None
    v = re.sub(r'^SAMPLE-', '', v) if k in ('campaign_id', 'drop_id', 'brand_id') else v
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', v): return dt.datetime.fromisoformat(v)
    if re.fullmatch(r'-?\d+(\.\d+)?', v): return float(v) if '.' in v else int(v)
    return v
drops = set()
for f, tab in (('campaign_actuals.csv', 'Actuals_Campaigns'), ('drop_costs.csv', 'Actuals_DropCosts'), ('hours_log.csv', 'Hours_Log')):
    rows = list(csv.reader(open(os.path.join(exp, f))))
    ws = wb[tab]
    hdr = [ws.cell(4, c).value for c in range(1, len(rows[0]) + 1)]
    assert hdr == rows[0], (tab, hdr, rows[0])
    for i, r in enumerate(rows[1:], start=5):
        for j, (k, v) in enumerate(zip(rows[0], r), start=1):
            x = conv(k, v)
            ws.cell(i, j, x)
            if isinstance(x, dt.datetime): ws.cell(i, j).number_format = 'yyyy-mm-dd'
            if k == 'drop_id' and x: drops.add(x)
    print(tab, len(rows) - 1, 'rows pasted')
ws = wb['Schedule']
for r in range(7, 19):
    if ws.cell(r, 1).value in drops:
        ws.cell(r, 14, 'Y')
wb.save(dst)
tmp = tempfile.mkdtemp(); env = dict(os.environ, HOME=tmp)
subprocess.run(['soffice', '--headless', '--norestore', '--calc', '--convert-to', 'xlsx', '--outdir', os.path.join(out, 'recalc'), dst], check=True, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=900)
rc = os.path.join(out, 'recalc', 'model-copy.xlsx')
wv = openpyxl.load_workbook(rc, data_only=True)
errs = []
for s in wv.worksheets:
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(('#N/A', '#VALUE', '#REF', '#DIV', '#NAME', '#NUM', 'Err:')):
                errs.append('%s!%s' % (s.title, c.coordinate))
print('formula errors:', len(errs), errs[:10])
ck = wv['Checks']
fails = [(ck.cell(r, 1).value, ck.cell(r, 3).value, ck.cell(r, 4).value) for r in range(5, ck.max_row + 1) if ck.cell(r, 5).value == 0]
print('Checks failing:', len(fails))
for f in fails: print('  ', f)
da = wv['Drop_Actuals']; hdr = [da.cell(6, c).value for c in range(1, 40)]
for r in range(7, 19):
    if da.cell(r, 1).value in drops:
        d = dict(zip(hdr, [da.cell(r, c).value for c in range(1, 40)]))
        print('Drop_Actuals', {k: d[k] for k in list(d)[:20] if k})
        print('   ', {k: v for k, v in list(d.items())[20:] if k})
print('H file unchanged:', hashlib.sha256(open(H, 'rb').read()).hexdigest() == sha0)
