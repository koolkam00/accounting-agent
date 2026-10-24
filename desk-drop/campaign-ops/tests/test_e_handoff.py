#!/usr/bin/env python3
"""Paste the outreach engine's own SAMPLE handoff files into a fresh tracker, the way the founder will.

  python3 tests/test_e_handoff.py

Reads desk-drop-outreach/handoff/SAMPLE-signed_offices.csv and SAMPLE-paid_brands.csv (E's published layout,
CONTRACTS 6.4 plus the columns this folder asked for, in E's column order), pastes each with its header row at A1
of in_signed_offices / in_paid_brands (numbers and ISO dates converted, as a Google Sheets import does),
recalculates with LibreOffice and requires:
  - 0 formula errors;
  - every offices row (only office_id typed) fills company, headcount band, owner phone, industry, naming
    consent, window and categories from E's row, read by column name;
  - tools/tracker.py reads the pantry owner's email and the requested columns the generators use;
  - the dashboard's two "in E's file but not on the tracker" checks read 0 when the IDs are typed, and the number
    of E's rows when they are not (second build);
  - the payment reconciliation (review round 1, R1-06): with no invoice on the tracker, "Brands whose payments in E's
    paid_brands.csv differ from the invoices tab" reads 1; with the same payment logged on the invoices tab (amount and
    date as E's payment_1), it reads 0; with a different amount, 1 again.
Never writes outside a temporary folder.
"""
import csv
import datetime as dt
import os
import re
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DESK = os.path.dirname(ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))
import build_tracker as B  # noqa: E402
import dryrun as DR  # noqa: E402
import recalc as RC  # noqa: E402
import tracker as TR  # noqa: E402
import trackerdef as T  # noqa: E402

E = os.path.join(DESK, "desk-drop-outreach", "handoff")
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
NUM = re.compile(r"^-?\d+(\.\d+)?$")


def cell(v):
    if v == "":
        return None
    if ISO.match(v):
        return dt.date.fromisoformat(v)
    if NUM.match(v) and not v.startswith("0"):
        return float(v) if "." in v else int(v)
    return v


def paste(wb, tab, path):
    ws = wb[tab]
    for row in ws.iter_rows():
        for c in row:
            c.value = None
    rows = list(csv.reader(open(path, encoding="utf-8")))
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v if i == 1 else cell(v))
    return rows


def ids(path, col):
    return [r[col] for r in csv.DictReader(open(path, encoding="utf-8")) if r.get(col)]


def build(tmp, with_ids, invoices=None, tag=""):
    p = os.path.join(tmp, "t-%s%s.xlsx" % (with_ids, tag))
    oids = ids(os.path.join(E, "SAMPLE-signed_offices.csv"), "office_id")
    bids = ids(os.path.join(E, "SAMPLE-paid_brands.csv"), "brand_id")
    sample = {"offices": [{"office_id": o} for o in oids], "brands": [{"brand_id": b} for b in bids]} if with_ids else {}
    if invoices:
        sample["invoices"] = invoices
    B.build(p, sample=sample, today=dt.date(2026, 12, 7), config=DR.SAMPLE_CONFIG)
    wb = openpyxl.load_workbook(p)
    so = paste(wb, "in_signed_offices", os.path.join(E, "SAMPLE-signed_offices.csv"))
    pb = paste(wb, "in_paid_brands", os.path.join(E, "SAMPLE-paid_brands.csv"))
    wb.save(p)
    out = RC.recalc(p, os.path.join(tmp, "r-%s%s.xlsx" % (with_ids, tag)))
    return out, so, pb


def dash(path):
    ws = openpyxl.load_workbook(path, data_only=True)["dashboard"]
    out, r = {}, 7
    while ws.cell(row=r, column=1).value:
        out[ws.cell(row=r, column=1).value] = ws.cell(row=r, column=2).value
        r += 1
    return out


def main():
    fails = []
    tmp = tempfile.mkdtemp(prefix="ops_e_")
    out, so, pb = build(tmp, True)
    errs, n = RC.errors(out)
    if n:
        fails.append("formula errors: %s" % errs[:5])
    hdr = so[0]
    for need in T.SIGNED_OFFICES_REQUESTED if hasattr(T, "SIGNED_OFFICES_REQUESTED") else []:
        if need not in hdr:
            fails.append("E's signed_offices.csv lacks requested column %s" % need)
    for need in T.PAID_BRANDS_REQUESTED if hasattr(T, "PAID_BRANDS_REQUESTED") else []:
        if need not in pb[0]:
            fails.append("E's paid_brands.csv lacks requested column %s" % need)
    tr = TR.Tracker(out)

    def band(h):
        h = int(h)
        return "under-40" if h < 40 else "40-75" if h <= 75 else "76-150" if h <= 150 else "151-250" if h <= 250 else "251-500" if h <= 500 else "over-500"
    for row in so[1:]:
        e = dict(zip(hdr, row))
        o = tr.office(e["office_id"])
        phone = e["pantry_owner_phone"] or (e["contact_phone"] if e["pantry_owner_is_contact"] == "Y" else "")
        want = {"company": e["company"], "industry": e["industry"], "public_naming_ok": e["public_naming_ok"],
                "headcount_band": band(e["tuesday_headcount"]) if e["tuesday_headcount"] else "unknown", "owner_phone": phone,
                "pantry_owner_email": e["pantry_owner_email"], "receiving_window": e["receiving_window"],
                "categories_wanted": e["categories_wanted"]}
        for k, v in want.items():
            if TR._s(o.get(k)) != v:
                fails.append("offices %s %s = %r, want %r (from E's file)" % (e["office_id"], k, o.get(k), v))
    d = dash(out)
    for k in ("Brands in E's paid_brands.csv not on brands", "Offices in E's signed_offices.csv not on offices"):
        if d.get(k) != 0:
            fails.append("dashboard %r = %r with the IDs typed, want 0" % (k, d.get(k)))
    out2, _, _ = build(tmp, False)
    d2 = dash(out2)
    for k, n in (("Brands in E's paid_brands.csv not on brands", len(pb) - 1), ("Offices in E's signed_offices.csv not on offices", len(so) - 1)):
        if d2.get(k) != n:
            fails.append("dashboard %r = %r with the IDs missing, want %d" % (k, d2.get(k), n))
    # Payments: E's payment_1/2 against the invoices tab, on the campaign E's row names (R1-06).
    PAY = "Brands whose payments in E's paid_brands.csv differ from the invoices tab"
    if d.get(PAY) != 1:
        fails.append("dashboard %r = %r with no invoice logged, want 1" % (PAY, d.get(PAY)))
    e = dict(zip(pb[0], pb[1]))
    inv = {"invoice_id": "SAMPLE-INV-901", "campaign_id": e["campaign_id"], "invoice_type": "founding_full",
           "issued_date": dt.date(2026, 12, 1), "amount_usd": 600, "sent_to": "SAMPLE-ap@example.com",
           "paid_date": dt.date.fromisoformat(e["payment_1_date"]), "paid_amount_usd": float(e["payment_1_usd"]), "method": "bank_transfer"}
    out3, _, _ = build(tmp, True, [inv], "-same")
    d3 = dash(out3)
    if d3.get(PAY) != 0:
        fails.append("dashboard %r = %r with the same payment on the invoices tab, want 0" % (PAY, d3.get(PAY)))
    out4, _, _ = build(tmp, True, [dict(inv, paid_amount_usd=300)], "-differs")
    d4 = dash(out4)
    if d4.get(PAY) != 1:
        fails.append("dashboard %r = %r with a different amount, want 1" % (PAY, d4.get(PAY)))
    print("E's files: %d office row(s), %d brand row(s), %d and %d columns; payment check %s / %s / %s (want 1 / 0 / 1)" % (
        len(so) - 1, len(pb) - 1, len(so[0]), len(pb[0]), d.get(PAY), d3.get(PAY), d4.get(PAY)))
    for f in fails:
        print("  FAIL " + f)
    print("PASS test_e_handoff" if not fails else "FAIL test_e_handoff")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
