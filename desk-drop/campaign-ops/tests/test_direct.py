#!/usr/bin/env python3
"""The direct-ship model and the other review round 2 findings, run through the tracker and the tools.

  python3 tests/test_direct.py

From the direct scenario (tools/scenarios.py direct(): SAMPLE-DRP-004, Tue 2027-03-23, direct-ship, SAMPLE Brightside
founding_lock and SAMPLE Tallgrass standard, ten offices; SAMPLE-O-003's Brightside carton 26 of 50; SAMPLE-O-008 refuses
the Tallgrass carton and two of its people answer that survey; SAMPLE-O-009's front desk will not hold food), each case
builds a tracker as the founder would have it on one day, recalculates it with LibreOffice and requires:

 1. T-5 as logged (the reserve and SAMPLE-O-009's cartons in): no short-shipment, units_override or reserve check.  (R2-02)
 2. T-5 with the Brightside reserve 12 short: "at your address: received ... of ..." asks for an issue.             (R2-02)
 3. T+1 without SAMPLE-O-003's issue: "delivered 26 of 50 allocated" on approvals and "received 543 of 567 due".   (R2-02)
 4. T+1 with the office cartons logged as delivered only (the old SOP): "reserve below zero: in direct-ship ...".     (R2-02)
 5. T+7 as logged: 0 errors, 0 disagreements, 0 dashboard checks. Tallgrass: responded 61 = the roster's sum, 2
    answers from SAMPLE-O-008 shown apart; both reports render; the reasons fit inside "did not try"; the page carries
    the roster line and the delivered line; trace.csv and the report email say what was not counted and how delivered
    was counted; a quote from the office off the drop is refused; a debrief sheet without a report says NO REPORT YET.
                                                                                                                (R2-01, R2-03, R2-06, R2-07)
 6. The documents: the tracking sheet asks for office counts and sends a refused carton back to the brand; the T-3
    email names the envelope and the count; the T-1 text promises no window to a carrier office; the hand-carried
    office is "YOU DELIVER" and its ship-date row says "with the reserve"; the kit checklist's reserve is not negative;
    a day-3 text with two brands under five says "have".                                                     (R2-03, R2-09)
 7. The final state: reserve 0, no check left; the export's units_received is the reserve plus every carton.       (R2-02)
 8. Roster rules: two stops where an office declined every brand raise the offices check, a one-brand refusal does
    not count; coworking and outside-the-boundary rows raise a check until confirmed; two campaigns of one category
    on one drop raise a check.                                                                                 (R2-05)
 9. The dashboard's "Tabs over 90% full" survives replacing the office_form tab (INDIRECT), and no formula refers to
    office_form or brand_form directly.                                                                         (R2-04)
Writes only to a temporary folder.
"""
import copy
import csv
import datetime as dt
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "tools"))
import dryrun as DR  # noqa: E402
import opscal  # noqa: E402
import scenarios as S  # noqa: E402
import tracker as TR  # noqa: E402
import trackerdef as T  # noqa: E402

DID, C7, C8 = "SAMPLE-DRP-004", "SAMPLE-CMP-007", "SAMPLE-CMP-008"
fails = []


def ok(cond, msg):
    if not cond:
        fails.append(msg)
    return cond


def snap(scn, day, tmp, tag):
    p = DR.build_snapshot(scn, day, os.path.join(tmp, "in-%s.xlsx" % tag))
    out = os.path.join(tmp, "snap-%s.xlsx" % tag)
    tr, errs, n = DR.recalc_and_read(p, out)
    ok(n == 0, "%s: %d formula errors %s" % (tag, n, errs[:3]))
    return out, tr


def rows(path, tab):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[tab]
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] not in (None, "")]


def row(path, tab, **key):
    return next((d for d in rows(path, tab) if all(d.get(k) == v for k, v in key.items())), {})


def tool(*args):
    r = subprocess.run([sys.executable] + list(args), capture_output=True, text=True, cwd=ROOT, timeout=600)
    return r.returncode, r.stdout + r.stderr


def read(p):
    return open(p, encoding="utf-8").read() if os.path.exists(p) else ""


def main():
    tmp = tempfile.mkdtemp(prefix="ops_direct_")
    base = S.direct()
    T0 = base["T"]
    sd = opscal.sop_dates(T0)

    # 1. T-5 as logged
    p1, _ = snap(base, sd["receive"], tmp, "t5")
    for cid in (C7, C8):
        ch = row(p1, "campaigns", campaign_id=cid).get("checks") or ""
        ok(not re.search(r"received \d+|reduce units_override|reserve below zero|at your address", ch), "T-5: %s false receiving check %r" % (cid, ch))
    ok(not DR.dashboard_checks(p1)[0], "T-5: dashboard checks %s" % (DR.dashboard_checks(p1)[0],))

    # 2. T-5, the Brightside reserve 12 short and no issue
    s2 = copy.deepcopy(base)
    for r in s2["rows"]["inventory"]:
        if r["campaign_id"] == C7 and r["movement"] == "received" and r["qty"] == 52:
            r["qty"] = 40
    p2, _ = snap(s2, sd["receive"], tmp, "t5-short")
    ch = row(p2, "campaigns", campaign_id=C7).get("checks") or ""
    ok("at your address: received 115 of 127" in ch and "log a short_shipment issue" in ch, "T-5 short reserve: checks %r" % ch)

    # 3. T+1 without the short-carton issue
    s3 = copy.deepcopy(base)
    s3["rows"]["issues"] = [r for r in s3["rows"]["issues"] if r.get("issue_id") != "SAMPLE-ISS-030"]
    p3, _ = snap(s3, sd["placement_confirm"], tmp, "t1-noissue")
    a = row(p3, "approvals", campaign_id=C7, office_id="SAMPLE-O-003").get("checks") or ""
    ok("delivered 26 of 50 allocated: a short carton?" in a, "T+1 short carton, no issue: approvals check %r" % a)
    ch = row(p3, "campaigns", campaign_id=C7).get("checks") or ""
    ok("received 543 of 567 due" in ch, "T+1 short carton, no issue: campaign check %r" % ch)
    ch8 = row(p3, "campaigns", campaign_id=C8).get("checks") or ""
    ok("due" not in ch8, "T+1: the carton refused at the door is still counted as due: %r" % ch8)

    # 4. T+1 with the office cartons logged as delivered only (no received row per carton)
    s4 = copy.deepcopy(base)
    s4["rows"]["inventory"] = [r for r in s4["rows"]["inventory"] if not (r["movement"] == "received" and r.get("office_id"))]
    p4, _ = snap(s4, sd["placement_confirm"], tmp, "t1-noreceived")
    ch = row(p4, "campaigns", campaign_id=C7).get("checks") or ""
    ok("reserve below zero: in direct-ship, log a received row" in ch, "T+1 without received rows: %r" % ch)

    # 5. T+7 as logged
    p5, tr5 = snap(base, sd["report_due"], tmp, "t7")
    bad = [b for b in TR.compare(tr5) if not b.startswith("no cached")]
    ok(not bad, "T+7: disagreements %s" % bad[:3])
    checks, msgs, rel = DR.dashboard_checks(p5)
    ok(not checks, "T+7: dashboard checks %s %s" % (checks, msgs[:3]))
    c8 = row(p5, "campaigns", campaign_id=C8)
    on = [r for r in rows(p5, "approvals") if r["campaign_id"] == C8 and r["on_drop"] == "Y"]
    ok(c8.get("survey_responses") == sum(r["survey_responses"] for r in on) == 61, "T+7: Tallgrass responded %r, roster sum %r" % (
        c8.get("survey_responses"), sum(r["survey_responses"] for r in on)))
    ok(c8.get("responses_office_not_on_drop") == 2 and c8.get("responses_unmatched_code") == 0,
       "T+7: off-drop %r unmatched %r" % (c8.get("responses_office_not_on_drop"), c8.get("responses_unmatched_code")))
    ok("2 in-window answers from an office not on the drop" in (c8.get("checks") or ""), "T+7: no info line for the off-drop answers")
    photos = os.path.join(tmp, "photos-root")
    DR.make_photos(base, photos)
    rep = os.path.join(tmp, "reports")
    for cid in (C7, C8):
        rc, out = tool("tools/make_report.py", p5, cid, rep, "--photos", os.path.join(photos, "photos"))
        ok(rc == 0 and os.path.exists(os.path.join(rep, cid, "report.pdf")), "T+7: %s report refused: %s" % (cid, out[-500:]))
        rj = json.load(open(os.path.join(rep, cid, "report.json"), encoding="utf-8"))
        total = sum(int(o["survey_responses"]) for o in rj["offices"])
        ok(total == tr5.ladder(cid)["survey_responses"], "T+7: %s headline responses %r, roster sum %r" % (cid, tr5.ladder(cid)["survey_responses"], total))
        reasons = sum(int(rj[k] or 0) for k in ("tried_no_didnt_see", "tried_no_not_for_me", "tried_no_gone"))
        ok(reasons <= total - int(rj["survey_tried_yes"]), "T+7: %s reasons %d over did-not-try %d" % (cid, reasons, total - int(rj["survey_tried_yes"])))
        ok("Delivered:" in rj["learned"] and "counted their own carton" in rj["learned"], "T+7: %s no delivered line on the page: %r" % (cid, rj["learned"][:200]))
    rj8 = json.load(open(os.path.join(rep, C8, "report.json"), encoding="utf-8"))
    ok(rj8["learned"].startswith("Roster: 9 of the 10 offices in the agreement.") and "its share was credited" in rj8["learned"],
       "T+7: no roster line first on the Tallgrass page: %r" % rj8["learned"][:200])
    trace = read(os.path.join(rep, C8, "trace.csv"))
    ok("not counted: answers from an office not on the drop" in trace and "SAMPLE-CMP-008.SAMPLE-O-008" in trace, "T+7: trace.csv does not say what was not counted")
    rc, out = tool("tools/make_messages.py", p5, DID, "report", os.path.join(tmp, "msg"), "--reports", rep)
    body = read(os.path.join(tmp, "msg", "report", C8 + ".md"))
    ok(rc == 0 and "2 survey answers were not counted" in body and "Delivered: counted on arrival by" in body, "T+7: report email %d %r" % (rc, body[-600:]))
    s5 = copy.deepcopy(base)
    for r in s5["rows"]["campaigns"]:
        if r["campaign_id"] == C8:
            r["comment_3"] = "SAMPLE never saw them"
    p5b, _ = snap(s5, sd["report_due"], tmp, "t7-offquote")
    rc, out = tool("tools/make_report.py", p5b, C8, os.path.join(tmp, "rep-q"), "--photos", os.path.join(photos, "photos"))
    ok(rc == 2 and "office on the drop" in out, "T+7: a quote from the office off the drop was not refused (%d) %s" % (rc, out[-300:]))
    deb = os.path.join(tmp, "deb")
    tool("tools/make_debrief.py", p5, DID, deb, "--reports", os.path.join(tmp, "no-reports"))
    ok(read(os.path.join(deb, "debrief-brand-%s.md" % C8)).count("NO REPORT YET") == 1 and "from the report" not in read(os.path.join(deb, "debrief-brand-%s.md" % C8)),
       "T+7: a brand debrief sheet without a report does not say so")
    tool("tools/make_debrief.py", p5, DID, os.path.join(tmp, "deb2"), "--reports", rep)
    ok("NO REPORT YET" not in read(os.path.join(tmp, "deb2", "debrief-brand-%s.md" % C8)), "T+7: NO REPORT YET although the report exists")
    ok("turned away at the door" in read(os.path.join(tmp, "deb2", "debrief-office-SAMPLE-O-008.md")), "T+7: SAMPLE-O-008's debrief does not mention the refused brand")

    # 6. The documents
    p6, _ = snap(base, sd["confirm"], tmp, "t-1")
    rc, out = tool("tools/make_route.py", p6, DID, os.path.join(tmp, "route"))
    html = read(os.path.join(tmp, "route", "route-%s.html" % DID))
    ok(rc == 0 and ("the office&#x27;s count" in html or "the office's count" in html), "tracking sheet: no office count box (%d)" % rc)
    ok("office_count" in html and "carrier returns it to the brand" in html and "founder_count or courier_count" not in html,
       "tracking sheet: the founder's block is not the direct-ship one")
    ok("YOU DELIVER" in html and html.count("YOU DELIVER") >= 2, "tracking sheet: the hand-carried office is not marked")
    tool("tools/make_messages.py", p6, DID, "t3", os.path.join(tmp, "msg6"))
    tool("tools/make_messages.py", p6, DID, "t1", os.path.join(tmp, "msg6"))
    t3, t1 = read(os.path.join(tmp, "msg6", "t3", "SAMPLE-O-008.md")), read(os.path.join(tmp, "msg6", "t1", "SAMPLE-O-008.md"))
    ok("envelope" in t3 and "count the units in each carton" in t3 and "courier" not in t3, "T-3 direct-ship email: %r" % t3[:600])
    ok("window" not in t1 and "as we agreed" not in t1 and "no set time" in t1, "T-1 direct-ship text: %r" % t1)
    t1h = read(os.path.join(tmp, "msg6", "t1", "SAMPLE-O-009.md"))
    ok("I'll bring" in t1h and "window" in t1h, "T-1 hand-carried text: %r" % t1h)
    s6 = copy.deepcopy(base)
    for r in s6["rows"]["campaigns"]:
        if r["drop_id"] == DID:
            r.pop("_fields_history", None)
    p6b, _ = snap(s6, sd["kit_print"], tmp, "release")
    tool("tools/make_messages.py", p6b, DID, "ship_release", os.path.join(tmp, "msg6b"))
    rel = read(os.path.join(tmp, "msg6b", "ship_release", C7 + ".md"))
    ok("with the reserve, to" in rel and "front desk will not hold food" in rel, "ship-date email: the hand-carried office row %r" % rel[-900:])
    rc, out = tool("tools/make_kit.py", p6, DID, os.path.join(tmp, "kit"))
    kit = re.sub(r"<[^>]+>", " ", read(os.path.join(tmp, "kit", "kit-checklist.html")))
    res = re.findall(r"SAMPLE (?:Brightside|Tallgrass)\s+(-?\d+)\s", kit.split("Reserve (stays with you")[1]) if "Reserve (stays with you" in kit else []
    ok(rc == 0 and res and all(int(x) == 52 for x in res), "kit checklist reserve %r (rc %d)" % (res, rc))
    p6c, _ = snap(base, T0 + dt.timedelta(days=0), tmp, "day1")
    tool("tools/make_messages.py", p6c, DID, "day3", os.path.join(tmp, "msg6c"))
    d3 = "\n".join(read(f) for f in glob.glob(os.path.join(tmp, "msg6c", "day3", "*.md")))
    ok("SAMPLE Brightside and SAMPLE Tallgrass have under five" in d3 and " and SAMPLE Tallgrass has under" not in d3, "day-3 grammar: %r" % d3[:400])

    # 7. Final state
    fin = DR.last_day(base) + dt.timedelta(days=1)
    p7, tr7 = snap(base, fin, tmp, "final")
    ok(not DR.dashboard_checks(p7)[0] and not DR.dashboard_checks(p7)[1], "final: checks %s %s" % DR.dashboard_checks(p7)[:2])
    for cid in (C7, C8):
        ok(tr7.ladder(cid)["reserve_on_hand"] == 0, "final: %s reserve %r" % (cid, tr7.ladder(cid)["reserve_on_hand"]))
    tool("tools/export_actuals.py", p7, os.path.join(tmp, "exp"))
    ex = {r["campaign_id"]: r for r in csv.DictReader(open(os.path.join(tmp, "exp", "campaign_actuals.csv"), encoding="utf-8"))}
    ok(ex[C7]["units_received"] == "543" and ex[C7]["units_delivered"] == "491", "export CMP-007 received %r delivered %r" % (ex[C7]["units_received"], ex[C7]["units_delivered"]))
    ok(ex[C8]["units_received"] == "522" and ex[C8]["offices"] == "9" and ex[C8]["survey_responses"] == "61",
       "export CMP-008 received %r offices %r responses %r" % (ex[C8]["units_received"], ex[C8]["offices"], ex[C8]["survey_responses"]))

    # 8. Roster rules
    s8 = copy.deepcopy(base)
    s8["rows"]["inventory"] = [r for r in s8["rows"]["inventory"] if not (r.get("office_id") == "SAMPLE-O-006" and r["campaign_id"] in (C7, C8)
                                                                          and r["movement"] in ("received", "delivered"))]
    s8["rows"]["counts"] = [r for r in s8["rows"]["counts"] if not (r["office_id"] == "SAMPLE-O-006" and r["campaign_id"] in (C7, C8))]
    s8["rows"]["photos"] = [r for r in s8["rows"]["photos"] if not (r["office_id"] == "SAMPLE-O-006" and r["campaign_id"] in (C7, C8))]
    for r in s8["rows"]["stops"]:
        if r["drop_id"] == DID and r["office_id"] == "SAMPLE-O-006":
            r["outcome"] = "declined"
    for r in s8["rows"]["in_signed_offices"]:
        if r["office_id"] == "SAMPLE-O-001":
            r["in_coworking"] = "Y"
        if r["office_id"] == "SAMPLE-O-002":
            r["in_boundary"] = "N"
    for r in s8["rows"]["offices"]:
        if r["office_id"] == "SAMPLE-O-002":
            r["eligibility_confirmed"] = "Y"
    for r in s8["rows"]["campaigns"]:
        if r["campaign_id"] == C8:
            r["brand_id"] = "SAMPLE-BR-005"   # SAMPLE Kettle Ridge: chips, like Brightside
    p8, _ = snap(s8, T0, tmp, "roster")
    o6, o8 = row(p8, "offices", office_id="SAMPLE-O-006"), row(p8, "offices", office_id="SAMPLE-O-008")
    ok(o6.get("declines") == 2 and "declined every brand at the door twice" in (o6.get("checks") or ""), "two full declines: %r %r" % (o6.get("declines"), o6.get("checks")))
    ok(o8.get("declines") == 0, "a one-brand refusal counted as a decline: %r" % o8.get("declines"))
    o1, o2 = row(p8, "offices", office_id="SAMPLE-O-001"), row(p8, "offices", office_id="SAMPLE-O-002")
    ok((o1.get("checks") or "").startswith("in a coworking space"), "coworking office: %r" % o1.get("checks"))
    ok((o2.get("checks") or "").startswith("info: outside the cluster boundary"), "outside the boundary, confirmed: %r" % o2.get("checks"))
    ok(DR.dashboard_checks(p8)[0].get("Offices with roster checks (declined twice, coworking, outside the boundary)") == 2,
       "dashboard roster line %s" % DR.dashboard_checks(p8)[0])
    ok("another campaign on this drop is also chips-crisps" in (row(p8, "campaigns", campaign_id=C7).get("checks") or ""), "category clash: no check")

    # 9. Replacing the office_form tab keeps the capacity warning alive
    import build_tracker as B
    import recalc as RC
    p9 = os.path.join(tmp, "forms.xlsx")
    B.build(p9, today=dt.date(2027, 1, 5), config=DR.SAMPLE_CONFIG)
    wb = openpyxl.load_workbook(p9)
    direct = []
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("="):
                    for m in re.finditer(r"(office_form|brand_form)!", c.value):
                        if c.value[m.start() - 1] != '"':
                            direct.append("%s!%s" % (ws.title, c.coordinate))
    ok(not direct, "formulas that refer to office_form or brand_form directly: %s" % direct[:5])
    hdr = [c.value for c in wb["office_form"][1]]
    idx = wb.sheetnames.index("office_form")
    wb.remove(wb["office_form"])
    ws = wb.create_sheet("Form Responses 1", idx)
    ws.append(hdr)
    for k in range(int(0.9 * T.N["office_form"]) + 1):
        ws.append([dt.datetime(2027, 1, 4, 9, 0)] + ["SAMPLE"] * (len(hdr) - 1))
    ws.title = "office_form"
    wb.save(p9)
    r9 = RC.recalc(p9, os.path.join(tmp, "forms-r.xlsx"))
    d9 = DR.dashboard_checks(r9)[0]
    ok(d9.get("Tabs over 90% full (rows below the last formula row are not counted)") == 1, "office_form replaced and over 90%%: dashboard %s" % d9)

    for f in fails:
        print("  FAIL " + f)
    print("PASS test_direct" if not fails else "FAIL test_direct")
    shutil.rmtree(tmp, ignore_errors=True)
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
