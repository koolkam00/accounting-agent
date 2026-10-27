#!/usr/bin/env python3
"""Load this folder's SAMPLE exports into a COPY of the financial model and check that its outputs move sensibly.

  python3 tests/test_h_import.py [exports/SAMPLE-zero exports/SAMPLE-standard exports/SAMPLE-exception exports/SAMPLE-direct]

For each SAMPLE export folder (campaign_actuals.csv, drop_costs.csv, hours_log.csv; the Pipeline tab is the outreach
workstream's file only since review round 1, R1-06, so nothing of this folder's is pasted there):
 1. copy financial-model/<workbook>.xlsx to _work/h-import/SAMPLE-model-<scenario>.xlsx (H's file is never written:
    its SHA-256 is checked before and after);
 2. paste each file with its header row at A4 of Actuals_Campaigns, Actuals_DropCosts and Hours_Log, dates
    as dates and amounts as numbers (financial-model/HANDOFF.md, R3-01), the SAMPLE drop IDs mapped to the model's
    Schedule keys (SAMPLE-DRP-001 -> DRP-001) in this copy only; set Actuals in = Y for those drops and the as-of date;
 3. recalculate the untouched copy and the loaded copy (LibreOffice), and compare:
    - H's own checks: 0 formula errors; the paste checks (headers, text, drop_id not on Schedule) read 0;
    - Drop_Actuals for each drop equals the CSVs: revenue, direct costs, founder hours, placed-within-1bd share,
      campaigns under 5 responses per office, late reports, brands that would buy again;
    - contribution: H subtracts its planned overhead allocation (column I), the tracker's dashboard does not; the
      log shows both and the difference;
    - the Dashboard and gate rows before and after.
 4. A negative control: the same files pasted with their SAMPLE- drop IDs unmapped are flagged ("not on Schedule")
    and move no drop.
Writes _work/h-import/h-import-log.md. Exit 1 on any failure.
"""
import csv
import datetime as dt
import hashlib
import os
import shutil
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DESK = os.path.dirname(ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))
import opscal  # noqa: E402
import recalc as RC  # noqa: E402

# The model's file name carries the working name; find it by pattern so no source file here spells the name.
H = sorted(p for p in __import__("glob").glob(os.path.join(DESK, "financial-model", "* financial model v*.xlsx")) if "SAMPLE" not in p)[-1]
OUT = os.path.join(ROOT, "_work", "h-import")
TABS = [("campaign_actuals.csv", "Actuals_Campaigns"), ("drop_costs.csv", "Actuals_DropCosts"), ("hours_log.csv", "Hours_Log")]
DATES = {"drop_date", "invoice_paid_date", "report_delivered_date", "deposit_paid_date", "date"}
TEXT = {"campaign_id", "drop_id", "delivery_model", "tier", "brand_id", "brand_would_buy_again", "renewed", "notes", "activity", "note",
        "side", "org_id", "org_name", "event", "value_text", "controls_budget"}
AS_OF = {"zero": dt.date(2027, 2, 12), "standard": dt.date(2027, 3, 12), "exception": dt.date(2027, 3, 26), "direct": dt.date(2027, 4, 26)}
DROP_ROW = {"DRP-001": 7, "DRP-002": 8, "DRP-003": 9, "DRP-004": 10}


def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def conv(col, v):
    if v == "":
        return None
    if col in DATES:
        return dt.date.fromisoformat(v)
    if col in TEXT:
        return v
    return float(v) if "." in v else int(v)


def load(folder):
    return {fn: list(csv.DictReader(open(os.path.join(folder, fn), encoding="utf-8"))) for fn, _ in TABS}, \
        {fn: next(csv.reader(open(os.path.join(folder, fn), encoding="utf-8"))) for fn, _ in TABS}


def mapped(v, keep):
    return v if keep or not v else v.replace("SAMPLE-DRP-", "DRP-")


def build(path, data, hdrs, as_of, keep_sample_ids=False):
    shutil.copy2(H, path)
    wb = openpyxl.load_workbook(path)
    wb["Dashboard"]["A4"] = "SAMPLE DATA. A scratch copy loaded with SAMPLE exports from campaign-ops. Illustrative placeholders, not results."
    for fn, tab in TABS:
        ws = wb[tab]
        for j, h in enumerate(hdrs[fn], start=1):
            ws.cell(row=4, column=j, value=h)
        for i, r in enumerate(data[fn], start=5):
            for j, h in enumerate(hdrs[fn], start=1):
                v = conv(h, r[h])
                if h == "drop_id":
                    v = mapped(v, keep_sample_ids)
                ws.cell(row=i, column=j, value=v)
    ref = wb["Inputs"]
    ref["C16"] = as_of
    drops = {mapped(r["drop_id"], keep_sample_ids) for r in data["drop_costs.csv"]}
    for d in drops:
        if d in DROP_ROW:
            wb["Schedule"]["N%d" % DROP_ROW[d]] = "Y"
    wb.save(path)


def recalc(path, out):
    RC.recalc(path, out)
    return openpyxl.load_workbook(out, data_only=True)


def checks(wb):
    ck = wb["Checks"]
    rows = {}
    for r in range(5, 200):
        if ck["A%d" % r].value is None:
            break
        rows[ck["C%d" % r].value] = (ck["D%d" % r].value, ck["E%d" % r].value)
    return rows


def errors(wb):
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(RC.ERRORS):
                    n += 1
    return n


def expected(data):
    exp = {}
    camps, costs, hours = data["campaign_actuals.csv"], data["drop_costs.csv"], data["hours_log.csv"]
    for dc in costs:
        d = mapped(dc["drop_id"], False)
        cs = [c for c in camps if mapped(c["drop_id"], False) == d]
        T = dt.date.fromisoformat(dc["drop_date"])
        due = opscal.add_bd(T, 7)
        deliv = sum(float(c["units_delivered"] or 0) for c in cs)
        exp[d] = {
            "revenue": sum(float(c["revenue_usd"] or 0) for c in cs),
            "processing": sum(float(c["processing_fee_usd"] or 0) for c in cs),
            "direct": sum(float(dc[k] or 0) for k in ("kit_cost_usd", "transport_cost_usd", "other_direct_cost_usd")),
            "hours": sum(float(h["hours"]) for h in hours if mapped(h["drop_id"], False) == d),
            "share": (sum(float(c["units_placed_within_1bd"] or 0) for c in cs) / deliv) if deliv else None,
            "under5": sum(1 for c in cs if c["survey_responses_min_office"] != "" and float(c["survey_responses_min_office"]) < 5),
            "late": sum(1 for c in cs if not c["report_delivered_date"] or dt.date.fromisoformat(c["report_delivered_date"]) > due),
            "rebuy": sum(1 for c in cs if c["brand_would_buy_again"] == "yes"),
        }
    return exp


def near(a, b, tol=0.006):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return a == b


def fmt(v):
    if isinstance(v, float):
        return ("%.4f" % v).rstrip("0").rstrip(".")
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    return "" if v is None else str(v)


def run(folder, log, fails):
    name = os.path.basename(folder.rstrip("/")).replace("SAMPLE-", "")
    data, hdrs = load(folder)
    as_of = AS_OF.get(name, dt.date(2027, 3, 12))
    base_p, load_p = os.path.join(OUT, "SAMPLE-model-%s-before.xlsx" % name), os.path.join(OUT, "SAMPLE-model-%s.xlsx" % name)
    shutil.copy2(H, base_p)
    wb = openpyxl.load_workbook(base_p)
    wb["Dashboard"]["A4"] = "SAMPLE DATA. Untouched scratch copy, as-of date set only."
    wb["Inputs"]["C16"] = as_of
    wb.save(base_p)
    before = recalc(base_p, base_p.replace(".xlsx", "-recalculated.xlsx"))
    build(load_p, data, hdrs, as_of)
    after = recalc(load_p, load_p.replace(".xlsx", "-recalculated.xlsx"))
    ck = checks(after)
    log += ["## %s (as of %s)" % (name, as_of), "",
            "Pasted from `%s`: %s." % (os.path.relpath(folder, ROOT), ", ".join("%s %d rows" % (fn, len(data[fn])) for fn, _ in TABS)), ""]
    nerr = errors(after)
    if nerr:
        fails.append("%s: %d formula error cells after loading" % (name, nerr))
    paste_checks = [k for k in ck if k and (k.startswith("Actuals_") or k.startswith("Hours_Log") or k.startswith("Pipeline"))]
    bad = [(k, ck[k]) for k in paste_checks if ck[k][1] not in (1, "info")]
    for k, v in bad:
        fails.append("%s: H check failing after load: %s = %s" % (name, k, v[0]))
    log.append("H's own checks after loading: %d formula errors; %d paste checks, %d failing%s." % (
        nerr, len(paste_checks), len(bad), "" if not bad else " (" + "; ".join("%s = %s" % (k, v[0]) for k, v in bad) + ")"))
    allfail = [(k, v) for k, v in ck.items() if v[1] not in (1, "info")]
    log.append("All H checks failing after loading: %s." % ("none" if not allfail else "; ".join("%s = %s" % (k, v[0]) for k, v in allfail)))
    log.append("")
    exp = expected(data)
    da_b, da_a = before["Drop_Actuals"], after["Drop_Actuals"]
    log += ["| Drop_Actuals | drop | before | after | from the CSVs | match |", "| --- | --- | --- | --- | --- | --- |"]
    for d, e in exp.items():
        r = DROP_ROW[d]
        cells = [("Actuals in", "C", None), ("Revenue", "D", e["revenue"]), ("Card processing", "E", e["processing"]),
                 ("Kit + transport + other direct", None, e["direct"]), ("Overhead allocations (planned)", "I", None),
                 ("Founder hours", "R", e["hours"]), ("Contribution (H: after allocations)", "S", None),
                 ("Contribution per founder hour (H)", "T", None), ("Placed within 1 bd, share", "AG", e["share"]),
                 ("Campaigns under 5 per office", "AH", e["under5"]), ("Reports late or missing", "AI", e["late"]),
                 ("Brands that would buy again", "AJ", e["rebuy"])]
        for label, col, want in cells:
            if col is None:
                b = sum(float(da_b["%s%d" % (x, r)].value or 0) for x in "FGH")
                a = sum(float(da_a["%s%d" % (x, r)].value or 0) for x in "FGH")
            else:
                b, a = da_b["%s%d" % (col, r)].value, da_a["%s%d" % (col, r)].value
            ok = "" if want is None else ("yes" if near(a, want) else "NO")
            if ok == "NO":
                fails.append("%s %s %s: H %r, CSVs %r" % (name, d, label, a, want))
            log.append("| %s | %s | %s | %s | %s | %s |" % (label, d, fmt(b), fmt(a), fmt(want), ok))
        alloc = float(da_a["I%d" % r].value or 0)
        f_contrib = e["revenue"] - e["direct"] - e["processing"]
        h_contrib = da_a["S%d" % r].value
        if not near(h_contrib, f_contrib - alloc, 0.01):
            fails.append("%s %s: H contribution %r is not revenue - direct - processing - allocation (%r)" % (name, d, h_contrib, f_contrib - alloc))
        log.append("| Tracker dashboard contribution (before overhead) | %s | | %s | %s per founder hour | H lower by its planned allocation %s |" % (
            d, fmt(f_contrib), fmt(round(f_contrib / e["hours"], 2) if e["hours"] else None), fmt(alloc)))
    log.append("")
    dash_b, dash_a = before["Dashboard"], after["Dashboard"]
    log += ["| Dashboard | before | after |", "| --- | --- | --- |"]
    for label, ref in (("Next scheduled drop", "C9"), ("Standard runs to Gate 3: contribution per founder hour (plan / actual)", "D13"),
                       ("Lowest point of cumulative cash", "C15"), ("Six months all-in per founder hour, roadmap drops", "C21"),
                       ("Cumulative cash at week 26, roadmap drops", "C23"), ("Brand renewal rate", "C64"), ("The rule says", "C68")):
        log.append("| %s | %s | %s |" % (label, fmt(dash_b[ref].value), fmt(dash_a[ref].value)))
    log += ["", "| Gate row | live value before | after | status before | after |", "| --- | --- | --- | --- | --- |"]
    for r in range(46, 56):
        log.append("| %s %s | %s | %s | %s | %s |" % (dash_a["A%d" % r].value, dash_a["B%d" % r].value, fmt(dash_b["F%d" % r].value),
                                                      fmt(dash_a["F%d" % r].value), fmt(dash_b["G%d" % r].value), fmt(dash_a["G%d" % r].value)))
    for r in range(58, 62):
        log.append("| %s overall | | | %s | %s |" % (dash_a["A%d" % r].value, fmt(dash_b["C%d" % r].value), fmt(dash_a["C%d" % r].value)))
    log.append("")
    # Sensible movement: the gate values must now be the CSVs' numbers.
    g2 = exp.get("DRP-001")
    if g2 and not near(dash_a["F49"].value, g2["share"]):
        fails.append("%s: Gate 2 placed share on the Dashboard %r, CSVs %r" % (name, dash_a["F49"].value, g2["share"]))
    if g2 and not near(dash_a["F52"].value, g2["rebuy"]):
        fails.append("%s: Gate 2 buy-again %r, CSVs %r" % (name, dash_a["F52"].value, g2["rebuy"]))
    # Gate 1 (brands paid in full) comes from the Pipeline tab, which is E's file alone: it must not move here.
    if not near(dash_a["F48"].value, dash_b["F48"].value):
        fails.append("%s: Gate 1 brands paid in full moved (%r to %r) although nothing was pasted into Pipeline" % (name, dash_b["F48"].value, dash_a["F48"].value))
    # Negative control: SAMPLE- drop IDs left as they are.
    neg_p = os.path.join(OUT, "SAMPLE-model-%s-unmapped.xlsx" % name)
    build(neg_p, data, hdrs, as_of, keep_sample_ids=True)
    neg = recalc(neg_p, neg_p.replace(".xlsx", "-recalculated.xlsx"))
    nck = checks(neg)
    flagged = nck.get("Actuals_Campaigns: rows whose drop_id is not on Schedule", (0,))[0]
    moved = [d for d, r in DROP_ROW.items() if neg["Drop_Actuals"]["D%d" % r].value not in (None, "", 0)]
    if not flagged or moved:
        fails.append("%s: unmapped SAMPLE drop IDs were not flagged (%r) or moved a drop (%s)" % (name, flagged, moved))
    log.append("Negative control (SAMPLE- drop IDs pasted as they are): H flags %s campaign rows, %s cost rows and %s hour rows as not on Schedule, and no drop moves (%s)." % (
        flagged, nck.get("Actuals_DropCosts: rows whose drop_id is not on Schedule", ("?",))[0],
        nck.get("Hours_Log: drop_ids not on Schedule", ("?",))[0], "none moved" if not moved else "moved: " + ", ".join(moved)))
    log.append("")
    for p in (neg_p, neg_p.replace(".xlsx", "-recalculated.xlsx")):
        os.remove(p)


def main(argv):
    folders = argv[1:] or [os.path.join(ROOT, "exports", "SAMPLE-zero"), os.path.join(ROOT, "exports", "SAMPLE-standard")]
    os.makedirs(OUT, exist_ok=True)
    h0 = sha(H)
    log = ["# SAMPLE exports loaded into a copy of the financial model", "",
           "SAMPLE DATA. Illustrative placeholders, not results. Made by `python3 tests/test_h_import.py` on %s." % dt.date.today(),
           "H's workbook (`financial-model/%s`) is copied, never written: SHA-256 before %s." % (os.path.basename(H), h0[:16]),
           "The SAMPLE drop IDs are mapped to the model's Schedule keys in the copy only (SAMPLE-DRP-001 to DRP-001).", ""]
    fails = []
    for f in folders:
        run(f if os.path.isabs(f) else os.path.join(ROOT, f), log, fails)
    h1 = sha(H)
    if h1 != h0:
        fails.append("H's workbook changed!")
    log.append("H's workbook SHA-256 after: %s (%s)." % (h1[:16], "unchanged" if h1 == h0 else "CHANGED"))
    log.append("")
    log.append("Result: %s" % ("pass" if not fails else "FAIL: " + "; ".join(fails)))
    open(os.path.join(OUT, "h-import-log.md"), "w", encoding="utf-8").write("\n".join(log) + "\n")
    for f in fails:
        print("  FAIL " + f)
    print("log: " + os.path.join(OUT, "h-import-log.md"))
    print("PASS test_h_import" if not fails else "FAIL test_h_import")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
