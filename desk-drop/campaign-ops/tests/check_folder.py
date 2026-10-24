#!/usr/bin/env python3
"""Folder rules (CONTRACTS sections 1 and 2, and the brief's non-negotiables).

  python3 tests/check_folder.py

1. Rename rule: the working name is in no file name, folder name or source file here. It may appear
   only in generated outputs (they read the one config cell) and in the tracker's config cell.
2. Ladder words: no field, column or sentence calls anything conversions, sales lift, ROI or
   acquisition, except where a sentence forbids exactly that.
3. SAMPLE rule: in every SAMPLE tracker, every ID and every free-text value starts with SAMPLE, and
   every SAMPLE output document carries the SAMPLE banner: the HTML pages, and (review round 1, R1-10) every text
   draft (.md, .txt) under dry-run/, whose first lines must hold the banner line. A CSV keeps CONTRACTS section 5's
   one header row (a banner row would break the paste into the financial model), so every data row of a dry-run or
   SAMPLE export CSV must instead carry a SAMPLE value.
4. Required files: README.md and HANDOFF.md exist (REPORT.md is saved by the orchestrator).
5. No email address other than example.com / placeholders in any source file.
"""
import glob
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DESK = os.path.dirname(ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))
import trackerdef as T  # noqa: E402

NAME = json.load(open(os.path.join(DESK, "desk-drop-identity", "brand.json")))["BRAND_NAME"]
SOURCE_DIRS = ["tools", "templates", "forms", "spec", "sop", "links", "runbook", "tests"]
SOURCE_TOP = ["README.md", "HANDOFF.md"]
GENERATED = ("dry-run", "_work", "tracker", "exports")
BANNED = re.compile(r"\b(conversions?|sales lift|ROI|acquisition)\b", re.I)
ALLOW = re.compile(r"never|not |no field|no column|do not|don't|refus|claims list|forbid|nothing here|without|instead of|BANNED|ALLOW", re.I)
EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
TEXT_COLS = re.compile(r"(_id$|^company$|^brand_name$|^note|^notes$|what_happened|action_taken|^comment_|^learned$|^recommendation$|card_sentences|allergen_line|may_contain|product_sku|retailer_line|file_name|delivery_point|reply_note|renewal_note|receiving_location|courier_or_van|backup_person|vendor|^lot$|bank_reference|^industry|contact_name|contact_title|pantry_owner_name|pantry_owner_title|office_address|floor_suite|brand_legal_name|billing_address|hq_city|nyc_retailers|q7_why_now|budget_probe_verbatim|^discount_code$|reported_by_role|sent_to)")


def main():
    fails = []
    for p in ("README.md", "HANDOFF.md"):
        if not os.path.exists(os.path.join(ROOT, p)):
            fails.append("missing " + p)
    low = NAME.lower()
    squash = low.replace(" ", "")
    for dirpath, dirs, files in os.walk(ROOT):
        if "node_modules" in dirpath:
            continue
        for n in dirs + files:
            if low in n.lower() or squash in n.lower().replace("-", "").replace("_", ""):
                fails.append("name in path: " + os.path.relpath(os.path.join(dirpath, n), ROOT))
    srcs = [os.path.join(ROOT, f) for f in SOURCE_TOP if os.path.exists(os.path.join(ROOT, f))]
    for d in SOURCE_DIRS:
        srcs += [p for p in glob.glob(os.path.join(ROOT, d, "**", "*"), recursive=True) if os.path.isfile(p)]
    for p in srcs:
        if p.endswith((".png", ".pdf", ".xlsx", ".pyc")) or "__pycache__" in p:
            continue
        rel = os.path.relpath(p, ROOT)
        txt = open(p, encoding="utf-8", errors="replace").read()
        for i, line in enumerate(txt.splitlines(), start=1):
            if low in line.lower() and rel != "tests/check_folder.py":
                fails.append("name in source %s:%d" % (rel, i))
            if BANNED.search(line) and not ALLOW.search(line):
                fails.append("ladder word in %s:%d: %s" % (rel, i, line.strip()[:100]))
            for m in EMAIL.findall(line):
                if not (m.lower().endswith("example.com") or m.startswith("noreply") or "@campaign-ops" in m or m.endswith("fda.gov")):
                    fails.append("email address in %s:%d: %s" % (rel, i, m))
    T.define()
    for x in glob.glob(os.path.join(ROOT, "dry-run", "*", "SAMPLE-tracker-*.xlsx")):
        wb = openpyxl.load_workbook(x, data_only=True)
        for tab, t in T.TABS.items():
            ws = wb[tab]
            for col in t.cols:
                if col.kind == "A" or not TEXT_COLS.search(col.name):
                    continue
                ci = t.index[col.name] + 1
                for r in range(t.first(), ws.max_row + 1):
                    v = ws.cell(row=r, column=ci).value
                    if isinstance(v, str) and v.strip() and not v.startswith("SAMPLE") and col.name not in ("form_version",):
                        fails.append("%s %s!%s%d not SAMPLE: %r" % (os.path.basename(x), tab, T.colletter(ci), r, v[:40]))
        for name in wb.sheetnames:
            if name.startswith("SAMPLE-CMP-"):
                ws = wb[name]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[0] and row[6] and not str(row[6]).startswith("SAMPLE"):
                        fails.append("%s %s comment not SAMPLE: %r" % (os.path.basename(x), name, row[6]))
                    if row[0] and row[1] and not str(row[1]).startswith("SAMPLE"):
                        fails.append("%s %s card code not SAMPLE" % (os.path.basename(x), name))
    for h in glob.glob(os.path.join(ROOT, "dry-run", "**", "*.html"), recursive=True):
        if "SAMPLE" not in open(h, encoding="utf-8").read()[:20000] or "class=\"sample\"" not in open(h, encoding="utf-8").read():
            fails.append("SAMPLE output without the banner: " + os.path.relpath(h, ROOT))
    banner = "SAMPLE. Test data only: illustrative placeholders, not results."
    texts = 0
    for t in glob.glob(os.path.join(ROOT, "dry-run", "**", "*.md"), recursive=True) + glob.glob(os.path.join(ROOT, "dry-run", "**", "*.txt"), recursive=True):
        texts += 1
        head = "\n".join(open(t, encoding="utf-8").read().splitlines()[:3])
        if banner not in head:
            fails.append("SAMPLE text draft without the banner line: " + os.path.relpath(t, ROOT))
    csvs = 0
    for c in (glob.glob(os.path.join(ROOT, "dry-run", "**", "*.csv"), recursive=True) +
              glob.glob(os.path.join(ROOT, "exports", "SAMPLE-*", "*.csv"))):
        csvs += 1
        for i, row in enumerate(open(c, encoding="utf-8").read().splitlines()[1:], start=2):
            if row.strip() and "SAMPLE" not in row:
                fails.append("CSV row without a SAMPLE value: %s:%d" % (os.path.relpath(c, ROOT), i))
                break
    print("%d text drafts and %d CSV files under dry-run/ and exports/SAMPLE-* checked for SAMPLE marking" % (texts, csvs))
    for f in fails[:60]:
        print("FAIL " + f)
    print("%d problems" % len(fails))
    print("PASS check_folder" if not fails else "FAIL check_folder")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
