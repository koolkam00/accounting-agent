#!/usr/bin/env python3
"""The exceptions the SAMPLE stories used to skip, run through the tracker and the tools (review round 1).

  python3 tests/test_exception.py

From the exception scenario (tools/scenarios.py exception(): SAMPLE Saltmarsh, 10 offices, courier, drop
Tue 2027-02-23, a short shipment and SAMPLE-O-006 declining at the door), each case builds a tracker as the
founder would have it on one day, recalculates it with LibreOffice and requires:

 1. T+7 as logged: 0 formula errors, 0 disagreements, 0 dashboard checks; the office that declined is off the drop
    (on_drop N, no letter); the report generates with a roster note; no day-3 or day-5 ask goes to that office;
    the report email is drafted only with --reports and a PDF, and refused without one; that office gets the
    "turned away" follow-up, not the usual debrief; the export says offices 9, units_depleted_day5 408 (not blank),
    survey_responses_min_office 5 or more.                                                        (R1-02)
 2. T+1 without the credit note: the campaign's checks say "credit due from T+1".                 (R1-02)
 3. T with a delivered row logged for the office that declined: the stop's checks say so.         (R1-03)
 4. T-5 with the short shipment not yet handled: "received 480 of ..." and "reduce units_override". (R1-09)
 5. A brand that fails a qualification check: ship release on HOLD (qualification), a campaign check, and the
    kit refuses; with its billing details blank, make_invoice.py refuses (exit 2).                 (R1-07)
 6. A rename in the tracker only (config!B2 changed, brand.json not): every generator refuses (exit 2). (R1-04)
 7. The founding tier: a fourth founding brand and a third founding_lock campaign each raise a check. (R1-14)
Writes only to a temporary folder.
"""
import copy
import csv
import datetime as dt
import glob
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "tools"))
import dryrun as DR  # noqa: E402
import scenarios as S  # noqa: E402
import tracker as TR  # noqa: E402
import trackerdef as T  # noqa: E402

CID, DID, OFF = "SAMPLE-CMP-006", "SAMPLE-DRP-003", "SAMPLE-O-006"
fails = []


def ok(cond, msg):
    if not cond:
        fails.append(msg)
    return cond


def snap(scn, day, tmp, tag, config=None):
    p = DR.build_snapshot(scn, day, os.path.join(tmp, "in-%s.xlsx" % tag))
    if config:
        wb = openpyxl.load_workbook(p)
        for k, v in config.items():
            wb["config"].cell(row=T.CFG_ROW[k], column=2, value=v)
        wb.save(p)
    out = os.path.join(tmp, "snap-%s.xlsx" % tag)
    tr, errs, n = DR.recalc_and_read(p, out)
    ok(n == 0, "%s: %d formula errors %s" % (tag, n, errs[:3]))
    return out, tr


def row(path, tab, **key):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[tab]
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        if all(d.get(k) == v for k, v in key.items()):
            return d
    return {}


def tool(*args):
    r = subprocess.run([sys.executable] + list(args), capture_output=True, text=True, cwd=ROOT, timeout=600)
    return r.returncode, r.stdout + r.stderr


def main():
    tmp = tempfile.mkdtemp(prefix="ops_exc_")
    base = S.exception()
    T0 = base["T"]
    import opscal
    sd = opscal.sop_dates(T0)

    # 1. T+7 as logged
    p, tr = snap(base, sd["report_due"], tmp, "t7")
    bad = [b for b in TR.compare(tr) if not b.startswith("no cached")]
    ok(not bad, "T+7: disagreements %s" % bad[:3])
    checks, msgs, rel = DR.dashboard_checks(p)
    ok(not checks, "T+7: dashboard checks %s %s" % (checks, msgs[:3]))
    a = row(p, "approvals", campaign_id=CID, office_id=OFF)
    ok(a.get("on_drop") == "N" and not a.get("office_label"), "T+7: declined office on_drop=%r label=%r" % (a.get("on_drop"), a.get("office_label")))
    c = row(p, "campaigns", campaign_id=CID)
    ok(c.get("offices_on_drop") == 9 and c.get("units_depleted_day5") == 408 and (c.get("survey_responses_min_office") or 0) >= 5,
       "T+7: campaign on_drop=%r depleted=%r min_office=%r" % (c.get("offices_on_drop"), c.get("units_depleted_day5"), c.get("survey_responses_min_office")))
    ok(c.get("gate2_responses") == "pass", "T+7: Gate 2 responses %r" % c.get("gate2_responses"))
    photos = os.path.join(tmp, "photos-root")
    DR.make_photos(base, photos)
    rc, out = tool("tools/make_report.py", p, CID, os.path.join(tmp, "reports"), "--photos", os.path.join(photos, "photos"))
    ok(rc == 0 and os.path.exists(os.path.join(tmp, "reports", CID, "report.pdf")), "T+7: report refused: %s" % out[-400:])
    trace = open(os.path.join(tmp, "reports", CID, "trace.csv"), encoding="utf-8").read()
    ok("roster note" in trace and "declined it at the door" in trace, "T+7: no roster note in trace.csv")
    ok(not os.path.exists(os.path.join(tmp, "reports", "office-summary-%s.md" % OFF)), "T+7: an office summary went to the office that declined")
    for step in ("day3", "day5"):
        tool("tools/make_messages.py", p, DID, step, os.path.join(tmp, "msg"))
        ok(not os.path.exists(os.path.join(tmp, "msg", step, OFF + ".md")), "T+7: a %s ask went to the office that declined" % step)
        ok(len(glob.glob(os.path.join(tmp, "msg", step, "*.md"))) == 9, "T+7: %s drafts for the nine offices on the drop" % step)
    rc, out = tool("tools/make_messages.py", p, DID, "report", os.path.join(tmp, "msg"), "--reports", os.path.join(tmp, "reports"))
    body = open(os.path.join(tmp, "msg", "report", CID + ".md"), encoding="utf-8").read() if rc == 0 else ""
    ok(rc == 0 and "declined it at the door" in body, "T+7: report email with the roster note (%d) %s" % (rc, out[-200:]))
    rc, out = tool("tools/make_messages.py", p, DID, "report", os.path.join(tmp, "msg2"), "--reports", os.path.join(tmp, "no-reports"))
    ok(rc == 2 and os.path.exists(os.path.join(tmp, "msg2", "report", CID + "-NO-REPORT.md")), "T+7: report email without a PDF not refused (%d)" % rc)
    rc, out = tool("tools/make_messages.py", p, DID, "report", os.path.join(tmp, "msg3"))
    ok(rc == 1, "T+7: report email without --reports should be a usage error (%d)" % rc)
    tool("tools/make_debrief.py", p, DID, os.path.join(tmp, "deb"))
    deb = open(os.path.join(tmp, "deb", "debrief-office-%s.md" % OFF), encoding="utf-8").read()
    ok("turned away at the door" in deb and "How did it land" not in deb, "T+7: the declined office got the usual debrief")
    tool("tools/export_actuals.py", p, os.path.join(tmp, "exp"))
    ex = {r["campaign_id"]: r for r in csv.DictReader(open(os.path.join(tmp, "exp", "campaign_actuals.csv"), encoding="utf-8"))}
    e = ex.get(CID, {})
    ok(e.get("offices") == "9" and e.get("units_depleted_day5") == "408" and int(e.get("survey_responses_min_office") or 0) >= 5,
       "T+7: export offices=%r depleted=%r min=%r" % (e.get("offices"), e.get("units_depleted_day5"), e.get("survey_responses_min_office")))
    rc, out = tool("tools/export_actuals.py", p, os.path.join(tmp, "exp2"), "--pipeline")
    ok(rc == 1 and "withdrawn" in out, "--pipeline should be refused as withdrawn (%d)" % rc)

    # 2. T+1 without the credit note
    s2 = copy.deepcopy(base)
    s2["rows"]["invoices"] = [r for r in s2["rows"]["invoices"] if r["invoice_id"] != "SAMPLE-INV-022"]
    p2, _ = snap(s2, sd["placement_confirm"], tmp, "t1-nocredit")
    ok("credit due from T+1" in (row(p2, "campaigns", campaign_id=CID).get("checks") or ""), "T+1 without the credit note: no credit check")
    p2b, _ = snap(base, sd["placement_confirm"], tmp, "t1-credit")
    ok("credit due" not in (row(p2b, "campaigns", campaign_id=CID).get("checks") or ""), "T+1 with the credit note: credit check still shows")

    # 3. T with a delivered row logged for the office that declined
    s3 = copy.deepcopy(base)
    s3["rows"]["inventory"].append({"_on": T0, "date": T0, "time": dt.time(8, 23), "campaign_id": CID, "office_id": OFF,
                                    "movement": "delivered", "qty": 40, "count_source": "courier_count"})
    p3, _ = snap(s3, T0, tmp, "t-wrongdelivered")
    st = row(p3, "stops", drop_id=DID, office_id=OFF)
    ok("every brand here is logged delivered" in (st.get("checks") or ""), "T: declined stop with a delivered row: no check (%r)" % st.get("checks"))

    # 4. T-5, short shipment not handled
    s4 = copy.deepcopy(base)
    s4["rows"]["issues"] = [r for r in s4["rows"]["issues"] if r.get("issue_type") != "short_shipment"]
    for r in s4["rows"]["approvals"]:
        if r["campaign_id"] == CID:
            r.pop("units_override", None)
            (r.get("_fields_on") or {}).pop("units_override", None)
    p4, _ = snap(s4, sd["receive"], tmp, "t5-short")
    ch = row(p4, "campaigns", campaign_id=CID).get("checks") or ""
    ok("reduce units_override" in ch and "log a short_shipment issue" in ch, "T-5 short shipment: checks %r" % ch)

    # 5. A brand that fails a check, and blank billing details
    s5 = copy.deepcopy(base)
    for r in s5["rows"]["brand_intake"]:
        if r.get("brand_name") == "SAMPLE Saltmarsh":
            r["q5_product_liability_insurance"] = "N"
    for r in s5["rows"]["brands"]:
        if r["brand_id"] == "SAMPLE-BR-006":
            r.pop("billing_address", None)
    p5, tr5 = snap(s5, sd["kit_print"], tmp, "t3-unqualified")
    c5 = row(p5, "campaigns", campaign_id=CID)
    ok(str(c5.get("ship_release")).startswith("HOLD") and "qualification" in str(c5.get("ship_release")), "unqualified brand: ship_release %r" % c5.get("ship_release"))
    ok("brand not qualified" in (c5.get("checks") or "") and "billing details missing" in (c5.get("checks") or ""), "unqualified brand: checks %r" % c5.get("checks"))
    ok(tr5.release_status(CID).startswith("HOLD"), "unqualified brand: the Python release says %r" % tr5.release_status(CID))
    rc, out = tool("tools/make_kit.py", p5, DID, os.path.join(tmp, "kit5"))
    ok(rc == 2 and "qualification" in out, "unqualified brand: the kit was not refused (%d)" % rc)
    rc, out = tool("tools/make_invoice.py", p5, "SAMPLE-INV-020", os.path.join(tmp, "inv5"))
    ok(rc == 2 and "billing_address" in out, "blank billing address: the invoice was not refused (%d) %s" % (rc, out[-200:]))

    # 6. A rename in the tracker only
    p6, _ = snap(base, sd["kit_print"], tmp, "rename", config={"company_name": "SAMPLE Newname"})
    for args in (["tools/make_kit.py", p6, DID, os.path.join(tmp, "kit6")], ["tools/make_route.py", p6, DID, os.path.join(tmp, "route6")],
                 ["tools/make_messages.py", p6, DID, "t3", os.path.join(tmp, "msg6")], ["tools/make_invoice.py", p6, "SAMPLE-INV-020", os.path.join(tmp, "inv6")]):
        rc, out = tool(*args)
        ok(rc == 2 and "the name differs" in out, "rename in one place only: %s was not refused (%d)" % (os.path.basename(args[0]), rc))

    # 7. Founding tier rules
    s7 = copy.deepcopy(base)
    extra = []
    for k, (bid, lock) in enumerate((("SAMPLE-BR-003", ""), ("SAMPLE-BR-004", ""), ("SAMPLE-BR-001", "Y"), ("SAMPLE-BR-001", "Y"), ("SAMPLE-BR-001", "Y"))):
        extra.append({"_on": dt.date(2027, 2, 1), "campaign_id": "SAMPLE-CMP-09%d" % k, "drop_id": DID, "brand_id": bid, "tier": "founding",
                      "founding_lock": lock, "offices_contracted": 5, "agreement_signed_date": dt.date(2027, 2, 1) + dt.timedelta(days=k)})
    s7["rows"]["campaigns"] += extra
    p7, _ = snap(s7, dt.date(2027, 2, 2), tmp, "founding")
    ok("first three brands only" not in (row(p7, "campaigns", campaign_id="SAMPLE-CMP-090").get("checks") or ""), "the third founding brand raises the four-brand check")
    ok("first three brands only" in (row(p7, "campaigns", campaign_id="SAMPLE-CMP-091").get("checks") or ""), "a fourth founding brand raises no check")
    ok("next two campaigns only" in (row(p7, "campaigns", campaign_id="SAMPLE-CMP-094").get("checks") or ""), "a third founding_lock campaign raises no check")

    for f in fails:
        print("  FAIL " + f)
    print("PASS test_exception" if not fails else "FAIL test_exception")
    shutil.rmtree(tmp, ignore_errors=True)
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
