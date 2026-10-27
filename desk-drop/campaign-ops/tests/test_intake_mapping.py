#!/usr/bin/env python3
"""The founder-side Google Form answers come out in exactly the CONTRACTS 6.3 values.

  python3 tests/test_intake_mapping.py dry-run/standard/SAMPLE-tracker-standard-final.xlsx

In the SAMPLE scenarios some offices and one brand use the founder-side forms (office_form, brand_form,
written as Google Forms would). The tracker's formula tabs office_intake_founder and brand_intake_founder
must turn them into the same values the scenario's handoff rows carry: keys not labels, Y/N not
"Yes, me", https:// added, the contact copied as pantry owner, categories as semicolon keys.
"""
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "tools"))
import scenarios as S  # noqa: E402
import trackerdef as T  # noqa: E402


def rows(ws):
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] not in (None, "")]


def main(argv):
    path = argv[1] if len(argv) > 1 else os.path.join(ROOT, "dry-run", "standard", "SAMPLE-tracker-standard-final.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    scn = S.standard()
    fails, n = [], 0
    want_off = {r["company"]: r for r in scn["rows"]["in_signed_offices"] if r["source"] == "founder"}
    # The form answer is what the office ticked at intake; the outreach file may later add a category
    # (the standard story adds one for SAMPLE-O-002), so categories are checked against the form answer.
    form_cats = {o[1]: o[8] for o in S.OFFICES}
    for got in rows(wb["office_intake_founder"]):
        w = want_off.get(got["company"])
        if not w:
            fails.append("unexpected founder office row %s" % got["company"])
            continue
        for col in T.OFFICE_INTAKE_COLS:
            a, b = got.get(col), w.get(col)
            if col == "categories_wanted":
                a, b = set(str(a).split(";")), set(str(form_cats.get(got["company"], b)).split(";"))
            if col == "submitted_at":
                b = w["submitted_at"]
            if (a in (None, "") and b in (None, "")) or str(a) == str(b) or a == b:
                n += 1
                continue
            fails.append("office %s %s: got %r want %r" % (got["company"], col, a, b))
    want_b = {S.BRANDS[k]["name"]: k for k in S.BRANDS if S.BRANDS[k]["via"] == "founder"}
    for got in rows(wb["brand_intake_founder"]):
        bid = want_b.get(got["brand_name"])
        if not bid:
            fails.append("unexpected founder brand row %s" % got["brand_name"])
            continue
        b = S.BRANDS[bid]
        checks = {"category": b["cat"], "website": b["site"], "source": "founder", "excluded_category_confirmed": "Y",
                  "budget_approver_is_contact": "Y", "q1_sealed_shelf_stable_60d": "Y", "q6_contact_can_approve_600": "Y",
                  "can_issue_discount_code": "Y", "submitted_at": b["submitted"], "form_version": "brand-gform-1.0",
                  "budget_approver_name_title": "SAMPLE Founder of %s, SAMPLE Co-founder" % b["name"][7:]}
        for k, v in checks.items():
            n += 1
            if str(got.get(k)) != v:
                fails.append("brand %s %s: got %r want %r" % (got["brand_name"], k, got.get(k), v))
    for f in fails:
        print("FAIL " + f)
    print("%d values checked" % n)
    print("PASS test_intake_mapping" if not fails else "FAIL test_intake_mapping")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
