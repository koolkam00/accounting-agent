#!/usr/bin/env python3
"""Write cluster-1-flatiron.xlsx: Companies, Buildings, Coworking, Retailers, Route, Walk,
Near-miss, Method. Route and Walk carry live formulas driven by editable parameter cells."""
import json
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build  # noqa: E402
import coldefs  # noqa: E402
from method_text import method_sections  # noqa: E402

HEAD = PatternFill("solid", fgColor="1F1D1A")
HEADF = Font(bold=True, color="FFFFFF")
PARAMF = PatternFill("solid", fgColor="FFF4C2")
NOTEF = Font(italic=True, color="6B675F")
TITLEF = Font(bold=True, size=14)
SECF = Font(bold=True, size=12, color="2E47C9")
URL_COLS = {"website", "headcount_source_url", "address_source_url", "second_source_url", "lobby_source_url", "coi_source_url",
            "pantry_source_url", "source_url", "operating_evidence_url", "linkedin_company_url"}
WIDTH = {"notes": 80, "verify_checklist": 70, "company_name": 34, "street_address": 30, "building_name": 34, "stops_here": 50,
         "headcount_unknown_reason": 50, "near_miss_reason": 45, "shelf_set": 40, "emerging_brand_evidence": 45,
         "pin_check": 40, "location_name": 34, "operator": 28, "retailer_name": 34, "start_point": 34, "end_point": 34}


def rd(f):
    return pd.read_csv(os.path.join(ROOT, f), dtype=str).fillna("")


def num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in str(v) else f
    except (TypeError, ValueError):
        return v


NUMERIC = {"lat", "lon", "blocks_outside", "candidate_company_count", "other_company_count", "candidates_within_5_blocks",
           "nearest_candidate_blocks", "pluto_numfloors", "pluto_yearbuilt", "pluto_office_sqft", "stop", "leg_metres_est",
           "avenue_crossings", "street_crossings", "leg_minutes_est", "car_minutes_est", "stop_minutes", "session", "order",
           "walk_minutes_from_previous_est", "dwell_minutes_est", "cumulative_minutes_est", "session_minutes_est"}


def table(ws, df, start_row=1, freeze="C2"):
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=start_row, column=j, value=c)
        cell.fill, cell.font = HEAD, HEADF
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(c, max(12, min(26, len(c) + 3)))
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, (c, v) in enumerate(zip(cols, row), 1):
            val = num(v) if c in NUMERIC and v != "" else v
            cell = ws.cell(row=i, column=j, value=val)
            if c in URL_COLS and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(color="2E47C9", underline="single")
            if c in ("notes", "verify_checklist", "headcount_unknown_reason", "near_miss_reason", "stops_here"):
                cell.alignment = Alignment(wrap_text=False, vertical="top")
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(df)}"
    return {c: get_column_letter(j) for j, c in enumerate(cols, 1)}


def main():
    S = json.load(open(os.path.join(HERE, "out", "stats.json")))
    c, n, b, w, r = rd("companies.csv"), rd("near_miss.csv"), rd("buildings.csv"), rd("coworking.csv"), rd("retailers.csv")
    route, walk = rd("route.csv"), rd("walk.csv")
    wb = Workbook()
    wsC = wb.active
    wsC.title = "Companies"
    L = {}
    L["Companies"] = table(wsC, c)
    red = PatternFill("solid", fgColor="FBE3E1")
    grey = PatternFill("solid", fgColor="EFECE6")
    lc = L["Companies"]
    last = len(c) + 1
    wsC.conditional_formatting.add(f"A2:{get_column_letter(len(c.columns))}{last}",
                                   FormulaRule(formula=[f'${lc["candidate_status"]}2<>"candidate"'], fill=grey))
    wsC.conditional_formatting.add(f'{lc["lease_status"]}2:{lc["move_in_date"]}{last}',
                                   FormulaRule(formula=[f'${lc["lease_status"]}2="signed-not-occupied"'], fill=red))
    wsB = wb.create_sheet("Buildings")
    L["Buildings"] = table(wsB, b)
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    wsB.add_data_validation(dv)
    dv.add(f'{L["Buildings"]["verified_on_foot"]}2:{L["Buildings"]["verified_on_foot"]}{len(b) + 1}')
    wsW = wb.create_sheet("Coworking")
    L["Coworking"] = table(wsW, w)
    wsR = wb.create_sheet("Retailers")
    L["Retailers"] = table(wsR, r)
    route_sheet(wb.create_sheet("Route"), route, S)
    walk_sheet(wb.create_sheet("Walk"), walk)
    wsN = wb.create_sheet("Near-miss")
    L["Near-miss"] = table(wsN, n)
    method_sheet(wb.create_sheet("Method"), S, L, {"Companies": c, "Near-miss": n, "Buildings": b, "Coworking": w,
                                                   "Retailers": r, "Route": route, "Walk": walk})
    wb.defined_names["company_name"] = DefinedName("company_name", attr_text="Method!$B$3")
    wb.defined_names["tagline"] = DefinedName("tagline", attr_text="Method!$B$4")
    wb.save(os.path.join(ROOT, "cluster-1-flatiron.xlsx"))
    print("saved")


# ------------------------------------------------------------------ Route
def route_sheet(ws, route, S):
    P = build.P
    ws["A1"] = '=company_name&": model ten-stop Tuesday route, 8 to 11 am, on foot from the first stop (all times are estimates)"'
    ws["A1"].font = TITLEF
    params = [
        ("hand_truck_pace_m_per_min", P["hand_truck_pace_m_per_min"], "MUTCD 2009 4E.06: 3.5 ft/s design walking speed (64 m/min). " + build.MUTCD_URL),
        ("wait_avenue_crossing_s", P["wait_avenue_crossing_s"], "Assumption: expected wait at an avenue crossing with a hand truck. Time it on the walk."),
        ("wait_street_crossing_s", P["wait_street_crossing_s"], "Assumption: expected wait crossing a side street. Time it on the walk."),
        ("minutes_per_stop", P["minutes_per_stop"], "plan.md Phase 3, per-stop procedure: 15 minutes."),
        ("same_building_transfer_min", P["same_building_transfer_min"], "Assumption: down one elevator and up another."),
        ("route_start", "08:00", "Window opens (plan.md)."),
        ("window_minutes", 180, "8 to 11 am."),
        ("car_overhead_min", P["car_overhead_min"], "Assumption: 5 min to hail and load the hand truck, 3 min to unload."),
        ("car_speed_mph", P["car_speed_mph"], "NBER Digest 2025-06-01: CBD average speed rose from 8.2 to 9.7 mph after congestion pricing. " + build.NBER_URL),
        ("car_detour_factor", P["car_detour_factor"], "Assumption: one-way streets and turns."),
    ]
    ws["A2"] = "Parameters (yellow cells are editable; every time below recalculates)"
    ws["A2"].font = SECF
    pos = {}
    for i, (k, v, src) in enumerate(params, 3):
        ws.cell(row=i, column=1, value=k)
        cell = ws.cell(row=i, column=2, value=v)
        cell.fill = PARAMF
        if k == "route_start":
            cell.value = 8 / 24
            cell.number_format = "hh:mm"
        ws.cell(row=i, column=3, value=src).font = NOTEF
        pos[k] = f"$B${i}"
    hdr = 15
    cols = list(route.columns)
    for j, cname in enumerate(cols, 1):
        cell = ws.cell(row=hdr, column=j, value=cname)
        cell.fill, cell.font = HEAD, HEADF
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(cname, max(11, min(24, len(cname) + 2)))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["C"].width = 30
    col = {cname: get_column_letter(j) for j, cname in enumerate(cols, 1)}
    for i, row in enumerate(route.itertuples(index=False), hdr + 1):
        for j, (cname, v) in enumerate(zip(cols, row), 1):
            ws.cell(row=i, column=j, value=num(v) if cname in NUMERIC and v != "" else v)
        R = i
        if i == hdr + 1:
            ws[f'{col["leg_minutes_est"]}{R}'] = 0
            ws[f'{col["arrive_est"]}{R}'] = f"={pos['route_start']}"
            ws[f'{col["car_minutes_est"]}{R}'] = ""
            ws[f'{col["car_faster"]}{R}'] = ""
        else:
            m, a, s_, sb = (f'{col[k]}{R}' for k in ("leg_metres_est", "avenue_crossings", "street_crossings", "same_building"))
            ws[f'{col["leg_minutes_est"]}{R}'] = (f'=ROUND(IF({sb}="Y",{pos["same_building_transfer_min"]},{m}/{pos["hand_truck_pace_m_per_min"]}'
                                                 f'+({a}*{pos["wait_avenue_crossing_s"]}+{s_}*{pos["wait_street_crossing_s"]})/60),1)')
            ws[f'{col["car_minutes_est"]}{R}'] = (f'=IF({sb}="Y","",ROUND({pos["car_overhead_min"]}+{m}*{pos["car_detour_factor"]}'
                                                 f'/({pos["car_speed_mph"]}*1609.344/60),1))')
            ws[f'{col["car_faster"]}{R}'] = f'=IF({col["car_minutes_est"]}{R}="","",IF({col["car_minutes_est"]}{R}<{col["leg_minutes_est"]}{R},"Y","N"))'
            ws[f'{col["arrive_est"]}{R}'] = f'={col["depart_est"]}{R - 1}+{col["leg_minutes_est"]}{R}/1440'
        ws[f'{col["stop_minutes"]}{R}'] = f"={pos['minutes_per_stop']}"
        ws[f'{col["depart_est"]}{R}'] = f'={col["arrive_est"]}{R}+{col["stop_minutes"]}{R}/1440'
        for k in ("arrive_est", "depart_est"):
            ws[f'{col[k]}{R}'].number_format = "hh:mm"
    last = hdr + len(route)
    t = last + 2
    ws[f"A{t}"] = "Totals (estimates)"
    ws[f"A{t}"].font = SECF
    rows = [
        ("walking minutes", f'=SUM({col["leg_minutes_est"]}{hdr + 1}:{col["leg_minutes_est"]}{last})'),
        ("walking metres", f'=SUM({col["leg_metres_est"]}{hdr + 1}:{col["leg_metres_est"]}{last})'),
        ("stop minutes", f'=SUM({col["stop_minutes"]}{hdr + 1}:{col["stop_minutes"]}{last})'),
        ("total minutes, 08:00 to last departure", f'=ROUND(({col["depart_est"]}{last}-{pos["route_start"]})*1440,0)'),
        ("finish time", f'={col["depart_est"]}{last}'),
        ("fits 8 to 11 am?", f'=IF(B{t + 4}<={pos["window_minutes"]},"fits, "&({pos["window_minutes"]}-B{t + 4})&" min spare","over by "&(B{t + 4}-{pos["window_minutes"]})&" min")'),
        ("legs where a car is faster", f'=COUNTIF({col["car_faster"]}{hdr + 1}:{col["car_faster"]}{last},"Y")'),
        ("total if stops took 12 min", f'=ROUND(B{t + 1}+12*{len(route)},0)'),
        ("total with the first nine stops only", f'=ROUND(SUM({col["leg_minutes_est"]}{hdr + 1}:{col["leg_minutes_est"]}{last - 1})+9*{pos["minutes_per_stop"]},0)'),
    ]
    for i, (k, f) in enumerate(rows, t + 1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=f)
    ws[f"B{t + 5}"].number_format = "hh:mm"
    notes = [
        "How to read this: stop 1 starts at 08:00. Each leg runs sidewalk to sidewalk along the street grid, to the corner and across (never mid-block), at hand-truck pace, plus an expected signal wait per crossing. The walk from the kerb to the door is inside each stop's 15 minutes.",
        (f"Result: {S['route']['stops']} stops = {S['route']['stop_minutes']} min at stops + {S['route']['walking_minutes']} min walking = {S['route']['total_minutes']} min, finishing about {S['route']['end']}. "
         + ("That fits the 3-hour window." if S['route']['over_minutes'] <= 0 else f"That is about {S['route']['over_minutes']} min over the 3-hour window, inside the model's error.")
         + " The stops: high-confidence, occupied offices with one 40-75 and one 76-150 office, two security desks last, and no food or drink brand."),
        "What makes ten fit: a compact roster. Allowing medium-confidence offices, the best set found (about three blocks, 21st to 23rd St between Fifth and Sixth Ave) takes about "
        + f"{S['route']['variants']['high_or_medium']['total_minutes']:.0f} min and ends about {S['route']['variants']['high_or_medium']['end']}. 12-minute stops at repeat offices give about {S['route']['at_12_min_stops']:.0f} min. Splitting a ten-office drop across Tuesday and Wednesday also works (the plan allows both days).",
        "Load check: ten 50-unit cases of a canned drink weigh about 200 kg (50 cans x about 0.37 kg each, by arithmetic). Check the hand truck's printed load rating; heavy products need a mid-route reload or two loops.",
        "8 am check: a walk-up office needs a person there at 08:00. Confirm every receiving window at intake and re-order the stops to match. Security desks often send deliveries to a freight elevator or loading dock with set hours.",
        "Lobby types: no walk-up lobby is sourced yet, so unknown-lobby buildings come first and the two sourced security-desk buildings come last. Re-run the order after the verification walk.",
        "This is a model. The real route uses the signed roster; company names are examples only and none has been contacted.",
    ]
    for i, tx in enumerate(notes, t + len(rows) + 2):
        ws.cell(row=i, column=1, value=tx).font = NOTEF
    ws.freeze_panes = ws.cell(row=hdr + 1, column=4)


# ------------------------------------------------------------------ Walk
def walk_sheet(ws, walk):
    cols = list(walk.columns)
    col = table(ws, walk)
    n = len(walk)
    for i in range(2, n + 2):
        ws[f'{col["cumulative_minutes_est"]}{i}'] = (f'=IF({col["session"]}{i}={col["session"]}{i - 1},{col["cumulative_minutes_est"]}{i - 1},0)'
                                                    f'+{col["walk_minutes_from_previous_est"]}{i}+{col["dwell_minutes_est"]}{i}') if i > 2 else \
            f'={col["walk_minutes_from_previous_est"]}{i}+{col["dwell_minutes_est"]}{i}'
        ws[f'{col["session_minutes_est"]}{i}'] = (f'=SUMIF(${col["session"]}$2:${col["session"]}${n + 1},{col["session"]}{i},${col["walk_minutes_from_previous_est"]}$2:${col["walk_minutes_from_previous_est"]}${n + 1})'
                                                 f'+SUMIF(${col["session"]}$2:${col["session"]}${n + 1},{col["session"]}{i},${col["dwell_minutes_est"]}$2:${col["dwell_minutes_est"]}${n + 1})')
    founder = PatternFill("solid", fgColor="E9E6DF")
    for k in ("observed_lobby_type", "security_desk", "loading_dock_hours", "coi_signage", "receiving_hours", "emerging_set_photo_taken", "founder_notes"):
        for i in range(2, n + 2):
            ws[f"{col[k]}{i}"].fill = founder
    dv1 = DataValidation(type="list", formula1='"walk-up,security desk,loading dock,other"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f'{col["observed_lobby_type"]}2:{col["observed_lobby_type"]}{n + 1}')
    for k in ("security_desk", "coi_signage", "emerging_set_photo_taken"):
        dv2.add(f'{col[k]}2:{col[k]}{n + 1}')
    # session summary to the right
    sc = len(cols) + 2
    heads = ["session_summary", "phase", "stops", "minutes_est", "start_point", "end_point"]
    for j, h in enumerate(heads):
        cell = ws.cell(row=1, column=sc + j, value=h)
        cell.fill, cell.font = HEAD, HEADF
        ws.column_dimensions[get_column_letter(sc + j)].width = 34 if "point" in h else 16
    sessions = sorted(set(int(s) for s in walk.session))
    rng = lambda k: f'${col[k]}$2:${col[k]}${n + 1}'
    for i, s in enumerate(sessions, 2):
        first = walk[walk.session == str(s)].iloc[0]
        ws.cell(row=i, column=sc, value=s)
        ws.cell(row=i, column=sc + 1, value=first.phase)
        ws.cell(row=i, column=sc + 2, value=f"=COUNTIF({rng('session')},{get_column_letter(sc)}{i})")
        ws.cell(row=i, column=sc + 3, value=f"=ROUND(SUMIF({rng('session')},{get_column_letter(sc)}{i},{rng('walk_minutes_from_previous_est')})+SUMIF({rng('session')},{get_column_letter(sc)}{i},{rng('dwell_minutes_est')}),0)")
        ws.cell(row=i, column=sc + 4, value=first.start_point)
        ws.cell(row=i, column=sc + 5, value=first.end_point)
    tr = len(sessions) + 2
    ws.cell(row=tr, column=sc, value="total")
    ws.cell(row=tr, column=sc + 2, value=f"=SUM({get_column_letter(sc + 2)}2:{get_column_letter(sc + 2)}{tr - 1})")
    ws.cell(row=tr, column=sc + 3, value=f"=SUM({get_column_letter(sc + 3)}2:{get_column_letter(sc + 3)}{tr - 1})")
    ws.cell(row=tr + 1, column=sc, value="Times are estimates: walking at 73 m/min (MUTCD 4 ft/s) plus signal waits, and the dwell rule in Method. Fill the grey founder columns on foot.").font = NOTEF
    ws.cell(row=tr + 2, column=sc, value=f"Do sessions 1 to {walk[walk.phase.str.startswith('core')].session.astype(int).max()} (the core, inside the boundary) on Tuesday or Wednesday mornings between 8 and 11 am, two sessions a morning, so lobby observations (security desk staffed? receiving hours?) match drop day.").font = Font(bold=True, color="B3261E")


# ------------------------------------------------------------------ Method
def method_sheet(ws, S, L, dfs):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    ws.column_dimensions["C"].width = 40
    ws["A1"] = '=company_name&": cluster 1 (Flatiron, NoMad and Union Square) method, definitions and sources"'
    ws["A1"].font = TITLEF
    ws["A2"] = "Config (rename here once; titles on Route and Method read it)"
    ws["A2"].font = SECF
    ws["A3"], ws["B3"] = "company_name", "Desk Drop"
    ws["A4"], ws["B4"] = "tagline", "Emerging brands. Delivered to your workday."
    ws["B3"].fill = ws["B4"].fill = PARAMF
    row = 6
    lc = L["Companies"]
    live = [
        ("Live counts (formulas)", None),
        ("candidate companies", f'=COUNTIF(Companies!{lc["candidate_status"]}:{lc["candidate_status"]},"candidate")'),
        ("  of which high confidence", f'=COUNTIFS(Companies!{lc["candidate_status"]}:{lc["candidate_status"]},"candidate",Companies!{lc["confidence"]}:{lc["confidence"]},"high")'),
        ("  of which medium", f'=COUNTIFS(Companies!{lc["candidate_status"]}:{lc["candidate_status"]},"candidate",Companies!{lc["confidence"]}:{lc["confidence"]},"medium")'),
        ("  of which low", f'=COUNTIFS(Companies!{lc["candidate_status"]}:{lc["candidate_status"]},"candidate",Companies!{lc["confidence"]}:{lc["confidence"]},"low")'),
        ("  signed, not yet occupied", f'=COUNTIFS(Companies!{lc["candidate_status"]}:{lc["candidate_status"]},"candidate",Companies!{lc["lease_status"]}:{lc["lease_status"]},"signed-not-occupied")'),
        ("coworking members (excluded)", f'=COUNTIF(Companies!{lc["candidate_status"]}:{lc["candidate_status"]},"excluded-coworking")'),
        ("near-miss rows", f'=COUNTA(\'Near-miss\'!A:A)-1'),
        ("buildings", "=COUNTA(Buildings!A:A)-1"),
        ("coworking sites", "=COUNTA(Coworking!A:A)-1"),
        ("stores", "=COUNTA(Retailers!A:A)-1"),
        ("verified on foot so far", f'=COUNTIF(Buildings!{L["Buildings"]["verified_on_foot"]}:{L["Buildings"]["verified_on_foot"]},"Y")'),
    ]
    for k, f in live:
        ws.cell(row=row, column=1, value=k).font = SECF if f is None else Font()
        if f:
            ws.cell(row=row, column=2, value=f).alignment = Alignment(horizontal="left")
        row += 1
    row += 1
    for title, items in method_sections(S):
        ws.cell(row=row, column=1, value=title).font = SECF
        row += 1
        for it in items:
            if isinstance(it, tuple):
                ws.cell(row=row, column=1, value=it[0]).alignment = Alignment(vertical="top", wrap_text=True)
                ws.cell(row=row, column=2, value=it[1]).alignment = Alignment(wrap_text=True, vertical="top")
                if len(it) > 2:
                    ws.cell(row=row, column=3, value=it[2]).alignment = Alignment(wrap_text=True, vertical="top")
            else:
                ws.cell(row=row, column=2, value=it).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1
    ws.cell(row=row, column=1, value="Column definitions").font = SECF
    row += 1
    for tab in ("Companies", "Buildings", "Coworking", "Retailers", "Route", "Walk", "Near-miss"):
        ws.cell(row=row, column=1, value=tab).font = Font(bold=True)
        row += 1
        for colname in dfs[tab].columns:
            ws.cell(row=row, column=1, value=colname)
            d = coldefs.define(tab, colname) or ("Same as the Companies tab." if tab == "Near-miss" else "")
            ws.cell(row=row, column=2, value=d).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1


if __name__ == "__main__":
    main()
