#!/usr/bin/env python3
"""Test the conversation tracker with SAMPLE rows and a LibreOffice recalculation.

    python3 tools/test_tracker.py            # builds, recalculates, compares, writes _work/partB/test_log.txt

Expected values are computed here in Python from tools/sample_data.py, independently of the
spreadsheet formulas, then compared with what LibreOffice calculates from the workbook.
"""
import datetime as dt
import os
import shutil
import statistics
import subprocess
import sys
sys.dont_write_bytecode = True   # no __pycache__ in the deliverable folder (review round 3, R3-11)
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_tracker as BT          # noqa: E402
import sample_data as S             # noqa: E402

ERRORS = ("#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "Err:")
LOG = []
FAILS = []


def say(msg):
    LOG.append(msg)
    print(msg)


def check(label, got, want):
    ok = (got == want) or (isinstance(got, (int, float)) and isinstance(want, (int, float)) and abs(got - want) < 1e-9)
    say(f"{'ok  ' if ok else 'FAIL'} {label}: got {got!r} want {want!r}")
    if not ok:
        FAILS.append(label)


def recalc(path):
    tmp = tempfile.mkdtemp(prefix="ddtr_")
    src = os.path.join(tmp, "in.xlsx")
    shutil.copy(path, src)
    out = os.path.join(tmp, "out")
    os.makedirs(out)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, src],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    res = os.path.join(out, "in.xlsx")
    return openpyxl.load_workbook(res, data_only=True), res


def errors(wb):
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRORS):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    return errs


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    return v


# ----------------------------------------------------------------------------- expected values
CUTS = {"B": dt.date(2026, 10, 30), "C": dt.date(2026, 11, 6), "D": dt.date(2026, 11, 13), "E": S.AS_OF}
QUAL = ("qual_q1_sealed_shelf_stable_60d", "qual_q2_labeled_ingredients_allergens", "qual_q3_buy_path_near_cluster",
        "qual_q4_can_ship_275_in_10bd", "qual_q5_product_liability_insurance", "qual_q6_contact_can_approve_600", "qual_not_excluded")


def org(r):
    """Review round 3 (R3-01): the organisation a row counts under, written out here, not imported from the builder:
    a typed org_key, else the company name, ignoring case and spaces (the SAMPLE Clay rows have no grp- company_group)."""
    return (r.get("org_key") or r["company"]).strip().lower()


def first_per_org(rows, pred, cut):
    """Earliest row (by call_date, then sheet order) per organisation satisfying pred, with date <= cut."""
    best = {}
    for i, r in enumerate(rows):
        if not pred(r):
            continue
        key = (r["call_date"], i)
        if org(r) not in best or key < best[org(r)][0]:
            best[org(r)] = (key, r)
    return [r for (k, r) in best.values() if k[0] <= cut]


def org_first_rid(rows, r):
    """The row_id of the organisation's first conversation: the pipeline's org_id (review round 3)."""
    return first_per_org(rows, lambda x: org(x) == org(r), dt.date(2100, 1, 1))[0]["row_id"]


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def median_or(vals):
    return statistics.median(vals) if vals else "none yet"


def standing(rows, key, answered, positive, cut):
    """Review round 2 (R2-02): the org's latest answered row (by call_date, then sheet order) on or before the cut-off
    decides whether its earlier yes still stands. Review round 3: across all the organisation's contacts."""
    best = None
    for i, r in enumerate(rows):
        if org(r) != key or not answered(r) or r["call_date"] > cut:
            continue
        pos = (r["call_date"], i)
        if best is None or pos > best[0]:
            best = (pos, r)
    return best is not None and positive(best[1])


def office_answered(r):
    return r["yes_in_principle"] in ("Y", "N", "maybe")


def office_standing(key, cut, rows=None):
    return standing(rows or S.OFFICE_LOG, key, office_answered, lambda r: r["yes_in_principle"] == "Y", cut)


def brand_pos(r):
    return r["controls_budget"] == "Y" and max(r["wtp_before_usd"] or 0, r["wtp_after_usd"] or 0) >= 500


def brand_neg(r):
    has = num(r["wtp_before_usd"]) or num(r["wtp_after_usd"])
    return r["controls_budget"] == "N" or (has and max(r["wtp_before_usd"] or 0, r["wtp_after_usd"] or 0) < 500)


def brand_standing(key, cut, rows=None):
    return standing(rows or S.BRAND_LOG, key, lambda r: brand_pos(r) or brand_neg(r), brand_pos, cut)


def office_hc(rows, r):
    """Review round 3 (R3-04): the office's latest Tuesday headcount on any of its rows (any date), or None."""
    hit = [(x["call_date"], i, x["tuesday_headcount"]) for i, x in enumerate(rows) if org(x) == org(r) and num(x.get("tuesday_headcount"))]
    return max(hit)[2] if hit else None


def in_band(rows, r):
    hc = office_hc(rows, r)
    return hc is None or 30 <= hc <= 250


def q_yes(r):
    return r["yes_in_principle"] == "Y" and r["pantry_owner_named"] == "Y" and r["receiving_window"].strip()


def expected_office(cut):
    L = S.OFFICE_LOG
    e = {}
    e["o_conv"] = len(first_per_org(L, lambda r: True, cut))
    e["o_yes"] = len([r for r in first_per_org(L, q_yes, cut) if office_standing(org(r), cut) and in_band(L, r)])
    e["o_yes_any"] = len([r for r in first_per_org(L, lambda r: r["yes_in_principle"] == "Y", cut) if office_standing(org(r), cut) and in_band(L, r)])
    e["o_wd"] = len([r for r in first_per_org(L, lambda r: r["yes_in_principle"] == "Y", cut) if not office_standing(org(r), cut)])
    e["o_yes_hcout"] = len([r for r in first_per_org(L, q_yes, cut) if office_standing(org(r), cut) and not in_band(L, r)])
    e["o_yes_3039"] = len([r for r in first_per_org(L, q_yes, cut) if office_standing(org(r), cut) and in_band(L, r)
                           and 30 <= (office_hc(L, r) or 0) < 40])
    e["o_maybe"] = len(first_per_org(L, lambda r: r["yes_in_principle"] == "maybe", cut))
    e["o_warm"] = len(first_per_org(L, lambda r: r["channel"] == "warm-intro", cut))
    e["o_short"] = len(first_per_org(L, lambda r: r.get("shortlist") == "Y", cut))
    bp = [r["budget_probe_low_usd"] for r in first_per_org(L, lambda r: num(r["budget_probe_low_usd"]), cut)]
    e["o_bp_med"] = median_or(bp)
    e["o_bp_n"] = len(bp)
    e["o_bp_zero"] = sum(1 for v in bp if v == 0)
    return e


def expected_office_asof():
    L, cut = S.OFFICE_LOG, S.AS_OF
    bp100 = len(first_per_org(L, lambda r: num(r["budget_probe_low_usd"]) and r["budget_probe_low_usd"] >= 100, cut))
    spend = [r["monthly_pantry_spend_usd"] for r in first_per_org(L, lambda r: num(r["monthly_pantry_spend_usd"]), cut)]
    return bp100, median_or(spend)


def expected_brand(cut):
    L = S.BRAND_LOG
    e = {}
    mx = lambda r: max(r["wtp_before_usd"] or 0, r["wtp_after_usd"] or 0)
    e["b_conv"] = len(first_per_org(L, lambda r: True, cut))
    e["b_yes"] = len([r for r in first_per_org(L, lambda r: r["controls_budget"] == "Y" and mx(r) >= 500, cut) if brand_standing(org(r), cut)])
    e["b_strict"] = len([r for r in first_per_org(L, lambda r: r["controls_budget"] == "Y" and num(r["wtp_before_usd"]) and r["wtp_before_usd"] >= 500, cut)
                         if brand_standing(org(r), cut)])
    e["b_wd"] = len([r for r in first_per_org(L, lambda r: r["controls_budget"] == "Y" and mx(r) >= 500, cut) if not brand_standing(org(r), cut)])
    e["b_noctrl"] = len(first_per_org(L, lambda r: r["controls_budget"] != "Y" and mx(r) >= 500, cut))
    wb = [r["wtp_before_usd"] for r in first_per_org(L, lambda r: num(r["wtp_before_usd"]), cut)]
    wa = [r["wtp_after_usd"] for r in first_per_org(L, lambda r: num(r["wtp_after_usd"]), cut)]
    e["b_wb_med"], e["b_wb_n"] = median_or(wb), len(wb)
    e["b_wa_med"], e["b_wa_n"] = median_or(wa), len(wa)
    e["b_short"] = len(first_per_org(L, lambda r: r.get("shortlist") == "Y", cut))
    e["b_qualified"] = len(first_per_org(L, lambda r: all(r[k] == "Y" for k in QUAL) and r["qual_q7_why_now"].strip(), cut))
    return e


def gate_status(yes, target, conv, cut, conv_target=15):
    """Review round 2 (R2-03): FAIL only when the plan's test has been run (15 conversations); else INCOMPLETE.
    Review round 3 (R3-09): from the final read, Fri 2026-11-27, a shortfall reads FAIL whatever the conversations."""
    if yes >= target:
        return "PASS"
    if cut < dt.date(2026, 11, 13):
        return "not yet"
    return "FAIL" if (conv >= conv_target or cut >= dt.date(2026, 11, 27)) else f"INCOMPLETE: {conv} of {conv_target} conversations"


def gate_overall(o, b):
    if o == "PASS" and b == "PASS":
        return "PASS"
    if "FAIL" in (o, b):
        return "FAIL"
    if o.startswith("INCOMPLETE") or b.startswith("INCOMPLETE"):
        return "INCOMPLETE"
    return "not yet"


def expected_top3(rows, codes, labels, cut):
    counts = {c: 0 for c in codes}
    for r in rows:
        if r["call_date"] > cut:
            continue
        for k in ("objection_1", "objection_2", "objection_3"):
            if r[k]:
                counts[r[k]] += 1
    order = sorted(codes, key=lambda c: (-counts[c], codes.index(c)))
    out = []
    for k, c in enumerate(order[:3], start=1):
        if counts[c] < 1:
            out.append((f"{k}. none logged yet", counts[c], ""))
            continue
        quote = ""
        for col in ("objection_1", "objection_2", "objection_3"):
            hit = [r for r in rows if r[col] == c]
            if hit:
                quote = hit[0]["objections_verbatim"]
                break
        out.append((f"{k}. {labels[codes.index(c)]} ({c})", counts[c], quote))
    return out


def expected_pace(clay, cut, conv):
    first = sum(1 for r in clay if r["touch_1_date"] and r["touch_1_date"] <= cut)
    rep = sum(1 for r in clay if r["last_reply_date"] and r["last_reply_date"] <= cut and r["reply_type"] not in ("bounce", "out_of_office"))
    # a reply that can become a conversation (review round 1, R1-06): written out here, not imported from the builder
    pos_types = ("yes", "maybe", "needs_approval", "send_info", "referral", "not_now",
                 "obj_seed_ourselves", "obj_price", "obj_free_trial", "obj_what_we_get", "obj_other")
    pos = sum(1 for r in clay if r["last_reply_date"] and r["last_reply_date"] <= cut and r["reply_type"] in pos_types)
    oo = sum(1 for r in clay if r["opt_out_date"] and r["opt_out_date"] <= cut)
    return {"first": first, "replies": rep, "rate": (rep / first) if first else "", "positive": pos,
            "posrate": (pos / first) if first else "", "convrate": (conv / first) if first else "", "optout": oo,
            "optrate": (oo / first) if first else ""}


def expected_events():
    ev = []   # (date, block, i, side, org_id, org_name, event, value_number, value_text, controls, notes)
    OL_, BL_ = S.OFFICE_LOG, S.BRAND_LOG
    oid = lambda rows, r: org_first_rid(rows, r)
    via_log = lambda rows, rid: next((oid(rows, x) for x in rows if x["row_id"] == rid), rid)
    for i, r in enumerate(OL_, start=1):
        ev.append((r["call_date"], 1, i, "office", oid(OL_, r), r["company"], "conversation", "", "", "", r["call_type"]))
        g0 = dt.date(2026, 11, 13)
        if q_yes(r) and in_band(OL_, r) and office_standing(org(r), g0 if r["call_date"] <= g0 else S.AS_OF):
            ev.append((r["call_date"], 2, i, "office", oid(OL_, r), r["company"], "yes_in_principle", "", "", "",
                       f"owner: {r['pantry_owner_title']}; window: {r['receiving_window']}"))
        if r["budget_probe_verbatim"].strip() or num(r["budget_probe_low_usd"]):
            lo, hi = r["budget_probe_low_usd"], r["budget_probe_high_usd"]
            note = f"range to {hi} per month" if (num(hi) and (hi or 0) > (lo or 0)) else ""
            ev.append((r["call_date"], 3, i, "office", oid(OL_, r), r["company"], "budget_probe", lo if num(lo) else "",
                       r["budget_probe_verbatim"], "", note))
    for i, r in enumerate(BL_, start=1):
        ev.append((r["call_date"], 4, i, "brand", oid(BL_, r), r["company"], "conversation", "", "", "", r["call_type"]))
        if r["yes_in_principle"] == "Y":
            ev.append((r["call_date"], 5, i, "brand", oid(BL_, r), r["company"], "yes_in_principle", "", "", r["controls_budget"], ""))
        if num(r["wtp_before_usd"]) or r["wtp_before_verbatim"].strip():
            ev.append((r["call_date"], 6, i, "brand", oid(BL_, r), r["company"], "budget_probe",
                       r["wtp_before_usd"] if num(r["wtp_before_usd"]) else "", r["wtp_before_verbatim"], r["controls_budget"], ""))
        if r["price_1500_mentioned"] == "Y" and (num(r["wtp_after_usd"]) or r["wtp_after_verbatim"].strip()):
            ev.append((r["call_date"], 7, i, "brand", oid(BL_, r), r["company"], "budget_probe",
                       r["wtp_after_usd"] if num(r["wtp_after_usd"]) else "", r["wtp_after_verbatim"], r["controls_budget"], "after_1500"))
    for i, r in enumerate(S.SIGNED_OFFICES, start=1):
        ev.append((r["agreement_signed_date"], 8, i, "office", via_log(OL_, r["clay_row_id"]), r["company"], "signed", r["tuesday_headcount"], "", "", r["office_id"]))
    for i, r in enumerate(S.PAID_BRANDS, start=1):
        for k, b in ((1, 9), (2, 10)):
            if r.get(f"payment_{k}_date") and num(r.get(f"payment_{k}_usd")):
                ev.append((r[f"payment_{k}_date"], b, i, "brand", via_log(BL_, r["clay_row_id"]), r["brand_name"], "paid", r[f"payment_{k}_usd"], "", "",
                           f"payment {k}; {r['brand_id']}; {r['tier']}; invoice {r['invoice_id']}"))
    for i, r in enumerate(S.LATER_EVENTS, start=1):
        ev.append((r["date"], 11, i, r["side"], via_log(OL_ if r["side"] == "office" else BL_, r["clay_row_id"]), r["org_name"], r["event"],
                   r["value_number"] if num(r["value_number"]) else "", r["value_text"], "", r["notes"]))
    b = 12
    for side, clay in (("office", S.CLAY_OFFICES), ("brand", S.CLAY_BRANDS)):
        for k in (1, 2, 3):
            for i, r in enumerate(clay, start=1):
                d = r.get(f"touch_{k}_date")
                if d:
                    ev.append((d, b, i, side, via_log(OL_ if side == "office" else BL_, r["row_id"]), r["company"], "outreach_sent", "", "",
                               "", f"touch_{k}; {r['sequence']}; {r['channel']}"))
            b += 1
    ev.sort(key=lambda x: (x[0], x[1], x[2]))
    return [(e[0], e[3], e[4], e[5], e[6], e[7], e[8], e[9], e[10]) for e in ev]


# ----------------------------------------------------------------------------- tests
def test_sample(b):
    path = os.path.join(ROOT, "tracking", "conversation-tracker-SAMPLE.xlsx")
    wb, recalced = recalc(path)
    errs = errors(wb)
    check("SAMPLE workbook: error cells after recalculation", len(errs), 0)
    for e in errs[:20]:
        say("   " + e)
    sc = wb["Scorecard"]
    v = lambda ref: sc[ref].value
    check("SAMPLE banner present", isinstance(v("A1"), str) and v("A1").startswith("SAMPLE: this file holds"), True)
    for col, cut in CUTS.items():
        eo, eb = expected_office(cut), expected_brand(cut)
        for k, want in eo.items():
            check(f"office {k} @ {cut}", v(f"{col}{b.sc[k]}"), want)
        for k, want in eb.items():
            check(f"brand {k} @ {cut}", v(f"{col}{b.sc[k]}"), want)
        o_pass = gate_status(eo["o_yes"], 5, eo["o_conv"], cut)
        b_pass = gate_status(eb["b_yes"], 3, eb["b_conv"], cut)
        check(f"office Gate 0 status @ {cut}", v(f"{col}{b.sc['o_status']}"), o_pass)
        check(f"brand Gate 0 status @ {cut}", v(f"{col}{b.sc['b_status']}"), b_pass)
        g = gate_overall(o_pass, b_pass)
        check(f"Gate 0 overall @ {cut}", v(f"{col}{b.sc['g0']}"), g)
        for side, clay, conv_key in (("office", S.CLAY_OFFICES, "o_conv"), ("brand", S.CLAY_BRANDS, "b_conv")):
            conv = (eo if side == "office" else eb)[conv_key]
            ep = expected_pace(clay, cut, conv)
            for k, want in ep.items():
                check(f"{side} pace {k} @ {cut}", v(f"{col}{b.sc[f'{side}_{k}']}"), want)
    bp100, spend = expected_office_asof()
    check("offices naming $100+/month (as of)", v(f"E{b.sc['o_bp100']}"), bp100)
    check("median pantry spend (as of)", v(f"E{b.sc['o_spend']}"), spend)
    q2 = [r["q2_sampling_budget_quarter_low_usd"] for r in first_per_org(S.BRAND_LOG, lambda r: num(r["q2_sampling_budget_quarter_low_usd"]), S.AS_OF)]
    check("median quarterly sampling budget (as of)", v(f"E{b.sc['b_q2']}"), statistics.median(q2))
    # Gate 1 block (Gate 1 date and as of)
    g1 = dt.date(2026, 12, 11)
    for col, cut in (("D", g1), ("E", S.AS_OF)):
        signed = [r for r in S.SIGNED_OFFICES if r["agreement_signed_date"] <= cut]
        check(f"Gate 1 offices signed @ {cut}", v(f"{col}{b.sc['g1_signed']}"), len(signed))
        check(f"Gate 1 headcount @ {cut}", v(f"{col}{b.sc['g1_headcount']}"), sum(r["tuesday_headcount"] for r in signed))
        full = 0
        for r in S.PAID_BRANDS:
            tot = (r["payment_1_usd"] or 0) + (r["payment_2_usd"] or 0)
            dd = max(d for d in (r["payment_1_date"], r["payment_2_date"]) if d)
            if tot >= r["price_usd"] and dd <= cut:
                full += 1
        check(f"Gate 1 brands paid in full @ {cut}", v(f"{col}{b.sc['g1_paid']}"), full)
        o_st = "PASS" if (len(signed) >= 5 and sum(r["tuesday_headcount"] for r in signed) >= 350) else ("not yet" if cut < g1 else "FAIL")
        b_st = "PASS" if full >= 2 else ("not yet" if cut < g1 else ("SOFT PASS" if full == 1 else "FAIL"))
        check(f"Gate 1 office status @ {cut}", v(f"{col}{b.sc['g1_o_status']}"), o_st)
        check(f"Gate 1 brand status @ {cut}", v(f"{col}{b.sc['g1_b_status']}"), b_st)
    # numbers for the Startup CPG posts (as of)
    L, Lb, cut = S.OFFICE_LOG, S.BRAND_LOG, S.AS_OF
    f1 = lambda rows, pred: len(first_per_org(rows, pred, cut))
    check("post: office conversations", v(f"E{b.sc['post_office_conv']}"), expected_office(cut)["o_conv"])
    check("post: brands reach out", v(f"E{b.sc['post_reach']}"), f1(L, lambda r: r["q3_brands_reach_out"] == "Y"))
    check("post: one person decides", v(f"E{b.sc['post_alone']}"), f1(L, lambda r: r["owner_can_say_yes_alone"] == "Y"))
    check("post: photo and survey yes", v(f"E{b.sc['post_q5']}"), f1(L, lambda r: r["q5_photo_and_survey"] == "Y"))
    check("post: use a snack service", v(f"E{b.sc['post_svc']}"), f1(L, lambda r: r["uses_snack_service"] == "Y"))
    hcs = [r["tuesday_headcount"] for r in first_per_org(L, lambda r: num(r["tuesday_headcount"]), cut)]
    check("post: Tuesday range", v(f"E{b.sc['post_hc']}"), f"{min(hcs)} to {max(hcs)}")
    from collections import Counter
    ap = Counter(r["other_approver"] for r in L if r["call_date"] <= cut and r["other_approver"])
    order = BT.LISTS["other_approver"]
    best = max(ap.values())
    check("post: most common other approver", v(f"E{b.sc['post_appr']}"), [a for a in order if ap.get(a, 0) == best][0])
    lo = [r["q2_sampling_budget_quarter_low_usd"] for r in first_per_org(Lb, lambda r: num(r["q2_sampling_budget_quarter_low_usd"]), cut)]
    hi = [r["q2_sampling_budget_quarter_high_usd"] for r in first_per_org(Lb, lambda r: num(r["q2_sampling_budget_quarter_low_usd"]), cut)]
    check("post: brand spend range", v(f"E{b.sc['post_spend']}"), f"${min(lo):,} to ${max(hi + lo):,}")
    check("post: seeded offices", v(f"E{b.sc['post_seeded']}"), f1(Lb, lambda r: r["q3_seeded_offices"] == "Y"))
    check("post: median number named", v(f"E{b.sc['post_median']}"), expected_brand(cut)["b_wb_med"])
    check("post: top blocker equals top brand objection", v(f"E{b.sc['post_blocker']}"), v(f"A{b.sc['top_brand_1']}"))
    check("post: quotes allowed", v(f"E{b.sc['post_quotes']}"), f1(L, lambda r: r["quote_ok_to_share"] == "Y") + f1(Lb, lambda r: r["quote_ok_to_share"] == "Y"))
    # copy-slot counts
    oy = sum(1 for r in S.CLAY_OFFICES if r["status"] in ("yes_in_principle", "materials_sent", "signed"))
    check("slot offices_yes", v(f"E{b.sc['slot_offices_yes']}"), oy)
    check("slot offices_signed", v(f"E{b.sc['slot_offices_signed']}"), sum(1 for r in S.CLAY_OFFICES if r["status"] == "signed"))
    check("slot tuesday_headcount", v(f"E{b.sc['slot_headcount']}"), sum(r["tuesday_headcount"] for r in S.SIGNED_OFFICES))
    check("slot founding_taken", v(f"E{b.sc['slot_founding']}"), sum(1 for r in S.PAID_BRANDS if r["tier"] == "founding"))
    check("slot brands_paid", v(f"E{b.sc['slot_brands_paid']}"), len(S.PAID_BRANDS))
    eb_asof = expected_brand(S.AS_OF)
    check("slot brand_selling_on", v(f"E{b.sc['slot_brand_selling']}"), "Y" if (eb_asof["b_yes"] >= 3 and oy >= 2) else "N")
    check("check: duplicate handoff rows", v(f"E{b.sc['chk_Duplicate clay_row_id on signed_offices or paid_brands']}"), 0)
    # top three objections
    for side, rows, lst in (("office", S.OFFICE_LOG, "office_objection"), ("brand", S.BRAND_LOG, "brand_objection")):
        exp = expected_top3(rows, BT.LISTS[lst], BT.LISTS[lst + "_label"], S.AS_OF)
        for k in (1, 2, 3):
            r = b.sc[f"top_{side}_{k}"]
            check(f"top {side} objection {k} label", v(f"A{r}"), exp[k - 1][0])
            check(f"top {side} objection {k} mentions", v(f"B{r}"), exp[k - 1][1])
            check(f"top {side} objection {k} quote", v(f"C{r}") or "", exp[k - 1][2])
    # sentence
    sent = v(f"B{b.sc['sentence']}")
    eo, eb = expected_office(S.AS_OF), expected_brand(S.AS_OF)
    want = (f"As of {S.AS_OF.isoformat()}: {eo['o_yes']} of {eo['o_conv']} offices said yes with a pantry owner and a receiving window (target 5); "
            f"{eb['b_yes']} of {eb['b_conv']} brands would pay $500 or more and control the budget (target 3). Gate 0: PASS.")
    check("Gate 0 sentence", sent, want)
    # reply-rate check text
    for side, clay, conv_key in (("office", S.CLAY_OFFICES, "o_conv"), ("brand", S.CLAY_BRANDS, "b_conv")):
        ep = expected_pace(clay, S.AS_OF, 0)
        want = (f"Too few first touches to judge ({ep['first']} of 30)." if ep["first"] < 30 else
                ("UNDER 20%: conversation replies are under 20% of first touches. Run sending/sending-setup.md section 6 before sending more."
                 if ep["posrate"] < 0.2 else "OK: conversation replies at or above 20%."))
        check(f"{side} reply-rate message", v(f"B{b.sc[f'{side}_check']}"), want)
        want = (f"Too few first touches to judge ({ep['first']} of 30)." if ep["first"] < 30 else
                ("OPT-OUTS OVER 5%: check targeting and copy (sending/sending-setup.md 6.3 to 6.4)." if (ep["optout"] >= 3 and ep["optrate"] > 0.05)
                 else "OK: fewer than 3 opt-outs, or 5% or under."))
        check(f"{side} opt-out message", v(f"B{b.sc[f'{side}_optcheck']}"), want)
    # sequence table: office-discovery-A
    r = b.sc["seq_first"]
    exp_sends = sum(1 for x in S.CLAY_OFFICES if x["sequence"] == "office-discovery-A" and x["touch_1_date"] and x["touch_1_date"] <= S.AS_OF)
    exp_rep = sum(1 for x in S.CLAY_OFFICES if x["sequence"] == "office-discovery-A" and x["last_reply_date"] and x["last_reply_date"] <= S.AS_OF
                  and x["reply_type"] not in ("bounce", "out_of_office"))
    check("sequence office-discovery-A first touches", v(f"C{r}"), exp_sends)
    check("sequence office-discovery-A replies", v(f"D{r}"), exp_rep)
    # sync flags on the logs
    OL, BL = b.C_office, b.C_brand
    ws = wb["Office_Log"]
    flags = {ws[f"A{i}"].value: ws[f"{OL.L('calc_sync_issue')}{i}"].value for i in range(2, 2 + len(S.OFFICE_LOG)) if ws[f"{OL.L('calc_sync_issue')}{i}"].value}
    check("office log sync flags (row_ids)", sorted(flags), ["SAMPLE-O-004", "SAMPLE-O-006", "SAMPLE-O-014"])
    check("office flag O-004 text", flags.get("SAMPLE-O-004"), "Clay status is replied: set call_done or later")
    check("office flag O-006 text", flags.get("SAMPLE-O-006"), "Clay yes_in_principle should be Y")
    check("office flag O-014 text", flags.get("SAMPLE-O-014"), "row_id not found in the Clay paste: add the Clay row or fix the row_id")
    ws = wb["Brand_Log"]
    flags = {ws[f"A{i}"].value: ws[f"{BL.L('calc_sync_issue')}{i}"].value for i in range(2, 2 + len(S.BRAND_LOG)) if ws[f"{BL.L('calc_sync_issue')}{i}"].value}
    check("brand log sync flags", flags, {"SAMPLE-B-009": "Clay status is call_booked: set call_done or later"})
    sy = wb["Sync"]
    top = b.sync_top
    off = {sy[f"C{r}"].value: sy[f"G{r}"].value for r in range(top, top + 60) if sy[f"G{r}"].value}
    brd = {sy[f"J{r}"].value: sy[f"N{r}"].value for r in range(top, top + 60) if sy[f"N{r}"].value}
    check("Clay office-side flags", off, {"SAMPLE-O-020": "status call_done but no tracker row",
                                         "SAMPLE-O-030": "do_not_contact = Y but status is contacted"})
    check("Clay brand-side flags", brd, {"SAMPLE-B-021": "status paid but no paid_brands row and no tracker row"})
    check("handoff status flags", [sy["R22"].value, sy["V22"].value], [None, None])
    # handoff formulas
    so = wb["signed_offices"]
    S_ = b.signed_cols
    first_probe = first_per_org([r for r in S.OFFICE_LOG if r["row_id"] == "SAMPLE-O-001"],
                                lambda r: bool(r["budget_probe_verbatim"].strip()) or num(r["budget_probe_low_usd"]), S.AS_OF)[0]
    check("signed_offices budget_probe_verbatim (from log)", so[f"{S_.L('budget_probe_verbatim')}2"].value, first_probe["budget_probe_verbatim"])
    check("signed_offices budget_probe_date", as_date(so[f"{S_.L('budget_probe_date')}2"].value), first_probe["call_date"])
    check("signed_offices budget_probe_monthly_usd", so[f"{S_.L('budget_probe_monthly_usd')}2"].value, first_probe["budget_probe_low_usd"])
    pb = wb["paid_brands"]
    check("paid_brands payment_received_date", as_date(pb[f"{b.paid_cols.L('payment_received_date')}2"].value), S.PAID_BRANDS[0]["payment_1_date"])
    # note entry tab: the row to paste equals the first office log row
    ne = wb["Note_Office"]
    nr = b.note_row_office
    got_row = [as_date(ne.cell(row=nr, column=j + 1).value) for j in range(len(BT.OFFICE_COLS))]
    want_row = [S.OFFICE_LOG[0].get(c[0]) for c in BT.OFFICE_COLS]
    want_row = ["" if v is None else v for v in want_row]
    got_row = ["" if v is None else v for v in got_row]
    check("Note_Office row to paste equals the log row", got_row, want_row)
    check("Note_Office check cell", ne.cell(row=nr - 2, column=2).value, "OK to paste")
    check("Note_Brand row to paste is blank when empty", all(wb["Note_Brand"].cell(row=b.note_row_brand, column=j + 1).value in (None, "") for j in range(len(BT.BRAND_COLS))), True)
    # pipeline export
    pl = wb["pipeline_log"]
    check("pipeline_log header", [pl.cell(row=1, column=j).value for j in range(1, 10)], BT.PIPELINE_COLS)
    got = []
    for rr in range(2, BT.EXPORT_N + 2):
        row = [pl.cell(row=rr, column=j).value for j in range(1, 10)]
        if all(x in (None, "") for x in row):
            continue
        row[0] = as_date(row[0])
        row = [("" if x is None else x) for x in row]
        got.append(tuple(row))
    exp = expected_events()
    check("pipeline_log event count", len(got), len(exp))
    mism = [(i, g, e) for i, (g, e) in enumerate(zip(got, exp)) if tuple(g) != tuple(e)]
    check("pipeline_log rows identical to independent computation", len(mism), 0)
    for i, g, e in mism[:5]:
        say(f"   row {i}: got {g}\n          want {e}")
    dates = [g[0] for g in got]
    check("pipeline_log in date order", dates == sorted(dates), True)
    check("pipeline_log has no outreach opt_out rows (roster exits only)", sum(1 for g in got if g[4] == "opt_out"), 0)
    # data checks block
    chk = lambda label: (v(f"E{b.sc['chk_' + label]}"), v(f"F{b.sc['chk_' + label]}"))
    # O-004 has two log rows (a yes, then the withdrawal on 11-12), each flagged for Clay status replied
    check("check: office sync issues", chk("Office_Log rows with a sync issue"), (4, "FIX"))
    check("check: number cells typed as text", chk('Number cells typed as text (for example "$600")'), (0, "OK"))
    check("check: call_date with a time", chk("call_date cells with a time of day"), (0, "OK"))
    check("check: withdrawn yes answers (as of)", chk("Withdrawn yes answers (offices and brands, as of)")[0],
          expected_office(S.AS_OF)["o_wd"] + expected_brand(S.AS_OF)["b_wd"])
    # the SAMPLE data obeys the gate (R2-16): five offices signed, all approving, before the paid brand's agreement
    pb0 = S.PAID_BRANDS[0]
    check("SAMPLE paid brand: five offices signed before its agreement",
          sum(1 for o in S.SIGNED_OFFICES if o["agreement_signed_date"] < pb0["agreement_signed_date"]), 5)
    check("SAMPLE paid brand: office_approvals lists the five signed office_ids",
          sorted(pb0["office_approvals"].split(";")), sorted(o["office_id"] for o in S.SIGNED_OFFICES))
    check("check: brand sync issues", chk("Brand_Log rows with a sync issue"), (1, "FIX"))
    check("check: Clay office issues", chk("Clay office rows with an issue"), (2, "FIX"))
    check("check: Clay brand issues", chk("Clay brand rows with an issue"), (1, "FIX"))
    check("check: missing Clay columns", chk("Clay columns missing from a paste"), (0, "OK"))
    check("check: clay paste age", chk("Days since the Clay offices paste (as of)"), (0, "OK"))
    n_over = sum(1 for r in S.OFFICE_LOG + S.BRAND_LOG if r["call_date"] <= S.AS_OF and r["note_minutes"] > 10)
    check("check: notes over ten minutes", chk("Notes over the ten-minute target (as of)")[0], n_over)
    med = statistics.median([r["note_minutes"] for r in S.OFFICE_LOG + S.BRAND_LOG if r["call_date"] <= S.AS_OF])
    check("check: median minutes per note", chk("Median minutes per note (as of)")[0], med)
    check("check: yes missing owner or window", chk("Offices that said Y but owner or window is missing")[0],
          sum(1 for r in S.OFFICE_LOG if r["yes_in_principle"] == "Y" and not (r["pantry_owner_named"] == "Y" and r["receiving_window"].strip())))
    check("check: SAMPLE rows flagged", chk("SAMPLE rows in this file")[1], "SAMPLE")
    return wb, recalced


def test_sample_ignored(b):
    """Same SAMPLE file with include_sample = N: every count must drop to zero."""
    src = os.path.join(ROOT, "tracking", "conversation-tracker-SAMPLE.xlsx")
    tmp = tempfile.mkdtemp(prefix="ddtr_n_")
    p = os.path.join(tmp, "sample_ignored.xlsx")
    wb0 = openpyxl.load_workbook(src)
    wb0["Config"]["B19"] = "N"
    wb0.save(p)
    wb, _ = recalc(p)
    sc = wb["Scorecard"]
    check("include_sample=N: error cells", len(errors(wb)), 0)
    for k in ("o_conv", "o_yes", "b_conv", "b_yes", "office_first", "brand_first", "g1_signed", "g1_paid", "slot_brands_paid", "slot_offices_yes"):
        check(f"include_sample=N: {k} as of", sc[f"E{b.sc[k]}"].value, 0)
    check("include_sample=N: median reads none yet", sc[f"E{b.sc['o_bp_med']}"].value, "none yet")
    pl = wb["pipeline_log"]
    n = sum(1 for rr in range(2, 50) if pl.cell(row=rr, column=1).value not in (None, ""))
    check("include_sample=N: pipeline_log empty", n, 0)
    check("include_sample=N: banner says ignored", "(SAMPLE rows are ignored)" in (sc["A1"].value or ""), True)
    check("include_sample=N: sync flags hidden", wb["Office_Log"][f"{b.C_office.L('calc_sync_issue')}2"].value in (None, ""), True)


def test_optout_heavy(b):
    """Review round 1 (R1-06): a week of "no thanks" replies must not keep the reply-rate check green.
    Rewrite the SAMPLE Clay_Offices paste so most replies are opt-outs: the all-reply rate stays high,
    the conversation-reply rate falls under 20%, and both checks must say so."""
    src = os.path.join(ROOT, "tracking", "conversation-tracker-SAMPLE.xlsx")
    tmp = tempfile.mkdtemp(prefix="ddtr_oo_")
    p = os.path.join(tmp, "optout_heavy.xlsx")
    wb0 = openpyxl.load_workbook(src)
    ws = wb0["Clay_Offices"]
    head = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    keep_positive, n_first, n_reply, n_pos, n_oo = 3, 0, 0, 0, 0
    pos_types = ("yes", "maybe", "needs_approval", "send_info", "referral", "not_now",
                 "obj_seed_ourselves", "obj_price", "obj_free_trial", "obj_what_we_get", "obj_other")
    for rr in range(2, ws.max_row + 1):
        if not ws.cell(row=rr, column=head["row_id"]).value:
            continue
        if ws.cell(row=rr, column=head["touch_1_date"]).value:
            n_first += 1
        lr = ws.cell(row=rr, column=head["last_reply_date"]).value
        rt = ws.cell(row=rr, column=head["reply_type"]).value
        if not lr or rt in ("bounce", "out_of_office"):
            continue
        n_reply += 1
        if rt in pos_types and keep_positive > 0:
            keep_positive -= 1
            n_pos += 1
            continue
        ws.cell(row=rr, column=head["reply_type"]).value = "opt_out"
        ws.cell(row=rr, column=head["opt_out_date"]).value = lr
        ws.cell(row=rr, column=head["do_not_contact"]).value = "Y"
        ws.cell(row=rr, column=head["status"]).value = "do_not_contact"
        n_oo += 1
    wb0.save(p)
    wb, _ = recalc(p)
    sc = wb["Scorecard"]
    say(f"   planted: {n_first} first touches, {n_reply} replies, {n_pos} conversation replies, opt-outs rewritten {n_oo}")
    check("opt-out-heavy: all-reply rate still 20% or more", (sc[f"E{b.sc['office_rate']}"].value or 0) >= 0.2, True)
    check("opt-out-heavy: conversation-reply rate", sc[f"E{b.sc['office_posrate']}"].value, n_pos / n_first)
    check("opt-out-heavy: reply-rate check turns UNDER 20%", sc[f"B{b.sc['office_check']}"].value,
          "UNDER 20%: conversation replies are under 20% of first touches. Run sending/sending-setup.md section 6 before sending more.")
    check("opt-out-heavy: opt-out check flags", sc[f"B{b.sc['office_optcheck']}"].value,
          "OPT-OUTS OVER 5%: check targeting and copy (sending/sending-setup.md 6.3 to 6.4).")


def scenario_book(name, offices, brands, as_of=dt.date(2026, 11, 20), clay_brands=None):
    """Write rows into a copy of the blank workbook (include_sample = Y) and recalculate it.
    clay_brands: optional (headers, rows) pasted on Clay_Brands."""
    src = os.path.join(ROOT, "tracking", "conversation-tracker.xlsx")
    tmp = tempfile.mkdtemp(prefix="ddsc_")
    p = os.path.join(tmp, name + ".xlsx")
    wb0 = openpyxl.load_workbook(src)
    cfg = wb0["Config"]
    for rr in range(1, 40):
        if cfg.cell(rr, 1).value == "as_of_date":
            cfg.cell(rr, 2).value = as_of
        if cfg.cell(rr, 1).value == "include_sample":
            cfg.cell(rr, 2).value = "Y"
    for ws, rows in ((wb0["Office_Log"], offices), (wb0["Brand_Log"], brands)):
        hdr = [c.value for c in ws[1]]
        for i, row in enumerate(rows):
            for k, v in row.items():
                ws.cell(2 + i, hdr.index(k) + 1).value = v
    if clay_brands:
        ws = wb0["Clay_Brands"]
        heads, rows = clay_brands
        for j, h in enumerate(heads, start=1):
            ws.cell(1, j).value = h
        for i, row in enumerate(rows, start=2):
            for j, h in enumerate(heads, start=1):
                ws.cell(i, j).value = row.get(h)
    wb0.save(p)
    wb, _ = recalc(p)
    return wb


def test_scenarios(b):
    """Review round 2: the reviewer's three Gate 0 scenarios (withdrawn yes, an incomplete test, text and time inputs)
    plus a full test that fails. Expected values are the plan's reading, counted by hand."""
    DT = dt.datetime
    # one company per row_id (review round 3: the tracker counts organisations, so a shared company name would join rows)
    O = lambda rid, d, **kw: dict(row_id=rid, company=f"SAMPLE co {rid}", call_date=d, **kw)
    B = lambda rid, d, **kw: dict(row_id=rid, company=f"SAMPLE brand {rid}", call_date=d, **kw)
    # A: edge cases (the reviewer's scenario A)
    offices = [O("SAMPLE-T-O1", DT(2026, 10, 27), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9-11"),
               O("SAMPLE-T-O2", DT(2026, 11, 3), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="   "),
               O("SAMPLE-T-O3", DT(2026, 11, 5), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 8-10"),
               O("SAMPLE-T-O3", DT(2026, 11, 12), yes_in_principle="N", pantry_owner_named="Y", receiving_window=""),
               O("SAMPLE-T-O4", DT(2026, 11, 13, 14, 30), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Wed 9"),
               O("SAMPLE-T-O5", DT(2026, 11, 13), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue"),
               O("SAMPLE-T-O6", DT(2026, 11, 6), yes_in_principle="y", pantry_owner_named="y", receiving_window="Tue"),
               O("SAMPLE-T-O7", DT(2026, 11, 10), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9"),
               O("SAMPLE-T-O8", DT(2026, 11, 16), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9"),
               O("SAMPLE-T-O9", DT(2026, 11, 9), yes_in_principle="maybe", pantry_owner_named="N", receiving_window="")]
    brands = [B("SAMPLE-T-B1", DT(2026, 10, 29), controls_budget="Y", wtp_before_usd=500),
              B("SAMPLE-T-B2", DT(2026, 11, 2), controls_budget="unsure", wtp_before_usd=900),
              B("SAMPLE-T-B3", DT(2026, 11, 4), controls_budget="Y", wtp_before_usd=450, wtp_after_usd=1500),
              B("SAMPLE-T-B4", DT(2026, 11, 6), controls_budget="Y", wtp_before_usd=600),
              B("SAMPLE-T-B4", DT(2026, 11, 11), controls_budget="Y", wtp_before_usd=0, wtp_before_verbatim="SAMPLE: budget cut"),
              B("SAMPLE-T-B5", DT(2026, 11, 12), controls_budget="Y", wtp_before_usd="$600"),
              B("SAMPLE-T-B6", DT(2026, 11, 13), controls_budget="Y", wtp_before_usd=1000)]
    sc = scenario_book("scenario-A", offices, brands)["Scorecard"]
    v = lambda k, col="D": sc[f"{col}{b.sc[k]}"].value
    check("scenario A: office conversations at Gate 0 (a call at 14:30 on Gate 0 day counts)", v("o_conv"), 8)
    check("scenario A: office Gate 0 yes (withdrawn O3 out, lower-case y in, time-of-day O4 in)", v("o_yes"), 5)
    check("scenario A: offices withdrawn at Gate 0", v("o_wd"), 1)
    check("scenario A: office yes, any, still standing", v("o_yes_any"), 6)
    check("scenario A: office yes as of 11-20 (O8 after Gate 0)", v("o_yes", "E"), 6)
    check("scenario A: brand Gate 0 yes (B4 withdrawn, B5 text)", v("b_yes"), 3)
    check("scenario A: brand strict read", v("b_strict"), 2)
    check("scenario A: brands withdrawn at Gate 0", v("b_wd"), 1)
    check("scenario A: Gate 0 result", v("g0"), "PASS")
    chk = lambda label: (sc[f"E{b.sc['chk_' + label]}"].value, sc[f"F{b.sc['chk_' + label]}"].value)
    check("scenario A: data check flags the number typed as text", chk('Number cells typed as text (for example "$600")'), (1, "FIX"))
    check("scenario A: data check flags the call_date with a time", chk("call_date cells with a time of day"), (1, "FIX"))
    # B: Gate 0 day with too few conversations to run the plan's test
    offices = [O(f"SAMPLE-T-P{i}", DT(2026, 11, 2) + dt.timedelta(days=i), yes_in_principle="Y" if i <= 3 else "N",
                 pantry_owner_named="Y", receiving_window="Tue 9") for i in range(1, 7)]
    brands = [B(f"SAMPLE-T-Q{i}", DT(2026, 11, 2) + dt.timedelta(days=i), controls_budget="Y", wtp_before_usd=600 if i <= 2 else 200) for i in range(1, 6)]
    sc = scenario_book("scenario-B", offices, brands)["Scorecard"]
    check("scenario B: office side INCOMPLETE, not FAIL", v("o_status"), "INCOMPLETE: 6 of 15 conversations")
    check("scenario B: brand side INCOMPLETE, not FAIL", v("b_status"), "INCOMPLETE: 5 of 15 conversations")
    check("scenario B: Gate 0 result", v("g0"), "INCOMPLETE")
    check("scenario B: sentence carries INCOMPLETE", (sc[f"B{b.sc['sentence']}"].value or "").endswith("Gate 0: INCOMPLETE."), True)
    # C: a withdrawn yes decides the gate (plan reading on 11-13: 4 offices, 2 brands)
    offices = [O(f"SAMPLE-T-C{i}", DT(2026, 11, 1) + dt.timedelta(days=i), yes_in_principle="Y", pantry_owner_named="Y",
                 receiving_window="Tue 9") for i in range(1, 6)]
    offices.append(O("SAMPLE-T-C5", DT(2026, 11, 12), yes_in_principle="N", pantry_owner_named="Y", receiving_window=""))
    brands = [B(f"SAMPLE-T-D{i}", DT(2026, 11, 1) + dt.timedelta(days=i), controls_budget="Y", wtp_before_usd=700,
                wtp_before_verbatim="SAMPLE: about 700") for i in range(1, 4)]
    brands.append(B("SAMPLE-T-D3", DT(2026, 11, 11), controls_budget="Y", wtp_before_usd=0, wtp_before_verbatim="SAMPLE: budget frozen"))
    wb = scenario_book("scenario-C", offices, brands)
    sc = wb["Scorecard"]
    check("scenario C: office Gate 0 yes", v("o_yes"), 4)
    check("scenario C: brand Gate 0 yes", v("b_yes"), 2)
    check("scenario C: office yes before the withdrawal (Fri wk 2)", v("o_yes", "C"), 5)
    check("scenario C: Gate 0 result is not PASS", v("g0"), "INCOMPLETE")
    pl = wb["pipeline_log"]
    yes_rows = [pl.cell(rr, 3).value for rr in range(2, 60) if pl.cell(rr, 5).value == "yes_in_principle" and pl.cell(rr, 2).value == "office"]
    check("scenario C: the pipeline export drops the withdrawn office's yes", sorted(set(yes_rows)), [f"SAMPLE-T-C{i}" for i in range(1, 5)])
    ol = wb["Office_Log"]
    hdr = [c.value for c in ol[1]]
    cy = hdr.index("calc_yes_any") + 1
    check("scenario C: Clay yes_in_principle should read N for the withdrawn office", ol.cell(7, cy).value, "N")
    # D: the full test is run and the office side falls short: FAIL
    offices = [O(f"SAMPLE-T-F{i}", DT(2026, 10, 26) + dt.timedelta(days=i), yes_in_principle="Y" if i <= 4 else "N",
                 pantry_owner_named="Y", receiving_window="Tue 9") for i in range(1, 16)]
    brands = [B(f"SAMPLE-T-G{i}", DT(2026, 10, 26) + dt.timedelta(days=i), controls_budget="Y", wtp_before_usd=600 if i <= 3 else 100) for i in range(1, 16)]
    sc = scenario_book("scenario-D", offices, brands)["Scorecard"]
    check("scenario D: 4 of 15 offices is FAIL", v("o_status"), "FAIL")
    check("scenario D: 3 of 15 brands is PASS", v("b_status"), "PASS")
    check("scenario D: Gate 0 result", v("g0"), "FAIL")


def test_scenarios_r3(b):
    """Review round 3: the reviewer's scenarios E (two contacts at one company, neutral follow-ups, maybe then yes) and
    G (an office under the hard floor), a company_group and org_key scenario, and the final read (R3-09).
    Expected values are the plan's reading ("5 of 15 offices", "3 of 15 brands"), counted by hand."""
    D = dt.datetime
    O = lambda rid, co, d, **kw: dict(row_id=rid, company=co, call_date=d, **kw)
    offices = [
        O("SAMPLE-R3-E1", "SAMPLE Alpha", D(2026, 11, 2), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9-11"),
        O("SAMPLE-R3-E1", "SAMPLE Alpha", D(2026, 11, 9), pantry_owner_named="Y"),
        O("SAMPLE-R3-E2", "SAMPLE Beta", D(2026, 11, 3), yes_in_principle="Y", pantry_owner_named="N", receiving_window=""),
        O("SAMPLE-R3-E2", "SAMPLE Beta", D(2026, 11, 10), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Wed 8"),
        O("SAMPLE-R3-E3a", "SAMPLE Gamma", D(2026, 11, 4), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 10"),
        O("SAMPLE-R3-E3b", " sample gamma ", D(2026, 11, 6), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 10"),
        O("SAMPLE-R3-E5", "SAMPLE Epsilon", D(2026, 11, 14), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9"),
        O("SAMPLE-R3-E8", "SAMPLE Theta", D(2026, 11, 4), yes_in_principle="maybe", pantry_owner_named="N", receiving_window=""),
        O("SAMPLE-R3-E8", "SAMPLE Theta", D(2026, 11, 11), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 8"),
    ]
    brands = [
        O("SAMPLE-R3-F1", "SAMPLE One", D(2026, 11, 2), controls_budget="Y", wtp_before_usd=700),
        O("SAMPLE-R3-F1", "SAMPLE One", D(2026, 11, 9)),
        O("SAMPLE-R3-F2", "SAMPLE Two", D(2026, 11, 3), controls_budget="Y", wtp_before_usd=700),
        O("SAMPLE-R3-F2", "SAMPLE Two", D(2026, 11, 10), wtp_before_usd=800),
        O("SAMPLE-R3-F3a", "SAMPLE Three", D(2026, 11, 4), controls_budget="Y", wtp_before_usd=600),
        O("SAMPLE-R3-F3b", "SAMPLE Three", D(2026, 11, 5), controls_budget="Y", wtp_before_usd=600),
        O("SAMPLE-R3-F4", "SAMPLE Four", D(2026, 11, 6), controls_budget="y", wtp_before_usd=500),
        O("SAMPLE-R3-F5", "SAMPLE Five", D(2026, 11, 9), controls_budget="Y", wtp_before_usd=300, wtp_after_usd=600),
        O("SAMPLE-R3-F7", "SAMPLE Seven", D(2026, 11, 12), controls_budget="Y", wtp_after_usd=500),
    ]
    wb = scenario_book("scenario-R3-E", offices, brands)
    sc = wb["Scorecard"]
    v = lambda k, col="D": sc[f"{col}{b.sc[k]}"].value
    check("scenario E: office conversations at Gate 0 (Gamma's two contacts are one office)", v("o_conv"), 4)
    check("scenario E: office Gate 0 yes", v("o_yes"), 4)
    check("scenario E: office side INCOMPLETE, not a false PASS", v("o_status"), "INCOMPLETE: 4 of 15 conversations")
    check("scenario E: office yes as of 11-20 (Epsilon on Saturday 11-14)", v("o_yes", "E"), 5)
    check("scenario E: brand conversations at Gate 0 (Three's two founders are one brand)", v("b_conv"), 6)
    check("scenario E: brand Gate 0 yes", v("b_yes"), 6)
    check("scenario E: brand strict read", v("b_strict"), 4)
    chk = lambda label: sc[f"E{b.sc['chk_' + label]}"].value
    check("scenario E: data check counts the two shared organisations", chk("Two or more row_ids share one org_key (counted once)"), 2)
    pl = wb["pipeline_log"]
    rows = [[pl.cell(rr, c).value for c in range(1, 10)] for rr in range(2, 80) if pl.cell(rr, 1).value not in (None, "")]
    gam = sorted({r[2] for r in rows if (r[3] or "").strip().lower() == "sample gamma"})
    check("scenario E: pipeline org_id is Gamma's first row_id on every Gamma event", gam, ["SAMPLE-R3-E3a"])
    three = sorted({r[2] for r in rows if r[3] == "SAMPLE Three"})
    check("scenario E: pipeline org_id is Three's first row_id on every Three event", three, ["SAMPLE-R3-F3a"])
    ol = wb["Office_Log"]
    hdr = [c.value for c in ol[1]]
    co_ = hdr.index("calc_org") + 1
    check("scenario E: calc_org shows the key", [ol.cell(6, co_).value, ol.cell(7, co_).value], ["sample gamma", "sample gamma"])
    cy = hdr.index("calc_yes_any") + 1
    check("scenario E: Clay yes_in_principle stays per person (row_id)", [ol.cell(6, cy).value, ol.cell(7, cy).value], ["Y", "Y"])
    # G: an office under the hard floor says yes (R3-04); one at 35 is the founder's call and counts
    og = [O(f"SAMPLE-R3-G{i}", f"SAMPLE G{i}", D(2026, 11, i + 1), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9",
            tuesday_headcount=60) for i in range(1, 4)]
    og.append(O("SAMPLE-R3-G4", "SAMPLE G4", D(2026, 11, 5), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9",
                tuesday_headcount=35))
    og.append(O("SAMPLE-R3-G5", "SAMPLE G5", D(2026, 11, 9), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9",
                tuesday_headcount=12, headcount_verbatim="SAMPLE: about twelve of us in on a Tuesday"))
    og.append(O("SAMPLE-R3-G6", "SAMPLE G6", D(2026, 11, 10), yes_in_principle="Y", pantry_owner_named="Y", receiving_window="Tue 9",
                tuesday_headcount=300))
    og.append(O("SAMPLE-R3-G6", "SAMPLE G6", D(2026, 11, 12), tuesday_headcount=240, headcount_verbatim="SAMPLE: 240, I checked the badge data"))
    wb = scenario_book("scenario-R3-G", og, [])
    sc = wb["Scorecard"]
    check("scenario G: Gate 0 count leaves out the 12-person office (G6's latest number, 240, is in the band)", v("o_yes"), 5)
    check("scenario G: office side PASS on the five in the band", v("o_status"), "PASS")
    check("scenario G: yes outside 30 to 250, not counted", v("o_yes_hcout"), 1)
    check("scenario G: 30 to 39 (your call), counted", v("o_yes_3039"), 1)
    check("scenario G: data check counts the out-of-band yes", chk("Gate 0 yes with tuesday_headcount under 30 or over 250 (as of)"), 1)
    pl = wb["pipeline_log"]
    ys = sorted(pl.cell(rr, 3).value for rr in range(2, 60) if pl.cell(rr, 5).value == "yes_in_principle")
    check("scenario G: the pipeline export leaves out the 12-person office's yes", ys,
          ["SAMPLE-R3-G1", "SAMPLE-R3-G2", "SAMPLE-R3-G3", "SAMPLE-R3-G4", "SAMPLE-R3-G6"])
    # H: A's company_group links two brands with different names; a third row with no Clay row shares a grouped name;
    # two different brands that share a name are split by org_key
    heads = ["row_id", "company", "company_group", "status", "do_not_contact"]
    clay = [dict(row_id="SAMPLE-R3-H1", company="SAMPLE Moshi", company_group="grp-SAMPLE-1", status="call_done", do_not_contact="N"),
            dict(row_id="SAMPLE-R3-H2", company="SAMPLE Bruce Cost", company_group="grp-SAMPLE-1", status="call_done", do_not_contact="N"),
            dict(row_id="SAMPLE-R3-H4", company="SAMPLE Same", company_group="SAMPLE-R3-H4", status="call_done", do_not_contact="N")]
    bh = [O("SAMPLE-R3-H1", "SAMPLE Moshi", D(2026, 11, 2), controls_budget="Y", wtp_before_usd=600),
          O("SAMPLE-R3-H2", "SAMPLE Bruce Cost", D(2026, 11, 3), controls_budget="Y", wtp_before_usd=600),
          O("SAMPLE-R3-H3", "SAMPLE Moshi", D(2026, 11, 4), controls_budget="Y", wtp_before_usd=900),
          O("SAMPLE-R3-H4", "SAMPLE Same", D(2026, 11, 5), org_key="SAMPLE same east", controls_budget="Y", wtp_before_usd=600),
          O("SAMPLE-R3-H5", "SAMPLE Same", D(2026, 11, 6), org_key="SAMPLE same west", controls_budget="Y", wtp_before_usd=600)]
    wb = scenario_book("scenario-R3-H", [], bh, clay_brands=(heads, clay))
    sc = wb["Scorecard"]
    check("scenario H: company_group joins Moshi, Bruce Cost and the second Moshi row; org_key splits the two Sames", v("b_conv"), 3)
    check("scenario H: brand Gate 0 yes", v("b_yes"), 3)
    bl = wb["Brand_Log"]
    hdr = [c.value for c in bl[1]]
    co_ = hdr.index("calc_org") + 1
    check("scenario H: calc_org values", [bl.cell(rr, co_).value for rr in range(2, 7)],
          ["grp-sample-1", "grp-sample-1", "grp-sample-1", "sample same east", "sample same west"])
    # R3-09: the final read turns a side still short of 15 conversations into FAIL
    offices = [O(f"SAMPLE-T-P{i}", f"SAMPLE P{i}", D(2026, 11, 2) + dt.timedelta(days=i), yes_in_principle="Y" if i <= 3 else "N",
                 pantry_owner_named="Y", receiving_window="Tue 9") for i in range(1, 7)]
    sc = scenario_book("scenario-R3-final", offices, [], as_of=dt.date(2026, 11, 27))["Scorecard"]
    check("final read: Gate 0 column still INCOMPLETE", v("o_status"), "INCOMPLETE: 6 of 15 conversations")
    check("final read: As of Fri 2026-11-27 reads FAIL", v("o_status", "E"), "FAIL")
    check("final read: overall FAIL as of", v("g0", "E"), "FAIL")


def test_note_template_order():
    """The post-call note template lists the log columns in the tracker's order (the README says it is checked)."""
    import re as _re
    t = open(os.path.join(ROOT, "calls", "post-call-note-template.md"), encoding="utf-8").read()
    for title, cols in (("## Office note", BT.OFFICE_COLS), ("## Brand note", BT.BRAND_COLS)):
        block = t.split(title, 1)[1].split("```text", 1)[1].split("```", 1)[0]
        fields = [m.group(1) for m in _re.finditer(r"^([a-z0-9_]+):", block, flags=_re.M)]
        check(f"note template {title[3:]}: fields equal the log columns in order", fields, [c[0] for c in cols])


def test_optout_small(b):
    """Review round 2 (R2-11): 2 opt-outs in about 40 first touches is noise and must not flag; the bounce rule works the same way."""
    src = os.path.join(ROOT, "tracking", "conversation-tracker-SAMPLE.xlsx")
    tmp = tempfile.mkdtemp(prefix="ddtr_os_")
    p = os.path.join(tmp, "optout_small.xlsx")
    wb0 = openpyxl.load_workbook(src)
    ws = wb0["Clay_Offices"]
    head = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    n_first, n_oo = 0, 0
    for rr in range(2, ws.max_row + 1):
        if not ws.cell(row=rr, column=head["row_id"]).value:
            continue
        if ws.cell(row=rr, column=head["touch_1_date"]).value:
            n_first += 1
        if ws.cell(row=rr, column=head["opt_out_date"]).value:
            n_oo += 1
            if n_oo > 2:       # keep exactly two opt-outs
                ws.cell(row=rr, column=head["opt_out_date"]).value = None
                ws.cell(row=rr, column=head["reply_type"]).value = "no"
                ws.cell(row=rr, column=head["do_not_contact"]).value = "N"
                ws.cell(row=rr, column=head["status"]).value = "no"
    wb0.save(p)
    wb, _ = recalc(p)
    sc = wb["Scorecard"]
    say(f"   planted: {n_first} first touches, 2 opt-outs ({2 / n_first:.1%})")
    check("two opt-outs: rate over 5%", (sc[f"E{b.sc['office_optrate']}"].value or 0) > 0.05, 2 / n_first > 0.05)
    check("two opt-outs: no flag", sc[f"B{b.sc['office_optcheck']}"].value, "OK: fewer than 3 opt-outs, or 5% or under.")


def test_blank():
    path = os.path.join(ROOT, "tracking", "conversation-tracker.xlsx")
    b = BT.Book(False)   # only for row positions
    wb, _ = recalc(path)
    errs = errors(wb)
    check("blank workbook: error cells after recalculation", len(errs), 0)
    for e in errs[:20]:
        say("   " + e)
    sc = wb["Scorecard"]
    check("blank workbook: no SAMPLE banner", sc["A1"].value in (None, ""), True)
    check("blank workbook: include_sample", wb["Config"]["B19"].value, "N")
    return wb


def test_f_columns():
    """Every column operations (F) reads from signed_offices.csv and paid_brands.csv is in the files (parsed, not imported)."""
    import ast
    f = os.path.join(os.path.dirname(ROOT), "campaign-ops", "tools", "trackerdef.py")
    if not os.path.exists(f):
        say("skip F column check: campaign-ops/tools/trackerdef.py not found")
        return
    tree = ast.parse(open(f, encoding="utf-8").read())
    lists = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and isinstance(node.value, ast.List) and isinstance(node.targets[0], ast.Name):
            try:
                lists[node.targets[0].id] = ast.literal_eval(node.value)
            except ValueError:
                pass
    import csv as _csv
    so = next(_csv.reader(open(os.path.join(ROOT, "handoff", "signed_offices.csv"), encoding="utf-8")))
    pb = next(_csv.reader(open(os.path.join(ROOT, "handoff", "paid_brands.csv"), encoding="utf-8")))
    need_so = lists.get("OFFICE_INTAKE_COLS", []) + lists.get("SIGNED_OFFICES_EXTRA", []) + lists.get("SIGNED_OFFICES_REQUESTED", [])
    need_pb = lists.get("BRAND_INTAKE_COLS", []) + lists.get("PAID_BRANDS_EXTRA", []) + lists.get("PAID_BRANDS_REQUESTED", [])
    check("F reads from signed_offices.csv: columns missing", [c for c in need_so if c not in so], [])
    check("F reads from paid_brands.csv: columns missing", [c for c in need_pb if c not in pb], [])
    check("signed_offices.csv starts with the office_intake columns in order", so[:21], BT.OFFICE_INTAKE)
    check("paid_brands.csv starts with the brand_intake columns in order", pb[:27], BT.BRAND_INTAKE)
    # review round 1 (R1-18): the live paths ship header-only; the SAMPLE rows sit in SAMPLE- files with the same header
    for d, f in (("handoff", "signed_offices.csv"), ("handoff", "paid_brands.csv"), ("exports", "pipeline_log.csv")):
        live = list(_csv.reader(open(os.path.join(ROOT, d, f), encoding="utf-8")))
        sample = list(_csv.reader(open(os.path.join(ROOT, d, "SAMPLE-" + f), encoding="utf-8")))
        check(f"{d}/{f} is header-only (no SAMPLE rows at the live path)", len(live), 1)
        check(f"{d}/SAMPLE-{f} has the same header", sample[0], live[0])
        check(f"{d}/SAMPLE-{f} rows are all SAMPLE", all(any(str(v).startswith("SAMPLE") for v in r) for r in sample[1:]) and len(sample) > 1, True)


def main():
    b = BT.build(True, os.path.join(ROOT, "tracking", "conversation-tracker-SAMPLE.xlsx"))
    BT.build(False, os.path.join(ROOT, "tracking", "conversation-tracker.xlsx"))
    say(f"tracker test run {dt.datetime.now().isoformat(timespec='seconds')}")
    say("== SAMPLE workbook, include_sample = Y")
    wb, recalced = test_sample(b)
    keep = os.path.join(ROOT, "_work", "partB", "recalc-SAMPLE.xlsx")
    shutil.copy(recalced, keep)
    say("== SAMPLE workbook, include_sample = N")
    test_sample_ignored(b)
    say("== SAMPLE workbook, opt-out-heavy Clay paste (review round 1)")
    test_optout_heavy(b)
    say("== SAMPLE workbook, two opt-outs (review round 2)")
    test_optout_small(b)
    say("== Gate 0 scenarios: withdrawn yes, incomplete test, text and time inputs, a full test that fails (review round 2)")
    test_scenarios(b)
    say("== Gate 0 scenarios, review round 3: one organisation counted once, the headcount band, the final read")
    test_scenarios_r3(b)
    say("== post-call note template field order")
    test_note_template_order()
    say("== blank workbook")
    test_blank()
    say("== operations (F) columns")
    test_f_columns()
    say(f"RESULT: {len(FAILS)} failing of {sum(1 for l in LOG if l.startswith(('ok  ', 'FAIL')))} checks")
    with open(os.path.join(ROOT, "_work", "partB", "test_log.txt"), "w") as f:
        f.write("\n".join(LOG) + "\n")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
