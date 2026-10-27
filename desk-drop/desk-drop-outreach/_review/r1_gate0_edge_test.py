"""Independent reviewer test of the tracker's Gate 0 formulas.
Writes edge-case SAMPLE rows into a COPY of the blank workbook, recalculates with LibreOffice,
and compares Scorecard values with expectations computed by hand."""
import datetime as dt, os, shutil, subprocess, sys, tempfile
import openpyxl

SRC = sys.argv[1]
WORK = sys.argv[2]
os.makedirs(WORK, exist_ok=True)
copy = os.path.join(WORK, 'rv-test.xlsx')
shutil.copy(SRC, copy)
wb = openpyxl.load_workbook(copy)
D = lambda s: dt.datetime.fromisoformat(s)

cfg = wb['Config']
for r in range(1, 30):
    if cfg.cell(r, 1).value == 'as_of_date':
        cfg.cell(r, 2).value = D('2026-11-20')
    if cfg.cell(r, 1).value == 'include_sample':
        cfg.cell(r, 2).value = 'Y'

def put(ws, rows):
    hdr = [c.value for c in ws[1]]
    for i, row in enumerate(rows):
        for k, v in row.items():
            ws.cell(2 + i, hdr.index(k) + 1).value = v

O = lambda rid, d, **kw: dict(row_id=rid, company='SAMPLE co', call_date=D(d) if d else None, **kw)
offices = [
    O('SAMPLE-RV-O1', '2026-10-28', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9-11', budget_probe_low_usd=0),
    O('SAMPLE-RV-O2', '2026-11-02', yes_in_principle='Y', pantry_owner_named='N', receiving_window='Tue 9-11', budget_probe_low_usd=200),
    O('SAMPLE-RV-O3', '2026-11-03', yes_in_principle='Y', pantry_owner_named='Y', receiving_window=''),
    O('SAMPLE-RV-O4', '2026-11-04', yes_in_principle='maybe', pantry_owner_named='Y', receiving_window='Wed 10'),
    O('SAMPLE-RV-O5', '2026-11-05', yes_in_principle='N', pantry_owner_named='N', budget_probe_low_usd=100),
    O('SAMPLE-RV-O5', '2026-11-12', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 8-10', budget_probe_low_usd=300),
    O('SAMPLE-RV-O6', '2026-11-09', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9', budget_probe_low_usd=150),
    O('SAMPLE-RV-O6', '2026-11-10', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9'),
    O('SAMPLE-RV-O7', '2026-11-16', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9', budget_probe_low_usd=500),
    O('SAMPLE-RV-O8', '2026-11-13', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9'),
    O('SAMPLE-RV-O9', '2026-11-06', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9'),
    O('', '2026-11-06', yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9'),
    O('SAMPLE-RV-O11', None, yes_in_principle='Y', pantry_owner_named='Y', receiving_window='Tue 9'),
]
B = lambda rid, d, **kw: dict(row_id=rid, company='SAMPLE brand', call_date=D(d), **kw)
brands = [
    B('SAMPLE-RV-B1', '2026-10-29', controls_budget='Y', wtp_before_usd=600),
    B('SAMPLE-RV-B2', '2026-11-02', controls_budget='N', wtp_before_usd=800),
    B('SAMPLE-RV-B3', '2026-11-03', controls_budget='Y', wtp_before_usd=400, wtp_after_usd=1500),
    B('SAMPLE-RV-B4', '2026-11-04', controls_budget='Y', wtp_before_usd=500),
    B('SAMPLE-RV-B5', '2026-11-05', controls_budget='Y', wtp_before_usd=499),
    B('SAMPLE-RV-B6', '2026-11-10', controls_budget='Y', wtp_after_usd=600, wtp_before_verbatim='declined to name a number'),
    B('SAMPLE-RV-B7', '2026-11-16', controls_budget='Y', wtp_before_usd=2000),
    B('SAMPLE-RV-B8', '2026-11-09', controls_budget='N', wtp_before_usd=700),
    B('SAMPLE-RV-B8', '2026-11-12', controls_budget='Y', wtp_after_usd=700),
]
put(wb['Office_Log'], offices)
put(wb['Brand_Log'], brands)
wb.save(copy)

out = os.path.join(WORK, 'out'); os.makedirs(out, exist_ok=True)
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', out, copy], check=True,
               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
r = openpyxl.load_workbook(os.path.join(out, 'rv-test.xlsx'), data_only=True)['Scorecard']
def row(label):
    for i in range(1, 146):
        if r.cell(i, 1).value == label:
            return [r.cell(i, c).value for c in range(2, 6)]
    return None
exp = {
    'Office conversations held': [1, 6, 9, 10],
    'Yes in principle, pantry owner named, receiving window named': [1, 2, 5, 6],
    'Gate 0, office side': ['not yet', 'not yet', 'PASS', 'PASS'],
    'Median budget-probe answer, USD a month': [0, 'W2', 125, 150],
    'Brand conversations held': [1, 5, 7, 8],
    'Would pay $500 or more and controls the budget': [1, 3, 5, 6],
    'Gate 0, brand side': ['not yet', 'PASS', 'PASS', 'PASS'],
    'Strict read: named $500 or more BEFORE hearing $1,500, and controls the budget': [1, 2, 2, 3],
    'Median brand number before $1,500 was mentioned, USD per campaign': [600, 'W2', 550, 600],
    'Gate 0 result': ['not yet', 'not yet', 'PASS', 'PASS'],
}
# W2 budget-probe: O1 0, O2 200, O5 first row 100 -> median 100. Brand W2: 600,800,400,500,499 -> 500.
exp['Median budget-probe answer, USD a month'][1] = 100
exp['Median brand number before $1,500 was mentioned, USD per campaign'][1] = 500
fails = 0
for k, v in exp.items():
    got = row(k)
    ok = got is not None and all((g == e) or (isinstance(g, (int, float)) and isinstance(e, (int, float)) and abs(g - e) < 1e-9) for g, e in zip(got, v))
    fails += 0 if ok else 1
    print(('ok  ' if ok else 'FAIL'), k, 'got', got, 'want', v)
for i in range(1, 146):
    if r.cell(i, 1).value and str(r.cell(i, 1).value).startswith('Sentence for the Gate 0'):
        print('sentence:', r.cell(i, 2).value)
    if r.cell(i, 1).value and 'not control' in str(r.cell(i, 1).value):
        print('noctrl row:', [r.cell(i, c).value for c in range(2, 6)])
print('banner:', r['A1'].value)
print('FAILS', fails)
