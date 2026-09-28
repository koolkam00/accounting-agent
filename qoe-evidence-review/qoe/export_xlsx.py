"""Excel workpaper export for QoE Evidence Review (SPEC §8).

Writes ``QoE_Evidence_Review_<deal_id>.xlsx``: Cover, EBITDA Bridge,
Adjustment Summary, one support sheet per adjustment, Open Questions,
GL-P&L Reconciliation, Data Quality, and Review Log.

Deals workpaper conventions used throughout:

- Arial everywhere. Amounts use ``#,##0;(#,##0);"-"``: whole dollars are
  displayed, cell values keep cents. Units are named in the headers.
- Blue font = hard-coded input carried from the workpaper; black = formula;
  green = link to another sheet.
- Subtotals, differences, totals, and checks are live Excel formulas built
  only from Excel-2007 functions, so the file recalculates cleanly in Excel
  and in LibreOffice.
- Final amounts follow the review log (the latest decision per adjustment
  wins). An adjustment without a decision carries the tool proposal and is
  marked UNREVIEWED; a REQUEST_INFO item is "Pending" and is excluded from
  diligence adjusted EBITDA.
- Diligence-identified items (``source == "diligence"``, SPEC §5.7: e.g. a
  duplicate posting to reverse) are not on management's schedule. They get
  their own block in the Adjustment Summary, the Bridge (after the diligence
  revisions to management's items) and the Cover counts, and a support sheet
  like any other item ("Adj D-1"). Their claimed amount is zero, so the final
  amount is the whole adjustment.
"""

from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from qoe.money import D, fmt, period_map, q2
from qoe.periods import labels_for_month, months_in
from qoe.review_store import bridge_display_rows, bridge_row_adj_id, is_item_row
from qoe.schemas import (
    AdjustmentAssessment,
    AdjustmentCategory,
    AdjustmentClaim,
    BridgeRow,
    DealPackage,
    DocFacts,
    EvidenceQuote,
    FlagCode,
    GLEntry,
    ManagementSchedule,
    OpenQuestion,
    QuestionStatus,
    ReviewDecision,
    Severity,
    TOOL_ERROR_CORRECTIONS,
    Treatment,
    Workpaper,
)

__all__ = [
    "SHEET_BRIDGE",
    "SHEET_COVER",
    "SHEET_DATA_QUALITY",
    "SHEET_QUESTIONS",
    "SHEET_RECON",
    "SHEET_REVIEW_LOG",
    "SHEET_SUMMARY",
    "build_workbook",
    "export_workpaper",
    "find_recalc_script",
    "latest_reviews",
    "recalc_and_check",
    "resolve_final_amounts",
    "sanitize_sheet_name",
    "support_sheet_names",
    "workbook_filename",
]

# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------

FONT_NAME = "Arial"
NUMBER_FORMAT = '#,##0;(#,##0);"-"'
CHECK_FORMAT = '#,##0.00;(#,##0.00);"-"'
DATE_FORMAT = "dd-mmm-yyyy"
MONTH_FORMAT = "mmm-yy"

INPUT_BLUE = "0000FF"
FORMULA_BLACK = "000000"
LINK_GREEN = "008000"
NAVY = "1F3864"
WHITE = "FFFFFF"
MUTED = "595959"
HYPERLINK_BLUE = "0563C1"
BANNER_RED = "C00000"
SUBTOTAL_FILL = "F2F2F2"
FINAL_FILL = "DDEBF7"
SECTION_FILL = "D9E1F2"
FACTS_BAND = "E2EFDA"
JUDGMENT_BAND = "FCE4D6"
CHECK_FAIL_FILL = "FFC7CE"

# (fill, font) per treatment: ACCEPT green, REVISE amber, REJECT red, REQUEST_INFO blue-grey.
TREATMENT_STYLES: dict[Treatment, tuple[str, str]] = {
    Treatment.ACCEPT: ("C6EFCE", "006100"),
    Treatment.REVISE: ("FFEB9C", "9C5700"),
    Treatment.REJECT: ("FFC7CE", "9C0006"),
    Treatment.REQUEST_INFO: ("D6DCE4", "333F4F"),
}
TREATMENT_MEANING: dict[Treatment, str] = {
    Treatment.ACCEPT: "Evidence supports management's amount in every period (within tolerance).",
    Treatment.REVISE: "Evidence supports a different amount; the bridge carries the diligence amount.",
    Treatment.REJECT: "Evidence does not support an adjustment; the diligence amount is zero.",
    Treatment.REQUEST_INFO: "Pending information from management; excluded from diligence adjusted EBITDA.",
}
SEVERITY_STYLES: dict[Severity, tuple[str, str]] = {
    Severity.CRITICAL: ("FFC7CE", "9C0006"),
    Severity.WARNING: ("FFEB9C", "9C5700"),
    Severity.INFO: ("F2F2F2", "595959"),
}
_SEVERITY_RANK = {Severity.CRITICAL: 0, Severity.WARNING: 1, Severity.INFO: 2}

# Same vocabulary as the reviewer app (qoe.review_store), upper-cased for the workpaper.
DILIGENCE_SOURCE = "diligence"
DILIGENCE_BLOCK = "Diligence-identified items (not on management's schedule)"

UNREVIEWED = "UNREVIEWED"
AGREED = "AGREED"
OVERRIDDEN = "OVERRIDDEN"
UNREVIEWED_STYLE = ("FFFF00", "000000")  # yellow = needs the reviewer's input
PENDING = "Pending"

TICKMARKS: list[tuple[str, str]] = [
    ("T", "Traced: GL entry is part of management's claimed amount."),
    ("X", "Linked for context only (comparable period, recovery, or excess activity); not part of the claim."),
    ("D", "Agreed to the supporting document(s) listed on the entry."),
    ("F", "Entry is cited by a flag; see the Flags block on the support sheet."),
    ("V", "Quote verified verbatim against the cited page text. Quotes that fail are dropped, never repaired."),
]

SHEET_COVER = "Cover"
SHEET_BRIDGE = "EBITDA Bridge"
SHEET_SUMMARY = "Adjustment Summary"
SHEET_QUESTIONS = "Open Questions"
SHEET_RECON = "GL-P&L Reconciliation"
SHEET_DATA_QUALITY = "Data Quality"
SHEET_REVIEW_LOG = "Review Log"
_FIXED_SHEETS = (
    SHEET_COVER,
    SHEET_BRIDGE,
    SHEET_SUMMARY,
    SHEET_QUESTIONS,
    SHEET_RECON,
    SHEET_DATA_QUALITY,
    SHEET_REVIEW_LOG,
)

_CATEGORY_LABELS = {
    AdjustmentCategory.NON_RECURRING: "Non-recurring",
    AdjustmentCategory.OWNER_DISCRETIONARY: "Owner / discretionary",
    AdjustmentCategory.NORMALIZATION: "Normalization",
    AdjustmentCategory.OUT_OF_PERIOD: "Out-of-period",
    AdjustmentCategory.PRO_FORMA: "Pro forma",
    AdjustmentCategory.OTHER: "Other",
}
_ACRONYMS = {"gl": "GL", "ebitda": "EBITDA", "doc": "Document"}

_THIN = Side(style="thin", color="808080")
_THIN_BLACK = Side(style="thin", color="000000")
_DOUBLE = Side(style="double", color="000000")
_HEADER_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TOP_BORDER = Border(top=_THIN_BLACK)
_FINAL_BORDER = Border(top=_THIN_BLACK, bottom=_DOUBLE)
_CENTER = Alignment(horizontal="center", vertical="top")

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\']")
_GL_ROW = re.compile(r"^GL-R(\d+)$")
_MAX_CELL_TEXT = 32000  # Excel's hard limit is 32,767 characters per cell

# LibreOffice recalculation uses the xlsx skill's recalc.py (found under ~/.claude/skills);
# QOE_RECALC_SCRIPT points at another copy.
_RECALC_ENV = "QOE_RECALC_SCRIPT"


# ---------------------------------------------------------------------------
# Public helpers
# ---------------------------------------------------------------------------


def workbook_filename(wp: Workpaper) -> str:
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", wp.deal.deal_id).strip("_") or "deal"
    return f"QoE_Evidence_Review_{safe}.xlsx"


def sanitize_sheet_name(name: str, used: Optional[set[str]] = None) -> str:
    """Excel-safe sheet name: no []:*?/\\', at most 31 chars, unique (case-insensitive) within ``used``."""
    s = ILLEGAL_CHARACTERS_RE.sub("", _BAD_SHEET_CHARS.sub("-", name)).strip() or "Sheet"
    s = s[:31].rstrip()
    if s.lower() == "history":  # reserved by Excel
        s = "History-"
    if used is not None:
        base, n = s, 2
        while s.lower() in used:
            suffix = f" ({n})"
            s = base[: 31 - len(suffix)].rstrip() + suffix
            n += 1
        used.add(s.lower())
    return s


def support_sheet_names(wp: Workpaper) -> dict[str, str]:
    """adj_id -> support sheet name ("Adj M-01"), unique and distinct from the fixed sheets."""
    used = {s.lower() for s in _FIXED_SHEETS}
    return {a.adj_id: sanitize_sheet_name(f"Adj {a.adj_id}", used) for a in wp.assessments}


def latest_reviews(wp: Workpaper) -> dict[str, ReviewDecision]:
    """Latest decision per adjustment; the review log is append-only, so list order decides."""
    latest: dict[str, ReviewDecision] = {}
    for rv in wp.reviews:
        latest[rv.adj_id] = rv
    return latest


def _fallback_final_amounts(wp: Workpaper, reviews: dict[str, ReviewDecision]) -> dict[str, dict[str, str]]:
    """Same semantics as qoe.review_store.final_amounts, used when that module is unavailable."""
    out: dict[str, dict[str, str]] = {}
    for a in wp.assessments:
        rv = reviews.get(a.adj_id)
        if rv is not None:
            out[a.adj_id] = {} if rv.treatment == Treatment.REQUEST_INFO else dict(rv.amounts)
        elif a.treatment == Treatment.REQUEST_INFO:
            out[a.adj_id] = {}
        else:
            out[a.adj_id] = dict(a.proposed)
    return out


def resolve_final_amounts(wp: Workpaper) -> dict[str, dict[str, str]]:
    """adj_id -> period -> final amount; ``{}`` means pending (excluded from diligence EBITDA)."""
    reviews = latest_reviews(wp)
    fallback = _fallback_final_amounts(wp, reviews)
    try:
        from qoe.review_store import final_amounts
    except ImportError:
        raw = fallback
    else:
        raw = final_amounts(wp, reviews)
    labels = _period_labels(wp)
    out: dict[str, dict[str, str]] = {}
    for a in wp.assessments:
        amounts = raw.get(a.adj_id, fallback[a.adj_id])
        out[a.adj_id] = period_map(amounts, labels) if amounts else {}
    return out


def find_recalc_script() -> Optional[Path]:
    """Locate the xlsx skill's LibreOffice recalc script, or None."""
    env = os.environ.get(_RECALC_ENV)
    candidates: list[Path] = [Path(env)] if env else []
    skills = Path.home() / ".claude" / "skills"
    if skills.is_dir():
        for pattern in ("*/xlsx/scripts/recalc.py", "*/*/xlsx/scripts/recalc.py"):
            candidates.extend(sorted(skills.glob(pattern)))
    for c in candidates:
        if c.is_file():
            return c
    return None


def recalc_and_check(path: Path, timeout: int = 90) -> dict[str, Any]:
    """Recalculate every formula with LibreOffice (rewrites the file in place) and report errors.

    Returns the recalc script's JSON (``status`` / ``total_errors`` / ``error_summary`` /
    ``total_formulas``), ``{"status": "skipped", "reason": ...}`` when LibreOffice or the
    script is unavailable, or ``{"error": ...}`` when recalculation failed.
    """
    script = find_recalc_script()
    if script is None:
        return {"status": "skipped", "reason": "recalc.py not found (set QOE_RECALC_SCRIPT)"}
    if shutil.which("soffice") is None:
        return {"status": "skipped", "reason": "LibreOffice (soffice) is not on PATH"}
    result = _run_recalc(script, Path(path), timeout)
    if "error" in result:
        # A concurrent LibreOffice start-up occasionally fails or stalls; one retry clears it.
        result = _run_recalc(script, Path(path), timeout)
    return result


def _run_recalc(script: Path, path: Path, timeout: int) -> dict[str, Any]:
    try:
        proc = subprocess.run(
            [sys.executable, str(script), str(path.resolve()), str(timeout)],
            capture_output=True,
            text=True,
            timeout=timeout + 60,
            cwd=str(script.parent),
        )
    except subprocess.TimeoutExpired:
        return {"error": f"recalc timed out after {timeout + 60}s"}
    try:
        result = json.loads(proc.stdout)
    except json.JSONDecodeError:
        return {"error": "recalc produced no JSON", "stdout": proc.stdout[-2000:], "stderr": proc.stderr[-2000:]}
    return result if isinstance(result, dict) else {"error": f"unexpected recalc output: {result!r}"}


def export_workpaper(wp: Workpaper, out_path: Path, *, pkg: Optional[DealPackage] = None) -> Path:
    """Write the Excel workpaper and return its path.

    ``out_path`` is either the .xlsx path or a directory (the file is then named
    ``QoE_Evidence_Review_<deal_id>.xlsx``). ``pkg`` is optional: when given, the
    support sheets show management's description and full GL detail (date,
    account, counterparty, doc #, memo) for linked entries, which the
    workpaper itself does not carry.
    """
    out_path = Path(out_path)
    if out_path.is_dir() or out_path.suffix.lower() != ".xlsx":
        out_path = out_path / workbook_filename(wp)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(wp, pkg=pkg)
    tmp = out_path.with_name(out_path.stem + ".partial.xlsx")
    wb.save(tmp)
    os.replace(tmp, out_path)  # a reviewer's open copy is never left half-written
    return out_path


# ---------------------------------------------------------------------------
# Context
# ---------------------------------------------------------------------------


@dataclass
class _Ctx:
    wp: Workpaper
    labels: list[str]
    latest: dict[str, ReviewDecision]
    final: dict[str, dict[str, str]]
    gl_by_id: dict[str, GLEntry]
    claims: dict[str, AdjustmentClaim]
    doc_facts: dict[str, DocFacts]
    adj_sheets: dict[str, str]
    questions: dict[str, list[OpenQuestion]]
    has_gl_detail: bool

    @property
    def currency(self) -> str:
        return self.wp.deal.currency or "USD"

    def status(self, adj_id: str) -> str:
        rv = self.latest.get(adj_id)
        if rv is None:
            return UNREVIEWED
        return OVERRIDDEN if _differs_from_tool(rv.treatment, rv.amounts, rv.tool_treatment, rv.tool_amounts,
                                                self.wp.deal.tolerance) else AGREED

    def stale(self, rv: ReviewDecision) -> bool:
        """The tool's current proposal differs from what the reviewer saw."""
        a = next((x for x in self.wp.assessments if x.adj_id == rv.adj_id), None)
        if a is None:
            return False
        return _differs_from_tool(a.treatment, a.proposed, rv.tool_treatment, rv.tool_amounts, self.wp.deal.tolerance)

    def final_treatment(self, a: AdjustmentAssessment) -> Treatment:
        rv = self.latest.get(a.adj_id)
        return rv.treatment if rv is not None else a.treatment

    @property
    def mgmt_items(self) -> list[AdjustmentAssessment]:
        return [a for a in self.wp.assessments if not _is_diligence(a)]

    @property
    def dil_items(self) -> list[AdjustmentAssessment]:
        return [a for a in self.wp.assessments if _is_diligence(a)]


def _is_diligence(a: AdjustmentAssessment) -> bool:
    return a.source == DILIGENCE_SOURCE


def _period_labels(wp: Workpaper) -> list[str]:
    return list(wp.bridge.period_labels) or [p.label for p in wp.deal.periods]


def _context(wp: Workpaper, pkg: Optional[DealPackage]) -> _Ctx:
    schedule: Optional[ManagementSchedule] = pkg.schedule if pkg is not None else getattr(wp, "schedule", None)
    claims = {c.adj_id: c for c in schedule.adjustments} if schedule is not None else {}
    gl_by_id = {e.entry_id: e for e in pkg.gl} if pkg is not None else {}
    return _Ctx(
        wp=wp,
        labels=_period_labels(wp),
        latest=latest_reviews(wp),
        final=resolve_final_amounts(wp),
        gl_by_id=gl_by_id,
        claims=claims,
        doc_facts={f.doc_id: f for f in wp.doc_facts},
        adj_sheets=support_sheet_names(wp),
        questions=_questions_with_updates(wp),
        has_gl_detail=pkg is not None,
    )


def _questions_with_updates(wp: Workpaper) -> dict[str, list[OpenQuestion]]:
    """Open questions per adjustment with reviewer status/response updates applied in log order.

    An update value is a status ("ANSWERED"), a status with a note ("ANSWERED: insurer
    confirmed"), or a free-text response.
    """
    statuses = {s.value for s in QuestionStatus}
    by_id: dict[str, OpenQuestion] = {}
    order: dict[str, list[str]] = {}
    for a in wp.assessments:
        order[a.adj_id] = []
        for q in a.open_questions:
            by_id[q.q_id] = q
            order[a.adj_id].append(q.q_id)
    for rv in wp.reviews:
        for q_id, note in rv.question_updates.items():
            q = by_id.get(q_id)
            if q is None:
                continue
            text = note.strip()
            head, sep, rest = text.partition(":")
            if text.upper() in statuses:
                by_id[q_id] = q.model_copy(update={"status": QuestionStatus(text.upper())})
            elif sep and head.strip().upper() in statuses:
                by_id[q_id] = q.model_copy(
                    update={"status": QuestionStatus(head.strip().upper()), "response": rest.strip()}
                )
            elif text:
                by_id[q_id] = q.model_copy(update={"response": text})
    return {adj: [by_id[q] for q in ids] for adj, ids in order.items()}


# ---------------------------------------------------------------------------
# Cell writing
# ---------------------------------------------------------------------------


def _clean(value: str) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", value)
    return text if len(text) <= _MAX_CELL_TEXT else text[: _MAX_CELL_TEXT - 3] + "..."


def _fill(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)


def _quoted(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def _xref(sheet: str, col: int, row: int) -> str:
    return f"{_quoted(sheet)}!{get_column_letter(col)}{row}"


def _a1(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _to_date(iso: Optional[str]) -> Optional[date]:
    if not iso:
        return None
    try:
        return date.fromisoformat(iso[:10])
    except ValueError:
        return None


def _month_date(month: str) -> Optional[date]:
    try:
        y, m = month.split("-")[:2]
        return date(int(y), int(m), 1)
    except (ValueError, AttributeError):
        return None


class _Sheet:
    """Thin writer over a worksheet: consistent Arial styling, merged spans, row-height fitting."""

    def __init__(self, ws: Worksheet, widths: Sequence[float]):
        self.ws = ws
        self.widths = list(widths)
        for i, w in enumerate(self.widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    @property
    def last_col(self) -> int:
        return len(self.widths)

    def span_width(self, col: int, span: int) -> float:
        return sum(self.widths[c - 1] if c - 1 < len(self.widths) else 9.0 for c in range(col, col + span))

    def put(
        self,
        row: int,
        col: int,
        value: Any = None,
        *,
        span: int = 1,
        bold: bool = False,
        italic: bool = False,
        color: str = FORMULA_BLACK,
        size: float = 10,
        fill: Optional[str] = None,
        num_fmt: Optional[str] = None,
        wrap: bool = False,
        halign: Optional[str] = None,
        valign: str = "top",
        border: Optional[Border] = None,
        underline: Optional[str] = None,
        indent: int = 0,
    ) -> Cell:
        cell = self.ws.cell(row=row, column=col)
        if isinstance(value, str):
            value = _clean(value)
            cell.value = value
            if value.startswith("="):  # text that merely looks like a formula stays text
                cell.data_type = "s"
        else:
            cell.value = value
        cell.font = Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=color, underline=underline)
        cell.alignment = Alignment(horizontal=halign, vertical=valign, wrap_text=wrap, indent=indent)
        if num_fmt:
            cell.number_format = num_fmt
        targets = [self.ws.cell(row=row, column=c) for c in range(col, col + span)]
        if fill:
            for t in targets:
                t.fill = _fill(fill)
        if border is not None:
            for t in targets:
                t.border = border
        if span > 1:
            self.ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
        if wrap and isinstance(value, str):
            self.fit(row, value, self.span_width(col, span))
        return cell

    def fit(self, row: int, text: str, width: float) -> None:
        # Excel never auto-fits merged cells, so wrapped rows get an explicit height.
        per_line = max(int(width * 1.1), 6)
        lines = sum(max(1, math.ceil(len(part) / per_line)) for part in text.split("\n"))
        if lines <= 1:
            return
        height = min(409.0, 12.75 * lines + 3)
        current = self.ws.row_dimensions[row].height or 0
        if height > current:
            self.ws.row_dimensions[row].height = height

    def text(self, row: int, col: int, value: Any, **kw: Any) -> Cell:
        kw.setdefault("wrap", True)
        return self.put(row, col, "" if value is None else value, **kw)

    def money(self, row: int, col: int, value: object, **kw: Any) -> Cell:
        kw.setdefault("color", INPUT_BLUE)
        return self.put(row, col, q2(D(value)), num_fmt=NUMBER_FORMAT, halign="right", **kw)

    def pending(self, row: int, col: int, **kw: Any) -> Cell:
        return self.put(row, col, PENDING, italic=True, color=MUTED, halign="right", **kw)

    def formula(
        self, row: int, col: int, text: str, *, link: bool = False, num_fmt: str = NUMBER_FORMAT, **kw: Any
    ) -> Cell:
        kw.setdefault("halign", "right")
        cell = self.put(row, col, None, num_fmt=num_fmt, color=LINK_GREEN if link else FORMULA_BLACK, **kw)
        cell.value = text
        return cell

    def count(self, row: int, col: int, value: int, **kw: Any) -> Cell:
        kw.setdefault("color", INPUT_BLUE)
        return self.put(row, col, value, num_fmt="#,##0", halign="right", **kw)

    def header(self, row: int, cols: Sequence[tuple[str, int]], start_col: int = 1) -> None:
        c = start_col
        for title, span in cols:
            self.put(
                row, c, title, span=span, bold=True, color=WHITE, fill=NAVY, wrap=True,
                halign="center", valign="center", border=_HEADER_BORDER,
            )
            c += span

    def header_tall(self, top: int, bottom: int, col: int, title: str) -> None:
        """A header cell merged down across a two-row header band."""
        self.header(top, [(title, 1)], start_col=col)
        for r in range(top + 1, bottom + 1):
            self.put(r, col, None, fill=NAVY, border=_HEADER_BORDER)
        self.ws.merge_cells(start_row=top, start_column=col, end_row=bottom, end_column=col)

    def section(self, row: int, title: str, *, band: Optional[str] = None, color: str = NAVY) -> None:
        self.put(row, 1, title, span=self.last_col if band else 1, bold=True, size=11, color=color, fill=band)
        if band:
            self.ws.row_dimensions[row].height = 18

    def link(self, row: int, col: int, text: str, sheet: str, *, span: int = 1, bold: bool = False) -> Cell:
        cell = self.put(row, col, text, span=span, bold=bold, color=HYPERLINK_BLUE, underline="single")
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{_quoted(sheet)}!A1", display=text)
        return cell

    def treatment(self, row: int, col: int, t: Optional[Treatment], **kw: Any) -> Cell:
        if t is None:
            return self.put(row, col, "", **kw)
        fill, font = TREATMENT_STYLES[t]
        return self.put(row, col, t.value, bold=True, color=font, fill=fill, halign="center", **kw)

    def status(self, row: int, col: int, status: str, **kw: Any) -> Cell:
        if status == UNREVIEWED:
            fill, font = UNREVIEWED_STYLE
            return self.put(row, col, status, bold=True, color=font, fill=fill, halign="center", **kw)
        return self.put(row, col, status, bold=status == OVERRIDDEN, color=NAVY if status == OVERRIDDEN else
                        FORMULA_BLACK, halign="center", **kw)

    def severity(self, row: int, col: int, sev: Severity, **kw: Any) -> Cell:
        fill, font = SEVERITY_STYLES[sev]
        return self.put(row, col, sev.value, bold=True, color=font, fill=fill, halign="center", **kw)


def _title_block(sh: _Sheet, ctx: _Ctx, title: str, subtitle: str = "") -> int:
    deal = ctx.wp.deal
    sh.put(1, 1, title, bold=True, size=14, color=NAVY)
    sh.put(2, 1, f"{deal.target_name}  |  Deal {deal.deal_id}  |  Run {ctx.wp.run_id}", bold=True)
    sh.put(3, 1, _banner_text(ctx), bold=True, color=BANNER_RED)
    if subtitle:
        sh.put(4, 1, subtitle, italic=True, color=MUTED)
    return 6


def _banner_text(ctx: _Ctx) -> str:
    if ctx.wp.deal.synthetic:
        return "SYNTHETIC: generated for QoE Evidence Review testing. Fictitious company and data; not for reliance."
    return "DRAFT: diligence workpaper subject to reviewer sign-off."


def _units(ctx: _Ctx) -> str:
    return (
        f"{ctx.currency}. Whole dollars displayed; cell values keep cents. Positive amounts increase EBITDA. "
        "Blue = input from the workpaper; black = formula; green = link to another sheet."
    )


# ---------------------------------------------------------------------------
# Formatting helpers for evidence
# ---------------------------------------------------------------------------


def _gl_row(entry_id: str) -> Optional[int]:
    m = _GL_ROW.match(entry_id)
    return int(m.group(1)) if m else None


def _gl_rows_text(entry_ids: Iterable[str], limit: int = 60) -> str:
    ids = list(dict.fromkeys(entry_ids))
    if not ids:
        return ""
    rows = sorted((r, e) for e in ids if (r := _gl_row(e)) is not None)
    others = [e for e in ids if _gl_row(e) is None]
    parts = [str(r) for r, _ in rows] + others
    more = ""
    if len(parts) > limit:
        more = f" (+{len(parts) - limit} more)"
        parts = parts[:limit]
    noun = "GL row" if len(parts) == 1 and not more else "GL rows"
    return f"{noun} {', '.join(parts)}{more}"


def _flag_label(code: FlagCode) -> str:
    words = [_ACRONYMS.get(w, w) for w in code.value.lower().split("_")]
    text = " ".join(words)
    return text[0].upper() + text[1:]


def _cite(q: EvidenceQuote) -> str:
    return f"{q.doc_id}, p. {q.page}"


def _quote_text(q: EvidenceQuote) -> str:
    return f'"{q.quote}"'


def _dedupe_quotes(quotes: Iterable[EvidenceQuote]) -> list[EvidenceQuote]:
    seen: set[tuple[str, int, str]] = set()
    out: list[EvidenceQuote] = []
    for q in quotes:
        key = (q.doc_id, q.page, q.quote)
        if key not in seen:
            seen.add(key)
            out.append(q)
    return out


def _key_flags(a: AdjustmentAssessment) -> str:
    seen: dict[FlagCode, Severity] = {}
    for f in sorted(a.flags, key=lambda f: _SEVERITY_RANK[f.severity]):
        seen.setdefault(f.code, f.severity)
    return "; ".join(f"{_flag_label(c)} ({s.value[0]})" for c, s in seen.items())


def _signed_status(facts: Optional[DocFacts]) -> str:
    if facts is None:
        return ""
    if facts.is_draft:
        return "DRAFT"
    if facts.is_signed is True:
        return "Signed"
    if facts.is_signed is False:
        return "Unsigned"
    return "Not stated"


def _differs_from_tool(
    treatment: Treatment, amounts: dict[str, str], tool_treatment: Treatment, tool_amounts: dict[str, str],
    tolerance: object,
) -> bool:
    """Same rule as the reviewer app: a different treatment, or amounts apart by more than the tolerance."""
    if treatment != tool_treatment:
        return True
    if treatment == Treatment.REQUEST_INFO:
        return False
    tol = D(tolerance)
    return any(abs(D(amounts.get(k)) - D(tool_amounts.get(k))) > tol for k in set(amounts) | set(tool_amounts))


def _review_pending(rv: ReviewDecision) -> bool:
    return rv.treatment == Treatment.REQUEST_INFO or not rv.amounts


# ---------------------------------------------------------------------------
# Support sheets (SPEC §8.4)
# ---------------------------------------------------------------------------

# A..M: id | row/page | label/account | D.. period columns (>= 14 wide) | text columns.
_SUPPORT_WIDTHS = [12, 9, 34, 14, 14, 14, 14, 26, 14, 44, 20, 32, 36]


@dataclass
class _TieOut:
    sheet: str
    first_col: int
    claimed_row: int
    proposed_row: int
    final_row: int
    revision_row: int
    check_row: int


def _write_support(ws: Worksheet, ctx: _Ctx, a: AdjustmentAssessment) -> _TieOut:
    n = len(ctx.labels)
    widths = list(_SUPPORT_WIDTHS)
    while len(widths) < 3 + n + 1:
        widths.append(14)
    for i in range(3, 3 + n):
        widths[i] = max(widths[i], 14)
    sh = _Sheet(ws, widths)
    last = sh.last_col
    pc = 4  # first period column (D)
    name = ctx.adj_sheets[a.adj_id]
    claim = ctx.claims.get(a.adj_id)
    rv = ctx.latest.get(a.adj_id)
    status = ctx.status(a.adj_id)

    _title_block(sh, ctx, f"{name}: {a.title}", "Support schedule. " + _units(ctx))
    sh.link(5, 1, "<< EBITDA Bridge", SHEET_BRIDGE, span=2)
    sh.link(5, 3, "<< Adjustment Summary", SHEET_SUMMARY)
    row = 7

    # Conclusion strip.
    sh.section(row, "CONCLUSION")
    row += 1
    sh.put(row, 1, "Tool treatment", span=2, bold=True)
    sh.treatment(row, 3, a.treatment)
    sh.put(row, 4, f"Confidence: {a.confidence}", span=2, italic=True)
    row += 1
    sh.put(row, 1, "Final treatment", span=2, bold=True)
    sh.treatment(row, 3, ctx.final_treatment(a))
    sh.status(row, 4, status, span=2)
    row += 1
    sh.put(row, 1, "Reviewer", span=2, bold=True)
    if rv is not None:
        who = f"{rv.reviewer} at {rv.timestamp}; correction type {rv.correction_type.value}"
        if ctx.stale(rv):
            who += ". NOTE: the tool proposal has changed since this decision; re-review."
        sh.text(row, 3, who, span=last - 2)
    else:
        sh.text(
            row, 3, "No reviewer decision recorded. The tool proposal is carried and marked UNREVIEWED.",
            span=last - 2, bold=True, color=BANNER_RED,
        )
    row += 1
    sh.put(row, 1, "Tool rationale", span=2, bold=True)
    sh.text(row, 3, a.rationale or "", span=last - 2)
    row += 2

    # Management's claim (or, for a diligence-identified item, the tool's basis).
    diligence = _is_diligence(a)
    sh.section(row, DILIGENCE_BLOCK.upper() if diligence else "MANAGEMENT'S CLAIM")
    row += 1
    for label, value in _claim_pairs(a, claim):
        sh.put(row, 1, label, span=2, bold=True)
        sh.text(row, 3, value, span=last - 2)
        row += 1
    row += 1

    # Tie-out.
    sh.section(row, "TIE-OUT BY PERIOD")
    row += 1
    sh.header(row, [(f"{ctx.currency}", 3)] + [(p, 1) for p in ctx.labels])
    row += 1
    rows: dict[str, int] = {}
    pending_tool = a.treatment == Treatment.REQUEST_INFO
    final = ctx.final.get(a.adj_id, {})
    final_label = "(e) Final: reviewer decision" if rv is not None else "(e) Final: tool proposal carried (UNREVIEWED)"
    for key, label, values in (
        ("a", "(a) Claimed by management" + (" (not on the schedule: zero)" if diligence else ""), a.claimed),
        ("b", "(b) Traced to GL (claimed entries)", a.traced_gl),
        ("c", "(c) Documented (traced GL with document support)", a.documented),
        ("d", "(d) Tool proposed", None if pending_tool else a.proposed),
        ("e", final_label, final or None),
    ):
        rows[key] = row
        sh.put(row, 1, label, span=3, bold=key in ("d", "e"))
        for i, p in enumerate(ctx.labels):
            if values is None:
                sh.pending(row, pc + i)
            else:
                sh.money(row, pc + i, values.get(p), bold=key in ("d", "e"))
        row += 1
    row += 1
    diffs = (
        ("(a) - (b) Claimed less traced to GL", "a", "b", False),
        ("(b) - (c) Traced to GL less documented", "b", "c", False),
        ("(d) - (a) Tool proposed less claimed", "d", "a", True),
        ("(e) - (a) Final less claimed (bridge revision)", "e", "a", True),
    )
    revision_row = row
    for label, left, right, guard in diffs:
        sh.put(row, 1, label, span=3, italic=True)
        for i in range(n):
            col = pc + i
            lhs = _a1(col, rows[left])
            # A pending amount is text; it counts as zero, so the revision reverses the claim.
            lhs_expr = f"IF(ISNUMBER({lhs}),{lhs},0)" if guard else lhs
            sh.formula(row, col, f"={lhs_expr}-{_a1(col, rows[right])}", border=_TOP_BORDER if left == "a" else None)
        if left == "e":
            revision_row = row
        row += 1
    check_row = row
    sh.put(row, 1, "Check: EBITDA Bridge revision less (e) - (a); should be zero", span=3, italic=True, color=MUTED)
    row += 2

    # Flags.
    sh.section(row, f"FLAGS ({len(a.flags)})")
    row += 1
    if a.flags:
        sh.header(row, [("Severity", 1), ("Flag", 2), ("Period", 1), (f"EBITDA impact ({ctx.currency})", 1),
                        ("Message", 6), ("Evidence", 1), ("Related adj.", 1)])
        row += 1
        for f in sorted(a.flags, key=lambda f: _SEVERITY_RANK[f.severity]):
            sh.severity(row, 1, f.severity)
            sh.text(row, 2, _flag_label(f.code), span=2, bold=True)
            sh.text(row, 4, f.period_label or "All periods")
            if f.amount_impact is not None:
                sh.money(row, 5, f.amount_impact)
            sh.text(row, 6, f.message, span=6, indent=1)
            evidence = "; ".join(x for x in (_gl_rows_text(f.entry_ids), "; ".join(f.doc_ids)) if x)
            sh.text(row, 12, evidence)
            sh.text(row, 13, ", ".join(f.related_adj_ids))
            row += 1
            row = _quote_rows(sh, row, f.quotes)
    else:
        sh.put(row, 1, "No flags raised.", italic=True, color=MUTED)
        row += 1
    row += 1

    # Documented facts and judgment questions are deliberately separate blocks.
    sh.section(row, "DOCUMENTED FACTS: what the evidence establishes (GL rows and verbatim quotes cited)",
               band=FACTS_BAND)
    row += 1
    if a.facts:
        sh.header(row, [("#", 1), ("Fact", 10), ("GL rows cited", 2)])
        row += 1
        for i, fact in enumerate(a.facts, start=1):
            sh.put(row, 1, f"F{i}", bold=True, halign="center")
            sh.text(row, 2, fact.text, span=10)
            sh.text(row, 12, _gl_rows_text(fact.entry_ids), span=2)
            row += 1
            row = _quote_rows(sh, row, fact.quotes)
    else:
        sh.put(row, 1, "No documented facts recorded.", italic=True, color=MUTED)
        row += 1
    row += 1
    sh.section(row, "JUDGMENT QUESTIONS: calls for the reviewer; not established by the evidence", band=JUDGMENT_BAND)
    row += 1
    if a.judgment_questions:
        sh.header(row, [("#", 1), ("Judgment question", last - 1)])
        row += 1
        for i, jq in enumerate(a.judgment_questions, start=1):
            sh.put(row, 1, f"J{i}", bold=True, halign="center")
            sh.text(row, 2, jq, span=last - 1)
            row += 1
    else:
        sh.put(row, 1, "No open judgment questions.", italic=True, color=MUTED)
        row += 1
    row += 1

    # Questions for management.
    questions = ctx.questions.get(a.adj_id, [])
    sh.section(row, f"QUESTIONS FOR MANAGEMENT ({len(questions)})")
    row += 1
    if questions:
        sh.header(row, [("Q id", 1), ("Question", 9), ("Priority", 1), ("Status / response", 1), ("Basis", 1)])
        row += 1
        for q in questions:
            sh.text(row, 1, q.q_id)
            sh.text(row, 2, q.text, span=9)
            high = q.priority == "high"
            sh.text(row, 11, q.priority, bold=high, color=BANNER_RED if high else FORMULA_BLACK)
            sh.text(row, 12, q.status.value + (f": {q.response}" if q.response else ""))
            sh.text(row, 13, q.basis)
            row += 1
    else:
        sh.put(row, 1, "No questions for management.", italic=True, color=MUTED)
        row += 1
    row += 1

    # Recurrence observations.
    sh.section(row, f"RECURRENCE OBSERVATIONS ({len(a.recurrence)})")
    row += 1
    if a.recurrence:
        note_col = pc + n
        sh.header(row, [("#", 1), ("Group", 2)] + [(f"{p} ({ctx.currency})", 1) for p in ctx.labels]
                  + [("Note / comparable GL rows", max(1, last - note_col + 1))])
        row += 1
        for i, obs in enumerate(a.recurrence, start=1):
            sh.put(row, 1, f"R{i}", bold=True, halign="center")
            sh.text(row, 2, obs.group, span=2)
            for j, p in enumerate(ctx.labels):
                sh.money(row, pc + j, obs.amounts_by_period.get(p))
            note = "; ".join(x for x in (obs.note, _gl_rows_text(obs.entry_ids)) if x)
            sh.text(row, note_col, note, span=max(1, last - note_col + 1), indent=1)
            row += 1
    else:
        sh.put(row, 1, "No recurrence observed.", italic=True, color=MUTED)
        row += 1
    row += 1

    row = _gl_table(sh, ctx, a, row)
    row = _documents_block(sh, ctx, a, row)
    row = _decision_block(sh, ctx, a, row)

    ws.freeze_panes = "A6"
    return _TieOut(
        sheet=name,
        first_col=pc,
        claimed_row=rows["a"],
        proposed_row=rows["d"],
        final_row=rows["e"],
        revision_row=revision_row,
        check_row=check_row,
    )


def _claim_pairs(a: AdjustmentAssessment, claim: Optional[AdjustmentClaim]) -> list[tuple[str, str]]:
    """Label / value lines describing the item. Management's schedule (from the deal package or the
    workpaper) is the source when available; otherwise the assessment carries the same fields."""
    category = _CATEGORY_LABELS.get(a.category, a.category.value)
    if _is_diligence(a):
        return [
            ("Item", f"{a.adj_id}: {a.title}"),
            ("Source", "Identified by diligence from the GL and data room; not on management's schedule. The "
                       "claimed amount is zero, so the final amount is the whole diligence adjustment."),
            ("Category", category),
            ("Basis", a.description or "No basis recorded in the workpaper; see the tool rationale and flags."),
            ("GL accounts", ", ".join(a.gl_accounts) or "None recorded"),
            ("Support", "; ".join(a.support_refs) or "See the documents block below"),
        ]
    description = (claim.description if claim is not None else "") or a.description
    if not description:
        description = ("No description provided in the schedule." if claim is not None else
                       "Not carried in the workpaper; see management's adjusted EBITDA schedule.")
    pairs = [
        ("Adjustment", f"{a.adj_id}: {claim.title if claim else a.title}"),
        ("Category", f"{category}" + (f" (as labelled: {claim.category_raw})" if claim and claim.category_raw else "")),
        ("Description", description),
        ("GL accounts", ", ".join(claim.gl_accounts if claim is not None else a.gl_accounts) or "None cited"),
        ("Support refs", "; ".join(claim.support_refs if claim is not None else a.support_refs) or "None cited"),
    ]
    if claim is not None:
        pairs.append(("Schedule row", str(claim.source_row)))
    return pairs


def _quote_rows(sh: _Sheet, row: int, quotes: Iterable[EvidenceQuote]) -> int:
    for q in _dedupe_quotes(quotes):
        sh.put(row, 1, "V", bold=True, color=LINK_GREEN, halign="center")
        sh.text(row, 2, _cite(q), span=2, color=MUTED)
        sh.text(row, 4, _quote_text(q), span=sh.last_col - 3, italic=True)
        row += 1
    return row


def _gl_table(sh: _Sheet, ctx: _Ctx, a: AdjustmentAssessment, row: int) -> int:
    links = a.gl_links
    sh.section(row, f"LINKED GL ENTRIES ({len(links)})")
    row += 1
    if not ctx.has_gl_detail and links:
        sh.put(row, 1, "Date, account, counterparty, doc # and memo are blank: the workpaper was exported "
                       "without the source GL.", italic=True, color=MUTED)
        row += 1
    if not links:
        sh.put(row, 1, "No GL entries linked.", italic=True, color=MUTED)
        return row + 2
    sh.header(row, [("Entry ID", 1), ("GL row", 1), ("Account", 1), ("Date", 1),
                    (f"Amount ({ctx.currency}, debit +)", 1), ("Tick", 1), ("Supports claim?", 1),
                    ("Counterparty", 1), ("Doc #", 1), ("Memo", 1), ("Analysis period(s)", 1),
                    ("Documents", 1), ("Flags / link reasons", 1)])
    header_row = row
    row += 1
    flagged: dict[str, list[str]] = {}
    for f in a.flags:
        for e in f.entry_ids:
            flagged.setdefault(e, []).append(_flag_label(f.code))
    periods = ctx.wp.deal.periods
    first = row
    ordered = sorted(links, key=lambda lk: (not lk.supports_claim, lk.period, _gl_row(lk.entry_id) or 0, lk.entry_id))
    for link in ordered:
        e = ctx.gl_by_id.get(link.entry_id)
        ticks = ["T" if link.supports_claim else "X"]
        if link.doc_ids:
            ticks.append("D")
        if link.entry_id in flagged:
            ticks.append("F")
        sh.text(row, 1, link.entry_id, wrap=False)
        gl_row = _gl_row(link.entry_id) if e is None else e.source_row
        if gl_row is not None:
            sh.put(row, 2, gl_row, num_fmt="0", halign="center", color=INPUT_BLUE)
        sh.text(row, 3, f"{e.account} {e.account_name}" if e else "")
        d = _to_date(e.date) if e else None
        if d is not None:
            sh.put(row, 4, d, num_fmt=DATE_FORMAT, halign="center", color=INPUT_BLUE)
        else:
            sh.put(row, 4, _month_date(link.period), num_fmt=MONTH_FORMAT, halign="center", color=INPUT_BLUE)
        sh.money(row, 5, link.amount)
        sh.put(row, 6, " ".join(ticks), bold=True, color=LINK_GREEN, halign="center")
        sh.put(row, 7, "Yes" if link.supports_claim else "No", halign="center")
        sh.text(row, 8, e.counterparty if e else "")
        sh.text(row, 9, e.doc_number if e else "")
        sh.text(row, 10, e.memo if e else "")
        sh.text(row, 11, ", ".join(labels_for_month(link.period, periods)) or "Outside analysis periods")
        sh.text(row, 12, "; ".join(link.doc_ids), indent=1)
        reasons = []
        if link.entry_id in flagged:
            reasons.append("Flags: " + ", ".join(dict.fromkeys(flagged[link.entry_id])))
        if link.reasons:
            reasons.append("Link: " + "; ".join(link.reasons))
        if link.group:
            reasons.append(f"Group: {link.group}")
        sh.text(row, 13, " | ".join(reasons))
        row += 1
    last_row = row - 1
    amt = f"E{first}:E{last_row}"
    sh.put(row, 3, "Total linked entries", bold=True, border=_TOP_BORDER)
    sh.formula(row, 5, f"=SUM({amt})", bold=True, border=_TOP_BORDER)
    row += 1
    sh.put(row, 3, "Of which supports the claim (T)", italic=True)
    sh.formula(row, 5, f'=SUMIFS({amt},G{first}:G{last_row},"Yes")', italic=True)
    row += 1
    sh.put(row, 3, "Linked for context only (X)", italic=True)
    sh.formula(row, 5, f"=E{row - 2}-E{row - 1}", italic=True)
    sh.ws.auto_filter.ref = f"A{header_row}:{get_column_letter(13)}{last_row}"
    return row + 2


def _documents_block(sh: _Sheet, ctx: _Ctx, a: AdjustmentAssessment, row: int) -> int:
    sh.section(row, f"DOCUMENTS AND VERBATIM QUOTES ({len(a.doc_links)})")
    row += 1
    if not a.doc_links:
        sh.put(row, 1, "No documents linked.", italic=True, color=MUTED)
        return row + 2
    sh.header(row, [("#", 1), ("Document", 2), ("Type", 1), ("Doc date", 1), ("Signed?", 1), ("Relation", 1),
                    ("Counterparty", 1), ("Reference #s", 1), ("Key terms", 1), ("Service period", 1),
                    ("Linked GL rows", 1), ("Link reasons", 1)])
    row += 1
    for i, link in enumerate(a.doc_links, start=1):
        facts = ctx.doc_facts.get(link.doc_id)
        sh.put(row, 1, f"D{i}", bold=True, halign="center")
        sh.text(row, 2, link.doc_id, span=2)
        sh.text(row, 4, facts.doc_type.replace("_", " ") if facts else "")
        d = _to_date(facts.doc_date) if facts else None
        if d is not None:
            sh.put(row, 5, d, num_fmt=DATE_FORMAT, halign="center", color=INPUT_BLUE)
        signed = _signed_status(facts)
        sh.text(row, 6, signed, bold=signed in ("DRAFT", "Unsigned"),
                color=BANNER_RED if signed in ("DRAFT", "Unsigned") else FORMULA_BLACK)
        sh.text(row, 7, link.relation.replace("_", " "))
        sh.text(row, 8, (facts.counterparty or "") if facts else "")
        sh.text(row, 9, ", ".join(facts.reference_numbers) if facts else "")
        sh.text(row, 10, "; ".join(f"{t.kind}: {t.text}" for t in facts.terms) if facts else "")
        span = ""
        if facts and (facts.service_period_start or facts.service_period_end):
            span = f"{facts.service_period_start or '?'} to {facts.service_period_end or '?'}"
        sh.text(row, 11, span)
        sh.text(row, 12, _gl_rows_text(link.entry_ids))
        sh.text(row, 13, "; ".join(link.reasons))
        row += 1
    quotes: list[EvidenceQuote] = []
    for link in a.doc_links:
        quotes.extend(link.quotes)
        facts = ctx.doc_facts.get(link.doc_id)
        if facts is not None:
            quotes.extend(t.quote for t in facts.terms)
    quotes = _dedupe_quotes(quotes)
    row += 1
    sh.put(row, 1, "Verbatim quotes (exact text of the cited page)", bold=True)
    row += 1
    if quotes:
        sh.header(row, [("Tick", 1), ("Document, page", 2), ("Quote", sh.last_col - 3)])
        row += 1
        row = _quote_rows(sh, row, quotes)
    else:
        sh.put(row, 1, "No quotes recorded for the linked documents.", italic=True, color=MUTED)
        row += 1
    return row + 1


def _decision_block(sh: _Sheet, ctx: _Ctx, a: AdjustmentAssessment, row: int) -> int:
    last = sh.last_col
    pc = 4
    sh.section(row, "REVIEWER DECISION")
    row += 1
    rv = ctx.latest.get(a.adj_id)
    if rv is None:
        sh.status(row, 1, UNREVIEWED, span=2)
        sh.text(row, 3, "No reviewer decision recorded; the tool proposal is carried in the bridge.", span=last - 2,
                bold=True)
        return row + 2
    sh.header(row, [("", 3)] + [(p, 1) for p in ctx.labels])
    row += 1
    sh.put(row, 1, "Reviewer treatment and amounts", span=3, bold=True)
    for i, p in enumerate(ctx.labels):
        if _review_pending(rv):
            sh.pending(row, pc + i)
        else:
            sh.money(row, pc + i, rv.amounts.get(p), bold=True)
    row += 1
    sh.put(row, 1, "Tool proposal at time of review", span=3, italic=True)
    for i, p in enumerate(ctx.labels):
        if rv.tool_amounts:
            sh.money(row, pc + i, rv.tool_amounts.get(p), italic=True)
        else:
            sh.pending(row, pc + i)
    row += 1
    for label, value in (
        ("Treatment", None),
        ("Reviewer", f"{rv.reviewer} at {rv.timestamp}"),
        ("Correction type", rv.correction_type.value
         + (" (tool error: becomes a regression case)" if rv.correction_type in TOOL_ERROR_CORRECTIONS else "")),
        ("Rationale", rv.rationale),
    ):
        sh.put(row, 1, label, span=2, bold=True)
        if value is None:
            sh.treatment(row, 3, rv.treatment)
            sh.text(row, 4, f"Tool treatment at review: {rv.tool_treatment.value}", span=3, italic=True)
        else:
            sh.text(row, 3, value, span=last - 2)
        row += 1
    history = [r for r in ctx.wp.reviews if r.adj_id == a.adj_id]
    if len(history) > 1:
        row += 1
        sh.put(row, 1, "Decision history (review log order; latest wins)", bold=True)
        row += 1
        sh.header(row, [("#", 1), ("Timestamp", 2), ("Reviewer", 1), ("Treatment", 1), ("Current?", 1),
                        ("Correction type", 2), ("Rationale", last - 8)])
        row += 1
        for i, h in enumerate(history, start=1):
            sh.put(row, 1, i, halign="center")
            sh.text(row, 2, h.timestamp, span=2)
            sh.text(row, 4, h.reviewer)
            sh.treatment(row, 5, h.treatment)
            sh.text(row, 6, "Current" if h is rv else "Superseded", italic=h is not rv)
            sh.text(row, 7, h.correction_type.value, span=2)
            sh.text(row, 9, h.rationale, span=last - 8)
            row += 1
    return row + 1


# ---------------------------------------------------------------------------
# Adjustment Summary
# ---------------------------------------------------------------------------


@dataclass
class _SummaryRefs:
    first_row: int  # management items (the Cover's status counts read these ranges)
    last_row: int
    total_row: int  # grand total: management + diligence-identified items
    claimed_col: int
    proposed_col: int
    final_col: int
    diff_col: int
    tool_col: int
    reviewer_col: int
    status_col: int
    rows: dict[str, int] = field(default_factory=dict)
    mgmt_total_row: int = 0  # management items only (== total_row when there are no diligence items)
    dil_first_row: Optional[int] = None
    dil_last_row: Optional[int] = None
    dil_total_row: Optional[int] = None


def _write_summary(ws: Worksheet, ctx: _Ctx, tie: dict[str, _TieOut]) -> _SummaryRefs:
    n = len(ctx.labels)
    claimed_col = 4
    proposed_col = claimed_col + n
    final_col = proposed_col + n
    diff_col = final_col + n
    tool_col = diff_col + n
    reviewer_col, status_col, corr_col, conf_col, flags_col = (tool_col + i for i in range(1, 6))
    links_col, supp_col, docs_col, qs_col = (flags_col + i for i in range(1, 5))
    widths = [9, 42, 20] + [13] * (4 * n) + [15, 15, 14, 24, 11, 52, 9, 11, 8, 11]
    sh = _Sheet(ws, widths)
    _title_block(sh, ctx, "Adjustment Summary", _units(ctx))
    cur = ctx.currency
    g, h = 6, 7
    # Two header rows: group band over the period columns, then period labels.
    for col, title in ((1, "Ref"), (2, "Title"), (3, "Category")):
        sh.header_tall(g, h, col, title)
    for col, title in (
        (claimed_col, f"Claimed by management ({cur})"),
        (proposed_col, f"Tool proposed ({cur})"),
        (final_col, f"Final: diligence ({cur})"),
        (diff_col, f"Final less claimed ({cur})"),
    ):
        sh.header(g, [(title, n)], start_col=col)
        sh.header(h, [(p, 1) for p in ctx.labels], start_col=col)
    for col, title in (
        (tool_col, "Tool treatment"), (reviewer_col, "Reviewer treatment"), (status_col, "Status"),
        (corr_col, "Correction type"), (conf_col, "Tool confidence"), (flags_col, "Key flags (C/W/I = severity)"),
        (links_col, "# GL links"), (supp_col, "# supporting GL links"), (docs_col, "# docs"),
        (qs_col, "# open questions"),
    ):
        sh.header_tall(g, h, col, title)
    ws.row_dimensions[g].height = 30
    refs_rows: dict[str, int] = {}

    def item_row(row: int, a: AdjustmentAssessment) -> None:
        t = tie[a.adj_id]
        rv = ctx.latest.get(a.adj_id)
        refs_rows[a.adj_id] = row
        sh.link(row, 1, a.adj_id, t.sheet)
        sh.text(row, 2, a.title)
        sh.text(row, 3, _CATEGORY_LABELS.get(a.category, a.category.value))
        for i in range(n):
            sc = t.first_col + i
            sh.formula(row, claimed_col + i, "=" + _xref(t.sheet, sc, t.claimed_row), link=True)
            sh.formula(row, proposed_col + i, "=" + _xref(t.sheet, sc, t.proposed_row), link=True)
            sh.formula(row, final_col + i, "=" + _xref(t.sheet, sc, t.final_row), link=True, bold=True)
            fin = _a1(final_col + i, row)
            sh.formula(row, diff_col + i, f"=IF(ISNUMBER({fin}),{fin},0)-{_a1(claimed_col + i, row)}")
        sh.treatment(row, tool_col, a.treatment)
        sh.treatment(row, reviewer_col, rv.treatment if rv else None)
        sh.status(row, status_col, ctx.status(a.adj_id))
        sh.text(row, corr_col, rv.correction_type.value if rv else "", halign="center")
        sh.text(row, conf_col, a.confidence, halign="center")
        sh.text(row, flags_col, _key_flags(a))
        sh.count(row, links_col, len(a.gl_links))
        sh.count(row, supp_col, sum(1 for lk in a.gl_links if lk.supports_claim))
        sh.count(row, docs_col, len(a.doc_links))
        sh.count(row, qs_col, sum(1 for q in ctx.questions.get(a.adj_id, []) if q.status == QuestionStatus.OPEN))

    total_cols = [*range(claimed_col, tool_col), links_col, supp_col, docs_col, qs_col]

    def total_line(row: int, label: str, formula: Any, border: Border) -> None:
        sh.put(row, 2, label, bold=True, border=border)
        sh.put(row, 1, None, border=border)
        sh.put(row, 3, None, border=border)
        for col in total_cols:
            count = col in (links_col, supp_col, docs_col, qs_col)
            sh.formula(row, col, formula(get_column_letter(col)), bold=True, border=border,
                       num_fmt="#,##0" if count else NUMBER_FORMAT)

    row = h + 1
    first = row
    mgmt, dil = ctx.mgmt_items, ctx.dil_items
    for a in mgmt:
        item_row(row, a)
        row += 1
    last_row = max(first, row - 1)
    dil_first = dil_last = dil_total = None
    if not dil:
        total_row = row if mgmt else first
        total_line(total_row, "Total", lambda L: f"=SUM({L}{first}:{L}{last_row})", _FINAL_BORDER)
        mgmt_total = total_row
    else:
        # Management items, their subtotal, then the diligence-identified block and its subtotal.
        mgmt_total = row if mgmt else first
        total_line(mgmt_total, "Total management adjustments", lambda L: f"=SUM({L}{first}:{L}{last_row})",
                   _TOP_BORDER)
        row = mgmt_total + 2
        sh.put(row, 1, DILIGENCE_BLOCK + ": management claimed nothing, so the final amount is the adjustment",
               span=qs_col, bold=True, color=NAVY, fill=SECTION_FILL)
        row += 1
        dil_first = row
        for a in dil:
            item_row(row, a)
            row += 1
        dil_last = row - 1
        dil_total = row
        total_line(dil_total, "Total diligence-identified items",
                   lambda L: f"=SUM({L}{dil_first}:{L}{dil_last})", _TOP_BORDER)
        total_row = dil_total + 2
        total_line(total_row, "Total", lambda L: f"={L}{mgmt_total}+{L}{dil_total}", _FINAL_BORDER)
    note = total_row + 2
    sh.put(note, 2, "Pending (REQUEST_INFO) items show \"Pending\": they are excluded from Final totals, and the "
                    "bridge reverses their claimed amounts.", italic=True, color=MUTED)
    sh.put(note + 1, 2, "Final = latest reviewer decision; an UNREVIEWED item carries the tool proposal until a "
                        "reviewer signs off.", italic=True, color=MUTED)
    ws.freeze_panes = _a1(4, h + 1)
    if mgmt:
        ws.auto_filter.ref = f"A{h}:{get_column_letter(qs_col)}{last_row}"
    return _SummaryRefs(
        first_row=first, last_row=last_row, total_row=total_row, claimed_col=claimed_col,
        proposed_col=proposed_col, final_col=final_col, diff_col=diff_col, tool_col=tool_col,
        reviewer_col=reviewer_col, status_col=status_col, rows=refs_rows, mgmt_total_row=mgmt_total,
        dil_first_row=dil_first, dil_last_row=dil_last, dil_total_row=dil_total,
    )


# ---------------------------------------------------------------------------
# EBITDA Bridge (SPEC §5.6)
# ---------------------------------------------------------------------------


def _is_total(r: BridgeRow) -> bool:
    return r.key.endswith("_total") or r.label.strip().lower().startswith("total")


def _plan_subtotals(rows: Sequence[BridgeRow], labels: Sequence[str]) -> dict[int, tuple[list[tuple[int, int]], bool]]:
    """For each subtotal row index: the (row index, sign) terms of its formula, and whether they foot.

    A "total" subtotal sums the adjustment rows since the previous subtotal. An
    EBITDA-level subtotal starts from the previous level and adds the
    intervening total (or rows). Memo rows between levels (the difference to
    management's reported EBITDA) are reconciling items: their sign is chosen
    so the formula reproduces the tool's figure, and an inconsistent bridge is
    reported rather than papered over.
    """
    plans: dict[int, tuple[list[tuple[int, int]], bool]] = {}
    prev_sub: Optional[int] = None
    prev_level: Optional[int] = None
    block: list[int] = []
    for i, r in enumerate(rows):
        if r.kind != "subtotal":
            block.append(i)
            continue
        plain = [(j, 1) for j in block if rows[j].kind != "memo"]
        memo = [j for j in block if rows[j].kind == "memo"]
        candidates: list[list[tuple[int, int]]] = []
        if _is_total(r):
            candidates.append(plain)
        else:
            base: list[tuple[int, int]] = []
            if prev_level is not None:
                base.append((prev_level, 1))
            if prev_sub is not None and prev_sub != prev_level:
                base.append((prev_sub, 1))
            for sign in (1, -1):
                candidates.append(base + plain + [(j, sign) for j in memo])
            if memo:
                candidates.append(base + plain)
        chosen, ok = candidates[0], False
        for terms in candidates:
            if all(
                abs(sum((D(rows[j].amounts.get(p)) * s for j, s in terms), Decimal(0)) - D(r.amounts.get(p)))
                <= Decimal("0.01")
                for p in labels
            ):
                chosen, ok = terms, True
                break
        plans[i] = (chosen, ok)
        prev_sub = i
        if not _is_total(r):
            prev_level = i
        block = []
    return plans


def _terms_formula(terms: list[tuple[int, int]], sheet_rows: dict[int, int], col: int) -> str:
    if not terms:
        return "=0"
    rows = sorted(sheet_rows[j] for j, _ in terms)
    signs = {s for _, s in terms}
    L = get_column_letter(col)
    # A run of added rows reads best as SUM(), but only when no other bridge row
    # (e.g. a memo line) sits inside the range; section and blank rows are harmless.
    others = set(sheet_rows.values()) - set(rows)
    if len(terms) > 2 and signs == {1} and not any(rows[0] < r < rows[-1] for r in others):
        return f"=SUM({L}{rows[0]}:{L}{rows[-1]})"
    parts = []
    for j, s in terms:
        ref = f"{L}{sheet_rows[j]}"
        parts.append(("+" if s > 0 else "-") + ref)
    expr = "".join(parts)
    return "=" + (expr[1:] if expr.startswith("+") else expr)


@dataclass
class _BridgeRefs:
    rows: dict[str, int]
    first_col: int
    footed: bool
    checks: list[str] = field(default_factory=list)  # every "should be zero" range, bridge and support sheets


def _write_bridge(ws: Worksheet, ctx: _Ctx, srefs: _SummaryRefs, tie: dict[str, _TieOut]) -> _BridgeRefs:
    n = len(ctx.labels)
    pc = 3
    treat_col, status_col, link_col = pc + n, pc + n + 1, pc + n + 2
    sh = _Sheet(ws, [10, 62] + [15] * n + [16, 14, 13])
    ws.sheet_view.showGridLines = False
    _title_block(sh, ctx, "EBITDA Bridge: reported to diligence adjusted", _units(ctx))
    row = 6
    sh.header(row, [("Ref", 1), ("Line item", 1)] + [(f"{p}\n{ctx.currency}", 1) for p in ctx.labels]
              + [("Final treatment", 1), ("Status", 1), ("Support", 1)])
    ws.row_dimensions[row].height = 28
    header_row = row
    row += 1
    item_ids = {a.adj_id for a in ctx.dil_items}
    rows = bridge_display_rows(ctx.wp.bridge.rows, item_ids)
    if not rows:
        sh.put(row, 2, "No bridge rows in the workpaper.", italic=True, color=MUTED)
        return _BridgeRefs(rows={}, first_col=pc, footed=True)

    plans = _plan_subtotals(rows, ctx.labels)
    sheet_rows: dict[int, int] = {}
    by_id = {a.adj_id: a for a in ctx.wp.assessments}
    seen_kinds: set[str] = set()
    sections = {
        "component": "Reported EBITDA",
        "mgmt_adjustment": "Management adjustments (as claimed)",
        "diligence_adjustment": "Diligence adjustments (final less claimed; claimed reversed when pending)",
    }
    prev: Optional[BridgeRow] = None
    item_block = False
    for i, r in enumerate(rows):
        if prev is not None and prev.kind == "subtotal" and not _is_total(prev) and r.kind != "subtotal":
            row += 1
        if r.kind in sections and r.kind not in seen_kinds:
            sh.put(row, 1, sections[r.kind], span=link_col, bold=True, color=NAVY, fill=SECTION_FILL)
            row += 1
        seen_kinds.add(r.kind)
        if is_item_row(r, item_ids) and not item_block:
            # A text-only row inside the diligence block; SUM ranges over it are unaffected.
            sh.put(row, 1, DILIGENCE_BLOCK + ": the final amount (management claimed nothing)", span=link_col,
                   bold=True, italic=True, color=NAVY)
            row += 1
            item_block = True
        sheet_rows[i] = row
        is_sub = r.kind == "subtotal"
        level = is_sub and not _is_total(r)
        final_line = r.key == "diligence_adjusted_ebitda"
        memo = r.kind == "memo"
        border = _FINAL_BORDER if final_line else (_TOP_BORDER if is_sub else None)
        fill = FINAL_FILL if final_line else (SUBTOTAL_FILL if level else None)
        sh.put(row, 1, r.adj_id or "", fill=fill, border=border, bold=is_sub)
        sh.text(row, 2, r.label, bold=is_sub, italic=memo, fill=fill, border=border,
                color=MUTED if memo else FORMULA_BLACK,
                indent=0 if (is_sub or r.kind == "component" or memo) else 1)
        for k, p in enumerate(ctx.labels):
            col = pc + k
            if is_sub:
                terms, ok = plans[i]
                cell = sh.formula(row, col, _terms_formula(terms, sheet_rows, col), bold=True, fill=fill, border=border)
                if not ok:
                    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=BANNER_RED)
            else:
                sh.money(row, col, r.amounts.get(p), italic=memo, fill=fill, border=border)
        for col in (treat_col, status_col, link_col):
            sh.put(row, col, None, fill=fill, border=border)
        a = by_id.get(bridge_row_adj_id(r) or "")
        if a is not None and r.kind == "diligence_adjustment":
            sh.treatment(row, treat_col, ctx.final_treatment(a), border=border)
            sh.status(row, status_col, ctx.status(a.adj_id), border=border)
        if a is not None and r.kind in ("mgmt_adjustment", "diligence_adjustment"):
            sh.link(row, link_col, ctx.adj_sheets[a.adj_id], ctx.adj_sheets[a.adj_id]).alignment = _CENTER
        elif r.key in ("dil_recon", "mgmt_recon_diff"):
            sh.link(row, link_col, "Data Quality", SHEET_DATA_QUALITY).alignment = _CENTER
        row += 1
        prev = r

    key_rows = {rows[i].key: sheet_rows[i] for i in sheet_rows}
    footed = all(ok for _, ok in plans.values())

    # Controls: the bridge must tie to the Adjustment Summary and to its own identity.
    row += 1
    sh.put(row, 1, "Checks (should be zero)", span=link_col, bold=True, color=NAVY, fill=SECTION_FILL)
    row += 1
    checks: list[tuple[str, Any]] = []
    if ctx.wp.assessments and {"diligence_adjusted_ebitda", "gl_ebitda"} <= key_rows.keys():
        checks.append((
            "Diligence adjusted EBITDA less (GL EBITDA + total final amounts, Adjustment Summary)",
            lambda col, k: (f"={_a1(col, key_rows['diligence_adjusted_ebitda'])}-{_a1(col, key_rows['gl_ebitda'])}"
                            f"-{_xref(SHEET_SUMMARY, srefs.final_col + k, srefs.total_row)}"),
        ))
    if ctx.mgmt_items and "mgmt_total" in key_rows:
        checks.append((
            "Total management adjustments less total claimed, Adjustment Summary",
            lambda col, k: (f"={_a1(col, key_rows['mgmt_total'])}"
                            f"-{_xref(SHEET_SUMMARY, srefs.claimed_col + k, srefs.mgmt_total_row)}"),
        ))
    item_rows = [key_rows[f"dil:{a.adj_id}"] for a in ctx.dil_items if f"dil:{a.adj_id}" in key_rows]
    if srefs.dil_total_row is not None and len(item_rows) == len(ctx.dil_items):
        checks.append((
            "Diligence-identified items less their total final amounts, Adjustment Summary",
            lambda col, k: ("=" + "+".join(_a1(col, r) for r in item_rows)
                            + f"-{_xref(SHEET_SUMMARY, srefs.final_col + k, srefs.dil_total_row)}"),
        ))
    check_ranges: list[str] = []
    for label, build in checks:
        check_ranges.append(
            f"{_quoted(SHEET_BRIDGE)}!${get_column_letter(pc)}${row}:${get_column_letter(pc + n - 1)}${row}"
        )
        sh.text(row, 2, label, italic=True, color=MUTED)
        for k in range(n):
            col = pc + k
            sh.formula(row, col, build(col, k), num_fmt=CHECK_FORMAT, italic=True)
        rng = f"{_a1(pc, row)}:{_a1(pc + n - 1, row)}"
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f"ABS({_a1(pc, row)})>=0.01"], fill=_fill(CHECK_FAIL_FILL))
        )
        row += 1
    if not footed:
        sh.put(row, 2, "WARNING: a subtotal in red does not foot from the rows above to the tool's bridge; "
                       "investigate before relying on it.", bold=True, color=BANNER_RED)
        row += 1
    row += 1
    sh.put(row, 2, "Subtotals are formulas over the rows above; adjustment rows are values from the workpaper. "
                   "Diligence adjusted EBITDA = GL EBITDA + final amounts (pending items excluded).",
           italic=True, color=MUTED)

    # Each support sheet ties its (e) - (a) revision to the bridge.
    for a in ctx.wp.assessments:
        t = tie[a.adj_id]
        dil_row = key_rows.get(f"dil:{a.adj_id}")
        support = ctx.adj_sheets[a.adj_id]
        sup = _Sheet(ws.parent[support], [])
        for k in range(n):
            col = t.first_col + k
            if dil_row is None:
                sup.put(t.check_row, col, "n/a", italic=True, color=MUTED, halign="right")
                continue
            sup.formula(t.check_row, col, f"={_xref(SHEET_BRIDGE, pc + k, dil_row)}-{_a1(col, t.revision_row)}",
                        num_fmt=CHECK_FORMAT, italic=True)
        if dil_row is not None and n:
            check_ranges.append(f"{_quoted(support)}!${get_column_letter(t.first_col)}${t.check_row}:"
                                f"${get_column_letter(t.first_col + n - 1)}${t.check_row}")
            rng = f"{_a1(t.first_col, t.check_row)}:{_a1(t.first_col + n - 1, t.check_row)}"
            ws.parent[support].conditional_formatting.add(
                rng, FormulaRule(formula=[f"ABS({_a1(t.first_col, t.check_row)})>=0.01"], fill=_fill(CHECK_FAIL_FILL))
            )

    ws.freeze_panes = _a1(pc, header_row + 1)
    return _BridgeRefs(rows=key_rows, first_col=pc, footed=footed, checks=check_ranges)


# ---------------------------------------------------------------------------
# Open Questions, Reconciliation, Data Quality, Review Log
# ---------------------------------------------------------------------------


def _write_questions(ws: Worksheet, ctx: _Ctx) -> tuple[int, int]:
    """Returns the (first, last) data rows, for the Cover's live counts."""
    sh = _Sheet(ws, [13, 9, 72, 10, 34, 12, 50])
    _title_block(sh, ctx, "Open Questions for Management", "Questions come from flags and document gaps; "
                 "status and responses follow the review log.")
    row = 6
    sh.header(row, [("Q id", 1), ("Ref", 1), ("Question", 1), ("Priority", 1), ("Basis", 1), ("Status", 1),
                    ("Response", 1)])
    header = row
    row += 1
    first = row
    rank = {"high": 0, "medium": 1, "low": 2}
    for a in ctx.wp.assessments:
        for q in sorted(ctx.questions.get(a.adj_id, []), key=lambda q: rank.get(q.priority, 3)):
            sh.text(row, 1, q.q_id, wrap=False)
            sh.link(row, 2, a.adj_id, ctx.adj_sheets[a.adj_id])
            sh.text(row, 3, q.text)
            high = q.priority == "high"
            sh.text(row, 4, q.priority, bold=high, color=BANNER_RED if high else FORMULA_BLACK, halign="center")
            sh.text(row, 5, q.basis)
            is_open = q.status == QuestionStatus.OPEN
            sh.text(row, 6, q.status.value, bold=is_open, halign="center",
                    fill=TREATMENT_STYLES[Treatment.REVISE][0] if is_open else None)
            sh.text(row, 7, q.response)
            row += 1
    last = max(first, row - 1)
    if row == first:
        sh.put(row, 3, "No open questions.", italic=True, color=MUTED)
    else:
        ws.auto_filter.ref = f"A{header}:G{last}"
    ws.freeze_panes = _a1(3, header + 1)
    return first, last


_RECON_COLS = [("Month", 1), ("GL", 1), ("P&L", 1), ("Variance", 1), ("Within tolerance?", 1), ("Account", 1),
               ("Account name", 1)]


def _write_recon(ws: Worksheet, ctx: _Ctx) -> None:
    recon = ctx.wp.reconciliation
    cur = ctx.currency
    sh = _Sheet(ws, [13, 17, 17, 17, 17, 12, 44])
    _title_block(sh, ctx, "GL to P&L Reconciliation", f"Management's monthly P&L compared with the GL by account and "
                 f"month. {cur}, debit-positive (expenses +, revenue -). Variance = P&L less GL.")
    row = 6
    for label, value in (
        ("Months compared", recon.months_compared),
        ("Accounts compared", recon.accounts_compared),
        ("Account-months outside tolerance", recon.variance_count),
    ):
        sh.put(row, 1, label, span=2, bold=True)
        sh.count(row, 3, value)
        row += 1
    sh.put(row, 1, f"Tolerance ({cur})", span=2, bold=True)
    sh.put(row, 3, q2(D(ctx.wp.deal.tolerance)), num_fmt="#,##0.00", color=INPUT_BLUE, halign="right")
    row += 2

    items = sorted(recon.items, key=lambda it: (it.month, it.account))
    months = sorted({it.month for it in items})
    variances = [it for it in items if not it.within_tolerance]

    sh.section(row, "SUMMARY BY MONTH")
    row += 1
    sh.header(row, [("Month", 1), (f"GL total ({cur})", 1), (f"P&L total ({cur})", 1),
                    (f"Variance: P&L less GL ({cur})", 1), ("# account variances", 1)])
    row += 1
    sum_first = row
    month_rows: list[int] = []
    for m in months:
        sh.put(row, 1, _month_date(m), num_fmt=MONTH_FORMAT, halign="center", color=INPUT_BLUE)
        month_rows.append(row)  # SUMIFS formulas are filled once the detail table's rows are known
        row += 1
    if not months:
        sh.put(row, 1, "No account-months compared.", italic=True, color=MUTED)
        row += 1
    sum_last = row - 1
    sh.put(row, 1, "Total", bold=True, border=_FINAL_BORDER)
    for col in range(2, 6):
        L = get_column_letter(col)
        formula = f"=SUM({L}{sum_first}:{L}{sum_last})" if months else "=0"
        sh.formula(row, col, formula, bold=True, border=_FINAL_BORDER, num_fmt="#,##0" if col == 5 else NUMBER_FORMAT)
    row += 2

    sh.section(row, f"VARIANCE DETAIL: OUTSIDE TOLERANCE ({len(variances)})")
    row += 1
    sh.header(row, [("Month", 1), (f"GL ({cur})", 1), (f"P&L ({cur})", 1), (f"Variance ({cur})", 1),
                    ("Within tolerance?", 1), ("Account", 1), ("Account name", 1)])
    row += 1
    if variances:
        for it in variances:
            _recon_row(sh, row, it)
            row += 1
    else:
        sh.put(row, 1, "None.", italic=True, color=MUTED)
        row += 1
    row += 1

    sh.section(row, f"ALL ACCOUNT-MONTHS COMPARED ({len(items)})")
    row += 1
    detail_header = row
    sh.header(row, [("Month", 1), (f"GL ({cur})", 1), (f"P&L ({cur})", 1), (f"Variance ({cur})", 1),
                    ("Within tolerance?", 1), ("Account", 1), ("Account name", 1)])
    row += 1
    detail_first = row
    for it in items:
        _recon_row(sh, row, it)
        row += 1
    detail_last = max(detail_first, row - 1)

    def rng(col: str) -> str:
        return f"${col}${detail_first}:${col}${detail_last}"

    for r in month_rows:
        sh.formula(r, 2, f"=SUMIFS({rng('B')},{rng('A')},$A{r})")
        sh.formula(r, 3, f"=SUMIFS({rng('C')},{rng('A')},$A{r})")
        sh.formula(r, 4, f"=C{r}-B{r}")
        sh.formula(r, 5, f'=COUNTIFS({rng("A")},$A{r},{rng("E")},"No")', num_fmt="#,##0")
    if items:
        ws.auto_filter.ref = f"A{detail_header}:G{detail_last}"
    ws.freeze_panes = "A6"


def _recon_row(sh: _Sheet, row: int, it: Any) -> None:
    sh.put(row, 1, _month_date(it.month), num_fmt=MONTH_FORMAT, halign="center", color=INPUT_BLUE)
    sh.money(row, 2, it.gl_amount)
    sh.money(row, 3, it.pl_amount)
    sh.formula(row, 4, f"=C{row}-B{row}")
    ok = it.within_tolerance
    sh.put(row, 5, "Yes" if ok else "No", halign="center", bold=not ok, color=BANNER_RED if not ok else FORMULA_BLACK)
    sh.text(row, 6, it.account, wrap=False)
    sh.text(row, 7, it.account_name)


def _write_data_quality(ws: Worksheet, ctx: _Ctx) -> tuple[int, int]:
    """Returns the (first, last) issue rows, for the Cover's live counts."""
    cur = ctx.currency
    sh = _Sheet(ws, [11, 34, 10, 10, 14, 15, 80, 30])
    _title_block(sh, ctx, "Data Quality", "Every reconciliation and ingest issue, most severe first.")
    row = 6
    issues = sorted(enumerate(ctx.wp.reconciliation.issues), key=lambda t: (_SEVERITY_RANK[t[1].severity], t[0]))
    sh.header(row, [("Severity", 1), ("Code", 1), ("Month", 1), ("Account", 1), ("Period", 1),
                    (f"Amount ({cur})", 1), ("Message", 1), ("GL rows", 1)])
    header = row
    row += 1
    first = row
    for _, issue in issues:
        sh.severity(row, 1, issue.severity)
        sh.text(row, 2, issue.code.value, wrap=False)
        md = _month_date(issue.month) if issue.month else None
        if md is not None:
            sh.put(row, 3, md, num_fmt=MONTH_FORMAT, halign="center", color=INPUT_BLUE)
        sh.text(row, 4, issue.account or "", halign="center")
        sh.text(row, 5, issue.period_label or "")
        if issue.amount is not None:
            sh.money(row, 6, issue.amount)
        sh.text(row, 7, issue.message, indent=1)
        sh.text(row, 8, _gl_rows_text(issue.entry_ids))
        row += 1
    issues_last = max(first, row - 1)
    if not issues:
        sh.put(row, 2, "No data quality issues.", italic=True, color=MUTED)
        row += 1
    else:
        ws.auto_filter.ref = f"A{header}:H{row - 1}"
    row += 1

    dropped = [f for f in ctx.wp.doc_facts if f.dropped_quotes]
    sh.section(row, "AI QUOTE VERIFICATION")
    row += 1
    sh.put(row, 1, f"{len(ctx.wp.doc_facts)} documents analysed ({ctx.wp.ai_mode}). Quotes that were not verbatim on "
                   "the cited page were dropped and counted, never repaired.", italic=True, color=MUTED)
    row += 1
    if dropped:
        sh.header(row, [("Document", 2), ("Extractor", 1), ("Dropped", 1)], start_col=1)
        row += 1
        f0 = row
        for f in dropped:
            sh.text(row, 1, f.doc_id, span=2)
            sh.text(row, 3, f.extractor)
            sh.count(row, 4, f.dropped_quotes)
            row += 1
        sh.put(row, 1, "Total dropped", span=3, bold=True, border=_TOP_BORDER)
        sh.formula(row, 4, f"=SUM(D{f0}:D{row - 1})", bold=True, num_fmt="#,##0", border=_TOP_BORDER)
        row += 1
    else:
        sh.put(row, 1, "No quotes dropped.", italic=True, color=MUTED)
        row += 1
    row += 1

    sh.section(row, f"INGEST NOTES ({len(ctx.wp.ingest_notes)})")
    row += 1
    for note in ctx.wp.ingest_notes:
        sh.text(row, 1, note, span=8)
        row += 1
    if not ctx.wp.ingest_notes:
        sh.put(row, 1, "None.", italic=True, color=MUTED)
    ws.freeze_panes = _a1(1, header + 1)
    return first, issues_last


def _write_review_log(ws: Worksheet, ctx: _Ctx) -> None:
    n = len(ctx.labels)
    cur = ctx.currency
    tool_amt, final_amt = 9, 9 + n
    rat_col = final_amt + n
    widths = [5, 22, 16, 9, 15, 15, 24, 10] + [13] * (2 * n) + [60, 40, 12, 14]
    sh = _Sheet(ws, widths)
    _title_block(sh, ctx, "Review Log", "Every reviewer decision in log order (append-only). The latest decision "
                 "per adjustment is Current.")
    g, h = 6, 7
    for col, title in enumerate(("#", "Timestamp", "Reviewer", "Ref", "Tool treatment", "Reviewer treatment",
                                 "Correction type", "Tool error?"), start=1):
        sh.header_tall(g, h, col, title)
    sh.header(g, [(f"Tool amounts at review ({cur})", n)], start_col=tool_amt)
    sh.header(h, [(p, 1) for p in ctx.labels], start_col=tool_amt)
    sh.header(g, [(f"Reviewer amounts ({cur})", n)], start_col=final_amt)
    sh.header(h, [(p, 1) for p in ctx.labels], start_col=final_amt)
    for col, title in ((rat_col, "Rationale"), (rat_col + 1, "Question updates"), (rat_col + 2, "Current?"),
                       (rat_col + 3, "Tool changed since?")):
        sh.header_tall(g, h, col, title)
    ws.row_dimensions[g].height = 30
    row = h + 1
    for i, rv in enumerate(ctx.wp.reviews, start=1):
        current = ctx.latest.get(rv.adj_id) is rv
        sh.put(row, 1, i, halign="center")
        sh.text(row, 2, rv.timestamp, wrap=False)
        sh.text(row, 3, rv.reviewer)
        if rv.adj_id in ctx.adj_sheets:
            sh.link(row, 4, rv.adj_id, ctx.adj_sheets[rv.adj_id])
        else:
            sh.text(row, 4, rv.adj_id)
        sh.treatment(row, 5, rv.tool_treatment)
        sh.treatment(row, 6, rv.treatment)
        sh.text(row, 7, rv.correction_type.value)
        tool_error = rv.correction_type in TOOL_ERROR_CORRECTIONS
        sh.text(row, 8, "Yes" if tool_error else "No", bold=tool_error, halign="center",
                color=BANNER_RED if tool_error else FORMULA_BLACK)
        for k, p in enumerate(ctx.labels):
            if rv.tool_amounts:
                sh.money(row, tool_amt + k, rv.tool_amounts.get(p))
            else:
                sh.pending(row, tool_amt + k)
            if _review_pending(rv):
                sh.pending(row, final_amt + k)
            else:
                sh.money(row, final_amt + k, rv.amounts.get(p), bold=True)
        sh.text(row, rat_col, rv.rationale, indent=1)
        sh.text(row, rat_col + 1, "; ".join(f"{q}: {v}" for q, v in rv.question_updates.items()))
        sh.text(row, rat_col + 2, "Current" if current else "Superseded", italic=not current, halign="center")
        changed = ctx.stale(rv)
        sh.text(row, rat_col + 3, "Yes: re-review" if changed else "No", bold=changed, halign="center",
                color=BANNER_RED if changed else FORMULA_BLACK)
        row += 1
    if not ctx.wp.reviews:
        sh.put(row, 2, "No reviewer decisions recorded: every adjustment is UNREVIEWED.", bold=True,
               color=BANNER_RED)
    else:
        ws.auto_filter.ref = f"A{h}:{get_column_letter(rat_col + 3)}{row - 1}"
    ws.freeze_panes = _a1(5, h + 1)


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------


_CONTENTS = [
    (SHEET_BRIDGE, "Reported EBITDA (GL) to management adjusted to diligence adjusted EBITDA, with checks."),
    (SHEET_SUMMARY, "One row per adjustment: claimed, tool proposed, final, difference, treatment, and status."),
    ("Adj <ref>", "One support sheet per adjustment: tie-out, flags, facts vs judgment, GL entries, quotes, decision."),
    (SHEET_QUESTIONS, "Questions for management, with priority, basis, status, and responses."),
    (SHEET_RECON, "GL vs management P&L by month and account, with variance detail."),
    (SHEET_DATA_QUALITY, "Reconciliation and ingest issues; AI quote verification."),
    (SHEET_REVIEW_LOG, "Every reviewer decision, in order, with correction types."),
]


def _write_cover(
    ws: Worksheet,
    ctx: _Ctx,
    srefs: _SummaryRefs,
    brefs: _BridgeRefs,
    q_rows: tuple[int, int],
    dq_rows: tuple[int, int],
) -> None:
    wp = ctx.wp
    deal = wp.deal
    n = len(ctx.labels)
    sh = _Sheet(ws, [40] + [18] * max(n, 4) + [40])
    last = sh.last_col
    ws.sheet_view.showGridLines = False
    sh.put(1, 1, "Quality of Earnings: Evidence Review Workpaper", bold=True, size=16, color=NAVY)
    sh.put(2, 1, deal.target_name, bold=True, size=13)
    sh.put(3, 1, _banner_text(ctx), span=last, bold=True, size=12, color=WHITE, fill=BANNER_RED,
           halign="center", valign="center")
    ws.row_dimensions[3].height = 24
    row = 5

    sh.section(row, "ENGAGEMENT AND RUN")
    row += 1
    mgmt_ids = {a.adj_id for a in ctx.mgmt_items}
    dil_ids = {a.adj_id for a in ctx.dil_items}
    reviewed = f"{len(mgmt_ids & ctx.latest.keys())} of {len(mgmt_ids)} management adjustments reviewed"
    if dil_ids:
        reviewed += f"; {len(dil_ids & ctx.latest.keys())} of {len(dil_ids)} diligence-identified items reviewed"
    info: list[tuple[str, Any]] = [
        ("Target", deal.target_name),
        ("Industry", deal.industry or ""),
        ("Deal ID", deal.deal_id),
        ("Currency", f"{ctx.currency}. Whole dollars displayed; cell values keep cents."),
        ("GL data range", f"{deal.data_start} to {deal.data_end}"),
        ("Tie-out tolerance", f"{fmt(deal.tolerance)} {ctx.currency}"),
        ("Run ID", wp.run_id),
        ("Run created", wp.created_at),
        ("Tool version", wp.tool_version),
        ("AI mode", f"{wp.ai_mode} (AI proposes facts, links and question wording; code computes every amount "
                    "and treatment; the reviewer decides)"),
        ("Documents analysed", f"{len(wp.doc_facts)}; quotes dropped as not verbatim: "
                               f"{sum(f.dropped_quotes for f in wp.doc_facts)}"),
        ("Adjustments", f"{len(mgmt_ids)} on management's schedule; {len(dil_ids)} identified by diligence "
                        "(not on the schedule)"),
        ("Reviewer decisions", f"{len(wp.reviews)} logged; {reviewed}"),
    ]
    for label, value in info:
        sh.put(row, 1, label, bold=True)
        sh.text(row, 2, value, span=last - 1)
        row += 1
    row += 1

    sh.section(row, "ANALYSIS PERIODS")
    row += 1
    sh.header(row, [("Period", 1), ("First month", 1), ("Last month", 1), ("Months", 1)])
    row += 1
    for p in deal.periods:
        sh.put(row, 1, p.label, bold=True)
        sh.put(row, 2, _month_date(p.start), num_fmt=MONTH_FORMAT, halign="center", color=INPUT_BLUE)
        sh.put(row, 3, _month_date(p.end), num_fmt=MONTH_FORMAT, halign="center", color=INPUT_BLUE)
        sh.count(row, 4, len(months_in(p)))
        row += 1
    if len({m for p in deal.periods for m in months_in(p)}) < sum(len(months_in(p)) for p in deal.periods):
        sh.put(row, 1, "Periods overlap (e.g. a TTM period shares months with a fiscal year): amounts are "
                       "computed by month and aggregated to each period, so period columns do not add across.",
               span=last, italic=True, color=MUTED, wrap=True)
        row += 1
    row += 1

    headline = [
        ("gl_ebitda", "Reported EBITDA (per GL)"),
        ("mgmt_reported_ebitda", "Reported EBITDA (per management)"),
        ("mgmt_adjusted_ebitda", "Management adjusted EBITDA"),
        ("diligence_adjusted_ebitda", "Diligence adjusted EBITDA"),
    ]
    present = [(k, lab) for k, lab in headline if k in brefs.rows]
    if present:
        sh.section(row, f"HEADLINE EBITDA ({ctx.currency})")
        row += 1
        sh.header(row, [("", 1)] + [(p, 1) for p in ctx.labels])
        row += 1
        hl_rows: dict[str, int] = {}
        for key, label in present:
            final_line = key == "diligence_adjusted_ebitda"
            sh.put(row, 1, label, bold=final_line, border=_FINAL_BORDER if final_line else None)
            for k in range(n):
                sh.formula(row, 2 + k, "=" + _xref(SHEET_BRIDGE, brefs.first_col + k, brefs.rows[key]), link=True,
                           bold=final_line, border=_FINAL_BORDER if final_line else None)
            hl_rows[key] = row
            row += 1
        if {"diligence_adjusted_ebitda", "mgmt_adjusted_ebitda"} <= hl_rows.keys():
            sh.put(row, 1, "Diligence less management adjusted", italic=True)
            for k in range(n):
                col = 2 + k
                sh.formula(row, col, f"={_a1(col, hl_rows['diligence_adjusted_ebitda'])}"
                                     f"-{_a1(col, hl_rows['mgmt_adjusted_ebitda'])}", italic=True)
            row += 1
        if not brefs.footed:
            sh.put(row, 1, "WARNING: the bridge does not foot to the tool's subtotals; see EBITDA Bridge.",
                   span=last, bold=True, color=BANNER_RED)
            row += 1
        row += 1
    checks = "+".join(f"SUMPRODUCT(ABS({r}))" for r in brefs.checks)
    if checks and len(checks) < 8000:  # Excel's formula length limit is 8,192 characters
        sh.put(row, 1, "Workbook checks (bridge and support ties)", bold=True)
        cell = sh.formula(row, 2, f'=IF({checks}<0.01,"OK","DIFFERENCE: see checks")', num_fmt="General",
                          bold=True, halign="left")
        ws.conditional_formatting.add(
            cell.coordinate, FormulaRule(formula=[f'{cell.coordinate}<>"OK"'], fill=_fill(CHECK_FAIL_FILL))
        )
        row += 2

    def rng(sheet: str, col: int, first: int, last_row: int) -> str:
        L = get_column_letter(col)
        return f"{_quoted(sheet)}!${L}${first}:${L}${last_row}"

    def summary_ranges(first: int, last_row: int) -> tuple[str, str, str]:
        return tuple(rng(SHEET_SUMMARY, c, first, last_row)  # type: ignore[return-value]
                     for c in (srefs.tool_col, srefs.reviewer_col, srefs.status_col))

    def status_table(row: int, title: str, ranges: tuple[str, str, str]) -> int:
        tool_r, rev_r, stat_r = ranges
        sh.section(row, title)
        row += 1
        sh.header(row, [("Treatment", 1), ("Tool proposal", 1), ("Reviewer decision", 1), ("Carried (final)", 1)])
        row += 1
        count_first = row
        for t in Treatment:
            sh.treatment(row, 1, t)
            crit = f'"{t.value}"'
            sh.formula(row, 2, f"=COUNTIF({tool_r},{crit})", num_fmt="#,##0")
            sh.formula(row, 3, f"=COUNTIF({rev_r},{crit})", num_fmt="#,##0")
            # Reviewed items carry the reviewer's treatment; unreviewed ones carry the tool's.
            sh.formula(row, 4, f'=COUNTIF({rev_r},{crit})+COUNTIFS({stat_r},"{UNREVIEWED}",{tool_r},{crit})',
                       num_fmt="#,##0")
            row += 1
        sh.put(row, 1, "Total", bold=True, border=_FINAL_BORDER)
        for col in (2, 3, 4):
            L = get_column_letter(col)
            sh.formula(row, col, f"=SUM({L}{count_first}:{L}{row - 1})", bold=True, num_fmt="#,##0",
                       border=_FINAL_BORDER)
        return row + 2

    mgmt_ranges = summary_ranges(srefs.first_row, srefs.last_row)
    dil_ranges = (summary_ranges(srefs.dil_first_row, srefs.dil_last_row)
                  if dil_ids and srefs.dil_first_row is not None and srefs.dil_last_row is not None else None)
    row = status_table(row, "ADJUSTMENT STATUS: MANAGEMENT ADJUSTMENTS" if dil_ranges else "ADJUSTMENT STATUS",
                       mgmt_ranges)
    if dil_ranges:
        row = status_table(row, DILIGENCE_BLOCK.upper(), dil_ranges)
        sh.header(row, [("Review status", 1), ("Management", 1), ("Diligence-identified", 1)])
        row += 1
    q_range = rng(SHEET_QUESTIONS, 6, *q_rows)
    dq_range = rng(SHEET_DATA_QUALITY, 1, *dq_rows)
    for label, status, style in (
        ("Reviewed: agreed with tool", AGREED, None),
        ("Reviewed: overridden", OVERRIDDEN, None),
        ("UNREVIEWED (tool proposal carried)", UNREVIEWED, UNREVIEWED_STYLE),
    ):
        if style:
            sh.put(row, 1, label, bold=True, fill=style[0], color=style[1])
        else:
            sh.put(row, 1, label, bold=True)
        sh.formula(row, 2, f'=COUNTIF({mgmt_ranges[2]},"{status}")', num_fmt="#,##0", bold=bool(style))
        if dil_ranges:
            sh.formula(row, 3, f'=COUNTIF({dil_ranges[2]},"{status}")', num_fmt="#,##0", bold=bool(style))
        row += 1
    for label, formula in (
        ("Open questions for management", f'=COUNTIF({q_range},"{QuestionStatus.OPEN.value}")'),
        ("Data quality issues: critical", f'=COUNTIF({dq_range},"{Severity.CRITICAL.value}")'),
        ("Data quality issues: warning", f'=COUNTIF({dq_range},"{Severity.WARNING.value}")'),
        ("Data quality issues: info", f'=COUNTIF({dq_range},"{Severity.INFO.value}")'),
    ):
        sh.put(row, 1, label, bold=True)
        sh.formula(row, 2, formula, num_fmt="#,##0")
        row += 1
    row += 1

    sh.section(row, "LEGEND")
    row += 1
    for t in Treatment:
        sh.treatment(row, 1, t)
        sh.text(row, 2, TREATMENT_MEANING[t], span=last - 1, indent=1)
        row += 1
    for status, meaning in (
        (UNREVIEWED, "No reviewer decision yet: the tool proposal is carried until a reviewer signs off."),
        (AGREED, "Reviewer agreed with the tool's treatment and amounts (within tolerance)."),
        (OVERRIDDEN, "Reviewer changed the tool's treatment or amounts; see the correction type."),
    ):
        sh.status(row, 1, status)
        sh.text(row, 2, meaning, span=last - 1, indent=1)
        row += 1
    for text, color, meaning in (
        ("1,234", INPUT_BLUE, "Blue: hard-coded input carried from the workpaper (GL, documents, tool, reviewer)."),
        ("1,234", FORMULA_BLACK, "Black: formula (subtotals, differences, checks)."),
        ("1,234", LINK_GREEN, "Green: link to another sheet."),
        ("(1,234)", FORMULA_BLACK, "Parentheses: negative (reduces EBITDA). A dash is zero."),
    ):
        sh.put(row, 1, text, color=color, halign="right")
        sh.text(row, 2, meaning, span=last - 1, indent=1)
        row += 1
    for mark, meaning in TICKMARKS:
        sh.put(row, 1, mark, bold=True, color=LINK_GREEN, halign="right")
        sh.text(row, 2, f"Tickmark: {meaning}", span=last - 1, indent=1)
        row += 1
    sh.put(row, 1, "C / W / I", bold=True, halign="right")
    sh.text(row, 2, "Flag severity: critical / warning / info.", span=last - 1, indent=1)
    row += 2

    sh.section(row, "CONTENTS")
    row += 1
    for sheet, desc in _CONTENTS:
        if sheet in ws.parent.sheetnames:
            sh.link(row, 1, sheet, sheet)
        else:
            sh.put(row, 1, sheet)
        sh.text(row, 2, desc, span=last - 1)
        row += 1
        if sheet.startswith("Adj "):
            for group, items in (("", ctx.mgmt_items), (DILIGENCE_BLOCK, ctx.dil_items)):
                if group and items:
                    sh.put(row, 1, group, span=last, italic=True, color=NAVY)
                    row += 1
                for a in items:
                    cell = sh.link(row, 1, ctx.adj_sheets[a.adj_id], ctx.adj_sheets[a.adj_id])
                    cell.alignment = Alignment(indent=2, vertical="top")
                    sh.text(row, 2, a.title, span=last - 1, color=MUTED)
                    row += 1
    row += 1

    sh.section(row, f"INPUT FILES ({len(wp.input_hashes)}), SHA-256")
    row += 1
    for relpath, digest in sorted(wp.input_hashes.items()):
        sh.text(row, 1, relpath)
        sh.put(row, 2, digest, span=last - 1, color=MUTED)
        row += 1


# ---------------------------------------------------------------------------
# Workbook assembly
# ---------------------------------------------------------------------------


def _use_arial_default(wb: Workbook) -> None:
    # Unstyled cells (and anything a reviewer types later) inherit the default font.
    arial = Font(name=FONT_NAME, size=10, family=2)
    wb._fonts = IndexedList([arial])
    wb._named_styles["Normal"].font = arial


def _page_setup(ws: Worksheet, ctx: _Ctx) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    label = "SYNTHETIC" if ctx.wp.deal.synthetic else "DRAFT"
    ws.oddHeader.left.text = f"{ctx.wp.deal.target_name}: &A"
    ws.oddFooter.left.text = f"{label}: QoE Evidence Review, run {ctx.wp.run_id}"
    ws.oddFooter.right.text = "Page &P of &N"


def build_workbook(wp: Workpaper, *, pkg: Optional[DealPackage] = None) -> Workbook:
    """Build the workpaper in memory (see ``export_workpaper``)."""
    ctx = _context(wp, pkg)
    wb = Workbook()
    _use_arial_default(wb)
    cover = wb.active
    cover.title = SHEET_COVER
    bridge = wb.create_sheet(SHEET_BRIDGE)
    summary = wb.create_sheet(SHEET_SUMMARY)
    support = {a.adj_id: wb.create_sheet(ctx.adj_sheets[a.adj_id]) for a in wp.assessments}
    questions = wb.create_sheet(SHEET_QUESTIONS)
    recon = wb.create_sheet(SHEET_RECON)
    dq = wb.create_sheet(SHEET_DATA_QUALITY)
    log = wb.create_sheet(SHEET_REVIEW_LOG)

    # Support sheets are the source; the summary links to them, the bridge checks against the summary.
    tie = {a.adj_id: _write_support(support[a.adj_id], ctx, a) for a in wp.assessments}
    srefs = _write_summary(summary, ctx, tie)
    brefs = _write_bridge(bridge, ctx, srefs, tie)
    q_rows = _write_questions(questions, ctx)
    _write_recon(recon, ctx)
    dq_rows = _write_data_quality(dq, ctx)
    _write_review_log(log, ctx)
    _write_cover(cover, ctx, srefs, brefs, q_rows, dq_rows)

    for ws in wb.worksheets:
        _page_setup(ws, ctx)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    label = "SYNTHETIC" if wp.deal.synthetic else "DRAFT"
    wb.properties.title = f"QoE Evidence Review: {wp.deal.target_name} ({label})"
    wb.properties.subject = f"Quality of earnings evidence review, deal {wp.deal.deal_id}, run {wp.run_id}"
    wb.properties.creator = f"QoE Evidence Review {wp.tool_version}"
    wb.properties.keywords = label
    return wb
