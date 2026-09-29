"""Excel outputs: management's monthly P&L (SPEC §3.5), the adjusted EBITDA
schedule (SPEC §3.6), and byte-reproducible xlsx saving.

Both workbooks hold values, not formulas: a QBO export has no formulas, and
readers open files with data_only semantics where uncalculated formulas read
as empty.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from qoe.money import ZERO

from .accounts import AccountInfo
from .ledger import MONTH_ABBR, MONTH_NAME, GenerationError
from .spec import DealSpec, MonthlyPLSpec, PLSection

_CORE_TS = re.compile(rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)")
_THIN = Side(style="thin")


def save_workbook(wb: Workbook, path: Path, package_date: str, creator: str = "") -> None:
    """Save with fixed metadata and zip timestamps so the same input gives the same bytes."""
    stamp = datetime.fromisoformat(package_date + "T09:00:00")
    wb.properties.creator = creator or None
    wb.properties.lastModifiedBy = None
    wb.properties.created = stamp
    buf = io.BytesIO()
    wb.save(buf)  # openpyxl stamps "modified" with the wall clock; rewritten below
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _CORE_TS.sub(lambda m: m.group(1) + iso + m.group(3), data)
            zi = zipfile.ZipInfo(info.filename, date_time=(stamp.year, stamp.month, stamp.day, 9, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(out.getvalue())


def _month_header(month: str, style: str) -> object:
    y, m = (int(x) for x in month.split("-"))
    if style == "mon_yyyy":
        return f"{MONTH_ABBR[m - 1]} {y}"
    if style == "month_yyyy":
        return f"{MONTH_NAME[m - 1]} {y}"
    if style == "iso":
        return month
    return datetime(y, m, 1)


def _credit_natural(section_name: str) -> bool:
    # Same rule the reader applies (SPEC §3.5), so displayed signs round-trip.
    name = section_name.lower()
    return "income" in name and "expense" not in name


def _check_sections(pl: MonthlyPLSpec) -> None:
    for s in pl.sections:
        credit = _credit_natural(s.name)
        if s.role in ("income", "other_income") and not credit:
            raise GenerationError(f"monthly_pl section {s.name!r} ({s.role}) must contain 'income' so readers treat it as credit-natural")
        if s.role in ("cogs", "expenses", "other_expenses") and credit:
            raise GenerationError(f"monthly_pl section {s.name!r} ({s.role}) would be read as credit-natural")


def _section_for(acct: AccountInfo, sections: list[PLSection]) -> PLSection:
    t = acct.source_type.strip().lower()
    for s in sections:
        if t in {x.strip().lower() for x in s.types}:
            return s
    raise GenerationError(f"account {acct.number} type {acct.source_type!r} is not listed in any monthly_pl section")


def write_monthly_pl(
    path: Path,
    spec: DealSpec,
    accounts: dict[str, AccountInfo],
    amounts: dict[str, dict[str, Decimal]],
    months: list[str],
) -> None:
    """Account-level 'Profit and Loss by Month' in QBO layout, natural signs by section."""
    pl = spec.monthly_pl
    _check_sections(pl)
    wb = Workbook()
    ws = wb.active
    ws.title = pl.sheet_name
    bold = Font(bold=True)
    y0, m0 = (int(x) for x in months[0].split("-"))
    y1, m1 = (int(x) for x in months[-1].split("-"))
    ws.append([spec.company.name])
    ws.append([pl.title])
    ws.append([f"{MONTH_NAME[m0 - 1]} {y0} - {MONTH_NAME[m1 - 1]} {y1}"])
    if pl.basis_line:
        ws.append([pl.basis_line])
    ws.append([])
    header = ["", *[_month_header(m, pl.month_format) for m in months]]
    if pl.total_column:
        header.append("Total")
    ws.append(header)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = bold
        cell.alignment = Alignment(horizontal="center")
        if isinstance(cell.value, datetime):
            cell.number_format = "mmm yyyy"
    for r in range(1, 4):
        ws.cell(row=r, column=1).font = bold

    ncols = len(months)

    def emit(label: str, values: list[Decimal], font: Optional[Font] = None, top_border: bool = False) -> None:
        cells: list[object] = [label, *[v if v != 0 else None for v in values]]
        if pl.total_column:
            total = sum(values, ZERO)
            cells.append(total if total != 0 else None)
        ws.append(cells)
        row = ws.max_row
        for c in range(2, len(cells) + 1):
            cell = ws.cell(row=row, column=c)
            cell.number_format = "#,##0.00"
            if top_border:
                cell.border = Border(top=_THIN)
        if font is not None:
            ws.cell(row=row, column=1).font = font

    by_section: dict[str, list[AccountInfo]] = {s.role: [] for s in pl.sections}
    for acct in accounts.values():
        if acct.is_pl and any(v != 0 for v in amounts.get(acct.number, {}).values()):
            by_section[_section_for(acct, pl.sections).role].append(acct)

    totals: dict[str, list[Decimal]] = {}
    for section in pl.sections:
        sign = Decimal(-1) if _credit_natural(section.name) else Decimal(1)
        total = [ZERO] * ncols
        accts = by_section[section.role]
        if accts:
            ws.append([section.name])
            ws.cell(row=ws.max_row, column=1).font = bold
            for acct in accts:
                values = [sign * amounts[acct.number].get(m, ZERO) for m in months]
                total = [a + b for a, b in zip(total, values)]
                emit(f"{acct.number} {acct.name}", values)
            emit(section.total_label or f"Total {section.name}", total, bold, top_border=True)
        totals[section.role] = total
        if section.role == "cogs":
            gp = [a - b for a, b in zip(totals.get("income", [ZERO] * ncols), total)]
            totals["gross_profit"] = gp
            emit(pl.gross_profit_label, gp, bold)
        elif section.role == "expenses":
            noi = [a - b for a, b in zip(totals.get("gross_profit", [ZERO] * ncols), total)]
            totals["noi"] = noi
            emit(pl.net_operating_income_label, noi, bold)
        elif section.role == "other_expenses":
            noth = [a - b for a, b in zip(totals.get("other_income", [ZERO] * ncols), total)]
            emit(pl.net_other_income_label, noth, bold)
            emit(pl.net_income_label, [a + b for a, b in zip(totals.get("noi", [ZERO] * ncols), noth)], bold, top_border=True)

    ws.column_dimensions["A"].width = 44
    for idx in range(2, ncols + 3):
        ws.column_dimensions[ws.cell(row=header_row, column=idx).column_letter].width = 14
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    save_workbook(wb, path, spec.package_date, creator=spec.company.name)


def write_schedule(
    path: Path,
    spec: DealSpec,
    components: dict[str, dict[str, Decimal]],
    claims: dict[str, dict[str, Decimal]],
) -> dict[str, dict[str, Decimal]]:
    """Management adjusted EBITDA schedule; returns the totals it printed."""
    sch = spec.schedule
    labels = [p.label for p in spec.periods]
    wb = Workbook()
    ws = wb.active
    ws.title = sch.sheet_name
    bold = Font(bold=True)
    for text in [spec.company.name, *sch.title_rows]:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    cols = sch.columns
    ws.append([cols.ref, cols.title, cols.category, cols.description, cols.accounts, cols.support, *labels])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = bold
        cell.border = Border(bottom=_THIN)
        cell.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)

    fmt = '#,##0;(#,##0);"-"'

    def emit(values: list[object], amounts: dict[str, Decimal], font: Optional[Font] = None, top: bool = False) -> None:
        ws.append([*values, *[amounts[l] for l in labels]])
        row = ws.max_row
        for c in range(7, 7 + len(labels)):
            cell = ws.cell(row=row, column=c)
            cell.number_format = fmt
            if top:
                cell.border = Border(top=_THIN)
        for c in range(1, 7):
            ws.cell(row=row, column=c).alignment = Alignment(vertical="top", wrap_text=c in (2, 4))
        if font is not None:
            for c in range(1, 7 + len(labels)):
                ws.cell(row=row, column=c).font = font

    lab = sch.labels
    ni = {l: components[l]["net_income"] for l in labels}
    interest = {l: components[l]["interest"] for l in labels}
    taxes = {l: components[l]["taxes"] for l in labels}
    da = {l: components[l]["depreciation"] + components[l]["amortization"] for l in labels}
    reported = {l: ni[l] + interest[l] + taxes[l] + da[l] for l in labels}
    emit(["", lab.net_income, "", "", "", ""], ni)
    emit(["", lab.interest, "", "", "", ""], interest)
    emit(["", lab.taxes, "", "", "", ""], taxes)
    emit(["", lab.depreciation_amortization, "", "", "", ""], da)
    emit(["", lab.reported_ebitda, "", "", "", ""], reported, bold, top=True)
    ws.append([])
    total = {l: ZERO for l in labels}
    for adj in sch.adjustments:
        amounts = claims[adj.ref]
        if adj.ref not in sch.total_excludes:  # planted mis-foot: the SUM range stops short of these rows
            total = {l: total[l] + amounts[l] for l in labels}
        emit([adj.ref, adj.title, adj.category, adj.description, adj.accounts, adj.support], amounts)
    emit(["", lab.total_adjustments, "", "", "", ""], total, bold, top=True)
    adjusted = {l: reported[l] + total[l] for l in labels}
    ws.append([])
    emit(["", lab.adjusted_ebitda, "", "", "", ""], adjusted, bold, top=True)

    for col, width in zip("ABCDEF", (8, 34, 20, 70, 16, 18)):
        ws.column_dimensions[col].width = width
    for idx in range(len(labels)):
        ws.column_dimensions[chr(ord("G") + idx)].width = 14
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    save_workbook(wb, path, spec.package_date, creator=spec.company.name)
    return {"net_income": ni, "interest": interest, "taxes": taxes, "depreciation_amortization": da,
            "reported_ebitda": reported, "total_adjustments": total, "adjusted_ebitda": adjusted}
