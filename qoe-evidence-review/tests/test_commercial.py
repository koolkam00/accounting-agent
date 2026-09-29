"""Tests for qoe.commercial: scenario math, the not-measured path, and the workbook contract."""

from __future__ import annotations

import importlib.util
import re
from dataclasses import replace
from decimal import Decimal
from pathlib import Path

import pytest

from qoe.money import D
from qoe.commercial import (
    ASSUMPTIONS_SHEET,
    CAPACITY_VS_VALUE,
    INPUT_SPECS,
    LINES,
    MEASURED_KEYS,
    SCENARIO_ILLUSTRATIVE_COL,
    SCENARIO_YOUR_COL,
    SCENARIOS_SHEET,
    STATUS_AWAITING,
    STATUS_COMPLETE,
    STATUS_INCOMPLETE,
    CommercialInputs,
    by_scenario,
    compare_with_model,
    compute,
    input_row,
    read_back,
    status,
    write_workbook,
)

ROOT = Path(__file__).resolve().parents[1]
EXCEL_2007_FUNCTIONS = {"IF", "COUNT", "MIN", "ROUNDDOWN", "ISNUMBER", "NOT"}


def _illustrative() -> dict[str, Decimal | None]:
    return compute(CommercialInputs.illustrative())


def _depends_on_measurement(key: str) -> bool:
    line = next((ln for ln in LINES if ln.key == key), None)
    if line is None:
        return key in MEASURED_KEYS
    return any(_depends_on_measurement(d) for d in line.deps)


# ---------------------------------------------------------------------------
# Scenario math (illustrative inputs: 12 adj, 6h + 1.5h, 40% / 15%, 20 eng/yr)
# ---------------------------------------------------------------------------


def test_scenario_a_capacity_only_is_hours():
    r = _illustrative()
    assert r["baseline_review_hours"] == Decimal("90")  # 12 x (6 + 1.5)
    assert r["preparer_hours_freed"] == Decimal("28.8")  # 12 x 6 x 0.40
    assert r["reviewer_hours_freed"] == Decimal("2.7")  # 12 x 1.5 x 0.15
    assert r["hours_freed_per_engagement"] == Decimal("31.5")
    assert r["hours_freed_per_year"] == Decimal("630")
    assert r["engagement_hours_with_tool"] == Decimal("558.5")  # 90 + 500 - 31.5
    assert r["equivalent_engagements"] == Decimal("630") / Decimal("558.5")
    # Capacity carries no dollar line of its own.
    assert set(by_scenario(r)["a"]) == {
        "baseline_review_hours",
        "preparer_hours_freed",
        "reviewer_hours_freed",
        "hours_freed_per_engagement",
        "hours_freed_per_year",
        "engagement_hours_with_tool",
        "equivalent_engagements",
    }
    assert not any("USD" in ln.label for ln in LINES if ln.section == "a")


def test_scenario_b_fixed_fee_margin_and_hourly_giveback():
    r = _illustrative()
    assert r["cost_avoided_per_engagement"] == Decimal("3064.5")  # 28.8 x 90 + 2.7 x 175
    assert r["margin_lift_per_fixed_fee_engagement"] == Decimal("3064.5") / Decimal("150000")
    assert r["fixed_fee_engagements"] == Decimal("12")
    assert r["fixed_fee_margin_per_year"] == Decimal("36774")
    assert r["hourly_engagements"] == Decimal("8")
    # Hourly work: the client keeps the saving, so fees fall faster than cost.
    assert r["hourly_fees_forgone_per_year"] == Decimal("-75600")  # 31.5 x 300 x 8
    assert r["hourly_cost_avoided_per_year"] == Decimal("24516")
    assert r["hourly_net_margin_per_year"] == Decimal("-51084")


def test_scenario_c_values_only_redeployed_share():
    r = _illustrative()
    assert r["hours_redeployed_per_year"] == Decimal("315")
    assert r["hours_not_redeployed_per_year"] == Decimal("315")
    assert r["redeployment_value_per_year"] == Decimal("94500")  # 315 x 300, not 630 x 300
    none_redeployed = compute(replace(CommercialInputs.illustrative(), redeployment_rate=Decimal("0")))
    assert none_redeployed["redeployment_value_per_year"] == 0
    assert none_redeployed["hours_freed_per_year"] == Decimal("630")  # capacity unchanged


def test_scenario_d_hire_capped_at_whole_ftes():
    r = _illustrative()
    assert r["fte_freed"] == Decimal("0.42")
    assert r["hires_avoided"] == 0  # 0.42 FTE does not avoid a hire
    assert r["hire_cost_avoided_per_year"] == 0
    big = compute(
        replace(CommercialInputs.illustrative(), engagements_per_year=Decimal("60"), hires_planned=Decimal("3"))
    )
    assert big["fte_freed"] == Decimal("1.26")  # 1,890 / 1,500
    assert big["hires_avoided"] == 1
    assert big["hire_cost_avoided_per_year"] == Decimal("140000")
    not_hiring = compute(
        replace(CommercialInputs.illustrative(), engagements_per_year=Decimal("60"), hires_planned=Decimal("0"))
    )
    assert not_hiring["hire_cost_avoided_per_year"] == 0


def test_zero_denominators_give_none():
    r = compute(
        replace(
            CommercialInputs.illustrative(),
            fixed_fee_per_engagement=Decimal("0"),
            chargeable_hours_per_fte=Decimal("0"),
        )
    )
    assert r["margin_lift_per_fixed_fee_engagement"] is None
    assert r["fte_freed"] is None
    assert r["hires_avoided"] is None and r["hire_cost_avoided_per_year"] is None
    assert r["fixed_fee_margin_per_year"] == Decimal("36774")


# ---------------------------------------------------------------------------
# Not measured -> no result
# ---------------------------------------------------------------------------


def test_default_inputs_produce_no_result():
    inputs = CommercialInputs()
    assert status(inputs) == STATUS_AWAITING
    assert all(v is None for v in compute(inputs).values())


def test_missing_measurement_blanks_every_dependent_line():
    inputs = replace(CommercialInputs.illustrative(), preparer_time_reduction=None, reviewer_time_reduction=None)
    assert status(inputs) == STATUS_AWAITING
    r = compute(inputs)
    for line in LINES:
        if _depends_on_measurement(line.key):
            assert r[line.key] is None, line.key
        else:
            assert r[line.key] is not None, line.key
    # Only descriptive, non-result lines survive.
    assert {k for k, v in r.items() if v is not None} == {
        "baseline_review_hours",
        "fixed_fee_engagements",
        "hourly_engagements",
    }
    # One measurement missing is enough to withhold the result.
    assert (
        compute(replace(CommercialInputs.illustrative(), reviewer_time_reduction=None))["hours_freed_per_year"] is None
    )


def test_incomplete_assumptions_blank_only_what_they_feed():
    inputs = replace(CommercialInputs.illustrative(), preparer_cost_per_hour=None)
    assert status(inputs) == STATUS_INCOMPLETE
    r = compute(inputs)
    assert r["hours_freed_per_year"] == Decimal("630")
    assert r["cost_avoided_per_engagement"] is None
    assert r["fixed_fee_margin_per_year"] is None
    assert r["redeployment_value_per_year"] == Decimal("94500")
    assert status(CommercialInputs.illustrative()) == STATUS_COMPLETE


# ---------------------------------------------------------------------------
# Input contract
# ---------------------------------------------------------------------------


def test_lines_declare_exactly_the_cells_they_use():
    known = {s.key for s in INPUT_SPECS}
    for line in LINES:
        placeholders = set(re.findall(r"\{(\w+)\}", line.excel))
        assert placeholders == set(line.deps), line.key
        assert set(line.deps) <= known, (line.key, set(line.deps) - known)
        known.add(line.key)


def test_every_input_is_labeled_and_noted():
    assert {s.key for s in INPUT_SPECS} == set(CommercialInputs().as_dict())
    for spec in INPUT_SPECS:
        assert spec.basis in ("ASSUMPTION", "MEASURED")
        assert spec.note.strip()
    assert set(MEASURED_KEYS) == {"preparer_time_reduction", "reviewer_time_reduction"}
    assert CommercialInputs().preparer_time_reduction is None


def test_input_validation():
    with pytest.raises(ValueError, match="fraction"):
        CommercialInputs(redeployment_rate=Decimal("35"))
    with pytest.raises(ValueError, match="negative"):
        CommercialInputs(engagements_per_year=Decimal("-1"))
    with pytest.raises(TypeError):
        CommercialInputs(fixed_fee_per_engagement=150000.0)  # type: ignore[arg-type]
    parsed = CommercialInputs.from_mapping(
        {"preparer_time_reduction": "0.35", "fixed_fee_per_engagement": "150,000", "reviewer_time_reduction": ""}
    )
    assert parsed.preparer_time_reduction == Decimal("0.35")
    assert parsed.fixed_fee_per_engagement == Decimal("150000")
    assert parsed.reviewer_time_reduction is None
    with pytest.raises(KeyError):
        CommercialInputs.from_mapping({"time_saved": 0.5})


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------


def test_workbook_structure(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "commercial_model.xlsx"
    rows = write_workbook(path)
    wb = load_workbook(path)
    assert wb.sheetnames == [ASSUMPTIONS_SHEET, SCENARIOS_SHEET]
    a, sc = wb[ASSUMPTIONS_SHEET], wb[SCENARIOS_SHEET]

    for spec in INPUT_SPECS:
        r = input_row(spec.key)
        yours, example = a[f"C{r}"], a[f"F{r}"]
        assert yours.value is None  # blank until the user supplies data
        assert yours.fill.start_color.rgb.endswith("FFFF00")  # yellow: must replace
        assert yours.font.color.rgb.endswith("0000FF")  # blue: input
        assert D(example.value) == spec.illustrative
        assert example.font.color.rgb.endswith("0000FF")
        assert yours.number_format == spec.number_format
        if spec.fraction:
            assert yours.number_format == "0.0%"
            assert D(example.value) <= 1  # stored as a fraction
    assert "Illustrative" in a["F12"].value
    legend = " ".join(str(a[f"C{r}"].value) for r in range(6, 10))
    assert "Yellow fill" in legend and "Blue text" in legend and "ILLUSTRATIVE" in legend

    texts = " ".join(str(c.value) for ws in (a, sc) for row in ws.iter_rows() for c in row if isinstance(c.value, str))
    assert CAPACITY_VS_VALUE in texts
    assert "do not add them together" in sc["A2"].value

    functions: set[str] = set()
    for key, r in rows.items():
        for col in (SCENARIO_YOUR_COL, SCENARIO_ILLUSTRATIVE_COL):
            value = sc[f"{col}{r}"].value
            assert isinstance(value, str) and value.startswith("="), (key, value)
            functions |= set(re.findall(r"([A-Z][A-Z0-9.]*)\(", value))
        # Each column reads only its own Assumptions column.
        assert "Assumptions!$F$" not in sc[f"C{r}"].value
        assert "Assumptions!$C$" not in sc[f"D{r}"].value
    assert functions <= EXCEL_2007_FUNCTIONS, functions

    for ws in (a, sc):
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    assert c.font.name == "Arial", (ws.title, c.coordinate)
    money_rows = [rows[ln.key] for ln in LINES if "USD" in ln.label]
    assert all(sc[f"D{r}"].number_format.startswith("#,##0") for r in money_rows)


def test_workbook_carries_user_inputs_and_sources(tmp_path):
    from openpyxl import load_workbook

    inputs = replace(CommercialInputs.illustrative(), preparer_time_reduction=Decimal("0.3"))
    path = tmp_path / "m.xlsx"
    write_workbook(path, inputs, {"preparer_time_reduction": "timed pilot, 3 engagements"})
    a = load_workbook(path)[ASSUMPTIONS_SHEET]
    r = input_row("preparer_time_reduction")
    assert D(a[f"C{r}"].value) == Decimal("0.3")
    assert "timed pilot" in a[f"G{r}"].value
    assert a[f"E{r}"].value.startswith("=IF(ISNUMBER(")  # basis label follows the cell


def test_compare_with_model_flags_disagreements():
    expected = {"a": Decimal("10"), "b": None}
    assert compare_with_model({"a": 10.0, "b": "n/a"}, expected) == []
    assert compare_with_model({"a": 11.0, "b": "n/a"}, expected)
    assert compare_with_model({"a": 10.0, "b": 0}, expected)


def _load_script():
    path = ROOT / "scripts" / "qoe_commercial_model.py"
    spec = importlib.util.spec_from_file_location("_test_qoe_commercial_model", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def test_libreoffice_recalc_has_no_errors_and_matches_model(tmp_path):
    script = _load_script()
    if not script.calc_available():
        pytest.skip("LibreOffice Calc is not installed")
    recalc_py = script.find_recalc_script(None)
    if recalc_py is None:
        pytest.skip("xlsx skill recalc.py not found")
    inputs = replace(CommercialInputs.illustrative(), preparer_time_reduction=None)
    path = tmp_path / "commercial_model.xlsx"
    rows = write_workbook(path, inputs)
    result = script.recalc(path, recalc_py, timeout=60)
    assert result.get("status") == "success", result
    assert result["total_errors"] == 0
    cached = read_back(path, rows)
    assert compare_with_model(cached["your"], compute(inputs)) == []
    assert compare_with_model(cached["illustrative"], compute(CommercialInputs.illustrative())) == []
    from openpyxl import load_workbook

    sc = load_workbook(path, data_only=True)[SCENARIOS_SHEET]
    assert sc["C6"].value.startswith("Awaiting measured time reduction")
    assert sc["D6"].value.startswith("ILLUSTRATIVE")


def test_script_writes_workbook_without_recalc(tmp_path, capsys):
    script = _load_script()
    out = tmp_path / "cm.xlsx"
    assert script.main(["--out", str(out), "--no-recalc"]) == 0
    assert out.is_file()
    printed = capsys.readouterr().out
    assert "AWAITING_MEASUREMENT" in printed
    assert "Capacity is not value" in printed
