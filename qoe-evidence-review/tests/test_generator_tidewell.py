"""Dev deal 2 (Tidewell Distribution, NetSuite) and the generator features it introduced.

The deal exercises every additive generator option: per-row NetSuite dimension noise,
normalization levels, a planted missing GL month kept through a diligence item, a
schedule whose total row does not foot, TTM-only verification for run-rate pro formas,
and related documents. Checks read the generated files back with small independent
readers, as tests/test_generator.py does for Meridian.
"""

from __future__ import annotations

import csv
import datetime
import re
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from qoe.money import D, ZERO
from qoe.schemas import GroundTruth, Treatment

ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT / "data" / "specs" / "tidewell_distribution.yaml"
COMMITTED = ROOT / "data" / "dev" / "tidewell_distribution"
LABELS = ["FY2025", "FY2026", "TTM Jun-26"]
PERIODS = {"FY2025": ("2024-04", "2025-03"), "FY2026": ("2025-04", "2026-03"), "TTM Jun-26": ("2025-07", "2026-06")}

pytest.importorskip("reportlab")
sys.path.insert(0, str(ROOT / "scripts"))
from qoe_synth import generate_deal, load_spec  # noqa: E402

pytestmark = pytest.mark.skipif(not SPEC.exists(), reason="Tidewell spec not present")

CLAIMS = {  # FY2025 / FY2026 / TTM Jun-26, as management presents them
    "1": (540000, 540000, 540000), "2": (48000, 48000, 48000), "3": (60000, 60000, 60000),
    "4": (0, 0, Decimal("88230.18")), "5": (0, 0, 48000), "6": (0, 90000, 90000), "7": (0, 100000, 100000),
    "8": (0, 148600, 148600), "9": (0, 51000, 34200), "10": (0, 128400, 144600), "11": (0, 54000, 54000),
    "12": (0, 46500, 46500), "13": (0, 38400, 38400), "14": (0, 69600, 69600), "15": (0, 18500, 18500),
    "16": (0, 95000, 95000),
}
TRUTH = {
    "1": ("ACCEPT", (540000, 540000, 540000)), "2": ("REVISE", (-60000, -60000, -60000)), "3": ("REQUEST_INFO", None),
    "4": ("ACCEPT", (0, 0, Decimal("88230.18"))), "5": ("REVISE", (0, 0, 43200)), "6": ("ACCEPT", (0, 90000, 90000)),
    "7": ("REVISE", (0, 72000, 42000)), "8": ("ACCEPT", (0, 148600, 148600)), "9": ("ACCEPT", (0, 51000, 34200)),
    "10": ("REVISE", (0, 64000, 64000)), "11": ("REVISE", (0, 84000, 84000)), "12": ("REVISE", (0, -46500, -46500)),
    "13": ("REVISE", (0, 0, 19200)), "14": ("REVISE", (0, 42000, 42000)), "15": ("ACCEPT", (0, 18500, 18500)),
    "16": ("ACCEPT", (0, 95000, 95000)),
}
DILIGENCE = {"D-1": (0, 32000, 32000), "D-2": (0, -186000, 0), "D-3": (Decimal("-363586.47"), 0, 0)}


def all_files(root: Path) -> dict[str, bytes]:
    return {str(p.relative_to(root)): p.read_bytes() for p in sorted(root.rglob("*")) if p.is_file()}


def month_of(mdy: str) -> str:
    m, _, y = mdy.split("/")
    return f"{y}-{int(m):02d}"


def read_gl(deal: Path) -> dict[int, dict]:
    with (deal / "gl" / "general_ledger.csv").open(newline="", encoding="utf-8") as f:
        return {line: rec for line, rec in enumerate(csv.DictReader(f), start=2)}


def dp(rec: dict) -> Decimal:
    return D(rec["Debit"] or "0") - D(rec["Credit"] or "0")


def read_pl(path: Path) -> dict[tuple[str, str], Decimal]:
    ws = load_workbook(path, data_only=True).worksheets[0]
    header, credit, out = None, False, defaultdict(lambda: ZERO)
    for rec in ws.iter_rows(values_only=True):
        if header is None:
            if sum(isinstance(v, datetime.datetime) for v in rec) >= 2:
                header = rec
            continue
        label = rec[0]
        if isinstance(label, str) and all(v is None for v in rec[1:]):
            credit = "income" in label.lower() and "expense" not in label.lower()
            continue
        if isinstance(label, str) and re.match(r"^\d{4} ", label):
            for h, v in zip(header, rec):
                if isinstance(h, datetime.datetime) and v is not None:
                    out[(label[:4], f"{h.year}-{h.month:02d}")] += -D(str(v)) if credit else D(str(v))
    return dict(out)


@pytest.fixture(scope="module")
def deal(tmp_path_factory):
    return generate_deal(load_spec(SPEC), tmp_path_factory.mktemp("tidewell"))


@pytest.fixture(scope="module")
def truth(deal):
    return GroundTruth.model_validate_json((deal.deal_dir / "ground_truth.json").read_text())


def test_committed_tidewell_package_is_current(deal):
    if not COMMITTED.exists():
        pytest.skip("data/dev/tidewell_distribution not generated")
    fresh, committed = all_files(deal.deal_dir), all_files(COMMITTED)
    assert fresh.keys() == committed.keys()
    assert [k for k in fresh if fresh[k] != committed[k]] == [], "regenerate with scripts/qoe_generate_deals.py --spec"


def test_truth_matches_catalog(truth):
    got = {a.adj_id: a for a in truth.adjustments}
    assert list(got) == list(TRUTH)
    for adj_id, (treatment, amounts) in TRUTH.items():
        a = got[adj_id]
        assert a.treatment == Treatment(treatment), adj_id
        expected = {} if amounts is None else {l: str(D(v).quantize(Decimal("0.01"))) for l, v in zip(LABELS, amounts)}
        assert a.amounts == expected, adj_id
        assert a.rationale and a.question_topics and a.case_type, adj_id
    adequate = [a for a in truth.adjustments if a.case_type == "ADEQUATE"]
    assert 13 <= len(truth.adjustments) <= 16 and len(adequate) / len(truth.adjustments) >= 0.4
    assert all(a.treatment == Treatment.ACCEPT for a in adequate)
    items = {d.adj_id: d for d in truth.diligence_items}
    assert {k: tuple(D(items[k].amounts[l]) for l in LABELS) for k in items} == {k: tuple(D(v) for v in vals) for k, vals in DILIGENCE.items()}
    assert got["10"].related_docs == ["5.1.3 Kestrel Systems Invoice KS-SUB-2511.pdf", "5.1.4 Kestrel Systems Customer Statement Dec 2025.pdf"]


def test_identity_and_totals(truth):
    finals = {l: ZERO for l in LABELS}
    for a in [*truth.adjustments, *truth.diligence_items]:
        for l in LABELS:
            finals[l] += D(a.amounts[l]) if a.amounts else ZERO
    for l in LABELS:
        assert D(truth.diligence_adjusted_ebitda[l]) == D(truth.gl_ebitda[l]) + finals[l]
    mgmt = [sum((D(a.amounts[l]) for a in truth.adjustments if a.amounts), ZERO) for l in LABELS]
    assert mgmt == [Decimal("480000"), Decimal("1098600"), Decimal("1202430.18")]


def test_amounts_tie_to_exported_gl(deal, truth):
    gl = read_gl(deal.deal_dir)

    def per(lines):
        out = {l: ZERO for l in LABELS}
        for n in lines:
            for l, (s, e) in PERIODS.items():
                if s <= month_of(gl[n]["Date"]) <= e:
                    out[l] += dp(gl[n])
        return out

    got = {a.adj_id: a for a in [*truth.adjustments, *truth.diligence_items]}
    owner, rent = per(got["1"].supporting_gl_rows), per(got["2"].supporting_gl_rows)
    for l in LABELS:  # 12-month periods: actual less 12 x the benchmark level
        assert owner[l] - 12 * 30000 == D(got["1"].amounts[l]) and rent[l] - 12 * 29000 == D(got["2"].amounts[l])
    assert per(got["4"].supporting_gl_rows)["TTM Jun-26"] == D(got["4"].amounts["TTM Jun-26"])
    kestrel = per(got["10"].supporting_gl_rows)
    assert kestrel == {"FY2025": 0, "FY2026": 64000, "TTM Jun-26": 64000}
    freight = per(got["13"].supporting_gl_rows)  # booked Nov-25; service Apr-Sep 25
    assert freight["FY2026"] - 38400 == 0 and freight["TTM Jun-26"] - 38400 * Decimal(3) / 6 == 19200


def test_data_quality_plants(deal, truth):
    codes = sorted((d.code.value, d.month, d.account) for d in truth.data_quality)
    assert codes == sorted([
        ("DUPLICATE_GL_ENTRY", "2025-10", "6600"), ("RECON_VARIANCE", "2026-03", "6040"), ("RECON_VARIANCE", "2026-06", "6040"),
        ("MISSING_PERIOD", "2024-08", None), ("MGMT_EBITDA_DIFFERS_FROM_GL", None, None), ("MGMT_SCHEDULE_ARITHMETIC", None, None),
    ])
    gl = read_gl(deal.deal_dir)
    pl = read_pl(deal.deal_dir / "financials" / "monthly_pl.xlsx")
    in_gl = {rec["Account"][:4] for rec in gl.values() if month_of(rec["Date"]) == "2024-08"}
    active = {a for (a, m), v in pl.items() if m == "2024-08" and v != 0}
    assert len(in_gl) < 0.5 * len(active) and in_gl < active
    dup = next(d for d in truth.data_quality if d.code.value == "DUPLICATE_GL_ENTRY")
    first, second = (gl[n] for n in dup.gl_rows)
    assert first["Document Number"] == second["Document Number"] == "KS-10442" and first["Internal ID"] != second["Internal ID"]
    # dimension noise: blank and mixed tags on the same streams
    assert {rec["Department"] for rec in gl.values() if rec["Account"].startswith("4000")} >= {"Sales", ""}
    assert {rec["Class"] for rec in gl.values() if rec["Account"].startswith("4100")} >= {"Janitorial & Sanitation", "Foodservice Packaging"}


def test_schedule_total_row_omits_ref_16(deal):
    ws = load_workbook(deal.deal_dir / "adjustments" / "management_adjusted_ebitda.xlsx", data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if "#" in r and all(l in r for l in LABELS))
    cols = {l: rows[hdr].index(l) for l in LABELS}
    by_key = {(r[0] or (r[1] or "").strip().lower()): r for r in rows[hdr + 1:] if r[0] or r[1]}
    for i, l in enumerate(LABELS):
        listed = sum((D(str(by_key[ref][cols[l]])) for ref in CLAIMS), ZERO)
        assert listed == sum((D(v[i]) for v in CLAIMS.values()), ZERO)
        printed = D(str(by_key["total adjustments"][cols[l]]))
        assert printed == listed - D(CLAIMS["16"][i])
        assert D(str(by_key["adjusted ebitda"][cols[l]])) == D(str(by_key["ebitda, as reported"][cols[l]])) + printed
