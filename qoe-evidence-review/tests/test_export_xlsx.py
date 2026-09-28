"""Tests for qoe.export_xlsx: workbook structure, Deals formatting rules, formulas, and LibreOffice recalculation."""

from __future__ import annotations

import re
import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from openpyxl.utils import get_column_letter

from qoe.export_xlsx import (
    NUMBER_FORMAT,
    SHEET_BRIDGE,
    SHEET_COVER,
    SHEET_DATA_QUALITY,
    SHEET_QUESTIONS,
    SHEET_RECON,
    SHEET_SUMMARY,
    TICKMARKS,
    _fallback_final_amounts,
    _management_text,
    _plain_basis,
    _plan_subtotals,
    _provisional_from_rationale,
    _vouch,
    _wrapped_lines,
    build_workbook,
    export_workpaper,
    find_recalc_script,
    latest_reviews,
    recalc_and_check,
    resolve_final_amounts,
    sanitize_sheet_name,
    support_sheet_names,
    workbook_filename,
)
from qoe.money import D, fmt
from qoe.review_store import bridge_display_rows
from qoe.periods import add_months, month_range
from qoe.schemas import (
    Account,
    AdjustmentAssessment,
    AmountFact,
    AdjustmentCategory,
    AdjustmentClaim,
    BridgeRow,
    CorrectionType,
    DataQualityCode,
    DataQualityIssue,
    DealFiles,
    DealMeta,
    DealPackage,
    DocFacts,
    DocLink,
    EbitdaBridge,
    EbitdaClass,
    EbitdaComponents,
    EvidenceQuote,
    Fact,
    Flag,
    FlagCode,
    GLEntry,
    GLLink,
    ManagementPL,
    ManagementSchedule,
    OpenQuestion,
    PeriodDef,
    ReconciliationItem,
    ReconciliationResult,
    RecurrenceObservation,
    ReviewDecision,
    Severity,
    TermFact,
    Treatment,
    Workpaper,
)

ROOT = Path(__file__).resolve().parents[1]
LABELS = ["FY2024", "FY2025", "TTM Jun-26"]
BANNED = re.compile(r"\b(XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SEQUENCE)\s*\(", re.I)


# ---------------------------------------------------------------------------
# Fixture: a realistic, fictitious workpaper
# ---------------------------------------------------------------------------


def _pm(*values: object) -> dict[str, str]:
    return {label: fmt(v) for label, v in zip(LABELS, values)}


def _q(doc: str, text: str, page: int = 1) -> EvidenceQuote:
    return EvidenceQuote(doc_id=doc, page=page, quote=text)


LIT_LETTER = "3.1 Harrow Legal Engagement Letter Matter 7781.pdf"
RETAINER_LETTER = "3.2 Harrow Legal General Retainer 2023.pdf"
LIT_INVOICE = "3.3 Harrow Legal Invoice 25-114.pdf"
ERP_MSA = "5.1 Corvid ERP Master Services Agreement.pdf"
ERP_EMAIL = "5.2 Email controller re ERP subscription.txt"
COMP_DRAFT = "6.1 DRAFT CEO Employment Agreement.pdf"
STORM_LETTER = "7.4 Keystone Mutual claim settlement letter.pdf"
SEARCH_LETTER = "4.1 Talbot Search engagement letter.pdf"


def _assessments() -> list[AdjustmentAssessment]:
    retainer_quote = _q(RETAINER_LETTER, "a monthly retainer of $2,500 continuing until terminated by either party")
    lit = AdjustmentAssessment(
        adj_id="A-1",
        title="Litigation legal fees (Matter 7781)",
        category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(0, 90000, "55000.50"),
        traced_gl=_pm(0, 90000, "55000.50"),
        documented=_pm(0, 90000, "55000.50"),
        proposed=_pm(0, 60000, "25000.50"),
        treatment=Treatment.REVISE,
        confidence="medium",
        gl_links=[
            GLLink(entry_id="GL-R1201", period="2025-02", amount="34999.50", score=4.5,
                   reasons=["account 6400", "reference 7781"], group="harrow legal|7781",
                   doc_ids=[LIT_LETTER, LIT_INVOICE], role="supporting", claimed=True),
            GLLink(entry_id="GL-R1410", period="2025-11", amount="25000.50", score=4.0,
                   reasons=["account 6400", "counterparty"], group="harrow legal|7781", doc_ids=[LIT_LETTER],
                   role="supporting", claimed=True),
            GLLink(entry_id="GL-R1500", period="2025-08", amount="30000.00", score=3.0, supports_claim=False,
                   reasons=["account 6400", "counterparty", "Removed (RECURRING_PATTERN): recurs"],
                   group="harrow legal|retainer", doc_ids=[RETAINER_LETTER], role="removed", claimed=True,
                   removed_by=FlagCode.RECURRING_PATTERN),
            GLLink(entry_id="GL-R880", period="2024-06", amount="2500.00", score=2.0,
                   reasons=["counterparty"], group="harrow legal|retainer", supports_claim=False, role="context"),
        ],
        doc_links=[
            DocLink(doc_id=LIT_LETTER, relation="agreement", entry_ids=["GL-R1201", "GL-R1410"], score=3.0,
                    reasons=["matter 7781"], quotes=[_q(LIT_LETTER, "Matter 7781: Pryor v. Harborview")]),
            DocLink(doc_id=RETAINER_LETTER, relation="agreement", entry_ids=["GL-R1500", "GL-R880"], score=2.5,
                    reasons=["counterparty"], quotes=[retainer_quote]),
            DocLink(doc_id=LIT_INVOICE, relation="invoice_for_entry", entry_ids=["GL-R1201"], score=4.0,
                    reasons=["doc number 25-114"], quotes=[_q(LIT_INVOICE, "Total due: $30,000.00", page=2)]),
        ],
        flags=[
            Flag(code=FlagCode.RECURRING_PATTERN, severity=Severity.WARNING,
                 message="General retainer activity also appears in FY2024 at a comparable level.",
                 effects={"FY2025": "-30000.00", "TTM Jun-26": "-30000.00"}, entry_ids=["GL-R1500", "GL-R880"],
                 doc_ids=[RETAINER_LETTER]),
            Flag(code=FlagCode.CONTINUING_OBLIGATION, severity=Severity.WARNING,
                 message="The retainer continues until terminated; it is part of the ongoing cost base.",
                 entry_ids=["GL-R1500"], doc_ids=[RETAINER_LETTER], quotes=[retainer_quote]),
        ],
        recurrence=[
            RecurrenceObservation(group="harrow legal|retainer", amounts_by_period=_pm(30000, 0, 0),
                                  entry_ids=["GL-R880"], note="Retainer billed every month."),
        ],
        facts=[
            Fact(text="Matter 7781 invoices in FY2025 total $60,000.00.", entry_ids=["GL-R1201", "GL-R1410"],
                 quotes=[_q(LIT_INVOICE, "Total due: $30,000.00", page=2)]),
            Fact(text="The general retainer is $2,500 per month and continues until terminated.",
                 entry_ids=["GL-R1500"], quotes=[retainer_quote]),
        ],
        judgment_questions=["Is any post-settlement monitoring cost part of the ongoing cost base?"],
        open_questions=[
            OpenQuestion(q_id="Q-A-1-1", adj_id="A-1", text="Confirm no further fees on Matter 7781 after settlement.",
                         priority="high", basis="CONTINUING_OBLIGATION"),
            OpenQuestion(q_id="Q-A-1-2", adj_id="A-1", priority="medium", basis="documented fact",
                         text="Provide the insurer's confirmation of the settlement payment."),
        ],
        rationale="Litigation invoices under Matter 7781 support $60,000 (FY2025); the general retainer recurs.",
    )
    search = AdjustmentAssessment(
        adj_id="A-2",
        title="Executive search fee",
        category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(0, 30000, 10000),
        traced_gl=_pm(0, 30000, 10000),
        documented=_pm(0, 30000, 10000),
        proposed=_pm(0, 30000, 10000),
        treatment=Treatment.ACCEPT,
        confidence="high",
        gl_links=[
            GLLink(entry_id=f"GL-R{row}", period=month, amount="10000.00", score=5.0, doc_ids=[SEARCH_LETTER],
                   role="supporting", claimed=True)
            for row, month in ((1601, "2025-04"), (1602, "2025-05"), (1603, "2025-07"))
        ],
        doc_links=[DocLink(doc_id=SEARCH_LETTER, relation="agreement", entry_ids=["GL-R1601"], score=4.0,
                           quotes=[_q(SEARCH_LETTER, "payable in three equal installments")])],
        facts=[Fact(text="Three installments of $10,000 agree to the engagement letter.",
                    entry_ids=["GL-R1601", "GL-R1602", "GL-R1603"])],
        rationale="Retained search fee agrees to the engagement letter and GL.",
    )
    erp = AdjustmentAssessment(
        adj_id="A-3",
        title='"One-time" ERP implementation',
        category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(0, 72000, 36000),
        traced_gl=_pm(0, 72000, 36000),
        documented=_pm(0, 72000, 36000),
        proposed=_pm(0, 0, 0),
        treatment=Treatment.REJECT,
        confidence="medium",
        gl_links=[GLLink(entry_id=f"GL-R{1700 + i}", period=add_months("2025-01", i), amount="6000.00", score=4.0,
                         doc_ids=[ERP_MSA], supports_claim=False, role="removed", claimed=True,
                         removed_by=FlagCode.CONTRADICTORY_EVIDENCE) for i in range(12)],
        doc_links=[
            DocLink(doc_id=ERP_MSA, relation="agreement", score=4.0,
                    quotes=[_q(ERP_MSA, "a monthly fee of $6,000"), _q(ERP_MSA, "renews automatically", page=3)]),
            DocLink(doc_id=ERP_EMAIL, relation="correspondence", score=2.0,
                    quotes=[_q(ERP_EMAIL, "our ERP subscription")]),
        ],
        flags=[
            Flag(code=FlagCode.CONTRADICTORY_EVIDENCE, severity=Severity.WARNING,
                 message="The controller calls the cost a subscription.", doc_ids=[ERP_EMAIL],
                 effects={"FY2025": "-72000.00", "TTM Jun-26": "-36000.00"},
                 entry_ids=[f"GL-R{1700 + i}" for i in range(12)], quotes=[_q(ERP_EMAIL, "our ERP subscription")]),
            Flag(code=FlagCode.CONTINUING_OBLIGATION, severity=Severity.WARNING,
                 message="36-month term with automatic renewal.", doc_ids=[ERP_MSA]),
        ],
        facts=[Fact(text="The MSA sets a monthly fee of $6,000.", quotes=[_q(ERP_MSA, "a monthly fee of $6,000")])],
        judgment_questions=["Does any part of the fee relate to one-time configuration work?"],
        open_questions=[OpenQuestion(q_id="Q-A-3-1", adj_id="A-3", text="=Provide the statement of work, if any.",
                                     priority="low", basis="CONTRADICTORY_EVIDENCE")],
        rationale="Documents describe a recurring subscription.",
    )
    comp = AdjustmentAssessment(
        adj_id="A-4",
        title="Owner compensation normalization",
        category=AdjustmentCategory.NORMALIZATION,
        claimed=_pm(250000, 250000, 250000),
        traced_gl=_pm(550000, 550000, 550000),
        documented=_pm(0, 0, 0),
        proposed={},
        treatment=Treatment.REQUEST_INFO,
        confidence="high",
        gl_links=[GLLink(entry_id=f"GL-R{row}", period=month, amount=amount, score=4.0, role="supporting",
                         claimed=True, group="payroll|owner")
                  for row, month, amount in ((300, "2024-06", "550000.00"), (1400, "2025-03", "275000.00"),
                                             (1450, "2025-09", "275000.00"), (2100, "2026-03", "275000.00"))],
        doc_links=[DocLink(doc_id=COMP_DRAFT, relation="agreement", score=2.0,
                           quotes=[_q(COMP_DRAFT, "DRAFT - FOR DISCUSSION ONLY")])],
        flags=[
            Flag(code=FlagCode.UNSIGNED_OR_DRAFT_SUPPORT, severity=Severity.WARNING,
                 message="Only support is an unsigned draft agreement.", doc_ids=[COMP_DRAFT]),
            Flag(code=FlagCode.NORMALIZATION_BENCHMARK_MISSING, severity=Severity.WARNING,
                 message="No market compensation benchmark."),
        ],
        open_questions=[OpenQuestion(q_id="Q-A-4-1", adj_id="A-4", text="Provide the executed agreement.",
                                     priority="high", basis="UNSIGNED_OR_DRAFT_SUPPORT")],
        rationale=("REQUEST_INFO: no executed agreement sets the normalized level. Provisional: at a normalized "
                   "level of 300,000 a year the GL supports FY2024 250,000 / FY2025 250,000 / TTM Jun-26 250,000, "
                   "before payroll taxes."),
    )
    storm = AdjustmentAssessment(
        adj_id="A-5",
        title="Hurricane repairs, net of insurance",
        category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(45000, 0, 0),
        traced_gl=_pm(45000, 0, 0),
        documented=_pm(45000, 0, 0),
        proposed=_pm(45000, -30000, 0),
        treatment=Treatment.REVISE,
        confidence="high",
        gl_links=[GLLink(entry_id="GL-R400", period="2024-10", amount="45000.00", score=4.0, role="supporting",
                         claimed=True),
                  GLLink(entry_id="GL-R1250", period="2025-02", amount="-30000.00", score=3.5, supports_claim=False,
                         reasons=["claim number KM-24-5521"], doc_ids=[STORM_LETTER], role="recovery")],
        doc_links=[DocLink(doc_id=STORM_LETTER, relation="recovery", entry_ids=["GL-R1250"], score=3.5,
                           quotes=[_q(STORM_LETTER, "net payment of $30,000.00")])],
        flags=[Flag(code=FlagCode.OFFSETTING_RECOVERY, severity=Severity.WARNING,
                    message="Insurance recovery booked to other income was not adjusted.",
                    period_label="FY2025", amount_impact="-30000.00", effects={"FY2025": "-30000.00"},
                    entry_ids=["GL-R1250"], doc_ids=[STORM_LETTER])],
        rationale="Recovery offsets the add-back in FY2025.",
    )
    refi = AdjustmentAssessment(
        adj_id="A-6",
        title="Refinancing costs",
        category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(0, 20000, 0),
        traced_gl=_pm(0, 20000, 0),
        documented=_pm(0, 20000, 0),
        proposed=_pm(0, 0, 0),
        treatment=Treatment.REJECT,
        confidence="high",
        gl_links=[GLLink(entry_id="GL-R1330", period="2025-06", amount="20000.00", score=4.0, supports_claim=False,
                         role="removed", claimed=True, removed_by=FlagCode.ALREADY_EXCLUDED_FROM_EBITDA)],
        flags=[Flag(code=FlagCode.ALREADY_EXCLUDED_FROM_EBITDA, severity=Severity.CRITICAL,
                    message="Booked to interest expense, which EBITDA already adds back.",
                    effects={"FY2025": "-20000.00"}, entry_ids=["GL-R1330"])],
        rationale="Already excluded from EBITDA.",
    )
    return [lit, search, erp, comp, storm, refi]


PREMIUM_NOTICE = "9.4 Keystone Risk premium installment notice May 2025.pdf"


def _diligence_item() -> AdjustmentAssessment:
    """SPEC §5.7: a duplicate posting the tool proposes to reverse (not on management's schedule)."""
    return AdjustmentAssessment(
        adj_id="D-1",
        title="Reverse duplicate premium posting (KRI-25-0507)",
        category=AdjustmentCategory.OTHER,
        source="diligence",
        description="Premium installment KRI-25-0507 is posted twice in May 2025; the P&L carries both.",
        gl_accounts=["6200"],
        support_refs=[PREMIUM_NOTICE],
        claimed=_pm(0, 0, 0),
        traced_gl=_pm(0, 18400, 0),
        documented=_pm(0, 18400, 0),
        proposed=_pm(0, 18400, 0),
        treatment=Treatment.REVISE,
        confidence="high",
        gl_links=[
            GLLink(entry_id="GL-R2210", period="2025-05", amount="18400.00", score=5.0, supports_claim=False,
                   reasons=["first posting: kept"], doc_ids=[PREMIUM_NOTICE], role="context"),
            GLLink(entry_id="GL-R2215", period="2025-05", amount="18400.00", score=5.0,
                   reasons=["duplicate of GL row 2210"], doc_ids=[PREMIUM_NOTICE], role="supporting"),
        ],
        doc_links=[DocLink(doc_id=PREMIUM_NOTICE, relation="invoice_for_entry", entry_ids=["GL-R2210", "GL-R2215"],
                           score=4.0, quotes=[_q(PREMIUM_NOTICE, "Installment due: $18,400.00")])],
        flags=[Flag(code=FlagCode.DUPLICATE_GL_ENTRY, severity=Severity.WARNING,
                    message="The same bill (KRI-25-0507) is posted twice, three days apart.",
                    effects={"FY2025": "18400.00"}, entry_ids=["GL-R2210", "GL-R2215"])],
        open_questions=[OpenQuestion(q_id="Q-D-1-1", adj_id="D-1", priority="medium", basis="DUPLICATE_GL_ENTRY",
                                     text="Was installment KRI-25-0507 paid twice, refunded, or applied to a later bill?")],
        rationale="One installment is due; the second posting overstates FY2025 insurance expense.",
    )


def _reviews() -> list[ReviewDecision]:
    return [
        ReviewDecision(adj_id="A-3", reviewer="k.osei", timestamp="2026-09-01T10:00:00Z", treatment=Treatment.REJECT,
                       amounts=_pm(0, 0, 0), rationale="Agree: subscription.", tool_treatment=Treatment.REJECT,
                       tool_amounts=_pm(0, 0, 0)),
        ReviewDecision(adj_id="A-2", reviewer="k.osei", timestamp="2026-09-01T10:05:00Z", treatment=Treatment.ACCEPT,
                       amounts=_pm(0, 30000, 10000), rationale="Agrees to the letter.", tool_treatment=Treatment.ACCEPT,
                       tool_amounts=_pm(0, 30000, 10000),
                       question_updates={"Q-A-1-2": "ANSWERED: insurer letter received"}),
        ReviewDecision(adj_id="A-3", reviewer="m.reyes", timestamp="2026-09-02T09:00:00Z", treatment=Treatment.REVISE,
                       amounts=_pm(0, "12000.25", 0), rationale="SOW shows two months of one-time configuration.",
                       tool_treatment=Treatment.REJECT, tool_amounts=_pm(0, 0, 0),
                       correction_type=CorrectionType.JUDGMENT_DIFFERENCE),
    ]


RECON_ACCOUNTS = [("4000", "Service Revenue", "-310000.00"), ("6000", "Salaries & Wages - Office", "118500.25"),
                  ("6200", "Insurance", "8200.00")]
PERIOD_MONTHS = {"FY2024": ("2024-01", "2024-12"), "FY2025": ("2025-01", "2025-12"),
                 "TTM Jun-26": ("2025-07", "2026-06")}


def _recon_gl(i: int, acct: str, base: str) -> Decimal:
    return D(base) + (D(-1000) * i if acct == "4000" else D(0))


def _net_income(label: str) -> str:
    """Net income per the GL = the negated debit-positive total of the reconciliation's GL side."""
    start, end = PERIOD_MONTHS[label]
    total = sum((_recon_gl(i, acct, base) for i, m in enumerate(month_range("2024-01", "2026-06"))
                 if start <= m <= end for acct, _, base in RECON_ACCOUNTS), Decimal(0))
    return fmt(-total)


GL_COMPONENTS = {
    "FY2024": (_net_income("FY2024"), "210000.00", "45000.00", "380000.00", "25000.00"),
    "FY2025": (_net_income("FY2025"), "195000.00", "52000.00", "395000.00", "25000.00"),
    "TTM Jun-26": (_net_income("TTM Jun-26"), "180000.00", "50000.00", "401000.00", "25000.00"),
}
MGMT_DIFF = _pm(0, -25000, -25000)


def _build_bridge(assessments: list[AdjustmentAssessment], final: dict[str, dict[str, str]],
                  diff: dict[str, str] = MGMT_DIFF) -> EbitdaBridge:
    """Mirror of SPEC §5.6 so the export test does not depend on qoe.bridge."""

    def per(fn) -> dict[str, str]:
        return {p: fmt(fn(p)) for p in LABELS}

    items = [a for a in assessments if a.source == "diligence"]
    assessments = [a for a in assessments if a.source != "diligence"]
    comp = {p: [D(x) for x in GL_COMPONENTS[p]] for p in LABELS}
    gl = per(lambda p: sum(comp[p], Decimal(0)))
    rows = [BridgeRow(key=k, label=lab, kind="component", amounts=per(lambda p, i=i: comp[p][i]))
            for i, (k, lab) in enumerate([("net_income", "Net income (per GL)"), ("interest", "Interest expense, net"),
                                          ("taxes", "Income taxes"), ("depreciation", "Depreciation"),
                                          ("amortization", "Amortization")])]
    rows.append(BridgeRow(key="gl_ebitda", label="Reported EBITDA (per GL)", kind="subtotal", amounts=gl))
    rows.append(BridgeRow(key="mgmt_recon_diff", label="Difference to management's reported EBITDA", kind="memo",
                          amounts=dict(diff)))
    mgmt_rep = per(lambda p: D(gl[p]) + D(diff[p]))
    rows.append(BridgeRow(key="mgmt_reported_ebitda", label="Reported EBITDA (per management)", kind="subtotal",
                          amounts=mgmt_rep))
    for a in assessments:
        rows.append(BridgeRow(key=f"mgmt:{a.adj_id}", label=f"{a.title} (as claimed)", kind="mgmt_adjustment",
                              adj_id=a.adj_id, amounts=dict(a.claimed)))
    mgmt_total = per(lambda p: sum((D(a.claimed[p]) for a in assessments), Decimal(0)))
    rows.append(BridgeRow(key="mgmt_total", label="Total management adjustments", kind="subtotal", amounts=mgmt_total))
    mgmt_adj = per(lambda p: D(mgmt_rep[p]) + D(mgmt_total[p]))
    rows.append(BridgeRow(key="mgmt_adjusted_ebitda", label="Management adjusted EBITDA", kind="subtotal",
                          amounts=mgmt_adj))
    rows.append(BridgeRow(key="dil_recon", label="Reverse unsupported reporting difference (to GL)",
                          kind="diligence_adjustment", amounts=per(lambda p: -D(diff[p]))))
    dil = {a.adj_id: per(lambda p, a=a: D(final[a.adj_id].get(p)) - D(a.claimed[p])) for a in assessments}
    for a in assessments:
        rows.append(BridgeRow(key=f"dil:{a.adj_id}", label=f"{a.title}: diligence revision",
                              kind="diligence_adjustment", adj_id=a.adj_id, amounts=dil[a.adj_id]))
    for a in items:  # diligence-identified items carry their final amount (claimed is zero)
        dil[a.adj_id] = per(lambda p, a=a: D(final[a.adj_id].get(p)))
        rows.append(BridgeRow(key=f"dil:{a.adj_id}", label=f"Diligence-identified: {a.title}",
                              kind="diligence_adjustment", adj_id=a.adj_id, amounts=dil[a.adj_id]))
    dil_total = per(lambda p: -D(diff[p]) + sum((D(v[p]) for v in dil.values()), Decimal(0)))
    rows.append(BridgeRow(key="dil_total", label="Total diligence adjustments", kind="subtotal", amounts=dil_total))
    rows.append(BridgeRow(key="diligence_adjusted_ebitda", label="Diligence adjusted EBITDA", kind="subtotal",
                          amounts=per(lambda p: D(mgmt_adj[p]) + D(dil_total[p]))))
    pending = [a for a in assessments if not final[a.adj_id]]
    rows.append(BridgeRow(key="pending", label="Memo: management adjustments pending information (excluded)",
                          kind="memo", amounts=per(lambda p: sum((D(a.claimed[p]) for a in pending), Decimal(0)))))
    return EbitdaBridge(period_labels=LABELS, rows=rows)


def _reconciliation() -> ReconciliationResult:
    items: list[ReconciliationItem] = []
    for i, month in enumerate(month_range("2024-01", "2026-06")):
        for acct, name, base in RECON_ACCOUNTS:
            gl = _recon_gl(i, acct, base)
            pl = gl
            if month == "2025-12" and acct == "6000":
                pl = gl + D(25000)  # top-side accrual not in the GL
            if month == "2024-03" and acct == "6200":
                pl = gl + D("0.40")  # rounding, within tolerance
            items.append(ReconciliationItem(month=month, account=acct, account_name=name, gl_amount=fmt(gl),
                                            pl_amount=fmt(pl), variance=fmt(pl - gl),
                                            within_tolerance=abs(pl - gl) <= D("1.00")))
    issues = [
        DataQualityIssue(code=DataQualityCode.GL_ACCOUNT_NOT_IN_PL, severity=Severity.INFO,
                         message="Account 8050 appears in the GL but not in the P&L.", account="8050"),
        DataQualityIssue(code=DataQualityCode.RECON_VARIANCE, severity=Severity.WARNING,
                         message="P&L exceeds GL by $25,000.00 for account 6000 in Dec 2025.",
                         month="2025-12", account="6000", amount="25000.00"),
        DataQualityIssue(code=DataQualityCode.MGMT_EBITDA_DIFFERS_FROM_GL, severity=Severity.CRITICAL,
                         message="Management's reported EBITDA is $25,000.00 below the GL.",
                         period_label="FY2025", amount="-25000.00"),
        DataQualityIssue(code=DataQualityCode.DUPLICATE_GL_ENTRY, severity=Severity.WARNING,
                         message="Premium bill posted twice.", month="2025-05", account="6200",
                         entry_ids=["GL-R2215", "GL-R2210"]),
    ]
    gl_ebitda = {}
    for p in LABELS:
        ni, i, t, d, am = GL_COMPONENTS[p]
        gl_ebitda[p] = EbitdaComponents(net_income=ni, interest=i, taxes=t, depreciation=d, amortization=am,
                                        ebitda=fmt(sum((D(x) for x in GL_COMPONENTS[p]), Decimal(0))))
    return ReconciliationResult(items=items, issues=issues, gl_ebitda=gl_ebitda,
                                mgmt_reported_ebitda={p: fmt(D(gl_ebitda[p].ebitda) + D(MGMT_DIFF[p])) for p in LABELS},
                                months_compared=30, accounts_compared=3,
                                variance_count=sum(1 for it in items if not it.within_tolerance))


def _meta() -> DealMeta:
    return DealMeta(
        deal_id="harborview_supply",
        target_name="Harborview Plumbing Supply, Inc.",
        industry="Plumbing distribution",
        periods=[PeriodDef(label="FY2024", start="2024-01", end="2024-12"),
                 PeriodDef(label="FY2025", start="2025-01", end="2025-12"),
                 PeriodDef(label="TTM Jun-26", start="2025-07", end="2026-06")],
        data_start="2024-01",
        data_end="2026-06",
        files=DealFiles(gl="gl/general_ledger.csv", chart_of_accounts="gl/chart_of_accounts.csv",
                        monthly_pl="financials/monthly_pl.xlsx", adjustments="adjustments/schedule.xlsx"),
    )


def make_workpaper(reviews: list[ReviewDecision] | None = None, diligence: bool = False) -> Workpaper:
    assessments = _assessments() + ([_diligence_item()] if diligence else [])
    reviews = _reviews() if reviews is None else reviews
    wp = Workpaper(
        run_id="run-test-001",
        tool_version="0.1.0",
        created_at="2026-09-28T12:00:00Z",
        ai_mode="rules",
        deal=_meta(),
        input_hashes={"gl/general_ledger.csv": "a" * 64, "documents/3.1 letter.pdf": "b" * 64},
        ingest_notes=["Detected GL format qbo_gl_csv."],
        reconciliation=_reconciliation(),
        doc_facts=[
            DocFacts(doc_id=RETAINER_LETTER, doc_type="engagement_letter", counterparty="Harrow Legal LLP",
                     doc_date="2023-03-01", is_signed=True,
                     terms=[TermFact(kind="retainer", text="$2,500 per month",
                                     quote=_q(RETAINER_LETTER, "a monthly retainer of $2,500 continuing until "
                                                               "terminated by either party"))]),
            DocFacts(doc_id=COMP_DRAFT, doc_type="contract", is_draft=True, is_signed=False, dropped_quotes=2,
                     extractor="rules"),
            DocFacts(doc_id=LIT_INVOICE, doc_type="invoice", counterparty="Harrow Legal LLP", doc_date="2025-02-10",
                     reference_numbers=["25-114", "7781"], service_period_start="2025-01-01",
                     service_period_end="2025-01-31"),
        ],
        assessments=assessments,
        reviews=reviews,
        bridge=EbitdaBridge(period_labels=LABELS, rows=[]),
    )
    final = _fallback_final_amounts(wp, latest_reviews(wp))
    return wp.model_copy(update={"bridge": _build_bridge(assessments, final)})


def make_package(wp: Workpaper) -> DealPackage:
    entries = [
        GLEntry(entry_id="GL-R1201", date="2025-02-14", period="2025-02", account="6400",
                account_name="Legal & Professional Fees", txn_type="Bill", doc_number="25-114",
                counterparty="Harrow Legal LLP", memo="Matter 7781 Pryor v. Harborview - Jan", amount="30000.00",
                source_file="gl/general_ledger.csv", source_row=1201),
        GLEntry(entry_id="GL-R880", date="2024-06-03", period="2024-06", account="6400",
                account_name="Legal & Professional Fees", counterparty="Harrow Legal LLP",
                memo="=General retainer - Jun", amount="2500.00", source_file="gl/general_ledger.csv", source_row=880),
    ]
    claims = [AdjustmentClaim(adj_id=a.adj_id, title=a.title, category_raw="Non-recurring",
                              category=a.category, description=f"Management narrative for {a.title}.",
                              gl_accounts=["6400"], support_refs=["DR 3.1", "DR 3.3"], amounts=dict(a.claimed),
                              source_row=10 + i) for i, a in enumerate(wp.assessments)]
    return DealPackage(
        deal_dir="data/dev/harborview_supply",
        meta=wp.deal,
        accounts={"6400": Account(number="6400", name="Legal & Professional Fees", source_type="Expense",
                                  ebitda_class=EbitdaClass.OPEX)},
        gl=entries,
        pl=ManagementPL(source_file="financials/monthly_pl.xlsx", months=[], lines=[]),
        schedule=ManagementSchedule(source_file="adjustments/schedule.xlsx", period_labels=LABELS, adjustments=claims),
        documents=[],
    )


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _find_row(ws, col: int, predicate) -> int:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if v is not None and predicate(v):
            return r
    raise AssertionError(f"row not found in {ws.title} col {col}")


def _bridge_rows(ws) -> dict[str, int]:
    """Map bridge row labels to sheet rows (column B, indentation stripped)."""
    return {str(ws.cell(row=r, column=2).value).strip(): r for r in range(7, ws.max_row + 1)
            if ws.cell(row=r, column=2).value}


@pytest.fixture(scope="module")
def wp() -> Workpaper:
    return make_workpaper()


@pytest.fixture(scope="module")
def exported(wp, tmp_path_factory) -> Path:
    return export_workpaper(wp, tmp_path_factory.mktemp("xlsx"))


@pytest.fixture(scope="module")
def recalculated(exported, tmp_path_factory) -> tuple[Path, dict]:
    if find_recalc_script() is None or shutil.which("soffice") is None:
        pytest.skip("LibreOffice recalc not available")
    copy = tmp_path_factory.mktemp("recalc") / exported.name
    shutil.copy(exported, copy)
    return copy, recalc_and_check(copy, timeout=120)


# ---------------------------------------------------------------------------
# Structure
# ---------------------------------------------------------------------------


def test_filename_and_sheet_order(wp, exported):
    assert exported.name == "QoE_Evidence_Review_harborview_supply.xlsx" == workbook_filename(wp)
    wb = load_workbook(exported)
    assert wb.sheetnames == [
        "Cover", "EBITDA Bridge", "Adjustment Summary",
        "Adj A-1", "Adj A-2", "Adj A-3", "Adj A-4", "Adj A-5", "Adj A-6",
        "Open Questions", "GL-P&L Reconciliation", "Data Quality", "Review Log",
    ]
    assert all(len(n) <= 31 for n in wb.sheetnames)


def test_explicit_file_path(wp, tmp_path):
    out = export_workpaper(wp, tmp_path / "nested" / "custom.xlsx")
    assert out == tmp_path / "nested" / "custom.xlsx" and out.is_file()
    assert not list(out.parent.glob("*.partial.xlsx"))


def test_sanitize_sheet_name():
    used: set[str] = set()
    assert sanitize_sheet_name("Adj A/1:[x]*?", used) == "Adj A-1--x---"
    long = sanitize_sheet_name("Adj " + "Z" * 40, used)
    assert len(long) == 31
    dup = sanitize_sheet_name("Adj " + "z" * 40, used)
    assert len(dup) <= 31 and dup.endswith("(2)") and dup.lower() != long.lower()
    assert sanitize_sheet_name("history") == "History-"
    assert "'" not in sanitize_sheet_name("O'Brien bonus")


def test_support_sheet_names_are_unique():
    wp = make_workpaper()
    a = wp.assessments[0]
    twins = [a.model_copy(update={"adj_id": "1/2"}), a.model_copy(update={"adj_id": "1-2"})]
    names = support_sheet_names(wp.model_copy(update={"assessments": twins}))
    assert len({n.lower() for n in names.values()}) == 2


# ---------------------------------------------------------------------------
# Final amounts
# ---------------------------------------------------------------------------


def test_fallback_final_amounts_semantics(wp):
    latest = latest_reviews(wp)
    assert latest["A-3"].reviewer == "m.reyes"  # later log entry wins
    final = _fallback_final_amounts(wp, latest)
    assert final["A-3"] == _pm(0, "12000.25", 0)  # reviewer override
    assert final["A-2"] == _pm(0, 30000, 10000)  # reviewed
    assert final["A-1"] == _pm(0, 60000, "25000.50")  # unreviewed: tool proposal
    assert final["A-4"] == {}  # REQUEST_INFO: pending
    pending_review = ReviewDecision(adj_id="A-1", reviewer="r", timestamp="t", treatment=Treatment.REQUEST_INFO,
                                    amounts={}, rationale="", tool_treatment=Treatment.REVISE, tool_amounts={})
    assert _fallback_final_amounts(wp, {"A-1": pending_review})["A-1"] == {}
    # A REQUEST_INFO decision is pending even if amounts were recorded with it.
    stray = pending_review.model_copy(update={"amounts": _pm(1, 2, 3)})
    assert _fallback_final_amounts(wp, {"A-1": stray})["A-1"] == {}


def test_resolve_final_amounts_matches_fallback(wp):
    resolved = resolve_final_amounts(wp)
    fallback = _fallback_final_amounts(wp, latest_reviews(wp))
    for adj_id, amounts in fallback.items():
        expected = {p: fmt(amounts.get(p)) for p in LABELS} if amounts else {}
        assert resolved[adj_id] == expected, adj_id


# ---------------------------------------------------------------------------
# Formatting rules
# ---------------------------------------------------------------------------


def test_arial_everywhere_and_no_banned_functions(exported):
    wb = load_workbook(exported)
    formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                assert cell.font.name == "Arial", (ws.title, cell.coordinate)
                if cell.data_type == "f":
                    formulas += 1
                    assert not BANNED.search(cell.value), (ws.title, cell.coordinate, cell.value)
                    # Every cross-sheet reference is quoted.
                    for m in re.finditer(r"('(?:[^']|'')+'|[A-Za-z0-9_.&\-]+)!", cell.value):
                        assert m.group(1).startswith("'"), cell.value
    assert formulas > 100
    assert wb._fonts[0].name == "Arial"  # default font for unstyled cells


def test_merged_ranges_never_overlap(wp):
    # Excel refuses to open ("repair" prompt) a sheet whose merged ranges overlap; LibreOffice does not care.
    wb = build_workbook(wp, pkg=make_package(wp))
    for ws in wb.worksheets:
        seen: set[tuple[int, int]] = set()
        for rng in ws.merged_cells.ranges:
            cells = {(r, c) for r in range(rng.min_row, rng.max_row + 1) for c in range(rng.min_col, rng.max_col + 1)}
            assert not cells & seen, (ws.title, str(rng))
            seen |= cells


def test_bridge_subtotals_are_formulas_and_rows_are_values(wp, exported):
    ws = load_workbook(exported)[SHEET_BRIDGE]
    rows = _bridge_rows(ws)
    by_key = {r.key: r for r in wp.bridge.rows}
    for key, bridge_row in by_key.items():
        r = rows[bridge_row.label]
        for k in range(3):
            cell = ws.cell(row=r, column=3 + k)
            if bridge_row.kind == "subtotal":
                assert cell.data_type == "f" and cell.value.startswith("="), (key, cell.value)
                assert cell.font.color.rgb.endswith("000000")
            else:
                assert isinstance(cell.value, (int, float)), (key, cell.value)
                assert D(cell.value) == D(bridge_row.amounts[LABELS[k]])
                assert cell.font.color.rgb.endswith("0000FF")
            assert cell.number_format == NUMBER_FORMAT
    gl = rows["Reported EBITDA (per GL)"]
    assert ws.cell(row=gl, column=3).value == f"=SUM(C{rows['Net income (per GL)']}:C{rows['Amortization']})"
    mgmt_rep = ws.cell(row=rows["Reported EBITDA (per management)"], column=4).value
    diff_row = rows["Difference to management's reported EBITDA"]
    assert mgmt_rep == f"=D{gl}+D{diff_row}"
    assert ws.freeze_panes == "C7"
    assert ws.cell(row=6, column=3).value == "FY2024\nUSD"


def test_treatment_fills_and_unreviewed_status(exported):
    ws = load_workbook(exported)[SHEET_SUMMARY]
    header = {ws.cell(row=6, column=c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row=6, column=c).value}
    tool_col, status_col, rev_col = header["Tool treatment"], header["Status"], header["Reviewer treatment"]
    fills = {"ACCEPT": "C6EFCE", "REVISE": "FFEB9C", "REJECT": "FFC7CE", "REQUEST_INFO": "D6DCE4"}
    status = {}
    for r in range(8, 14):
        ref = ws.cell(row=r, column=1).value
        t = ws.cell(row=r, column=tool_col)
        assert t.fill.start_color.rgb.endswith(fills[t.value])
        status[ref] = ws.cell(row=r, column=status_col)
    assert {k: v.value for k, v in status.items()} == {
        "A-1": "UNREVIEWED", "A-2": "AGREED", "A-3": "OVERRIDDEN",
        "A-4": "UNREVIEWED", "A-5": "UNREVIEWED", "A-6": "UNREVIEWED",
    }
    assert status["A-1"].font.bold and status["A-1"].fill.start_color.rgb.endswith("FFFF00")
    assert ws.cell(row=10, column=rev_col).value == "REVISE"  # A-3 latest decision
    assert ws.cell(row=8, column=rev_col).value in (None, "")


def test_summary_links_and_difference_formulas(exported):
    ws = load_workbook(exported)[SHEET_SUMMARY]
    # A-1: claimed FY2025 links to the support sheet tie-out, green.
    claimed = ws.cell(row=8, column=5)
    assert claimed.value.startswith("='Adj A-1'!E") and claimed.font.color.rgb.endswith("008000")
    diff = ws.cell(row=8, column=4 + 9)  # first "Final less claimed" column
    assert diff.value == "=IF(ISNUMBER(J8),J8,0)-D8"
    total = _find_row(ws, 2, lambda v: v == "Total")
    assert ws.cell(row=total, column=4).value == "=SUM(D8:D13)"
    assert ws.freeze_panes == "D8"


def test_support_sheet_blocks(exported):
    ws = load_workbook(exported)["Adj A-1"]
    col_a = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    text = [str(v) for v in col_a if v]
    facts = next(i for i, v in enumerate(text) if v.startswith("DOCUMENTED FACTS"))
    judgment = next(i for i, v in enumerate(text) if v.startswith("JUDGMENT QUESTIONS"))
    assert facts < judgment
    for heading in ("CONCLUSION", "MANAGEMENT'S CLAIM", "TIE-OUT BY PERIOD", "WALK: CLAIMED (a) TO TOOL PROPOSED",
                    "FLAGS (2)", "LINKED GL ENTRIES (4)", "DOCUMENTS AND VERBATIM QUOTES (3)",
                    "RECURRENCE OBSERVATIONS (1)", "REVIEWER DECISION"):
        assert any(v.startswith(heading) for v in text), heading
    # The ticked GL listing is the last block, so its header can repeat on every printed page.
    def at(prefix: str) -> int:
        return next(i for i, v in enumerate(text) if v.startswith(prefix))

    assert at("REVIEWER DECISION") < at("LINKED GL ENTRIES") and text[-1].startswith("GL-R")
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert '"Total due: $30,000.00"' in values
    assert f"{LIT_INVOICE}, p. 2" in values
    assert "GL rows 1201, 1410" in values
    assert "UNREVIEWED" in values
    # Tie-out differences are formulas; inputs are blue values.
    tie = _find_row(ws, 1, lambda v: str(v).startswith("(a) Claimed"))
    assert ws.cell(row=tie, column=5).font.color.rgb.endswith("0000FF")
    diff = _find_row(ws, 1, lambda v: str(v).startswith("(a) - (b)"))
    assert ws.cell(row=diff, column=4).value == f"=D{tie}-D{tie + 1}"
    check = _find_row(ws, 1, lambda v: str(v).startswith("Check: EBITDA Bridge"))
    assert ws.cell(row=check, column=4).value.startswith("='EBITDA Bridge'!C")


def test_pending_item_and_decision_history(exported):
    wb = load_workbook(exported)
    ws = wb["Adj A-4"]
    proposed = _find_row(ws, 1, lambda v: str(v).startswith("(d) Tool proposed"))
    assert [ws.cell(row=proposed, column=4 + k).value for k in range(3)] == ["Pending"] * 3
    ws3 = wb["Adj A-3"]
    values = [str(c.value) for row in ws3.iter_rows() for c in row if c.value is not None]
    assert "Superseded" in values and "Current" in values
    log = wb["Review Log"]
    assert [log.cell(row=r, column=4).value for r in (8, 9, 10)] == ["A-3", "A-2", "A-3"]
    assert log.cell(row=8, column=log.max_column - 1).value == "Superseded"


def test_stale_review_is_flagged(tmp_path):
    stale = ReviewDecision(adj_id="A-1", reviewer="k.osei", timestamp="2026-08-01T09:00:00Z",
                           treatment=Treatment.ACCEPT, amounts=_pm(0, 90000, 50000), rationale="Looked fine.",
                           tool_treatment=Treatment.ACCEPT, tool_amounts=_pm(0, 90000, 50000))
    wb = build_workbook(make_workpaper(reviews=[stale]))
    log = wb["Review Log"]
    assert log.cell(row=8, column=log.max_column).value == "Yes: re-review"
    values = [str(c.value) for row in wb["Adj A-1"].iter_rows() for c in row if c.value is not None]
    assert any("tool proposal has changed since this decision" in v for v in values)
    summary = wb[SHEET_SUMMARY]
    status_col = next(c for c in range(1, summary.max_column + 1) if summary.cell(row=6, column=c).value == "Status")
    assert summary.cell(row=8, column=status_col).value == "AGREED"


def test_open_questions_apply_review_updates_and_keep_formula_like_text(exported):
    ws = load_workbook(exported)["Open Questions"]
    headers = [ws.cell(row=6, column=c).value for c in range(1, 8)]
    assert headers == ["Q id", "Ref", "Question", "Priority", "Basis", "Status", "Response"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(7, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    assert ws.cell(row=rows["Q-A-1-2"], column=6).value == "ANSWERED"
    assert ws.cell(row=rows["Q-A-1-2"], column=7).value == "insurer letter received"
    q = ws.cell(row=rows["Q-A-3-1"], column=3)
    assert q.data_type == "s" and q.value.startswith("=Provide")


def test_gl_detail_with_package(wp, tmp_path):
    out = export_workpaper(wp, tmp_path, pkg=make_package(wp))
    ws = load_workbook(out)["Adj A-1"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "6400 Legal & Professional Fees" in values
    assert "Management narrative for Litigation legal fees (Matter 7781)." in values
    assert "=General retainer - Jun" in values  # memo stays text, never a formula
    without = load_workbook(export_workpaper(wp, tmp_path / "plain"))["Adj A-1"]
    plain = [str(c.value) for row in without.iter_rows() for c in row if c.value is not None]
    assert any("without the source GL" in v for v in plain)


def test_cover_banner_and_legend(exported):
    ws = load_workbook(exported)[SHEET_COVER]
    banner = ws["A3"]
    assert banner.value.startswith("SYNTHETIC") and banner.font.bold
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    for t in ("ACCEPT", "REVISE", "REJECT", "REQUEST_INFO", "UNREVIEWED"):
        assert t in values
    assert sum(v.startswith("Tickmark (") for v in values) == len(TICKMARKS)
    assert "run-test-001" in values and "a" * 64 in values


# ---------------------------------------------------------------------------
# Bridge planning
# ---------------------------------------------------------------------------


def test_plan_subtotals_standard_bridge(wp):
    plans = _plan_subtotals(wp.bridge.rows, LABELS)
    assert all(ok for _, ok in plans.values())
    keys = [r.key for r in wp.bridge.rows]
    terms = {keys[i]: sorted(keys[j] for j, _ in t) for i, (t, _) in plans.items()}
    assert terms["mgmt_adjusted_ebitda"] == ["mgmt_reported_ebitda", "mgmt_total"]
    assert terms["diligence_adjusted_ebitda"] == ["dil_total", "mgmt_adjusted_ebitda"]
    assert "pending" not in terms


def test_plan_subtotals_handles_opposite_memo_sign(wp):
    rows = list(wp.bridge.rows)
    i = next(k for k, r in enumerate(rows) if r.key == "mgmt_recon_diff")
    rows[i] = rows[i].model_copy(update={"amounts": {p: fmt(-D(v)) for p, v in rows[i].amounts.items()}})
    plans = _plan_subtotals(rows, LABELS)
    j = next(k for k, r in enumerate(rows) if r.key == "mgmt_reported_ebitda")
    terms, ok = plans[j]
    assert ok and (i, -1) in terms


def test_bridge_that_does_not_foot_is_flagged(wp, tmp_path):
    rows = list(wp.bridge.rows)
    i = next(k for k, r in enumerate(rows) if r.key == "gl_ebitda")
    rows[i] = rows[i].model_copy(update={"amounts": {p: fmt(D(v) + 5) for p, v in rows[i].amounts.items()}})
    bad = wp.model_copy(update={"bridge": wp.bridge.model_copy(update={"rows": rows})})
    ws = build_workbook(bad)[SHEET_BRIDGE]
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any(v.startswith("WARNING: a subtotal in red") for v in values)


def test_empty_workpaper_exports(wp, tmp_path):
    empty = wp.model_copy(update={
        "assessments": [], "reviews": [], "doc_facts": [],
        "bridge": EbitdaBridge(period_labels=LABELS, rows=[]),
        "reconciliation": wp.reconciliation.model_copy(update={"items": [], "issues": []}),
    })
    out = export_workpaper(empty, tmp_path)
    wb = load_workbook(out)
    assert wb.sheetnames[:3] == ["Cover", "EBITDA Bridge", "Adjustment Summary"]
    assert wb.sheetnames[3] == "Open Questions"


# ---------------------------------------------------------------------------
# Diligence-identified items (SPEC §5.7) and assessment-carried narrative
# ---------------------------------------------------------------------------

D1_TITLE = "Reverse duplicate premium posting (KRI-25-0507)"


@pytest.fixture(scope="module")
def dwp() -> Workpaper:
    agree = ReviewDecision(adj_id="D-1", reviewer="k.osei", timestamp="2026-09-02T10:00:00Z",
                           treatment=Treatment.REVISE, amounts=_pm(0, 18400, 0), rationale="AP shows one payment.",
                           tool_treatment=Treatment.REVISE, tool_amounts=_pm(0, 18400, 0))
    return make_workpaper(reviews=[*_reviews(), agree], diligence=True)


@pytest.fixture(scope="module")
def dexported(dwp, tmp_path_factory) -> Path:
    return export_workpaper(dwp, tmp_path_factory.mktemp("dxlsx"))


def _values(ws) -> list[str]:
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]


def _col_by_header(ws, title: str) -> int:
    return next(c for c in range(1, ws.max_column + 1) if ws.cell(row=6, column=c).value == title)


def test_diligence_item_sheet_and_summary_block(dexported):
    wb = load_workbook(dexported)
    assert wb.sheetnames[3:11] == ["Adj A-1", "Adj A-2", "Adj A-3", "Adj A-4", "Adj A-5", "Adj A-6", "Adj D-1",
                                   "Open Questions"]
    ws = wb[SHEET_SUMMARY]
    labels = {ws.cell(row=r, column=2).value: r for r in range(8, ws.max_row + 1) if ws.cell(row=r, column=2).value}
    mgmt_total, dil_total, total = (labels["Total management adjustments"], labels["Total diligence-identified items"],
                                    labels["Total"])
    d1 = _find_row(ws, 1, lambda v: v == "D-1")
    assert mgmt_total == 14 and ws.cell(row=mgmt_total, column=4).value == "=SUM(D8:D13)"
    assert mgmt_total < d1 < dil_total < total
    assert ws.cell(row=d1 - 1, column=1).value.startswith("Diligence-identified items (not on management's schedule)")
    assert ws.cell(row=d1, column=1).hyperlink.location == "'Adj D-1'!A1"
    assert ws.cell(row=dil_total, column=4).value == f"=SUM(D{d1}:D{d1})"
    assert ws.cell(row=total, column=4).value == f"=D{mgmt_total}+D{dil_total}"
    assert ws.cell(row=d1, column=_col_by_header(ws, "Status")).value == "AGREED"
    assert ws.cell(row=d1, column=_col_by_header(ws, "Tool treatment")).value == "REVISE"


def test_diligence_item_support_sheet_uses_assessment_fields(dexported):
    values = _values(load_workbook(dexported)["Adj D-1"])
    assert any(v.startswith("DILIGENCE-IDENTIFIED ITEMS (NOT ON MANAGEMENT'S SCHEDULE)") for v in values)
    assert "MANAGEMENT'S CLAIM" not in values
    assert "Premium installment KRI-25-0507 is posted twice in May 2025; the P&L carries both." in values
    assert "6200" in values and PREMIUM_NOTICE in values
    assert "(a) Claimed by management (not on the schedule: zero)" in values
    assert "Q-D-1-1" in values


def test_management_claim_falls_back_to_assessment_fields():
    wp = make_workpaper()
    a0 = wp.assessments[0].model_copy(update={"description": "Litigation narrative carried by the run.",
                                               "gl_accounts": ["6400"], "support_refs": ["DR 3.1", "DR 3.3"]})
    wp = wp.model_copy(update={"assessments": [a0, *wp.assessments[1:]]})
    values = _values(build_workbook(wp)["Adj A-1"])
    assert "Litigation narrative carried by the run." in values
    assert "6400" in values and "DR 3.1; DR 3.3" in values
    assert "Schedule row" not in values
    # Without a description anywhere, the sheet says where to look.
    values = _values(build_workbook(make_workpaper())["Adj A-2"])
    assert "Not carried in the workpaper; see management's adjusted EBITDA schedule." in values
    # A schedule stored on the workpaper is the source when present.
    sched = ManagementSchedule(source_file="adjustments/schedule.xlsx", period_labels=LABELS, adjustments=[
        AdjustmentClaim(adj_id="A-1", title=a0.title, description="From the schedule.", gl_accounts=["6410"],
                        amounts=dict(a0.claimed), source_row=12)])
    values = _values(build_workbook(wp.model_copy(update={"schedule": sched}))["Adj A-1"])
    assert "From the schedule." in values and "6410" in values and "Schedule row" in values


def test_bridge_shows_diligence_items_after_management_revisions(dwp, dexported):
    ws = load_workbook(dexported)[SHEET_BRIDGE]
    rows = _bridge_rows(ws)
    d1 = rows[f"Diligence-identified: {D1_TITLE}"]
    last_revision = max(rows[f"{a.title}: diligence revision"] for a in dwp.assessments if a.source != "diligence")
    assert last_revision < d1 < rows["Total diligence adjustments"]
    assert ws.cell(row=d1 - 1, column=1).value.startswith("Diligence-identified items")
    assert ws.cell(row=d1, column=1).value == "D-1"
    assert D(ws.cell(row=d1, column=4).value) == D("18400")
    assert ws.cell(row=d1, column=6).value == "REVISE" and ws.cell(row=d1, column=7).value == "AGREED"
    assert ws.cell(row=d1, column=8).value == "Adj D-1"
    assert not any(k.endswith("(as claimed)") and D1_TITLE in k for k in rows)
    total = ws.cell(row=rows["Total diligence adjustments"], column=3).value
    assert total.startswith("=SUM(C") and total.endswith(f":C{d1})")
    assert "Diligence-identified items less their total final amounts, Adjustment Summary" in rows


def test_bridge_display_rows_reorders_and_drops_zero_claims(dwp):
    rows = list(dwp.bridge.rows)
    # An older bridge: D-1 ahead of the management revisions, plus an all-zero "as claimed" row for it.
    d1 = next(r for r in rows if r.key == "dil:D-1")
    rows.remove(d1)
    rows.insert(next(k for k, r in enumerate(rows) if r.key == "dil:A-1"), d1)
    rows.insert(next(k for k, r in enumerate(rows) if r.key == "mgmt_total"),
                BridgeRow(key="mgmt:D-1", label="D-1 (as claimed)", kind="mgmt_adjustment", amounts=_pm(0, 0, 0)))
    out = bridge_display_rows(rows, {"D-1"})
    keys = [r.key for r in out]
    assert "mgmt:D-1" not in keys
    assert keys.index("dil:D-1") == keys.index("dil_total") - 1
    assert keys.index("dil:A-6") < keys.index("dil:D-1")
    assert all(ok for _, ok in _plan_subtotals(out, LABELS).values())
    # Without diligence items the order is untouched.
    assert bridge_display_rows(dwp.bridge.rows, set()) == list(dwp.bridge.rows)


def test_cover_counts_diligence_items_separately(dexported):
    ws = load_workbook(dexported)[SHEET_COVER]
    values = _values(ws)
    assert "ADJUSTMENT STATUS: MANAGEMENT ADJUSTMENTS" in values
    assert "DILIGENCE-IDENTIFIED ITEMS (NOT ON MANAGEMENT'S SCHEDULE)" in values
    assert "6 on management's schedule; 1 identified by diligence (not on the schedule)" in values
    assert any("2 of 6 management adjustments reviewed; 1 of 1 diligence-identified items reviewed" in v
               for v in values)
    # The diligence table counts the summary's diligence block (row 17), the management table rows 8-13; the
    # tool treatment column follows the four amount groups and the pending-memo group (5 x 3 periods).
    tool = get_column_letter(4 + 5 * 3)
    header = _find_row(ws, 1, lambda v: v == "DILIGENCE-IDENTIFIED ITEMS (NOT ON MANAGEMENT'S SCHEDULE)")
    revise = header + 3
    assert ws.cell(row=revise, column=1).value == "REVISE"
    assert ws.cell(row=revise, column=2).value == f"=COUNTIF('Adjustment Summary'!${tool}$17:${tool}$17,\"REVISE\")"
    mgmt_revise = _find_row(ws, 1, lambda v: v == "REVISE")
    assert ws.cell(row=mgmt_revise, column=2).value == f"=COUNTIF('Adjustment Summary'!${tool}$8:${tool}$13,\"REVISE\")"


@pytest.fixture(scope="module")
def drecalculated(dexported, tmp_path_factory) -> tuple[Path, dict]:
    if find_recalc_script() is None or shutil.which("soffice") is None:
        pytest.skip("LibreOffice recalc not available")
    copy = tmp_path_factory.mktemp("drecalc") / dexported.name
    shutil.copy(dexported, copy)
    return copy, recalc_and_check(copy, timeout=120)


def test_recalc_with_diligence_item_ties(dwp, drecalculated):
    path, result = drecalculated
    assert result.get("status") == "success", result
    assert result["total_errors"] == 0
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_BRIDGE]
    rows = _bridge_rows(ws)
    for r in dwp.bridge.rows:
        for k, p in enumerate(LABELS):
            got = ws.cell(row=rows[r.label], column=3 + k).value
            assert abs(D(got) - D(r.amounts[p])) <= D("0.01"), (r.key, p, got)
    final = resolve_final_amounts(dwp)
    assert final["D-1"] == _pm(0, 18400, 0)
    for k, p in enumerate(LABELS):
        gl = D(ws.cell(row=rows["Reported EBITDA (per GL)"], column=3 + k).value)
        dil = D(ws.cell(row=rows["Diligence adjusted EBITDA"], column=3 + k).value)
        assert abs(dil - (gl + sum((D(v[p]) for v in final.values() if v), Decimal(0)))) <= D("0.01")
    checks = [r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=2).value or "").endswith(
        ("Adjustment Summary", "Adjustment Summary)"))]
    assert len(checks) == 3
    for r in checks:
        for k in range(3):
            assert abs(D(ws.cell(row=r, column=3 + k).value)) < D("0.01")
    cover = wb[SHEET_COVER]
    assert cover.cell(row=_find_row(cover, 1, lambda v: str(v).startswith("Workbook checks")), column=2).value == "OK"
    header = _find_row(cover, 1, lambda v: v == "DILIGENCE-IDENTIFIED ITEMS (NOT ON MANAGEMENT'S SCHEDULE)")
    counts = {cover.cell(row=r, column=1).value: tuple(cover.cell(row=r, column=c).value for c in (2, 3, 4))
              for r in range(header + 2, header + 6)}
    assert counts == {"ACCEPT": (0, 0, 0), "REVISE": (1, 1, 1), "REJECT": (0, 0, 0), "REQUEST_INFO": (0, 0, 0)}
    agreed = _find_row(cover, 1, lambda v: v == "Reviewed: agreed with tool")
    unreviewed = _find_row(cover, 1, lambda v: str(v).startswith("UNREVIEWED ("))
    assert (cover.cell(row=agreed, column=2).value, cover.cell(row=agreed, column=3).value) == (1, 1)
    assert (cover.cell(row=unreviewed, column=2).value, cover.cell(row=unreviewed, column=3).value) == (4, 0)
    summary = wb[SHEET_SUMMARY]
    total = _find_row(summary, 2, lambda v: v == "Total")
    fy25_final = 4 + 6 + 1
    assert D(summary.cell(row=total, column=fy25_final).value) == D("60000") + D("30000") + D("12000.25") - D(
        "30000") + D("18400")
    sws = wb["Adj D-1"]
    check = _find_row(sws, 1, lambda v: str(v).startswith("Check: EBITDA Bridge"))
    assert [abs(D(sws.cell(row=check, column=4 + k).value)) < D("0.01") for k in range(3)] == [True] * 3


# ---------------------------------------------------------------------------
# LibreOffice recalculation
# ---------------------------------------------------------------------------


def test_recalc_has_zero_errors(recalculated):
    _, result = recalculated
    assert result.get("status") == "success", result
    assert result["total_errors"] == 0
    assert result["total_formulas"] > 100


def test_recalculated_bridge_matches_workpaper(wp, recalculated):
    path, _ = recalculated
    ws = load_workbook(path, data_only=True)[SHEET_BRIDGE]
    rows = _bridge_rows(ws)
    for r in wp.bridge.rows:
        for k, p in enumerate(LABELS):
            got = ws.cell(row=rows[r.label], column=3 + k).value
            assert abs(D(got) - D(r.amounts[p])) <= D("0.01"), (r.key, p, got, r.amounts[p])
    # Identity: diligence adjusted = GL EBITDA + final amounts (pending excluded).
    final = resolve_final_amounts(wp)
    for k, p in enumerate(LABELS):
        gl = D(ws.cell(row=rows["Reported EBITDA (per GL)"], column=3 + k).value)
        dil = D(ws.cell(row=rows["Diligence adjusted EBITDA"], column=3 + k).value)
        expected = gl + sum((D(v[p]) for v in final.values() if v), Decimal(0))
        assert abs(dil - expected) <= D("0.01")
    checks = [r for r in range(1, ws.max_row + 1)
              if str(ws.cell(row=r, column=2).value or "").startswith(("Diligence adjusted EBITDA less",
                                                                        "Total management adjustments less"))]
    assert len(checks) == 2
    for r in checks:
        for k in range(3):
            assert abs(D(ws.cell(row=r, column=3 + k).value)) < D("0.01")


def test_recalculated_summary_cover_and_support_checks(wp, recalculated):
    path, _ = recalculated
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_SUMMARY]
    total = _find_row(ws, 2, lambda v: v == "Total")
    # Final totals exclude pending A-4; Final less claimed reverses its claim.
    assert D(ws.cell(row=total, column=4 + 6 + 1).value) == D("60000") + D("30000") + D("12000.25") - D("30000")
    a4 = 11
    assert ws.cell(row=a4, column=4 + 6).value == "Pending"
    assert D(ws.cell(row=a4, column=4 + 9).value) == D("-250000")
    cover = wb[SHEET_COVER]
    counts = {}
    for r in range(1, cover.max_row + 1):
        label = cover.cell(row=r, column=1).value
        if label in ("ACCEPT", "REVISE", "REJECT", "REQUEST_INFO") and cover.cell(row=r, column=2).value is not None \
                and label not in counts:
            counts[label] = tuple(cover.cell(row=r, column=c).value for c in (2, 3, 4))
    assert counts == {"ACCEPT": (1, 1, 1), "REVISE": (2, 1, 3), "REJECT": (2, 0, 1), "REQUEST_INFO": (1, 0, 1)}
    unreviewed = _find_row(cover, 1, lambda v: str(v).startswith("UNREVIEWED ("))
    assert cover.cell(row=unreviewed, column=2).value == 4
    open_q = _find_row(cover, 1, lambda v: v == "Open questions for management")
    assert cover.cell(row=open_q, column=2).value == 3
    critical = _find_row(cover, 1, lambda v: v == "Data quality issues: critical")
    assert cover.cell(row=critical, column=2).value == 1
    checks = _find_row(cover, 1, lambda v: str(v).startswith("Workbook checks"))
    assert cover.cell(row=checks, column=2).value == "OK"
    for a in wp.assessments:
        sws = wb[f"Adj {a.adj_id}"]
        check = _find_row(sws, 1, lambda v: str(v).startswith("Check: EBITDA Bridge"))
        for k in range(3):
            assert abs(D(sws.cell(row=check, column=4 + k).value)) < D("0.01"), a.adj_id


def test_recalculated_reconciliation_summary(wp, recalculated):
    path, _ = recalculated
    ws = load_workbook(path, data_only=True)["GL-P&L Reconciliation"]
    first = _find_row(ws, 1, lambda v: v == "Month") + 1
    items = wp.reconciliation.items
    months = sorted({it.month for it in items})
    for i, m in enumerate(months):
        r = first + i
        gl = sum((D(it.gl_amount) for it in items if it.month == m), Decimal(0))
        pl = sum((D(it.pl_amount) for it in items if it.month == m), Decimal(0))
        assert abs(D(ws.cell(row=r, column=2).value) - gl) <= D("0.01"), m
        assert abs(D(ws.cell(row=r, column=4).value) - (pl - gl)) <= D("0.01"), m
        assert ws.cell(row=r, column=5).value == (1 if m == "2025-12" else 0), m


def test_recalculated_checks_catch_a_stale_bridge(wp, tmp_path):
    """A decision logged after the bridge was built must trip the workbook checks, not pass silently."""
    if find_recalc_script() is None or shutil.which("soffice") is None:
        pytest.skip("LibreOffice recalc not available")
    late = ReviewDecision(adj_id="A-1", reviewer="m.reyes", timestamp="2026-09-03T08:00:00Z",
                          treatment=Treatment.REVISE, amounts=_pm(0, 70000, "25000.50"), rationale="Added a fee.",
                          tool_treatment=Treatment.REVISE, tool_amounts=_pm(0, 60000, "25000.50"),
                          correction_type=CorrectionType.TOOL_MISSED_EVIDENCE)
    stale = wp.model_copy(update={"reviews": [*wp.reviews, late]})  # bridge not rebuilt
    out = export_workpaper(stale, tmp_path)
    result = recalc_and_check(out, timeout=120)
    assert result.get("status") == "success" and result["total_errors"] == 0, result
    wb = load_workbook(out, data_only=True)
    cover = wb[SHEET_COVER]
    checks = _find_row(cover, 1, lambda v: str(v).startswith("Workbook checks"))
    assert cover.cell(row=checks, column=2).value.startswith("DIFFERENCE")
    bridge = wb[SHEET_BRIDGE]
    row = _find_row(bridge, 2, lambda v: str(v).startswith("Diligence adjusted EBITDA less"))
    assert abs(D(bridge.cell(row=row, column=4).value) + D("10000")) < D("0.01")


# ---------------------------------------------------------------------------
# Review findings (excel-*): audit trail, ties, and presentation
# ---------------------------------------------------------------------------


def _row_of(ws, text: str, col: int = 1) -> int:
    """First row whose column-``col`` text starts with ``text``."""
    return _find_row(ws, col, lambda v: str(v).startswith(text))


def _period_values(ws, row: int, first_col: int = 4) -> list:
    return [ws.cell(row=row, column=first_col + k).value for k in range(len(LABELS))]


def _listing(ws) -> dict[str, dict[str, object]]:
    """The linked GL listing: entry id -> {header: value}."""
    header = _find_row(ws, 1, lambda v: v == "Entry ID")
    titles = {c: str(ws.cell(row=header, column=c).value) for c in range(1, ws.max_column + 1)
              if ws.cell(row=header, column=c).value}
    out = {}
    for r in range(header + 1, ws.max_row + 1):
        eid = ws.cell(row=r, column=1).value
        if eid:
            out[eid] = {t: ws.cell(row=r, column=c).value for c, t in titles.items()} | {"_row": r}
    return out


def _recalc_copy(wp: Workpaper, tmp_path: Path, pkg: DealPackage | None = None):
    if find_recalc_script() is None or shutil.which("soffice") is None:
        pytest.skip("LibreOffice recalc not available")
    out = export_workpaper(wp, tmp_path, pkg=pkg)
    result = recalc_and_check(out, timeout=120)
    assert result.get("status") == "success" and result["total_errors"] == 0, result
    return load_workbook(out, data_only=True)


def _sheet_check(ws) -> tuple[object, object]:
    """(check total rounded to the cent, status) of a support sheet."""
    r = _row_of(ws, "Sheet checks")
    return round(float(ws.cell(row=r, column=3).value), 2), ws.cell(row=r, column=4).value


def _cover_status(cover, label: str) -> object:
    return cover.cell(row=_row_of(cover, label), column=2).value


# -- excel-flag-impact-column -----------------------------------------------


def test_flags_show_effect_by_period_and_context_amounts_apart():
    wp = make_workpaper()
    search = wp.assessments[1].model_copy(update={"flags": [
        Flag(code=FlagCode.EXCESS_GL_ACTIVITY, severity=Severity.INFO, period_label="FY2025", amount_impact="533490.00",
             message="Linked activity exceeds the claim; the claimed subset ties exactly."),
    ]})
    wb = build_workbook(wp.model_copy(update={"assessments": [wp.assessments[0], search, *wp.assessments[2:]]}))
    ws = wb["Adj A-1"]
    header = _row_of(ws, "Severity")
    assert _period_values(ws, header) == [f"Effect on (d)\n{p}" for p in LABELS]
    recurring = _find_row(ws, 3, lambda v: v == "Recurring pattern")
    assert _period_values(ws, recurring) == [None, D("-30000.00"), D("-30000.00")]
    assert ws.cell(row=recurring, column=2).value == "F1"
    # A multi-period removal is never labelled "All periods" and never left blank.
    assert not any(v == "All periods" for v in _values(ws))
    ws2 = wb["Adj A-2"]
    excess = _find_row(ws2, 3, lambda v: v == "Excess GL activity")
    assert _period_values(ws2, excess) == [None, None, None]  # not an EBITDA effect
    noted = ws2.cell(row=excess, column=4 + len(LABELS)).value
    what = ws2.cell(row=excess, column=5 + len(LABELS)).value
    assert D(noted) == D("533490") and what.startswith("Unclaimed context activity (FY2025)")
    assert ws2.cell(row=_row_of(ws2, "Severity"), column=4 + len(LABELS)).value == "Amount noted (not an effect)"


def test_walk_takes_the_claim_to_the_proposal_one_flag_per_line(wp, recalculated):
    path, _ = recalculated
    wb = load_workbook(path, data_only=True)
    for a in wp.assessments:
        ws = wb[f"Adj {a.adj_id}"]
        if a.treatment == Treatment.REQUEST_INFO:
            assert any(v.startswith("Pending (REQUEST_INFO): the tool proposes no amount") for v in _values(ws))
            continue
        walk = _row_of(ws, "WALK: CLAIMED (a) TO TOOL PROPOSED")
        total = _find_row(ws, 1, lambda v: v == "Proposed per the walk")
        assert [D(x) for x in _period_values(ws, total)] == [D(a.proposed[p]) for p in LABELS], a.adj_id
        check = _row_of(ws, "Check: walk less (d) Tool proposed")
        assert walk < total < check and all(abs(D(x)) < D("0.01") for x in _period_values(ws, check)), a.adj_id
    ws = wb["Adj A-1"]
    line = _row_of(ws, "F1 Recurring pattern")
    assert [D(x) for x in _period_values(ws, line)] == [D(0), D("-30000"), D("-30000")]


def _moved_and_mismatch_items() -> list[AdjustmentAssessment]:
    """An out-of-period move, and a period mismatch whose flag effect is measured from the traced GL (A-8)
    or from the claim (A-9): the walk must close under either convention."""
    moved = AdjustmentAssessment(
        adj_id="A-7", title="Prior-year subcontractor true-up", category=AdjustmentCategory.OUT_OF_PERIOD,
        claimed=_pm(0, 42000, 0), traced_gl=_pm(0, 42000, 0), documented=_pm(0, 42000, 0),
        proposed=_pm(-42000, 42000, 0), treatment=Treatment.REVISE,
        gl_links=[GLLink(entry_id="GL-R1350", period="2025-03", amount="42000.00", score=4.0, role="moved",
                         claimed=True)],
        flags=[Flag(code=FlagCode.OUT_OF_PERIOD, severity=Severity.WARNING, entry_ids=["GL-R1350"],
                    message="The invoice dates the service to Jul-Dec 2024.", effects={"FY2024": "-42000.00"})],
    )

    def mismatch(adj_id: str, effects: dict[str, str]) -> AdjustmentAssessment:
        return AdjustmentAssessment(
            adj_id=adj_id, title=f"Warehouse relocation ({adj_id})", category=AdjustmentCategory.NON_RECURRING,
            claimed=_pm(0, 80000, 80000), traced_gl=_pm(0, 80000, 0), documented=_pm(0, 80000, 0),
            proposed=_pm(0, 80000, 0), treatment=Treatment.REVISE,
            gl_links=[GLLink(entry_id=f"GL-R{row}", period=m, amount="40000.00", score=4.0, role="supporting",
                             claimed=True) for row, m in ((1360 + int(adj_id[-1]), "2025-02"),
                                                           (1370 + int(adj_id[-1]), "2025-03"))],
            flags=[Flag(code=FlagCode.PERIOD_MISMATCH, severity=Severity.WARNING, period_label="TTM Jun-26",
                        amount_impact="-80000.00", effects=effects,
                        message="No claimed activity falls in TTM Jun-26.")],
        )

    return [moved, mismatch("A-8", {}), mismatch("A-9", {"TTM Jun-26": "-80000.00"})]


def test_walk_and_listing_tie_moves_and_period_mismatches(tmp_path):
    items = _moved_and_mismatch_items()
    wp = make_workpaper(reviews=[])
    wp = wp.model_copy(update={"assessments": [*wp.assessments, *items]})
    final = _fallback_final_amounts(wp, {})
    wp = wp.model_copy(update={"bridge": _build_bridge(list(wp.assessments), final)})
    wb = _recalc_copy(wp, tmp_path)
    for adj in ("A-7", "A-8", "A-9"):
        assert _sheet_check(wb[f"Adj {adj}"]) == (0, "OK"), adj
    a7 = wb["Adj A-7"]
    assert _listing(a7)["GL-R1350"]["Tick"] == "M"
    assert "Moved by F1 Out of period" in _listing(a7)["GL-R1350"]["Flags"]
    oop = _row_of(a7, "Out-of-period flag effects")
    assert [D(x) for x in _period_values(a7, oop)] == [D("-42000"), D(0), D(0)]
    listing_d = _row_of(a7, "Tool proposed per the listing")
    assert [D(x) for x in _period_values(a7, listing_d)] == [D("-42000"), D("42000"), D(0)]
    # Effects measured from the traced GL: the walk shows the claim-to-GL step as its own line.
    assert any(v.startswith("(b) - (a) Claimed amount not traced") for v in _values(wb["Adj A-8"]))
    assert not any(v.startswith("(b) - (a) Claimed amount not traced") for v in _values(wb["Adj A-9"]))
    assert _cover_status(wb[SHEET_COVER], "Workbook checks") == "OK"


def _edge_items() -> list[AdjustmentAssessment]:
    """Cases the second dev deal exposed: a proposal capped at the claim, an entry claimed in one of two
    overlapping periods only, a normalization item carried at the claim's level, and a gap within the
    tie-out tolerance that no flag carries."""

    def link(row: int, month: str, amount: str, role: str = "supporting", hint: str = "", **kw) -> GLLink:
        return GLLink(entry_id=f"GL-R{row}", period=month, amount=amount, score=4.0, role=role, claimed=True,
                      reasons=[hint] if hint else [], supports_claim=role == "supporting", **kw)

    capped = AdjustmentAssessment(
        adj_id="A-10", title="Relocation (no exact fit)", category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(0, 50000, 0), traced_gl=_pm(0, 60000, 0), documented=_pm(0, 0, 0), proposed=_pm(0, 50000, 0),
        treatment=Treatment.ACCEPT,
        gl_links=[link(2400 + i, f"2025-0{i}", "20000.00") for i in (2, 3, 4)],
    )
    overlap = AdjustmentAssessment(
        adj_id="A-11", title="Consulting (claimed FY only)", category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(0, 30000, 20000), traced_gl=_pm(0, 30000, 20000), documented=_pm(0, 0, 0),
        proposed=_pm(0, 20000, 20000), treatment=Treatment.REVISE,
        gl_links=[
            link(2501, "2025-07", "10000.00", hint="Claimed in FY2025"),
            link(2502, "2025-08", "10000.00", hint="Claimed in FY2025, TTM Jun-26"),
            link(2503, "2025-09", "10000.00", role="removed", removed_by=FlagCode.RECURRING_PATTERN),
            link(2504, "2026-01", "10000.00", hint="Claimed in TTM Jun-26"),
        ],
        flags=[Flag(code=FlagCode.RECURRING_PATTERN, severity=Severity.WARNING, message="Recurs.",
                    entry_ids=["GL-R2503"], effects={"FY2025": "-10000.00"})],
    )
    level = AdjustmentAssessment(
        adj_id="A-12", title="Owner pay at the signed level", category=AdjustmentCategory.NORMALIZATION,
        claimed=_pm(250000, 250000, 250000), traced_gl=_pm(550000, 550000, 550000), documented=_pm(0, 0, 0),
        proposed=_pm(250000, 250000, 250000), treatment=Treatment.ACCEPT,
        gl_links=[link(2600 + i, m, amt) for i, (m, amt) in enumerate(
            (("2024-06", "550000.00"), ("2025-03", "275000.00"), ("2025-09", "275000.00"), ("2026-03", "275000.00")))],
    )
    rounding = AdjustmentAssessment(
        adj_id="A-13", title="Repairs (partial)", category=AdjustmentCategory.NON_RECURRING,
        claimed=_pm(0, 10000, 10000), traced_gl=_pm(0, "9318.27", "9999.22"), documented=_pm(0, 0, 0),
        proposed=_pm(0, "9318.27", "9999.22"), treatment=Treatment.REVISE,
        gl_links=[link(2701, "2025-07", "9318.27", hint="Claimed in FY2025, TTM Jun-26"),
                  link(2702, "2026-01", "680.95", hint="Claimed in TTM Jun-26")],
        flags=[Flag(code=FlagCode.PARTIAL_GL_SUPPORT, severity=Severity.WARNING, period_label="FY2025",
                    amount_impact="-681.73", effects={"FY2025": "-681.73"}, message="Part of the claim is not in the GL.")],
    )
    return [capped, overlap, level, rounding]


def test_listing_and_walk_tie_caps_overlaps_normalization_and_rounding(tmp_path):
    wp = make_workpaper(reviews=[])
    wp = wp.model_copy(update={"assessments": [*wp.assessments, *_edge_items()]})
    wp = wp.model_copy(update={"bridge": _build_bridge(list(wp.assessments), _fallback_final_amounts(wp, {}))})
    wb = _recalc_copy(wp, tmp_path)
    for adj in ("A-10", "A-11", "A-12", "A-13"):
        assert _sheet_check(wb[f"Adj {adj}"]) == (0, "OK"), adj
    a10 = wb["Adj A-10"]
    assert [D(x) for x in _period_values(a10, _row_of(a10, "Cap at the claim"))] == [D(0), D(-10000), D(0)]
    a11 = wb["Adj A-11"]
    rows = _listing(a11)
    assert rows["GL-R2501"]["TTM Jun-26\n(USD, debit +)"] is None  # claimed in FY2025 only
    assert rows["GL-R2503"]["FY2025\n(USD, debit +)"] == 10000 and rows["GL-R2503"]["TTM Jun-26\n(USD, debit +)"] is None
    a12 = wb["Adj A-12"]
    assert [D(x) for x in _period_values(a12, _row_of(a12, "Normalized level deducted in (d)"))] == [D(300000)] * 3
    assert [D(x) for x in _period_values(a12, _row_of(a12, "Level deducted less the level the claim"))] == [D(0)] * 3
    assert not any(v.startswith("Check: listing less (d)") for v in _values(a12))
    a13 = wb["Adj A-13"]
    gap = _row_of(a13, "Difference within the tie-out tolerance")
    assert [round(D(x), 2) for x in _period_values(a13, gap)] == [D(0), D(0), D("-0.78")]
    assert _cover_status(wb[SHEET_COVER], "Workbook checks") == "OK"


def test_exact_subset_prefers_the_largest_fit():
    from qoe.export_xlsx import _exact_subset

    assert _exact_subset([100, 200, 300], 600) == {0, 1, 2}
    assert _exact_subset([100, 200, 300], 300) == {0, 1}  # two entries beat one
    assert _exact_subset([100, 200, 300], 0) == set()
    assert _exact_subset([100, 200], 50) is None


# -- excel-tick-x-on-claimed-entries ----------------------------------------


def test_listing_ticks_each_entry_by_its_role_and_cites_the_flag(wp):
    wb = build_workbook(wp, pkg=make_package(wp))
    rows = _listing(wb["Adj A-1"])
    assert rows["GL-R1201"]["Tick"] == "T" and rows["GL-R1201"]["Claimed by mgmt?"] == "Yes"
    removed = rows["GL-R1500"]
    assert removed["Tick"] == "R" and removed["Claimed by mgmt?"] == "Yes"
    assert removed["Flags"].startswith("Removed by F1 Recurring pattern") and "Also cited by F2" in removed["Flags"]
    assert rows["GL-R880"]["Tick"] == "X" and rows["GL-R880"]["Claimed by mgmt?"] == "No"
    # Claimed and driving entries first: T, then R, then X.
    order = sorted(rows, key=lambda e: rows[e]["_row"])
    assert [rows[e]["Tick"] for e in order] == ["T", "T", "R", "X"]
    recovery = _listing(wb["Adj A-5"])["GL-R1250"]
    assert recovery["Tick"] == "O" and recovery["Claimed by mgmt?"] == "No"
    assert recovery["Flags"].startswith("Offset under F1 Offsetting recovery")
    erp = _listing(wb["Adj A-3"])
    assert {r["Tick"] for r in erp.values()} == {"R"} and all(r["Claimed by mgmt?"] == "Yes" for r in erp.values())
    legend = _values(wb[SHEET_COVER])
    assert any(v.startswith("Tickmark (entry role): Removed: management claimed the entry") for v in legend)
    assert any(v.startswith("Tickmark (entry role): Context only") and "not part of management's claim" in v
               for v in legend)


def _legacy(wp: Workpaper) -> Workpaper:
    """A workpaper from before GLLink.role / Flag.effects."""
    out = []
    for a in wp.assessments:
        links = [lk.model_copy(update={"role": "", "claimed": False, "removed_by": None}) for lk in a.gl_links]
        flags = [f.model_copy(update={"effects": {}}) for f in a.flags]
        out.append(a.model_copy(update={"gl_links": links, "flags": flags}))
    return wp.model_copy(update={"assessments": out})


def test_legacy_workpaper_rebuilds_ticks_and_keeps_its_checks_honest(wp, tmp_path):
    old = _legacy(wp)
    wb = build_workbook(old)
    rows = _listing(wb["Adj A-1"])
    assert rows["GL-R1500"]["Tick"] == "R" and rows["GL-R1500"]["Claimed by mgmt?"] == "Yes"  # from its reason
    assert _listing(wb["Adj A-5"])["GL-R1250"]["Tick"] == "O"  # from the recovery flag
    ws = wb["Adj A-1"]
    assert any("does not record entry roles" in v for v in _values(ws))
    assert any("does not record flag effects" in v for v in _values(ws))
    recalced = _recalc_copy(old, tmp_path)
    assert _cover_status(recalced[SHEET_COVER], "Workbook checks") == "OK"


# -- excel-gl-listing-does-not-tie-to-tieout --------------------------------


def test_listing_ties_by_period_to_traced_and_proposed(wp, recalculated):
    path, _ = recalculated
    wb = load_workbook(path, data_only=True)
    formulas = load_workbook(path)
    for a in wp.assessments:
        ws = wb[f"Adj {a.adj_id}"]
        claimed = _row_of(ws, "Claimed by management (Claimed? = Yes)")
        assert [D(x) for x in _period_values(ws, claimed)] == [D(a.traced_gl[p]) for p in LABELS], a.adj_id
        if a.treatment != Treatment.REQUEST_INFO:
            listing_d = _row_of(ws, "Tool proposed per the listing")
            assert [D(x) for x in _period_values(ws, listing_d)] == [D(a.proposed[p]) for p in LABELS], a.adj_id
        assert _sheet_check(ws) == (0, "OK"), a.adj_id
        f = formulas[f"Adj {a.adj_id}"].cell(row=_row_of(ws, "Supporting: claimed and carried"), column=5).value
        assert f.startswith("=SUMIFS($E$") and ',"T")' in f
    a1 = wb["Adj A-1"]
    assert [D(x) for x in _period_values(a1, _row_of(a1, "Removed by a flag (R)"))] == [D(0), D("30000"), D("30000")]
    a5 = wb["Adj A-5"]
    assert [D(x) for x in _period_values(a5, _row_of(a5, "Recovery or offset applied"))] == [D(0), D("-30000"), D(0)]
    # Every support sheet's check total rolls up to the Cover.
    cover = formulas[SHEET_COVER]
    area = cover.cell(row=_row_of(cover, "Support sheets:"), column=3).value
    assert all(f"'Adj {a.adj_id}'!$C$" in area for a in wp.assessments)


def test_listing_check_catches_a_traced_amount_the_entries_do_not_support(wp, tmp_path):
    search = wp.assessments[1].model_copy(update={"traced_gl": _pm(0, 35000, 10000)})
    bad = wp.model_copy(update={"assessments": [wp.assessments[0], search, *wp.assessments[2:]]})
    wb = _recalc_copy(bad, tmp_path)
    ws = wb["Adj A-2"]
    check = _row_of(ws, "Check: claimed entries less (b) Traced to GL")
    assert [D(x) for x in _period_values(ws, check)] == [D(0), D("-5000"), D(0)]
    assert _sheet_check(ws)[1].startswith("DIFFERENCE")
    assert _cover_status(wb[SHEET_COVER], "Workbook checks").startswith("DIFFERENCE")


# -- excel-doc-vouching-overstated ------------------------------------------


def test_document_ticks_vouch_each_entry_to_its_own_document(wp):
    from qoe.export_xlsx import _context

    docs = [
        DocFacts(doc_id="inv-0212", doc_type="invoice", counterparty="Harrow Legal LLP", doc_date="2025-02-12",
                 reference_numbers=["25-0212", "7781"], amounts=[AmountFact(label="total_due", amount="14500.00",
                                                                             quote=_q("inv-0212", "Total due"))]),
        DocFacts(doc_id="inv-0418", doc_type="invoice", counterparty="Harrow Legal LLP", doc_date="2025-04-18",
                 reference_numbers=["25-0418", "7781"], amounts=[AmountFact(label="total_due", amount="22000.00",
                                                                             quote=_q("inv-0418", "Total due"))]),
        DocFacts(doc_id="retainer-0701", doc_type="invoice", counterparty="Harrow Legal LLP", doc_date="2025-07-01",
                 reference_numbers=["25-0701"], amounts=[AmountFact(label="total_due", amount="2500.00",
                                                                    quote=_q("retainer-0701", "Total due"))]),
        DocFacts(doc_id="letter", doc_type="engagement_letter", is_signed=True),
        DocFacts(doc_id="draft", doc_type="contract", is_draft=True, is_signed=False),
        DocFacts(doc_id="count memo", doc_type="memo"),
        DocFacts(doc_id="claim letter", doc_type="insurance", reference_numbers=["KM-24-5521"],
                 amounts=[AmountFact(label="net", amount="30000.00", quote=_q("claim letter", "net payment"))]),
        DocFacts(doc_id="statement", doc_type="invoice", counterparty="Lakeside Leasing", doc_date="2025-11-28",
                 amounts=[AmountFact(label="due", amount="1500.00", quote=_q("statement", "Amount due"))]),
    ]
    ctx = _context(wp.model_copy(update={"doc_facts": docs}), None)

    def entry(row: int, month: str, amount: str, number: str = "", memo: str = "", cp: str = "Harrow Legal LLP"):
        e = GLEntry(entry_id=f"GL-R{row}", date=f"{month}-15", period=month, account="6400", account_name="Legal",
                    doc_number=number, counterparty=cp, memo=memo, amount=amount, source_file="gl.csv",
                    source_row=row)
        return e, GLLink(entry_id=e.entry_id, period=month, amount=amount, score=1.0)

    lit = entry(1, "2025-02", "14500.00", "25-0212")
    assert [_vouch(ctx, *lit, d) for d in ("inv-0212", "inv-0418", "letter", "draft", "count memo")] == \
        ["D", "", "A", "U", "C"]
    # One month's retainer invoice is only a sample for the other months.
    assert _vouch(ctx, *entry(2, "2025-08", "2500.00", "25-0801"), "retainer-0701") == "S"
    assert _vouch(ctx, *entry(3, "2025-07", "2500.00", "25-0701"), "retainer-0701") == "D"
    # A recovery with the claim number in its memo agrees to the settlement letter.
    recovery = entry(4, "2025-02", "-30000.00", memo="Insurance proceeds - claim KM-24-5521", cp="Keystone Mutual")
    assert _vouch(ctx, *recovery, "claim letter") == "D"
    # Same party and amount: vouched in the statement's month, a sample in any other month.
    assert _vouch(ctx, *entry(5, "2025-12", "1500.00", cp="Lakeside Leasing Co."), "statement") == "D"
    assert _vouch(ctx, *entry(6, "2025-03", "1500.00", cp="Lakeside Leasing Co."), "statement") == "S"
    # An extractor amount that is not a number never breaks the export.
    odd = DocFacts(doc_id="odd", doc_type="invoice",
                   amounts=[AmountFact(label="x", amount="n/a", quote=_q("odd", "x"))])
    odd_ctx = _context(wp.model_copy(update={"doc_facts": [odd]}), None)
    assert _vouch(odd_ctx, *entry(7, "2025-03", "1500.00"), "odd") == ""


def test_documented_rows_follow_the_ticks(wp):
    ws = build_workbook(wp, pkg=make_package(wp))["Adj A-1"]
    rows = _listing(ws)
    assert rows["GL-R1201"]["Doc tick"] == "D" and "Doc 3 D" in rows["GL-R1201"]["Documents (tick)"]
    assert rows["GL-R1500"]["Doc tick"] == "A"
    gap = _row_of(ws, "(c) Documented less the two lines above")
    formula = ws.cell(row=gap, column=5).value
    assert formula.startswith("=E") and formula.count("-E") == 2


# -- excel-bridge-mgmt-reported-derived -------------------------------------


def test_bridge_agrees_to_schedule_line_by_line_and_to_net_income(wp, tmp_path):
    rows = list(wp.bridge.rows)
    i2 = next(k for k, r in enumerate(rows) if r.key == "mgmt:A-2")
    i3 = next(k for k, r in enumerate(rows) if r.key == "mgmt:A-3")
    rows[i2], rows[i3] = (rows[i2].model_copy(update={"amounts": rows[i3].amounts}),
                          rows[i3].model_copy(update={"amounts": rows[i2].amounts}))  # totals unchanged
    mgmt_adj = next(r for r in wp.bridge.rows if r.key == "mgmt_adjusted_ebitda").amounts
    schedule = ManagementSchedule(
        source_file="adjustments/schedule.xlsx", period_labels=LABELS,
        reported_ebitda=wp.reconciliation.mgmt_reported_ebitda,
        adjusted_ebitda={p: fmt(D(v) + (1000 if p == "FY2025" else 0)) for p, v in mgmt_adj.items()},
    )
    bad = wp.model_copy(update={"bridge": wp.bridge.model_copy(update={"rows": rows}), "schedule": schedule})
    wb = _recalc_copy(bad, tmp_path)
    ws = wb[SHEET_BRIDGE]

    def checks(label: str) -> list:
        return [D(x) for x in _period_values(ws, _row_of(ws, label, col=2), first_col=3)]

    assert checks("Total management adjustments less total claimed") == [D(0)] * 3  # the total still ties
    assert checks("Each management line less its claimed amount") == [D(0), D(84000), D(52000)]
    assert checks("Check: Reported EBITDA (per management) above less the schedule") == [D(0)] * 3
    assert checks("Check: Management adjusted EBITDA above less the schedule") == [D(0), D(-1000), D(0)]
    assert checks("Check: Net income (per GL) above less the reconciliation") == [D(0)] * 3
    recon = wb[SHEET_RECON]
    ni = _find_row(recon, 1, lambda v: v == "FY2025")
    assert abs(D(recon.cell(row=ni, column=5).value) - D(GL_COMPONENTS["FY2025"][0])) < D("0.01")
    cover = wb[SHEET_COVER]
    assert _cover_status(cover, "Workbook checks").startswith("DIFFERENCE")
    assert _cover_status(cover, "Agreement to source data").startswith("DIFFERENCE")


# -- excel-sign-conventions-mixed -------------------------------------------


def test_every_amount_states_its_sign_basis(wp):
    wb = build_workbook(wp)
    support = wb["Adj A-5"]
    assert "EBITDA-signed" in support["A4"].value and "debit +" in support["A4"].value
    header = _find_row(support, 1, lambda v: v == "Entry ID")
    assert all("debit +" in str(support.cell(row=header, column=c).value) for c in range(4, 8))
    assert "EBITDA-signed" in wb[SHEET_BRIDGE]["A4"].value and "EBITDA-signed" in wb[SHEET_SUMMARY]["A4"].value
    assert "debit-positive" in wb[SHEET_RECON]["A4"].value
    dq = wb[SHEET_DATA_QUALITY]
    basis_col = _col_by_header(dq, "Sign basis of the amount")
    by_code = {dq.cell(row=r, column=2).value: dq.cell(row=r, column=basis_col).value for r in range(7, 11)}
    assert by_code["RECON_VARIANCE"].startswith("P&L less GL, debit +")
    assert by_code["MGMT_EBITDA_DIFFERS_FROM_GL"].startswith("EBITDA: management less GL")
    legend = next(v for v in _values(wb[SHEET_COVER]) if v.startswith("Parentheses"))
    assert "reduces EBITDA" in legend and "credit" in legend


# -- excel-print-unreadable -------------------------------------------------


def test_print_setup_is_legible_with_repeating_headers(wp):
    from openpyxl.utils import column_index_from_string

    from qoe.export_xlsx import _MIN_SCALE, _PAPERS, _inches

    wb = build_workbook(wp)
    printable = dict(_PAPERS)
    for ws in wb.worksheets:
        ps = ws.page_setup
        assert ps.orientation == "landscape" and ps.fitToHeight == 0, ws.title
        last_col = column_index_from_string(re.findall(r"\$([A-Z]+)\$\d+", ws.print_area)[-1])
        widths = [_inches(ws.column_dimensions[get_column_letter(c)].width or 9.0) for c in range(1, last_col + 1)]
        pages = int(ps.fitToWidth)
        repeated = 0.0
        if ws.print_title_cols:
            repeated = sum(widths[:column_index_from_string(ws.print_title_cols.split(":")[-1].strip("$"))])
        per_page = (sum(widths) + repeated * (pages - 1)) / pages
        # Fitted to the page width, body text prints at 75% or more (it printed at 3-4pt before).
        assert per_page <= printable[int(ps.paperSize)] / _MIN_SCALE + 0.01, (ws.title, per_page)
        if ws.title != SHEET_COVER:
            assert ws.print_title_rows, ws.title
    support = wb["Adj A-1"]
    header = _find_row(support, 1, lambda v: v == "Entry ID")
    assert support.print_title_rows == f"${header}:${header}"
    assert support.row_breaks.brk[-1].id == _row_of(support, "LINKED GL ENTRIES") - 1
    reasons = next(c for c in range(1, support.max_column + 1)
                   if str(support.cell(row=header, column=c).value).startswith("Link reasons"))
    printed = column_index_from_string(re.findall(r"\$([A-Z]+)\$\d+", support.print_area)[-1])
    assert reasons > printed
    assert wb[SHEET_SUMMARY].print_title_rows == "$6:$7" and wb[SHEET_BRIDGE].print_title_rows == "$6:$6"
    recon = wb[SHEET_RECON]
    detail = _row_of(recon, "ALL ACCOUNT-MONTHS COMPARED")
    assert int(re.findall(r"\d+", recon.print_area.split(":")[-1])[0]) < detail


# -- excel-normalization-proforma-tieout ------------------------------------


def test_normalization_tieout_shows_actual_cost_and_implied_level(wp):
    wb = build_workbook(wp)
    ws = wb["Adj A-4"]
    values = _values(ws)
    a = _row_of(ws, "(a) Claimed by management: actual cost less the normalized level")
    b = _row_of(ws, "(b) Actual cost in the GL (linked entries)")
    level = _row_of(ws, "Normalized level implied by the claim: (b) - (a)")
    assert ws.cell(row=level, column=4).value == f"=D{b}-D{a}"
    note = ws.cell(row=level, column=4 + len(LABELS)).value
    assert COMP_DRAFT in note and "DRAFT" in note
    assert not any(v.startswith("(a) - (b) Claimed less traced") for v in values)
    summary = wb[SHEET_SUMMARY]
    supp = _col_by_header(summary, "# supporting GL links (T)")
    assert summary.cell(row=_find_row(summary, 1, lambda v: v == "A-4"), column=supp).value == "n/a"


def test_pro_forma_tieout_says_the_cost_is_still_in_the_gl():
    wp = make_workpaper()
    pf = wp.assessments[3].model_copy(update={"category": AdjustmentCategory.PRO_FORMA, "title": "Pro forma savings"})
    swapped = wp.model_copy(update={"assessments": [*wp.assessments[:3], pf, *wp.assessments[4:]]})
    ws = build_workbook(swapped)["Adj A-4"]
    values = _values(ws)
    assert "(b) Cost still in the GL (run-rate saving not yet realized)" in values
    assert "(a) - (b) Claimed saving less cost still in the GL" in values
    assert "(b) - (c) Cost still in the GL without document support" in values


# -- excel-pending-tool-diff-misleading -------------------------------------


def test_pending_item_shows_pending_difference_and_a_provisional_memo(wp, recalculated):
    ws = build_workbook(wp)["Adj A-4"]
    diff = _row_of(ws, "(d) - (a) Tool proposed less claimed")
    assert _period_values(ws, diff) == ["Pending"] * 3
    memo = _row_of(ws, "Memo: provisional amount while pending")
    assert _period_values(ws, memo) == [D("250000.00")] * 3
    assert "tool rationale" in ws.cell(row=memo, column=4 + len(LABELS)).value
    path, _ = recalculated
    summary = load_workbook(path, data_only=True)[SHEET_SUMMARY]
    prov = 4 + 4 * len(LABELS)
    a4 = _find_row(summary, 1, lambda v: v == "A-4")
    assert [D(summary.cell(row=a4, column=prov + k).value) for k in range(3)] == [D(250000)] * 3
    total = _find_row(summary, 2, lambda v: v == "Total")
    assert D(summary.cell(row=total, column=prov).value) == D(250000)
    cover = load_workbook(path)[SHEET_COVER]
    assert any(v.startswith("Final treatment: the reviewer's decision") for v in _values(cover))
    assert "Final treatment" in _values(cover) and "Carried (final)" not in _values(cover)


def test_provisional_amount_parsing():
    labels = ["FY2024", "FY2025", "TTM Jun-26"]
    text = ("REQUEST_INFO: x. Provisional amount the evidence would support: FY2024 0 / FY2025 (42,000) / "
            "TTM Jun-26 165,000.")
    assert _provisional_from_rationale(text, labels) == {"FY2024": D(0), "FY2025": D(-42000), "TTM Jun-26": D(165000)}
    assert _provisional_from_rationale("REQUEST_INFO: FY2024 1 / FY2025 2 / TTM Jun-26 3", labels) is None
    assert _provisional_from_rationale("Provisional: FY2024 1 / FY2025 2", labels) is None


# -- excel-pending-without-question -----------------------------------------


def test_pending_item_without_an_open_question_gets_a_request(tmp_path):
    hold = ReviewDecision(adj_id="A-6", reviewer="m.reyes", timestamp="2026-09-03T10:00:00Z",
                          treatment=Treatment.REQUEST_INFO, amounts={}, rationale="Need the lender's payoff statement.",
                          tool_treatment=Treatment.REJECT, tool_amounts=_pm(0, 0, 0),
                          correction_type=CorrectionType.NEW_INFORMATION)
    answered = ReviewDecision(adj_id="A-2", reviewer="k.osei", timestamp="2026-09-03T11:00:00Z",
                              treatment=Treatment.ACCEPT, amounts=_pm(0, 30000, 10000), rationale="Agrees.",
                              tool_treatment=Treatment.ACCEPT, tool_amounts=_pm(0, 30000, 10000),
                              question_updates={"Q-A-4-1": "ANSWERED: sent the draft again"})
    wp = make_workpaper(reviews=[hold, answered])
    wb = build_workbook(wp)
    qs = wb[SHEET_QUESTIONS]
    rows = {qs.cell(row=r, column=1).value: r for r in range(7, qs.max_row + 1) if qs.cell(row=r, column=1).value}
    assert "Q-A-6-P" in rows and "Q-A-4-P" in rows
    text = qs.cell(row=rows["Q-A-6-P"], column=3).value
    assert "Need the lender's payoff statement." in text and qs.cell(row=rows["Q-A-6-P"], column=6).value == "OPEN"
    assert "REQUEST_INFO" not in qs.cell(row=rows["Q-A-4-P"], column=3).value  # drafted from the tool rationale
    summary = wb[SHEET_SUMMARY]
    qcol = _col_by_header(summary, "# open questions")
    a6 = summary.cell(row=_find_row(summary, 1, lambda v: v == "A-6"), column=qcol)
    assert a6.value == 1 and a6.fill.start_color.rgb.endswith("FFC7CE")
    cover = wb[SHEET_COVER]
    r = _row_of(cover, "Pending items with no open question")
    assert cover.cell(row=r, column=2).value == 2 and "A-4, A-6" in cover.cell(row=r, column=3).value
    held = _values(wb["Adj A-6"])
    assert any(v.startswith("The item is pending but no question to management is open") for v in held)
    # A pending item that already has an open question gets nothing extra.
    assert "Q-A-4-P" not in _values(build_workbook(make_workpaper())[SHEET_QUESTIONS])


# -- excel-open-questions-internal-text -------------------------------------


def test_open_questions_read_as_plain_requests(wp):
    extra = [
        OpenQuestion(q_id="Q-A-3-10", adj_id="A-3", basis="ai:draft_questions",
                     text="Which entries make up the TTM Jun-26 claim of 36,000? 924 combinations tie; we assumed "
                          "Jul 2025-Dec 2025."),
        OpenQuestion(q_id="Q-A-3-2", adj_id="A-3", basis="RECURRING_PATTERN; CONTINUING_OBLIGATION",
                     text="The memo describes the cost as recurring. Removed ERP subscription: (72,000) in FY2025. "
                          "How does management reconcile this (NORMALIZATION_BENCHMARK_MISSING)?"),
    ]
    erp = wp.assessments[2].model_copy(update={"open_questions": [*wp.assessments[2].open_questions, *extra]})
    ws = build_workbook(wp.model_copy(update={"assessments": [*wp.assessments[:2], erp, *wp.assessments[3:]]}))[
        SHEET_QUESTIONS]
    rows = [r for r in range(7, ws.max_row + 1) if str(ws.cell(row=r, column=1).value or "").startswith("Q-A-3")]
    assert [ws.cell(row=r, column=1).value for r in rows] == ["Q-A-3-1", "Q-A-3-2", "Q-A-3-10"]
    texts = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=3).value for r in rows}
    assert texts["Q-A-3-10"] == "Which entries make up the TTM Jun-26 claim of 36,000?"
    assert texts["Q-A-3-2"] == ("The memo describes the cost as recurring. How does management reconcile this "
                                "(basis for the normalized level)?")
    bases = [str(ws.cell(row=r, column=5).value) for r in range(7, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert not any("_" in b or "ai:" in b for b in bases)
    assert ws.cell(row=rows[2], column=5).value == "Follow-up from the document review"
    assert ws.cell(row=rows[1], column=5).value == "Similar costs in other periods; Ongoing contract terms"


def test_plain_basis_and_management_text_helpers():
    assert _plain_basis("documented fact") == "Documented fact"
    assert _plain_basis("SOME_NEW_CODE") == "Some new code"
    assert _management_text("Only a request.") == "Only a request."
    assert _management_text("Removed everything.") == "Removed everything."  # never emptied


# -- excel-summary-diff-total-vs-bridge -------------------------------------


def test_summary_difference_reconciles_to_the_bridge(wp, recalculated):
    path, _ = recalculated
    ws = load_workbook(path, data_only=True)[SHEET_SUMMARY]
    diff = 4 + 3 * len(LABELS)
    rows = {k: _find_row(ws, 2, lambda v, k=k: str(v).startswith(k)) for k in (
        "Final less claimed, all items", "Reverse unsupported reporting difference", "Total diligence adjustments per",
        "Check: the two lines above")}
    items, recon, bridge, check = (rows[k] for k in rows)
    bridge_rows = {r.key: r for r in wp.bridge.rows}
    for k, p in enumerate(LABELS):
        assert D(ws.cell(row=recon, column=diff + k).value) == D(bridge_rows["dil_recon"].amounts[p])
        assert D(ws.cell(row=bridge, column=diff + k).value) == D(bridge_rows["dil_total"].amounts[p])
        assert abs(D(ws.cell(row=check, column=diff + k).value)) < D("0.01")
    total = _find_row(ws, 2, lambda v: v == "Total")
    for title in ("# GL links", "# docs", "# supporting GL links (T)"):
        assert ws.cell(row=total, column=_col_by_header(ws, title)).value is None  # would double count
    assert ws.cell(row=total, column=_col_by_header(ws, "# open questions")).value == 3


# -- excel-recurrence-table-misleading --------------------------------------


def test_recurrence_shows_the_claim_beside_comparable_activity(wp):
    ws = build_workbook(wp)["Adj A-1"]
    comparable = _row_of(ws, "Comparable activity outside the claim", col=2)
    claimed = _row_of(ws, "Claimed by management in the same group", col=2)
    assert _period_values(ws, comparable) == [D(30000), D(0), D(0)]
    assert _period_values(ws, claimed) == [D(0), D(30000), D(30000)]
    note = ws.cell(row=comparable - 1, column=4 + len(LABELS)).value
    assert note == "Retainer billed every month. GL row 880" and ".;" not in note


# -- excel-truncated-text ---------------------------------------------------


def test_row_heights_allow_for_long_unbroken_text():
    path = "documents/5.1 Castellano Executive Employment Agreement DRAFT.pdf"
    assert _wrapped_lines(path, 30) == 3  # word wrap, not characters / width
    assert _wrapped_lines("x" * 95, 30) == 4  # an unbroken token breaks by characters
    assert _wrapped_lines("a\nb", 30) == 2
    wp = make_workpaper()
    long_path = "documents/" + "Barrow Search Partners Invoice BSP-25-0412 " * 3 + "final.pdf"
    ws = build_workbook(wp.model_copy(update={"input_hashes": {long_path: "c" * 64}}))[SHEET_COVER]
    r = _find_row(ws, 1, lambda v: v == long_path)
    span = next(m for m in ws.merged_cells.ranges if m.min_row == r and m.min_col == 1)
    width = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(span.min_col, span.max_col + 1))
    assert span.max_col >= 3 and (ws.row_dimensions[r].height or 15) >= 12.75 * _wrapped_lines(long_path, int(width))


# -- excel-no-signoff-draft-marking -----------------------------------------


def test_signoff_cells_and_draft_marking(wp):
    wb = build_workbook(wp)
    cover = wb[SHEET_COVER]
    assert "SIGN-OFF" in _values(cover)
    for role in ("Prepared by", "Reviewed by", "Approved by (engagement manager)"):
        r = _find_row(cover, 1, lambda v, role=role: v == role)
        assert cover.cell(row=r, column=2).fill.start_color.rgb.endswith("FFF2CC")
        assert cover.cell(row=r, column=4).fill.start_color.rgb.endswith("FFF2CC")
    for ws in wb.worksheets[1:]:
        row5 = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Prepared by" in row5 and "Reviewed by" in row5, ws.title
        assert "DRAFT: 4 of 6 items UNREVIEWED" in ws["A3"].value, ws.title
    assert cover["A3"].value.startswith("SYNTHETIC") and "DRAFT: 4 of 6 items UNREVIEWED" in cover["A3"].value
    bridge = wb[SHEET_BRIDGE]
    final = _bridge_rows(bridge)["Diligence adjusted EBITDA"]
    assert "DRAFT" in [bridge.cell(row=final, column=c).value for c in range(3, bridge.max_column + 1)]
    assert "DRAFT" in bridge.oddFooter.left.text
    everyone = [ReviewDecision(adj_id=a.adj_id, reviewer="k.osei", timestamp="2026-09-04T09:00:00Z",
                               treatment=a.treatment, amounts=dict(a.proposed), rationale="Agree.",
                               tool_treatment=a.treatment, tool_amounts=dict(a.proposed)) for a in wp.assessments]
    done = build_workbook(make_workpaper(reviews=everyone))
    assert "DRAFT until signed off: all 6 items have a reviewer decision" in done[SHEET_COVER]["A3"].value
    real = make_workpaper().model_copy(update={"deal": _meta().model_copy(update={"synthetic": False})})
    banner = build_workbook(real)[SHEET_COVER]["A3"].value
    assert banner.startswith("DRAFT:") and "SYNTHETIC" not in banner


def test_integration_dev_deal_exports_and_ties(tmp_path):
    """End to end on the reference dev deal (other data/dev packages may be mid-authoring)."""
    engine = pytest.importorskip("qoe.engine")
    ingest = pytest.importorskip("qoe.ingest")
    deal_dir = ROOT / "data" / "dev" / "meridian_mechanical"
    if not (deal_dir / "deal.yaml").is_file():
        pytest.skip("no generated dev deal at data/dev/meridian_mechanical")
    wp = engine.run_review(deal_dir, run_id="export-test", created_at="2026-01-01T00:00:00Z")
    out = export_workpaper(wp, tmp_path, pkg=ingest.load_deal(deal_dir))
    result = recalc_and_check(out, timeout=180)
    if result.get("status") == "skipped":
        return
    assert result.get("status") == "success" and result["total_errors"] == 0, result
    ws = load_workbook(out, data_only=True)[SHEET_BRIDGE]
    rows = _bridge_rows(ws)
    labels = wp.bridge.period_labels
    for r in wp.bridge.rows:
        if r.kind != "subtotal":
            continue
        for k, p in enumerate(labels):
            got = ws.cell(row=rows[r.label], column=3 + k).value
            assert abs(D(got) - D(r.amounts.get(p))) <= D("0.01"), (r.key, p, got)
